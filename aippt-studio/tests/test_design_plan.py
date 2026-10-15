"""Fresh independent Work fixtures, never allocated in the production work tree."""
import json
from pathlib import Path

import pytest
from typer.testing import CliRunner

from aippt_studio.design_plan import PREFIX, author_design_plan, design_plan_status
from aippt_studio.maintenance import initialize_authoring_context
from aippt_studio.models import ChangeMode, DeckPlan, JobBrief, JobState, PresentationForm, PresentationIntent, ProductionRoute, SlideSpec
from aippt_studio.storage import job_dir, read_json, save_job, write_json
from aippt_studio.workflow import create_job

SCENARIOS = {
    PresentationForm.GENERAL: ('兴趣分享', '初次接触者', '理解主题', '主题引入—具体例子—理解回顾', '中心观点与例子'),
    PresentationForm.GOVERNMENT_BRIEFING: ('防汛项目调度', '协调部门', '明确责任与进度', '现状—区域风险—责任事项—时序', '区域风险与责任对应'),
    PresentationForm.BUSINESS_REPORT: ('供应链经营复盘', '运营负责人', '判断改善行动', '经营结果—差异证据—原因—行动', '目标与实际的同口径比较'),
    PresentationForm.COMPETITION_DEFENSE: ('职业规划大赛', '评委', '理解成长路径', '职业目标—自我与岗位证据—行动路径', '经历与岗位能力要求对应'),
    PresentationForm.ACADEMIC_DEFENSE: ('材料实验答辩', '学术评审', '判断方法与贡献', '问题—方法—对照结果—局限', '对照条件与实验结果对应'),
    PresentationForm.TEACHING: ('儿童水循环课程', '小学生', '理解变化过程', '观察现象—示例解释—练习—回顾', '现象到解释的学习步骤'),
    PresentationForm.TECHNICAL_EXPLANATION: ('分布式架构说明', '工程师', '理解接口和约束', '边界—组件—数据流—限制', '组件之间的数据传递'),
    PresentationForm.CEREMONIAL_STORY: ('校庆回顾', '师生校友', '建立共同经历的认同', '记忆场景—人物—共同经历—祝愿', '人物与场景的叙事呼应'),
    PresentationForm.PREMIUM_BRAND_PORTFOLIO: ('摄影作品集', '潜在合作方', '理解作品方向与能力', '代表作品—创作意图—系列联系—合作信息', '作品主视觉与创作说明'),
}


def fresh(tmp_path, monkeypatch, form=PresentationForm.BUSINESS_REPORT, number=1, mode=ChangeMode.NEW_COMPOSITION):
    work = tmp_path / 'work' / str(number)
    (work / 'input').mkdir(parents=True)
    (work / 'workspace').mkdir()
    topic, audience, _, _, relation = SCENARIOS[form]
    (work / 'input/request.md').write_text(f'{topic}；听众：{audience}；独立任务 {number}')
    write_json(work / 'work.json', {'id': str(number), 'name': f'独立{number}', 'storage_profile': 'minimal'})
    initialize_authoring_context(tmp_path, work)
    monkeypatch.setenv('AIPPT_JOBS_DIR', str(work / 'workspace/studio-jobs'))
    monkeypatch.delenv('AIPPT_WORK_DIR', raising=False)
    monkeypatch.setattr('aippt_studio.native_scene.WORKSPACE_ROOT', tmp_path)
    job = create_job('隔离验证', job_id='production')
    job.state = JobState.PROTOTYPE_REVIEW
    job.route = ProductionRoute.NATIVE
    save_job(job)
    root = job_dir(job.id)
    brief = JobBrief(goal=topic, audience=audience, presentation_intent=PresentationIntent(presentation_form=form, change_mode=mode))
    plan = DeckPlan(title=f'{topic}{number}', slides=[SlideSpec(id='p1', title=f'{topic}核心页', purpose=relation, layout='content')])
    write_json(root / 'planning/brief.json', brief)
    write_json(root / 'planning/deck-plan.json', plan)
    return job, root, plan, work


def decisions(form, plan):
    _, _, need, story, relation = SCENARIOS[form]
    return {'narrative_intent': {'task_specific_type': form.value, 'audience_need': need,
            'core_message': plan.title, 'storyline': story, 'evidence_boundary': '仅本次材料'},
            'whole_deck_reference_decision': {'mode': 'new_composition', 'reason': '本次结构由现有材料和听众需要决定'},
            'page_design': {s.id: {'page_job': s.purpose, 'content_relationship': relation,
              'reference_decision': {'mode': 'new_composition', 'reason': '只需表达此页关系，不强行匹配模板'}} for s in plan.slides}}


@pytest.mark.parametrize('form', list(PresentationForm))
def test_each_fresh_form_requires_its_own_plan_and_native_binding(tmp_path, monkeypatch, form):
    from aippt_studio.native_scene import set_native_scene
    job, root, plan, work = fresh(tmp_path, monkeypatch, form)
    scene = {'slides': [{'id': 'p1', 'elements': [{'type': 'text', 'text': plan.slides[0].title, 'options': {}}]}]}
    assert design_plan_status(root)['status'] == 'needs_design_plan'
    with pytest.raises(ValueError, match=PREFIX):
        set_native_scene(job.id, scene)
    assert not (root / 'planning/native-scene.json').exists()
    result = author_design_plan(job.id, decisions(form, plan))
    assert result['status'] == 'recorded_for_build_not_visual_approval'
    assert set_native_scene(job.id, scene)['slides'] == 1
    context = read_json(root / 'planning/design-context.json')
    assert context['form'] == form.value
    assert context['narrative_intent']['task_specific_type'] == form.value
    assert context['narrative_intent']['storyline'] == SCENARIOS[form][3]
    assert 'Work 68' not in json.dumps(context['narrative_intent'], ensure_ascii=False)


def test_new_work_does_not_inherit_other_work_decisions(tmp_path, monkeypatch):
    first, root1, plan1, _ = fresh(tmp_path, monkeypatch, number=1)
    author_design_plan(first.id, decisions(PresentationForm.BUSINESS_REPORT, plan1))
    _, root2, _, _ = fresh(tmp_path, monkeypatch, PresentationForm.TEACHING, number=2)
    assert design_plan_status(root1)['status'] == 'recorded_for_build_not_visual_approval'
    assert design_plan_status(root2)['status'] == 'needs_design_plan'
    assert not read_json(root2 / 'planning/deck-plan.json')['planning_strategy']


@pytest.mark.parametrize('changed', ['request', 'brief', 'copy', 'decision'])
def test_current_task_changes_invalidate_recorded_plan(tmp_path, monkeypatch, changed):
    job, root, plan, work = fresh(tmp_path, monkeypatch)
    author_design_plan(job.id, decisions(PresentationForm.BUSINESS_REPORT, plan))
    if changed == 'request':
        (work / 'input/request.md').write_text('改为当前不同目标')
    elif changed == 'brief':
        path = root / 'planning/brief.json'; data = read_json(path); data['audience'] = '小学生'; write_json(path, data)
    else:
        path = root / 'planning/deck-plan.json'; data = read_json(path)
        if changed == 'copy': data['slides'][0]['points'] = ['当前更新的材料']
        else: data['planning_strategy']['narrative_intent']['core_message'] = '被静默改写的主张'
        write_json(path, data)
    assert design_plan_status(root)['status'] == 'needs_design_plan'


@pytest.mark.parametrize('mode', [ChangeMode.LOCAL_PATCH, ChangeMode.FAITHFUL_RECONSTRUCTION])
def test_bounded_existing_work_does_not_require_whole_deck_replanning(tmp_path, monkeypatch, mode):
    _, root, _, _ = fresh(tmp_path, monkeypatch, mode=mode)
    assert not design_plan_status(root)['required']


def test_cli_records_plan_and_only_resumes_design_pause(tmp_path, monkeypatch):
    from aippt_studio.cli import app
    from aippt_studio.storage import load_job
    job, root, plan, _ = fresh(tmp_path, monkeypatch)
    job.state = JobState.BLOCKED; job.resume_state = JobState.PROTOTYPE_REVIEW; job.blocked_reason = PREFIX+'待规划'; save_job(job)
    request = tmp_path / 'author-plan.json'; write_json(request, decisions(PresentationForm.BUSINESS_REPORT, plan))
    result = CliRunner().invoke(app, ['design-plan', job.id, '--request', str(request), '--json'])
    assert result.exit_code == 0, result.output
    assert load_job(job.id).state == JobState.PROTOTYPE_REVIEW
    assert not load_job(job.id).approvals['prototype']
    job = load_job(job.id); job.state = JobState.BLOCKED; job.resume_state = JobState.PROTOTYPE_REVIEW; job.blocked_reason = 'source corrupted'; save_job(job)
    with pytest.raises(ValueError, match='non-design'):
        author_design_plan(job.id, decisions(PresentationForm.BUSINESS_REPORT, plan))


def test_normal_run_stops_before_renderer_without_author_plan(tmp_path, monkeypatch):
    from aippt_studio import workflow
    from aippt_studio.ingest import ingest_sources
    from aippt_studio.storage import load_job
    job, root, _, _ = fresh(tmp_path, monkeypatch)
    job.state = JobState.INTAKE; save_job(job)
    (root / 'planning/brief.json').unlink()
    (root / 'planning/deck-plan.json').unlink()
    source = tmp_path / 'source.md'; source.write_text('# 教学材料\n' + '\n'.join([
        '水循环包含蒸发、凝结与降水，课堂先观察杯壁的水滴，再比较水的三种状态。',
        '学习者用生活中的晾衣、烧水和雨后积水说明水会发生哪些变化。',
        '一组连续图像展示水从地面进入空气、形成云，再返回地面的过程。',
        '小组活动先描述观察到的现象，再根据记录解释温度变化产生的影响。',
        '练习题让学生给步骤排序，并说明为什么不能把云看成气态的水。',
        '课程最后用一张关系图回顾过程，保留仍需通过实验解决的问题。']))
    ingest_sources(job.id, [str(source)])
    initial = workflow.run_job(job.id)
    assert initial['state'] == 'brief_review', initial
    workflow.approve_gate(job.id, 'brief')
    monkeypatch.setattr(workflow, 'render_html', lambda *a, **k: pytest.fail('rendered before author planning'))
    result = workflow.run_job(job.id)
    assert result['state'] == 'blocked'
    assert load_job(job.id).blocked_reason.startswith(PREFIX)
    assert result['authoring_status']['status'] == 'needs_design_plan'


def test_old_work_is_compatible(tmp_path, monkeypatch):
    _, root, _, work = fresh(tmp_path, monkeypatch)
    record = read_json(work / 'work.json'); record.pop('authoring_workflow_version'); write_json(work / 'work.json', record)
    assert design_plan_status(root)['status'] == 'legacy_or_unmanaged'
    before = (root / 'planning/deck-plan.json').read_bytes()
    assert author_design_plan('production')['status'] == 'legacy_or_unmanaged'
    with pytest.raises(ValueError, match='versioned new Work'):
        author_design_plan('production', {})
    assert (root / 'planning/deck-plan.json').read_bytes() == before


def test_identical_plan_preserves_current_approvals_and_build(tmp_path, monkeypatch):
    from aippt_studio.storage import load_job
    job, root, plan, _ = fresh(tmp_path, monkeypatch)
    request = decisions(PresentationForm.BUSINESS_REPORT, plan)
    author_design_plan(job.id, request)
    current = load_job(job.id); current.approvals['prototype'] = True; save_job(current)
    candidate = root / 'build/prototype.pptx'; candidate.write_bytes(b'previously inspected fixture')
    assert author_design_plan(job.id, request)['changed'] is False
    assert candidate.read_bytes() == b'previously inspected fixture'
    assert load_job(job.id).approvals['prototype']
    current = load_job(job.id); current.state = JobState.BLOCKED; current.resume_state = JobState.PROTOTYPE_REVIEW
    current.blocked_reason = PREFIX+'此前需求漂移，现已恢复'; save_job(current)
    assert author_design_plan(job.id, request)['resumed']
    assert load_job(job.id).state == JobState.PROTOTYPE_REVIEW
    assert candidate.is_file() and load_job(job.id).approvals['prototype']


@pytest.mark.parametrize('route', [ProductionRoute.NATIVE, ProductionRoute.TEMPLATE, ProductionRoute.REBUILD])
def test_all_build_routes_stop_before_engine_for_missing_design(tmp_path, monkeypatch, route):
    from types import SimpleNamespace
    from aippt_studio import workflow
    job, root, plan, _ = fresh(tmp_path, monkeypatch)
    job.route = route; job.state = JobState.FULL_BUILD; save_job(job)
    monkeypatch.setattr(workflow, 'audit_job', lambda *_: SimpleNamespace(status='consistent'))
    monkeypatch.setattr(workflow, '_approval_is_current', lambda *_: True)
    monkeypatch.setattr(workflow, 'route_ready', lambda *_: (True, 'fixture'))
    monkeypatch.setattr(workflow, '_template_requirement_reason', lambda *_: None)
    result = workflow.run_job(job.id)
    assert result['state'] == 'blocked', result
    assert result['blocked_reason'].startswith(PREFIX), result['blocked_reason']
    assert result['resume_state'] == 'full_build'


def test_reference_identity_page_coverage_and_changes(tmp_path, monkeypatch):
    import hashlib
    from pptx import Presentation
    from PIL import Image
    job, root, plan, _ = fresh(tmp_path, monkeypatch)
    source = tmp_path / '模板库/test.pptx'; source.parent.mkdir()
    pptx = Presentation(); pptx.slides.add_slide(pptx.slide_layouts[6]); pptx.save(source)
    preview = tmp_path / 'source-page.png'; Image.new('RGB', (160, 90), 'white').save(preview)
    sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    request = decisions(PresentationForm.BUSINESS_REPORT, plan)
    reference = {'mode': 'reference', 'reason': '本次观察测试', 'source_path': 'test.pptx', 'source_sha256': sha(source),
                 'source_slide': 1, 'borrowed_relationship': '测试页对应关系', 'observation': '测试页图身份绑定，不是审美认可',
                 'adaptation': '按当前内容重排', 'previews': [{'path': str(preview), 'sha256': sha(preview),
                                                        'source_sha256': sha(source), 'source_slide': 1}]}
    request['page_design']['p1']['reference_decision'] = reference
    assert author_design_plan(job.id, request)['status'] == 'recorded_for_build_not_visual_approval'
    reference['source_slide'] = 9
    with pytest.raises(ValueError, match='超出|身份'):
        author_design_plan(job.id, request)
    reference['source_slide'] = 1
    reference['previews'][0]['source_sha256'] = '0'*64
    with pytest.raises(ValueError, match='身份|覆盖'):
        author_design_plan(job.id, request)
    reference['previews'][0]['source_sha256'] = sha(source)
    preview.write_bytes(b'changed picture')
    assert any('页图已变化' in x for x in design_plan_status(root)['issues'])


def test_preservation_requires_an_existing_bound_source(tmp_path, monkeypatch):
    job, root, plan, _ = fresh(tmp_path, monkeypatch)
    request = decisions(PresentationForm.BUSINESS_REPORT, plan)
    request['page_design']['p1']['reference_decision'] = {'mode': 'preserve_existing', 'reason': '保留'}
    with pytest.raises(ValueError, match='未绑定'):
        author_design_plan(job.id, request)
