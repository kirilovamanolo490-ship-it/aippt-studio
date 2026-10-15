from __future__ import annotations

import hashlib
import json
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from typer.testing import CliRunner

from aippt_studio import cli, pptx_file_ops


def _sample(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "file ops sample"
    deck.save(path)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_pptx_new_creates_blank_editable_widescreen_deck_and_refuses_overwrite(tmp_path: Path):
    output = tmp_path / "new folder" / "blank.pptx"
    result = CliRunner().invoke(cli.app, ["pptx-new", str(output), "--json"])
    assert result.exit_code == 0, result.output
    report = json.loads(result.output)
    assert report["status"] == "created"
    deck = Presentation(output)
    assert len(deck.slides) == 1
    assert round(deck.slide_width / deck.slide_height, 3) == 1.778
    assert len(deck.slides[0].shapes) == 0
    before = _sha(output)
    rejected = CliRunner().invoke(cli.app, ["pptx-new", str(output)])
    assert rejected.exit_code == 2
    assert _sha(output) == before


def test_pptx_save_as_copies_exact_source_bytes_and_preserves_source(tmp_path: Path):
    source = tmp_path / "source.pptx"
    output = tmp_path / "delivery" / "copy.pptx"
    _sample(source)
    source_hash = _sha(source)
    result = CliRunner().invoke(
        cli.app, ["pptx-save-as", str(source), "--output", str(output), "--json"]
    )
    assert result.exit_code == 0, result.output
    report = json.loads(result.output)
    assert report["status"] == "saved_as"
    assert _sha(output) == source_hash
    assert _sha(source) == source_hash
    assert len(Presentation(output).slides) == 1
    rejected = CliRunner().invoke(
        cli.app, ["pptx-save-as", str(source), "--output", str(output)]
    )
    assert rejected.exit_code == 2
    assert _sha(output) == source_hash
    changed = Presentation(source)
    changed.slides[0].shapes[0].text = "updated source"
    changed.save(source)
    replaced = CliRunner().invoke(
        cli.app, ["pptx-save-as", str(source), "--output", str(output), "--replace", "--json"]
    )
    assert replaced.exit_code == 0, replaced.output
    assert json.loads(replaced.output)["replaced"] is True
    assert _sha(output) == _sha(source)
    assert Presentation(output).slides[0].shapes[0].text == "updated source"


def test_pptx_open_calls_only_system_open_and_reports_request_not_acceptance(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    source = tmp_path / "open me.pptx"
    _sample(source)
    calls: list[tuple[list[str], dict[str, object]]] = []

    def fake_run(command, **kwargs):
        if command[0] == "/usr/sbin/sysctl":
            return subprocess.CompletedProcess(command, 0, "1\n", "")
        calls.append((command, kwargs))
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(pptx_file_ops.sys, "platform", "darwin")
    monkeypatch.setattr(pptx_file_ops.subprocess, "run", fake_run)
    result = CliRunner().invoke(cli.app, ["pptx-open", str(source), "--json"])
    assert result.exit_code == 0, result.output
    report = json.loads(result.output)
    assert report["status"] == "open_requested"
    assert report["native_review"] == "not_verified"
    assert calls[0][0] == ["/usr/bin/open", "-a", "Microsoft PowerPoint", str(source.resolve())]
    assert calls[0][1]["check"] is True


def test_pptx_open_rejects_non_mac_without_running_open(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    source = tmp_path / "source.pptx"
    _sample(source)
    monkeypatch.setattr(pptx_file_ops.sys, "platform", "linux")
    monkeypatch.setattr(pptx_file_ops.subprocess, "run", lambda *args, **kwargs: pytest.fail("must not run"))
    with pytest.raises(RuntimeError, match="only on macOS"):
        pptx_file_ops.request_powerpoint_open(source)
