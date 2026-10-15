from pathlib import Path
from zipfile import ZipFile
import json
import math
import shutil
import subprocess
import pytest
from pptx import Presentation
from pptx.util import Inches
from aippt_studio import native_powerpoint as native, mac_native, mac_video
from aippt_studio.animation_edit import _package, _slides, _xml, NS


@pytest.fixture(autouse=True)
def isolate_system_osascript(monkeypatch):
    # Unit tests replace the subprocess and must not depend on the CI host
    # having macOS' bridge, while production still resolves only /usr/bin.
    monkeypatch.setattr(mac_video, 'system_osascript', lambda: Path('/usr/bin/osascript'))


def fixture_deck(path):
    p=Presentation()
    for i in range(3):
        s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1)).text=str(i)
    p.save(path)


@pytest.mark.parametrize('field', ['duration','delay','repeat','motion_by_x','scale_percent'])
def test_nonfinite_native_values_are_rejected(field):
    with pytest.raises(ValueError):
        native.validate_native_plan({'effects':[{'slide':1,'shape':'X','effect':'fade',field:float('nan')}]})


def test_mac_backend_dispatch_bypasses_windows_script(tmp_path,monkeypatch):
    source=tmp_path/'source.pptx';source.write_bytes(b'fixture')
    monkeypatch.setattr(native.sys,'platform','darwin')
    calls=[]
    monkeypatch.setattr(mac_native,'apply_mac_native_plan',lambda src,out,plan:calls.append(plan) or {'status':'passed','backend':'mac'})
    result=native.apply_native_plan(source,tmp_path/'result.pptx',{})
    assert result['backend']=='mac' and len(calls)==1


def test_movie_subset_uses_slide_relationships_and_only_changes_copy(tmp_path):
    source=tmp_path/'source.pptx';fixture_deck(source);before=source.read_bytes()
    spec={'slides':[1,3],'use_timings_and_narrations':False,'seconds_per_slide':7}
    dest=tmp_path/'preview.pptx';mac_video._subset(source,dest,spec)
    parts=_package(dest);names=_slides(parts)
    assert names==['ppt/slides/slide1.xml','ppt/slides/slide3.xml']
    for name in names:
        trans=_xml(parts[name]).find('p:transition',NS)
        assert trans.get('advTm')=='7000' and trans.get('advClick')=='0'
    assert source.read_bytes()==before


def test_recorded_timing_preview_keeps_original_slide_bytes(tmp_path):
    source=tmp_path/'source.pptx';fixture_deck(source);before=_package(source)
    dest=tmp_path/'preview.pptx';mac_video._subset(source,dest,{'slides':[2],'use_timings_and_narrations':True})
    assert _package(dest)['ppt/slides/slide2.xml']==before['ppt/slides/slide2.xml']


def test_invalid_preview_slide_fails_before_office(tmp_path):
    source=tmp_path/'source.pptx';fixture_deck(source)
    with pytest.raises(ValueError,match='out of range'):
        mac_video._subset(source,tmp_path/'preview.pptx',{'slides':[7]})


def test_failed_mac_video_does_not_publish_new_native_deck(tmp_path,monkeypatch):
    source=tmp_path/'source.pptx';fixture_deck(source)
    output=tmp_path/'output.pptx';output.write_bytes(b'previous-valid-result')
    from aippt_studio import upstream_tools
    def fake_invoke(provider,operation,request,folder):
        folder.mkdir();shutil.copy2(source,folder/'candidate.pptx');return {'status':'succeeded'}
    monkeypatch.setattr(upstream_tools,'invoke',fake_invoke)
    monkeypatch.setattr(mac_native,'powerpoint_roundtrip',lambda *a:{'status':'passed'})
    monkeypatch.setattr(mac_native,'officecli_validate',lambda *a:{'status':'passed'})
    def fail(*args):raise ValueError('movie could not be decoded')
    monkeypatch.setattr(mac_video,'export_movie',fail)
    plan=native.validate_native_plan({'previews':[{'id':'preview','slides':[1]}]})
    with pytest.raises(native.NativePowerPointError,match='decoded'):
        mac_native.apply_mac_native_plan(source,output,plan)
    assert output.read_bytes()==b'previous-valid-result'


def test_mac_probe_classifies_candidate_failure(tmp_path,monkeypatch):
    base=tmp_path/'base.pptx';base.write_bytes(b'baseline');candidate=tmp_path/'candidate.pptx';candidate.write_bytes(b'candidate')
    def rt(path,out):
        if path.read_bytes()==b'candidate':raise ValueError('cannot open')
        return {'status':'passed'}
    monkeypatch.setattr(mac_native,'powerpoint_roundtrip',rt)
    monkeypatch.setattr(mac_native,'officecli_validate',lambda *a:{'status':'passed'})
    report=mac_native.probe_mac_powerpoint(base,candidate,tmp_path/'probe')
    assert report['status']=='failed' and report['failed_stage']=='candidate'


def test_movie_script_compiles_on_supported_mac(tmp_path):
    import sys
    import subprocess
    from aippt_studio.config import PROJECT_ROOT
    if sys.platform != 'darwin' or not Path('/Applications/Microsoft PowerPoint.app').exists():
        pytest.skip('requires the supported Mac PowerPoint scripting dictionary')
    result = subprocess.run(['osacompile', '-o', str(tmp_path/'movie.scpt'), str(PROJECT_ROOT/'scripts/powerpoint_movie.applescript')], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_export_failure_uses_office_container_and_preserves_output(tmp_path, monkeypatch):
    import subprocess
    from aippt_studio import mac_powerpoint
    source = tmp_path/'source.pptx'; fixture_deck(source)
    output = tmp_path/'preview.mp4'; output.write_bytes(b'previous-valid-video')
    container = tmp_path/'office-container'
    monkeypatch.setattr(mac_video, 'media_executable', lambda: Path('/fake/ffmpeg'))
    monkeypatch.setattr(mac_powerpoint, '_office_temporary_root', lambda: container)
    calls = []
    def failed_export(command, **kwargs):
        calls.append(command)
        assert command[0] == '/usr/bin/osascript'
        assert Path(command[2]).is_relative_to(container)
        assert Path(command[2]).is_file()
        return subprocess.CompletedProcess(command, 1, '', 'native export refused')
    monkeypatch.setattr(mac_video.subprocess, 'run', failed_export)
    spec = native.validate_native_plan({'previews':[{'id':'preview','slides':[1]}]})['previews'][0]
    with pytest.raises(ValueError, match='native export refused'):
        mac_video.export_movie(source, output, spec)
    assert len(calls) == 1
    assert output.read_bytes() == b'previous-valid-video'
    assert 'native export refused' in output.with_suffix('.export.log').read_text()
    assert not list(container.iterdir())


def test_mac_video_falls_back_to_work_owned_cache_when_office_container_is_protected(tmp_path, monkeypatch):
    source = tmp_path / 'source.pptx'; fixture_deck(source)
    output = tmp_path / 'preview.mp4'; output.write_bytes(b'previous-valid-video')
    protected = tmp_path / 'protected-container'; protected.write_bytes(b'not a directory')
    from aippt_studio import mac_powerpoint
    monkeypatch.setattr(mac_video, 'media_executable', lambda: Path('/fake/ffmpeg'))
    monkeypatch.setattr(mac_powerpoint, '_office_temporary_root', lambda: protected)
    calls = []

    def failed_export(command, **kwargs):
        calls.append(command)
        assert Path(command[2]).is_relative_to(tmp_path / '.powerpoint-movie-cache')
        return subprocess.CompletedProcess(command, 1, '', 'native export refused')

    monkeypatch.setattr(mac_video.subprocess, 'run', failed_export)
    spec = native.validate_native_plan({'previews':[{'id':'preview','slides':[1]}]})['previews'][0]
    with pytest.raises(ValueError, match='native export refused'):
        mac_video.export_movie(source, output, spec)
    assert len(calls) == 1
    assert output.read_bytes() == b'previous-valid-video'
    assert not list((tmp_path / '.powerpoint-movie-cache').iterdir())


def test_mac_movie_receipt_cannot_claim_object_animation_or_native_timing_control(tmp_path, monkeypatch):
    from aippt_studio import mac_powerpoint

    source = tmp_path / 'source.pptx'; fixture_deck(source)
    output = tmp_path / 'preview.mp4'
    ffmpeg = tmp_path / 'ffmpeg'; ffmpeg.write_bytes(b'host-native-fixture')
    container = tmp_path / 'office-container'
    monkeypatch.setattr(mac_video, 'PROJECT_ROOT', tmp_path)
    monkeypatch.setattr(mac_video, 'media_executable', lambda: ffmpeg)
    monkeypatch.setattr(mac_powerpoint, '_office_temporary_root', lambda: container)
    calls = []
    export_runs = [0]

    def fake_run(command, **kwargs):
        calls.append(command)
        if command[0] == '/usr/bin/osascript':
            export_runs[0] += 1
            assert len(command) == 5  # only input, output, and staging identity
            Path(command[3]).write_bytes(b'native movie fixture' * 16)
        elif '-frames:v' in command:
            Path(command[-1].replace('%03d', '001')).write_bytes(b'frame fixture')
            if export_runs[0] == 1:
                Path(command[-1].replace('%03d', '002')).write_bytes(b'old tail frame')
        elif '-f' in command and 'null' in command:
            pass
        else:
            Path(command[-1]).write_bytes(b'encoded movie fixture')
        return subprocess.CompletedProcess(command, 0, '', '')

    monkeypatch.setattr(mac_video.subprocess, 'run', fake_run)
    spec = native.validate_native_plan({'previews': [{'id': 'preview', 'slides': [1],
        'use_timings_and_narrations': True}]})['previews'][0]

    first = mac_video.export_movie(source, output, spec)
    report = mac_video.export_movie(source, output, spec)

    assert len(calls) == 8
    assert report['decoded'] is True and output.is_file()
    assert first['frames'] != report['frames']
    assert first['frame_count'] == 2 and report['frame_count'] == 1
    assert set(report['frame_sha256s']) == {'frame-001.png'}
    assert sorted(path.name for path in Path(report['frames']).glob('*.png')) == ['frame-001.png']
    assert report['preview_role'] == 'slide_sequence_only'
    assert report['playback_verified'] is False
    assert report['animation_export_support'] == 'object_animations_not_played_by_powerpoint_mac_movie'
    assert report['timing_source'] == 'not_verified'
    assert report['timing_request'] == 'recorded_timings_and_narrations'
    assert report['native_export_options']['timings_and_narrations'] == 'not_configured_by_applescript'
    assert report['audio_policy'] == 'carry_native_movie_audio_if_present_unverified'
    assert json.loads(output.with_suffix('.video.json').read_text()) == report
