import hashlib
import base64
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree as ET
from pptx import Presentation
from pptx.util import Inches

from aippt_studio.motion_clone import MotionCloneError, _motion_signature, clone_motion_deck, text_sha256

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"


def _make_deck(path):
    ppt = Presentation()
    for label in ("One", "Two", "Three"):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = label
    ppt.save(path)
    with ZipFile(path) as archive:
        parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
    for num in range(1, 4):
        root = ET.fromstring(parts[f"ppt/slides/slide{num}.xml"])
        sp_tree = root.find(f"{{{P}}}cSld/{{{P}}}spTree")
        shape = next(x for x in sp_tree if x.tag == f"{{{P}}}sp" and x.find(f".//{{{A}}}t") is not None)
        shape_id = shape.find(f"{{{P}}}nvSpPr/{{{P}}}cNvPr").get("id")
        timing = ET.SubElement(root, f"{{{P}}}timing")
        par = ET.SubElement(timing, f"{{{P}}}par")
        ctn = ET.SubElement(par, f"{{{P}}}cTn", id="1")
        target = ET.SubElement(ctn, f"{{{P}}}spTgt", spid=shape_id)
        transition = ET.SubElement(root, f"{{{P}}}transition", spd="slow")
        ET.SubElement(transition, f"{{{P}}}fade")
        if num == 2:
            ext_list = ET.SubElement(root, f"{{{P}}}extLst")
            ext = ET.SubElement(ext_list, f"{{{P}}}ext", uri="{A4A8B984-8BB5-495C-AF93-5E8A8B4F4D35}")
            ET.SubElement(ext, f"{{{P159}}}morph", option="byObject")
        parts[f"ppt/slides/slide{num}.xml"] = ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return path


def _parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _add_layout_motion(path, *, target_master=False):
    parts = _parts(path)
    slide_rel = ET.fromstring(parts["ppt/slides/_rels/slide1.xml.rels"])
    layout_target = next(x.get("Target") for x in slide_rel if x.get("Type", "").endswith("/slideLayout"))
    layout_part = "ppt/slides/" + layout_target
    layout_part = layout_part.replace("ppt/slides/../", "ppt/")
    layout = ET.fromstring(parts[layout_part])
    tree = layout.find(f"{{{P}}}cSld/{{{P}}}spTree")
    layout_shape = next(x for x in tree if x.tag == f"{{{P}}}sp")
    shape_id = layout_shape.find(f"{{{P}}}nvSpPr/{{{P}}}cNvPr").get("id")
    if target_master:
        rels_part = f"ppt/slideLayouts/_rels/{layout_part.rsplit('/', 1)[-1]}.rels"
        rels = ET.fromstring(parts[rels_part])
        master_target = next(x.get("Target") for x in rels if x.get("Type", "").endswith("/slideMaster"))
        master_part = "ppt/slideLayouts/" + master_target
        master_part = master_part.replace("ppt/slideLayouts/../", "ppt/")
        master = ET.fromstring(parts[master_part])
        master_tree = master.find(f"{{{P}}}cSld/{{{P}}}spTree")
        master_shape = next(x for x in master_tree if x.tag == f"{{{P}}}sp")
        master_nv = master_shape.find(f"{{{P}}}nvSpPr/{{{P}}}cNvPr")
        master_nv.set("id", "9999")
        shape_id = master_nv.get("id")
        parts[master_part] = ET.tostring(master, xml_declaration=True, encoding="UTF-8", standalone=True)
    timing = ET.SubElement(layout, f"{{{P}}}timing")
    par = ET.SubElement(timing, f"{{{P}}}par")
    ctn = ET.SubElement(par, f"{{{P}}}cTn", id="1")
    ET.SubElement(ctn, f"{{{P}}}spTgt", spid=shape_id)
    anim = ET.SubElement(ctn, f"{{{P}}}animMotion", path="M 0 0 L 1 1")
    transition = ET.SubElement(layout, f"{{{P}}}transition", spd="med")
    ET.SubElement(transition, f"{{{P}}}fade")
    parts[layout_part] = ET.tostring(layout, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return layout_part, ET.tostring(anim, method="c14n"), shape_id


def _shape_and_text(parts, slide=2):
    root = ET.fromstring(parts[f"ppt/slides/slide{slide}.xml"])
    shape = next(x for x in root.iter(f"{{{P}}}sp") if x.find(f".//{{{A}}}t") is not None)
    return shape.find(f"{{{P}}}nvSpPr/{{{P}}}cNvPr").get("id"), shape.find(f".//{{{A}}}t").text


def _make_picture_deck(path, image_path):
    source = _make_deck(path)
    ppt = Presentation(source)
    for slide_num in (2, 3):
        slide = ppt.slides[slide_num - 1]
        pic = slide.shapes.add_picture(str(image_path), Inches(4), Inches(1), width=Inches(2), height=Inches(1))
        pic.crop_left = 0.1
        pic.crop_right = 0.2
    ppt.save(source)
    return source


def _make_layout_picture_deck(path, image_path, shared=False):
    ppt = Presentation()
    layout = ppt.slide_layouts[6 if shared else 5]
    from pptx.oxml.shapes.picture import CT_Picture
    _, image_rid = layout.part.get_or_add_image_part(str(image_path))
    pic = CT_Picture.new_pic(99, "Layout Picture", "", image_rid, Inches(6), Inches(4), Inches(1), Inches(1))
    layout.shapes._spTree.insert_element_before(pic, "p:extLst")
    for i in range(3):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6] if shared and i else layout if not shared and i == 0 else ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = str(i)
    ppt.save(path)
    return path


def _picture_info(parts, slide):
    root = ET.fromstring(parts[f"ppt/slides/slide{slide}.xml"])
    pic = root.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}pic")
    shape_id = pic.find("{http://schemas.openxmlformats.org/presentationml/2006/main}nvPicPr/{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").get("id")
    rel_id = pic.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip").get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    rels_root = ET.fromstring(parts[f"ppt/slides/_rels/slide{slide}.xml.rels"])
    target = next(rel.get("Target") for rel in rels_root if rel.get("Id") == rel_id)
    media = "ppt/slides/" + target
    media = media.replace("ppt/slides/../", "ppt/")
    return shape_id, media


def _layout_picture_info(parts, layout_part):
    root = ET.fromstring(parts[layout_part])
    pic = root.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}pic")
    shape_id = pic.find("{http://schemas.openxmlformats.org/presentationml/2006/main}nvPicPr/{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").get("id")
    rel_id = pic.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip").get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    rels_name = f"{layout_part.rsplit('/', 1)[0]}/_rels/{layout_part.rsplit('/', 1)[-1]}.rels"
    rels = ET.fromstring(parts[rels_name])
    target = next(rel.get("Target") for rel in rels if rel.get("Id") == rel_id)
    media = layout_part.rsplit("/", 1)[0] + "/" + target
    while "/../" in media:
        prefix, tail = media.split("/../", 1)
        media = prefix.rsplit("/", 1)[0] + "/" + tail
    return shape_id, media


def test_clone_contiguous_slides_and_single_run_text_preserves_motion_and_other_parts(tmp_path):
    source = _make_deck(tmp_path / "source.pptx")
    original = _parts(source)
    shape_id, old_text = _shape_and_text(original)
    target = tmp_path / "cloned.pptx"
    result = clone_motion_deck(
        source,
        target,
        source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        slides=[2, 3],
        replacements=[{"slide": 2, "shape_id": shape_id, "original_text_sha256": text_sha256(old_text), "text": "Reused"}],
    )
    cloned = _parts(target)
    assert result["timing_transition_readback"] == "passed"
    assert result["playback_verified"] is False
    assert original["ppt/slides/slide2.xml"] != cloned["ppt/slides/slide2.xml"]
    assert original["ppt/slides/slide3.xml"] == cloned["ppt/slides/slide3.xml"]
    for name, data in original.items():
        if name not in {"ppt/presentation.xml", "ppt/slides/slide2.xml"}:
            assert cloned[name] == data
    root = ET.fromstring(cloned["ppt/slides/slide2.xml"])
    assert root.xpath("string(.//a:t)", namespaces={"a": A}) == "Reused"
    assert root.find(f"{{{P}}}timing") is not None and root.find(f"{{{P}}}transition") is not None
    assert root.find(f".//{{{P159}}}morph") is not None


def test_motion_signature_includes_morph_extension():
    root = ET.Element(f"{{{P}}}sld")
    ext_list = ET.SubElement(root, f"{{{P}}}extLst")
    ext = ET.SubElement(ext_list, f"{{{P}}}ext", uri="test")
    morph = ET.SubElement(ext, f"{{{P159}}}morph", option="byObject")
    before = _motion_signature(root)
    morph.set("option", "byWord")
    assert _motion_signature(root) != before


def test_related_layout_motion_is_validated_and_read_back_without_mutating_shared_layer(tmp_path):
    source = _make_deck(tmp_path / "source.pptx")
    layout_part, _, _ = _add_layout_motion(source)
    original = _parts(source)
    target = tmp_path / "clone.pptx"
    result = clone_motion_deck(source, target, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[1])
    cloned = _parts(target)
    assert result["related_motion_readback"] == "passed"
    assert cloned[layout_part] == original[layout_part]
    assert result["cross_layer_animation_targets"] == []


def test_layout_animation_target_found_on_master_is_reported_as_cross_layer(tmp_path):
    source = _make_deck(tmp_path / "source.pptx")
    layout_part, _, master_shape_id = _add_layout_motion(source, target_master=True)
    target = tmp_path / "clone.pptx"
    result = clone_motion_deck(source, target, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[1])
    assert result["related_motion_readback"] == "passed"
    assert {"part": layout_part, "shape_id": master_shape_id} in result["cross_layer_animation_targets"]


def test_image_replacement_preserves_picture_xml_motion_and_unrelated_parts(tmp_path):
    original_png = tmp_path / "original.png"
    replacement_png = tmp_path / "replacement.png"
    # Two distinct valid 1x1 PNG files, kept local and deterministic.
    original_png.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/aWQAAAAASUVORK5CYII="))
    replacement_png.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADUlEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC"))
    source = _make_picture_deck(tmp_path / "source.pptx", original_png)
    source_bytes_before = source.read_bytes()
    original = _parts(source)
    id2, media2 = _picture_info(original, 2)
    id3, media3 = _picture_info(original, 3)
    assert media2 == media3
    old_media = original[media2]
    target = tmp_path / "clone.pptx"
    result = clone_motion_deck(
        source, target, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[2, 3],
        image_replacements=[
            {"slide": 2, "shape_id": id2, "original_media_sha256": hashlib.sha256(old_media).hexdigest(), "path": replacement_png},
            {"slide": 3, "shape_id": id3, "original_media_sha256": hashlib.sha256(old_media).hexdigest(), "path": replacement_png},
        ],
    )
    cloned = _parts(target)
    assert result["timing_transition_readback"] == "passed"
    assert result["image_replacements"] == 2
    assert cloned[media2] == replacement_png.read_bytes()
    assert cloned["ppt/slides/slide2.xml"] == original["ppt/slides/slide2.xml"]
    assert cloned["ppt/slides/slide3.xml"] == original["ppt/slides/slide3.xml"]
    assert source.read_bytes() == source_bytes_before
    for name, data in original.items():
        if name not in {"ppt/presentation.xml", media2}:
            assert cloned[name] == data


def test_shared_media_requires_all_selected_references_and_matching_bytes(tmp_path):
    original_png = tmp_path / "original.png"
    replacement_png = tmp_path / "replacement.png"
    original_png.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/aWQAAAAASUVORK5CYII="))
    replacement_png.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADUlEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC"))
    source = _make_picture_deck(tmp_path / "source.pptx", original_png)
    original = _parts(source)
    id2, media = _picture_info(original, 2)
    id3, _ = _picture_info(original, 3)
    common = {"original_media_sha256": hashlib.sha256(original[media]).hexdigest(), "path": replacement_png}
    with pytest.raises(MotionCloneError, match="all selected-slide references"):
        clone_motion_deck(source, tmp_path / "partial.pptx", source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[2, 3], image_replacements=[{"slide": 2, "shape_id": id2, **common}])
    other = tmp_path / "different.png"
    other.write_bytes(original_png.read_bytes())
    with pytest.raises(MotionCloneError, match="same replacement bytes"):
        clone_motion_deck(source, tmp_path / "different.pptx", source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[2, 3], image_replacements=[
            {"slide": 2, "shape_id": id2, **common},
            {"slide": 3, "shape_id": id3, "original_media_sha256": common["original_media_sha256"], "path": other},
        ])
    assert not (tmp_path / "partial.pptx").exists()
    assert not (tmp_path / "different.pptx").exists()


def test_layout_image_replacement_is_hash_bound_and_blocks_unselected_layout_users(tmp_path):
    original_png = tmp_path / "original.png"
    replacement_png = tmp_path / "replacement.png"
    original_png.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/aWQAAAAASUVORK5CYII="))
    replacement_png.write_bytes(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADUlEQVR4nGP4z8AAAAMBAQDJ/pLvAAAAAElFTkSuQmCC"))
    shared_source = _make_layout_picture_deck(tmp_path / "shared.pptx", original_png, shared=True)
    shared_parts = _parts(shared_source)
    layout_part = "ppt/slideLayouts/slideLayout7.xml"
    shape_id, media_part = _layout_picture_info(shared_parts, layout_part)
    record = {"layer": "layout", "layout_part": layout_part, "shape_id": shape_id,
              "original_media_sha256": hashlib.sha256(shared_parts[media_part]).hexdigest(), "path": replacement_png}
    with pytest.raises(MotionCloneError, match="unselected slides"):
        clone_motion_deck(shared_source, tmp_path / "blocked.pptx", source_sha256=hashlib.sha256(shared_source.read_bytes()).hexdigest(), slides=[1], image_replacements=[record])

    unique_source = _make_layout_picture_deck(tmp_path / "unique.pptx", original_png, shared=False)
    unique_parts = _parts(unique_source)
    layout_part = "ppt/slideLayouts/slideLayout6.xml"
    shape_id, media_part = _layout_picture_info(unique_parts, layout_part)
    target = tmp_path / "layout-clone.pptx"
    result = clone_motion_deck(unique_source, target, source_sha256=hashlib.sha256(unique_source.read_bytes()).hexdigest(), slides=[1], image_replacements=[{
        "layer": "layout", "layout_part": layout_part, "shape_id": shape_id,
        "original_media_sha256": hashlib.sha256(unique_parts[media_part]).hexdigest(), "path": replacement_png,
    }])
    cloned = _parts(target)
    assert result["image_replacements"] == 1
    assert cloned[media_part] == replacement_png.read_bytes()
    assert cloned[layout_part] == unique_parts[layout_part]


@pytest.mark.parametrize("slides", [[], [1, 3], [0, 1], [1, 4]])
def test_rejects_noncontiguous_or_out_of_range_selection_without_output(tmp_path, slides):
    source = _make_deck(tmp_path / "source.pptx")
    target = tmp_path / "clone.pptx"
    with pytest.raises(MotionCloneError):
        clone_motion_deck(source, target, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=slides)
    assert not target.exists()


def test_stale_source_or_text_binding_rejected(tmp_path):
    source = _make_deck(tmp_path / "source.pptx")
    sha = hashlib.sha256(source.read_bytes()).hexdigest()
    with pytest.raises(MotionCloneError, match="source SHA"):
        clone_motion_deck(source, tmp_path / "a.pptx", source_sha256="stale", slides=[1])
    parts = _parts(source)
    shape_id, _ = _shape_and_text(parts, 1)
    with pytest.raises(MotionCloneError, match="text hash mismatch"):
        clone_motion_deck(source, tmp_path / "b.pptx", source_sha256=sha, slides=[1], replacements=[
            {"slide": 1, "shape_id": shape_id, "original_text_sha256": "stale", "text": "Changed"}
        ])
    assert not (tmp_path / "a.pptx").exists() and not (tmp_path / "b.pptx").exists()


def test_unresolved_animation_target_is_blocked(tmp_path):
    source = _make_deck(tmp_path / "source.pptx")
    with ZipFile(source, "r") as old:
        parts = {name: old.read(name) for name in old.namelist()}
    root = ET.fromstring(parts["ppt/slides/slide1.xml"])
    root.find(f".//{{{P}}}spTgt").set("spid", "999")
    parts["ppt/slides/slide1.xml"] = ET.tostring(root)
    with ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    with pytest.raises(MotionCloneError, match="unresolved animation"):
        clone_motion_deck(source, tmp_path / "clone.pptx", source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[1])


def test_streams_package_larger_than_100mb_uncompressed(tmp_path):
    source = _make_deck(tmp_path / "source.pptx")
    payload_size = 101 * 1024 * 1024
    with ZipFile(source, "a", compression=ZIP_DEFLATED, compresslevel=1) as archive:
        with archive.open("ppt/media/large-media.bin", "w") as media:
            block = b"M" * (1024 * 1024)
            for _ in range(101):
                media.write(block)
    target = tmp_path / "clone.pptx"
    result = clone_motion_deck(source, target, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[1])
    with ZipFile(target) as archive:
        info = archive.getinfo("ppt/media/large-media.bin")
        assert info.file_size == payload_size
        digest = hashlib.sha256()
        with archive.open(info) as media:
            for chunk in iter(lambda: media.read(1024 * 1024), b""):
                digest.update(chunk)
    expected = hashlib.sha256()
    block = b"M" * (1024 * 1024)
    for _ in range(101):
        expected.update(block)
    assert digest.hexdigest() == expected.hexdigest()
    assert result["timing_transition_readback"] == "passed"


def test_duplicate_entry_and_existing_output_are_rejected(tmp_path):
    source = _make_deck(tmp_path / "source.pptx")
    with ZipFile(source, "a") as archive:
        archive.writestr("ppt/presentation.xml", archive.read("ppt/presentation.xml"))
    target = tmp_path / "clone.pptx"
    with pytest.raises(MotionCloneError, match="duplicate ZIP member"):
        clone_motion_deck(source, target, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(), slides=[1])
    assert not target.exists()

    clean_source = _make_deck(tmp_path / "clean.pptx")
    target.write_bytes(b"keep existing output")
    with pytest.raises(MotionCloneError, match="refusing to overwrite"):
        clone_motion_deck(clean_source, target, source_sha256=hashlib.sha256(clean_source.read_bytes()).hexdigest(), slides=[1])
    assert target.read_bytes() == b"keep existing output"
