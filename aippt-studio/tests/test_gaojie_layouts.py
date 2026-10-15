from __future__ import annotations
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches

from aippt_studio.gaojie_layouts import (analysis_root, build_gaojie_layout_candidate,
                                       record_layout_application, search_gaojie_layouts)
from aippt_studio.motion_clone import text_sha256
from aippt_studio.style_library import inspect_style_deck
from aippt_studio.style_search import search_style_library


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def jsonl(path, values):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(value, ensure_ascii=False) + "\n" for value in values))


def fixture(library: Path, category="文本组件"):
    source = library / "高阶文化" / category / "1 - 样本.pptx"
    source.parent.mkdir(parents=True)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    groups = []
    for col in range(2):
        group = slide.shapes.add_group_shape()
        title = group.shapes.add_textbox(Inches(1 + col * 4), Inches(1), Inches(3), Inches(.5))
        title.text = f"示例标题{col}"
        body = group.shapes.add_textbox(Inches(1 + col * 4), Inches(1.7), Inches(3), Inches(1))
        body.text = f"示例正文{col}"
        groups.append((str(group.shape_id), str(title.shape_id), str(body.shape_id)))
    prs.save(source)
    digest = sha(source)
    relative = source.relative_to(library).as_posix()
    detail = inspect_style_deck(source, relative)
    detail_file = analysis_root(library) / "逐份结构" / f"{digest}.json"
    detail_file.parent.mkdir(parents=True)
    detail_file.write_text(json.dumps(detail, ensure_ascii=False))
    catalog = {"path": relative, "category": category, "source_sha256": digest, "status": "parsed",
               "source_format": "pptx", "slide_count": 1, "detail_ref": detail_file.relative_to(library).as_posix(),
               "detail_sha256": sha(detail_file), "style_profile": {}}
    page = {"canonical_path": relative, "source_sha256": digest, "slide": 1, "page_status": "parsed",
            "title": "示例", "text_excerpt": "社区共创", "role_evidence": [{"role": "text_component"}],
            "layout_measurements": {"structure_signature": "fixture"}, "layout_pattern": {"kind": "text_and_native_shapes"}}
    jsonl(analysis_root(library) / "STYLE-CATALOG.jsonl", [catalog])
    jsonl(analysis_root(library) / "STYLE-PAGES.jsonl", [page])
    jsonl(library / "CONTENT-CURATION.jsonl", [{"path": relative, "sha256": digest, "asset_type": "template", "category": category}])
    preview = library.parent / "source.png"
    preview.write_bytes(b"bound preview fixture")
    receipt = library.parent / "preview.json"
    receipt.write_text(json.dumps({"status": "rendered_advisory", "source_sha256": digest,
                                  "slide_number": 1, "output_png": str(preview), "output_sha256": sha(preview)}))
    request = {"source_path": relative, "source_sha256": digest, "source_slide": 1,
               "selected_object_ids": [groups[0][0]], "selection_note": "以一个标题和正文组成完整的信息单元",
               "reviewed_source_page": {"slide": 1, "finding": "所选组保持标题和正文的明确上下层级", "preview_receipt": str(receipt)},
               "text_replacements": [{"shape_id": groups[0][1], "original_text_sha256": text_sha256("示例标题0"), "new_text": "社区共创"},
                                     {"shape_id": groups[0][2], "original_text_sha256": text_sha256("示例正文0"), "new_text": "以共同参与组织社区活动"}]}
    return source, digest, detail_file, request, groups


def test_search_rejects_changed_source_and_changed_structural_detail(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    assert search_style_library(library, "文本组件 社区", collection="gaojie")["candidate_count"] == 1
    original = source.read_bytes()
    source.write_bytes(original + b"changed")
    assert not search_gaojie_layouts(library, "社区", category="文本组件")["candidates"]
    source.write_bytes(original)
    detail.write_text("{}")
    assert not search_gaojie_layouts(library, "社区", category="文本组件")["candidates"]


def test_single_component_copy_retains_group_geometry_and_maps_all_selected_text(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    output = tmp_path / "candidate.pptx"
    built = build_gaojie_layout_candidate(library, request, output)
    assert sha(source) == digest and built["mapped_text_objects"] == 2
    deck = inspect_style_deck(output, "candidate.pptx")
    texts = [obj["text"] for obj in deck["slides"][0]["objects"] if obj["text"]]
    assert texts == ["社区共创", "以共同参与组织社区活动"]
    assert [obj["id"] for obj in deck["slides"][0]["objects"] if obj["type"] == "grpSp"] == [groups[0][0]]
    assert deck["slides"][0]["objects"][0]["geometry"] == json.loads(detail.read_text())["slides"][0]["objects"][0]["geometry"]


def test_component_sheet_and_incomplete_fact_mapping_cannot_be_published(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    output = tmp_path / "candidate.pptx"
    missing = dict(request)
    missing.pop("selected_object_ids")
    with pytest.raises(ValueError, match="component sheets"):
        build_gaojie_layout_candidate(library, missing, output)
    incomplete = dict(request, text_replacements=request["text_replacements"][:1])
    with pytest.raises(ValueError, match="every retained text"):
        build_gaojie_layout_candidate(library, incomplete, output)
    assert not output.exists() and sha(source) == digest


def test_corrupt_reference_is_never_a_template_candidate(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    jsonl(library / "CONTENT-CURATION.jsonl", [{"path": request["source_path"], "sha256": digest, "asset_type": "reference",
                                              "evidence": {"availability_status": "unavailable_corrupt_source"}}])
    assert not search_gaojie_layouts(library, "社区", category="文本组件")["candidates"]
    with pytest.raises(ValueError, match="stale"):
        build_gaojie_layout_candidate(library, request, tmp_path / "bad.pptx")


def test_visual_rejection_survives_query_and_source_identity(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    jsonl(analysis_root(library) / "VISUAL-LEARNING.jsonl", [{"source_sha256": digest, "source_slide": 1, "decision": "avoid"}])
    assert not search_gaojie_layouts(library, "社区", category="文本组件")["candidates"]


def test_learning_event_requires_matching_output_preview_and_deduplicates(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    output = tmp_path / "candidate.pptx"
    build_gaojie_layout_candidate(library, request, output)
    source_receipt = request["reviewed_source_page"]["preview_receipt"]
    output_png = tmp_path / "output.png"
    output_png.write_bytes(b"observed output fixture")
    output_receipt = tmp_path / "output-preview.json"
    output_receipt.write_text(json.dumps({"status": "rendered_advisory", "source_sha256": sha(output), "slide_number": 1,
                                          "output_png": str(output_png), "output_sha256": sha(output_png)}))
    lesson = {"source_path": request["source_path"], "source_sha256": digest, "source_slide": 1,
              "output_pptx": str(output), "output_sha256": sha(output), "output_slide": 1,
              "source_preview_receipt": source_receipt, "output_preview_receipt": str(output_receipt),
              "finding": "替换后的标题和正文保持明确层级", "what_reused": "复用一个完整的标题正文组件组",
              "input_constraints": "限短标题与精简正文，长段需调整版面", "outcome_issues": "当前映射没有观察到文字重叠问题",
              "decision": "use_with_constraints"}
    assert record_layout_application(library, lesson)["events"] == 1
    assert record_layout_application(library, lesson)["events"] == 1
    output_png.write_bytes(b"changed preview")
    with pytest.raises(ValueError, match="preview"):
        record_layout_application(library, lesson)


def test_optional_larger_xml_budget_remains_explicitly_bounded(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    parsed = inspect_style_deck(source, request["source_path"], xml_part_limit=40_000_000)
    assert parsed["status"] == "parsed" and parsed["limits"]["xml_part_bytes"] == 40_000_000
    assert inspect_style_deck(source, request["source_path"])["limits"]["xml_part_bytes"] == 32_000_000
    assert inspect_style_deck(source, request["source_path"], xml_part_limit=100_000_000)["status"] == "unsupported"


def test_widening_selected_group_text_preserves_fonts_and_group_scale(tmp_path):
    library = tmp_path / "模板库"
    source, digest, detail, request, groups = fixture(library)
    before = Presentation(source).slides[0].shapes[0]
    scale = before.width / before.element.grpSpPr.xfrm.chExt.cx
    request["component_text_widths"] = [{"shape_id": groups[0][1], "width_fraction": .70},
                                        {"shape_id": groups[0][2], "width_fraction": .70}]
    output = tmp_path / "wide-component.pptx"
    build_gaojie_layout_candidate(library, request, output)
    result = Presentation(output).slides[0].shapes[0]
    assert result.width > before.width
    assert abs(result.width / result.element.grpSpPr.xfrm.chExt.cx - scale) < 0.00001
    assert result.shapes[0].width > before.shapes[0].width
    assert result.shapes[0].text_frame.paragraphs[0].runs[0].font.size == before.shapes[0].text_frame.paragraphs[0].runs[0].font.size
    assert sha(source) == digest
