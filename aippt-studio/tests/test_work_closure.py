from __future__ import annotations

import json
import hashlib
import os
import shutil
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches


BUNDLED_PWSH = Path(__file__).resolve().parents[1] / ".tools/pwsh/pwsh"
POWERSHELL = shutil.which("pwsh") or shutil.which("powershell") or (str(BUNDLED_PWSH) if BUNDLED_PWSH.is_file() else None)


def _make_work(tmp_path: Path) -> Path:
    work = tmp_path / "work" / "1"
    for directory in ("input/source", "input/reference", "workspace/qa", "workspace/studio-jobs/job/build", "workspace/studio-jobs/job/assets/acquired", "output"):
        (work / directory).mkdir(parents=True, exist_ok=True)
    (work / "input" / "request.md").write_text("fixture", encoding="utf-8")
    (work / "workspace" / "worklog.md").write_text("fixture", encoding="utf-8")
    (work / "workspace" / "knowledge-context.json").write_text(
        '{"status":"no_match","advisory_only":true,"matches":[]}', encoding="utf-8"
    )
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(work / "output" / "fixture.pptx")
    (work / "workspace" / "studio-jobs" / "job" / "build" / "adapter-last-result.json").write_text(
        '{"status":"succeeded"}', encoding="utf-8"
    )
    (work / "workspace" / "studio-jobs" / "job" / "planning").mkdir(parents=True, exist_ok=True)
    (work / "workspace" / "studio-jobs" / "job" / "planning" / "asset-plan.json").write_text(
        '{"needs":[]}', encoding="utf-8"
    )
    (work / "workspace" / "studio-jobs" / "job" / "assets" / "asset-ledger.json").write_text(
        '{"assets":[]}', encoding="utf-8"
    )
    (work / "workspace" / "studio-jobs" / "job" / "assets" / "acquired" / "icon.svg").write_text(
        '<svg/>', encoding="utf-8"
    )
    (work / "work.json").write_text(
        json.dumps({"id": "1", "work_id": "1", "name": "fixture", "client": "", "status": "in_progress"}), encoding="utf-8"
    )
    knowledge_entry = tmp_path / "任务知识库" / "entries" / "2026" / "work-1-fixture.md"
    knowledge_entry.parent.mkdir(parents=True, exist_ok=True)
    knowledge_entry.write_text(
        "# Work 1 fixture\n\n## 元数据\n\n- Work：1\n- 状态：review\n\n任务目标、反馈、修正、验证与证据均已记录。",
        encoding="utf-8",
    )
    (tmp_path / "任务知识库" / "INDEX.md").write_text(
        "# Index\n\n| Work | 日期 | 任务 | 状态 | 条目 |\n| ---: | --- | --- | --- | --- |\n| 1 | 2026-09-02 | fixture | review | [Work 1](entries/2026/work-1-fixture.md) |\n",
        encoding="utf-8",
    )
    return work


def _make_minimal_work(tmp_path: Path) -> Path:
    work = tmp_path / "work" / "7"
    for directory in ("input/source", "input/reference", "workspace/scripts", "workspace/qa", "output"):
        (work / directory).mkdir(parents=True, exist_ok=True)
    (work / "work.json").write_text(json.dumps({
        "work_id": "7", "name": "minimal fixture", "status": "in_progress",
        "storage_profile": "minimal", "conversation_retention": "disabled",
    }), encoding="utf-8")
    (work / "input/request.md").write_text("制作一个交付文件", encoding="utf-8")
    (work / "input/source/source.txt").write_text("provided", encoding="utf-8")
    (work / "workspace/scripts/build.py").write_text("print('build')\n", encoding="utf-8")
    (work / "workspace/settings.json").write_text('{"format":"fixture"}', encoding="utf-8")
    (work / "workspace/runtime-environment.json").write_text('{"host":"fixture"}', encoding="utf-8")
    (work / "workspace/.env.local").write_text("SECRET=must-not-be-retained", encoding="utf-8")
    (work / "workspace/qa/quality-receipt.json").write_text('{"status":"passed"}', encoding="utf-8")
    (work / "output/final.txt").write_text("deliverable", encoding="utf-8")
    return work


def _run_script(name: str, tmp_path: Path, *args: str) -> subprocess.CompletedProcess[str]:
    script = Path(__file__).resolve().parents[2] / "scripts" / name
    shell = POWERSHELL
    assert shell is not None
    command = [shell, "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script), *args]
    if name == "finish-work.ps1":
        command.extend(["-TaskKnowledgeRoot", str(tmp_path / "任务知识库")])
    command.extend(["-WorkRoot", str(tmp_path / "work")])
    return subprocess.run(
        command,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
        check=False,
    )


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_prepare_work_close_captures_adapter_evidence(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    completed = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1")

    assert completed.returncode == 0, completed.stderr or completed.stdout
    manifest = json.loads((work / "workspace" / "qa" / "final-output-manifest.json").read_text(encoding="utf-8-sig"))
    assert len(manifest["adapter_evidence"]) == 1
    assert manifest["adapter_evidence"][0]["path"].endswith("adapter-last-result.json")
    assert len(manifest["asset_evidence"]) == 3
    assert len(manifest["knowledge_context"]) == 1
    record = json.loads((work / "work.json").read_text(encoding="utf-8-sig"))
    assert record["required_confirmation_phrase"] == "确认 Work 1 彻底完成并结项"


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_minimal_close_creates_one_bound_receipt_and_finishes_only_after_approval(tmp_path: Path) -> None:
    work = _make_minimal_work(tmp_path)
    prepared = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "7")
    assert prepared.returncode == 0, prepared.stderr or prepared.stdout
    receipt_path = work / "workspace/production-receipt.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8-sig"))
    assert receipt["status"] == "review" and receipt["user_approved"] is False
    assert receipt["qa"]["status"] == "captured_not_assessed"
    assert [row["path"] for row in receipt["outputs"]] == ["output/final.txt"]
    assert "workspace/scripts/build.py" in {row["path"] for row in receipt["production_files"]}
    assert "workspace/settings.json" in {row["path"] for row in receipt["production_files"]}
    assert "workspace/runtime-environment.json" in {row["path"] for row in receipt["production_files"]}
    assert "workspace/.env.local" not in {row["path"] for row in receipt["production_files"]}
    assert not (work / "output/工作完成报告.md").exists()
    assert not (work / "workspace/knowledge-candidate.md").exists()
    assert not (work / "workspace/qa/closure-evidence.json").exists()
    assert not (work / "workspace/conversations").exists()

    prepared_again = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "7")
    assert prepared_again.returncode == 0, prepared_again.stderr or prepared_again.stdout
    denied = _run_script("finish-work.ps1", tmp_path, "-WorkId", "7", "-ApprovalNote", "未给出确认", "-KnowledgeStatus", "no-new-knowledge")
    assert denied.returncode != 0
    assert json.loads((work / "work.json").read_text(encoding="utf-8-sig"))["status"] == "review"

    completed = _run_script("finish-work.ps1", tmp_path, "-WorkId", "7", "-UserApproved", "-ApprovalNote", "用户明确同意结束", "-KnowledgeStatus", "no-new-knowledge")
    assert completed.returncode == 0, completed.stderr or completed.stdout
    record = json.loads((work / "work.json").read_text(encoding="utf-8-sig"))
    assert record["status"] == "completed" and record["user_approved"] is True
    assert record["knowledge_status"] == "no-new-knowledge"
    assert record["maintenance_status"] == "skipped_minimal_profile"
    assert receipt_path.is_file()


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
@pytest.mark.parametrize("mutation", ["output_bytes", "output_set", "receipt", "input_bytes", "input_set", "production_bytes", "production_set", "runtime"])
def test_minimal_finish_rejects_changed_review_evidence(tmp_path: Path, mutation: str) -> None:
    work = _make_minimal_work(tmp_path)
    prepared = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "7")
    assert prepared.returncode == 0, prepared.stderr or prepared.stdout
    if mutation == "output_bytes":
        (work / "output/final.txt").write_text("changed", encoding="utf-8")
    elif mutation == "output_set":
        (work / "output/extra.txt").write_text("extra", encoding="utf-8")
    elif mutation == "input_bytes":
        (work / "input/source/source.txt").write_text("changed", encoding="utf-8")
    elif mutation == "input_set":
        (work / "input/reference/extra.txt").write_text("new reference", encoding="utf-8")
    elif mutation == "production_bytes":
        (work / "workspace/settings.json").write_text('{"format":"changed"}', encoding="utf-8")
    elif mutation == "production_set":
        (work / "workspace/new-build.py").write_text("print('new code')\n", encoding="utf-8")
    elif mutation == "runtime":
        (work / "workspace/runtime-environment.json").write_text('{"host":"changed"}', encoding="utf-8")
    else:
        (work / "workspace/production-receipt.json").write_text("{}", encoding="utf-8")
    finished = _run_script("finish-work.ps1", tmp_path, "-WorkId", "7", "-UserApproved", "-ApprovalNote", "已验收", "-KnowledgeStatus", "no-new-knowledge")
    assert finished.returncode != 0
    assert json.loads((work / "work.json").read_text(encoding="utf-8-sig"))["status"] == "review"


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_minimal_close_blocks_pptx_with_unreviewed_overlap(tmp_path: Path) -> None:
    work = _make_minimal_work(tmp_path)
    (work / "output/final.txt").unlink()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "主要结论"
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "遮挡主要结论"
    deck.save(work / "output/final.pptx")
    prepared = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "7")
    assert prepared.returncode != 0
    assert not (work / "workspace/production-receipt.json").exists()
    assert json.loads((work / "work.json").read_text(encoding="utf-8-sig"))["status"] == "in_progress"
    reports = list((work / "workspace/qa/output-preflight").glob("*.json"))
    assert len(reports) == 1
    assert json.loads(reports[0].read_text())["quality"]["summary"]["codes"]["TEXT_OVERLAP"] >= 1


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_minimal_reopen_clears_receipt_binding_without_creating_history(tmp_path: Path) -> None:
    work = _make_minimal_work(tmp_path)
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "7").returncode == 0
    assert _run_script("finish-work.ps1", tmp_path, "-WorkId", "7", "-UserApproved", "-ApprovalNote", "已验收", "-KnowledgeStatus", "no-new-knowledge").returncode == 0
    before = {path.relative_to(work).as_posix() for path in work.rglob("*") if path.is_file()}
    receipt_before = (work / "workspace/production-receipt.json").read_bytes()
    reopened = _run_script("reopen-work.ps1", tmp_path, "-WorkId", "7", "-Reason", "用户要求修改")
    assert reopened.returncode == 0, reopened.stderr or reopened.stdout
    record = json.loads((work / "work.json").read_text(encoding="utf-8-sig"))
    assert record["status"] == "in_progress" and record["user_approved"] is False
    assert "production_receipt" not in record and "production_receipt_sha256" not in record
    assert (work / "workspace/production-receipt.json").read_bytes() == receipt_before
    after = {path.relative_to(work).as_posix() for path in work.rglob("*") if path.is_file()}
    assert before == after
    assert not (work / "workspace/review-history").exists()
    prepared_again = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "7")
    assert prepared_again.returncode == 0, prepared_again.stderr or prepared_again.stdout


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_prepare_work_close_blocks_current_pptx_with_unreviewed_overlap(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    target = work / "output" / "fixture.pptx"
    deck = Presentation(target)
    slide = deck.slides[0]
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "第一页正文"
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "遮挡的另一段"
    deck.save(target)

    prepared = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1")

    assert prepared.returncode != 0
    assert "current PPTX quality needs review" in (prepared.stderr + prepared.stdout)
    assert not (work / "workspace" / "qa" / "final-output-manifest.json").exists()
    reports = list((work / "workspace" / "qa" / "output-preflight").glob("*.json"))
    assert len(reports) == 1
    assert json.loads(reports[0].read_text())["quality"]["summary"]["codes"]["TEXT_OVERLAP"] >= 1
    assert json.loads(reports[0].read_text())["observation"]["status"] == "recorded"
    assert list((work / "workspace" / "quality-observations").glob("*.json"))
    assert json.loads((work / "work.json").read_text())["status"] == "in_progress"


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_review_can_be_reopened_and_requires_new_knowledge_evaluation(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0

    reopened = _run_script(
        "reopen-work.ps1",
        tmp_path,
        "-WorkId",
        "1",
        "-Reason",
        "用户要求继续调整版式",
    )

    assert reopened.returncode == 0, reopened.stderr or reopened.stdout
    record = json.loads((work / "work.json").read_text(encoding="utf-8-sig"))
    assert record["status"] == "in_progress"
    assert record["user_approved"] is False
    assert "需重新评估" in (work / "workspace" / "knowledge-candidate.md").read_text(encoding="utf-8-sig")
    blocked_finish = _run_script(
        "finish-work.ps1",
        tmp_path,
        "-WorkId",
        "1",
        "-UserApproved",
        "-ApprovalNote",
        "确认 Work 1 彻底完成并结项",
        "-KnowledgeStatus",
        "no-new-knowledge",
    )
    assert blocked_finish.returncode != 0


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_finish_requires_explicit_approval_and_unchanged_review_outputs(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0
    (work / "workspace" / "knowledge-candidate.md").write_text(
        "# Knowledge candidate\n\n状态：已评估\n\n本次无新的通用知识；现有规则已经覆盖。",
        encoding="utf-8",
    )

    missing_approval = _run_script(
        "finish-work.ps1",
        tmp_path,
        "-WorkId",
        "1",
        "-UserApproved:$false",
        "-ApprovalNote",
        "尚未确认",
        "-KnowledgeStatus",
        "no-new-knowledge",
    )
    assert missing_approval.returncode != 0

    (work / "output" / "fixture.pptx").write_bytes(b"changed")
    changed_output = _run_script(
        "finish-work.ps1",
        tmp_path,
        "-WorkId",
        "1",
        "-UserApproved",
        "-ApprovalNote",
        "确认 Work 1 彻底完成并结项",
        "-KnowledgeStatus",
        "no-new-knowledge",
    )
    assert changed_output.returncode != 0
    assert "changed after review" in (changed_output.stderr + changed_output.stdout)

    restored = Presentation()
    restored.slides.add_slide(restored.slide_layouts[6])
    restored.save(work / "output" / "fixture.pptx")
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0
    finished = _run_script(
        "finish-work.ps1",
        tmp_path,
        "-WorkId",
        "1",
        "-UserApproved",
        "-ApprovalNote",
        "确认 Work 1 彻底完成并结项",
        "-KnowledgeStatus",
        "no-new-knowledge",
    )
    assert finished.returncode == 0, finished.stderr or finished.stdout
    record = json.loads((work / "work.json").read_text(encoding="utf-8-sig"))
    assert record["status"] == "completed"
    assert record["user_approved"] is True
    assert record["task_knowledge_entry"].endswith("work-1-fixture.md")
    assert "- 状态：completed" in (
        tmp_path / "任务知识库" / "entries" / "2026" / "work-1-fixture.md"
    ).read_text(encoding="utf-8-sig")
    assert "| completed |" in (tmp_path / "任务知识库" / "INDEX.md").read_text(encoding="utf-8-sig")


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_finish_requires_independent_task_knowledge_entry(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0
    (work / "workspace" / "knowledge-candidate.md").write_text(
        "# Knowledge candidate\n\n状态：已评估\n\n本次无新的通用知识；现有规则已经覆盖。",
        encoding="utf-8",
    )
    (tmp_path / "任务知识库" / "entries" / "2026" / "work-1-fixture.md").unlink()

    blocked = _run_script(
        "finish-work.ps1",
        tmp_path,
        "-WorkId",
        "1",
        "-UserApproved",
        "-ApprovalNote",
        "确认 Work 1 彻底完成并结项",
        "-KnowledgeStatus",
        "no-new-knowledge",
    )

    assert blocked.returncode != 0
    assert "requires exactly one independent task knowledge entry" in (blocked.stderr + blocked.stdout)


def _evaluated_candidate(work: Path) -> None:
    (work / "workspace/knowledge-candidate.md").write_text(
        "已评估：已有通用规则覆盖，本次没有新的可复用结论。", encoding="utf-8")


def _finish(tmp_path: Path):
    return _run_script("finish-work.ps1", tmp_path, "-WorkId", "1", "-UserApproved",
                       "-ApprovalNote", "这个任务完成了，归档", "-KnowledgeStatus", "no-new-knowledge")


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_output_added_after_review_and_source_changed_are_detected(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    source = work / "input/source/source.txt"
    source.write_text("source")
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0
    _evaluated_candidate(work)
    extra = work / "output/unreviewed.pdf"
    extra.write_text("new unreviewed deliverable")
    result = _finish(tmp_path)
    assert result.returncode != 0
    assert "file set changed" in result.stderr
    extra.unlink()
    source.write_text("changed")
    result = _finish(tmp_path)
    assert result.returncode != 0
    assert "Source file changed" in result.stderr


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_new_closure_contract_rejects_empty_retrospective_and_preserves_old_review(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    record = json.loads((work / "work.json").read_text())
    record["closure_schema_version"] = 2
    (work / "work.json").write_text(json.dumps(record))
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0
    _evaluated_candidate(work)
    first_report = (work / "workspace/completion-report.md").read_bytes()
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0
    history = list((work / "workspace/review-history").glob("*/completion-report.md"))
    assert any(p.read_bytes() == first_report for p in history)
    result = _finish(tmp_path)
    assert result.returncode != 0
    assert "substantive section" in result.stderr


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_completed_work_can_reopen_and_withdraw_confirmed_case(tmp_path: Path) -> None:
    work = _make_work(tmp_path)
    assert _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1").returncode == 0
    _evaluated_candidate(work)
    result = _finish(tmp_path)
    assert result.returncode == 0, result.stderr
    reopened = _run_script("reopen-work.ps1", tmp_path, "-WorkId", "1", "-Reason", "用户要求继续修改")
    assert reopened.returncode == 0, reopened.stderr
    assert json.loads((work / "work.json").read_text(encoding="utf-8-sig"))["status"] == "in_progress"
    entry = tmp_path / "任务知识库/entries/2026/work-1-fixture.md"
    assert "- 状态：in_progress" in entry.read_text(encoding="utf-8-sig")
    assert "| in_progress |" in (tmp_path / "任务知识库/INDEX.md").read_text(encoding="utf-8-sig")


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_finish_binds_completed_retrospective_bytes_and_later_edits_lose_approval(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from aippt_studio import context_intelligence as ci

    monkeypatch.setattr(ci, "recommend_upstream_tools", lambda query: [])
    work = _make_work(tmp_path)
    record = json.loads((work / "work.json").read_text())
    record["closure_schema_version"] = 2
    (work / "work.json").write_text(json.dumps(record))
    entry = tmp_path / "任务知识库/entries/2026/work-1-fixture.md"
    entry.write_text("# 企业模板经营汇报\n\n- Work：1\n- 状态：review\n\n" + "\n".join(
        f"## {heading}\n\n企业模板经营汇报测试：保留模板母版、版式和主题；已保存本次核验的证据。\n"
        for heading in ("任务目标与完成结果", "输入语义", "用户反馈与对应修改", "失败、诊断与恢复", "最终验证", "可复用经验", "不应推广的任务事实", "证据位置")
    ), encoding="utf-8")
    prior_hash = hashlib.sha256(entry.read_bytes()).hexdigest()
    prepared = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1")
    assert prepared.returncode == 0, prepared.stderr
    _evaluated_candidate(work)
    finished = _finish(tmp_path)
    assert finished.returncode == 0, finished.stderr
    record = json.loads((work / "work.json").read_text(encoding="utf-8-sig"))
    expected = hashlib.sha256(entry.read_bytes()).hexdigest()
    assert record["task_knowledge_entry_sha256"] == expected
    assert expected != prior_hash, "approval must bind the completed bytes, after the status conversion"
    assert json.loads(finished.stdout)["task_knowledge_entry_sha256"] == expected
    result = ci.recall_context("企业模板经营汇报", tmp_path)
    assert result["completed_lessons"] == []
    assert result["unfinished_cases"][0]["trust"] == "confirmed_completed"
    assert "reusable_experience" not in result["unfinished_cases"][0]
    with entry.open("a", encoding="utf-8") as stream:
        stream.write("\n## 未审阅新增\n\n新增经验没有经过再次完成确认。\n")
    result = ci.recall_context("企业模板经营汇报", tmp_path)
    assert result["completed_lessons"] == []
    assert result["unfinished_cases"][0]["trust"] == "pending_or_unverified"
