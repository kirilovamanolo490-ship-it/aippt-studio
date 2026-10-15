import hashlib
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.shapes.picture import CT_Picture
from pptx.util import Inches
import pytest

from aippt_studio.motion_clone import _read_xml, _safe_entries, _slide_order, _slide_objects, text_sha256
from aippt_studio.template_content_replace import (
    TemplateContentReplaceError, _image_refs, inspect_page_images, replace_template_content,
)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _deck(path, image_path):
    ppt = Presentation()
    for label in ("第一段旧事实", "第二段旧事实"):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.text_frame.paragraphs[0].add_run().text = label
        shape.text_frame.paragraphs[0].add_run().text = "；补充旧数据"
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(2), width=Inches(2))
    ppt.save(path)
    return path


def _parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _shape_bindings(path, slide=1):
    with ZipFile(path) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        root = _read_xml(archive, entries, slides[slide - 1])
        objects = _slide_objects(root)
        text_id = next(i for i, shape in objects.items() if shape.xpath(".//a:t", namespaces={"a": A}))
        pic_id = next(i for i, shape in objects.items() if shape.tag.endswith("}pic"))
        shape = objects[text_id]
        old_text = "".join(shape.xpath(".//a:t/text()", namespaces={"a": A}))
        pic = objects[pic_id]
        rid = pic.find(f".//{{{A}}}blip").get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        rels = _read_xml(archive, entries, f"ppt/slides/_rels/slide{slide}.xml.rels")
        target = next(rel.get("Target") for rel in rels if rel.get("Id") == rid)
        from posixpath import normpath
        media_part = normpath("ppt/slides/" + target)
        return text_id, old_text, pic_id, media_part, archive.read(media_part)


@pytest.fixture
def sample(tmp_path):
    image = tmp_path / "old.png"
    Image.new("RGB", (30, 20), "red").save(image)
    source = _deck(tmp_path / "source.pptx", image)
    replacement = tmp_path / "new.png"
    Image.new("RGB", (30, 20), "blue").save(replacement)
    return source, replacement


def test_multirun_text_and_shared_picture_copy_on_write_preserves_other_consumers(sample, tmp_path):
    source, replacement = sample
    original = _parts(source)
    text_id, old_text, pic_id, media_part, media = _shape_bindings(source)
    output = tmp_path / "candidate.pptx"
    receipt = replace_template_content(
        source, output, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        text_replacements=[{"slide": 1, "shape_id": text_id,
                            "original_text_sha256": text_sha256(old_text),
                            "text": "全新且完整的目标页文案"}],
        image_replacements=[{"slide": 1, "shape_id": pic_id,
                             "original_media_sha256": hashlib.sha256(media).hexdigest(),
                             "path": str(replacement)}],
    )
    changed = _parts(output)
    with ZipFile(output) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        root = _read_xml(archive, entries, slides[0])
        shape = _slide_objects(root)[text_id]
        actual = "".join(shape.xpath(".//a:t/text()", namespaces={"a": A}))
        assert actual == "全新且完整的目标页文案"
        assert len(shape.findall(f".//{{{A}}}p")) == 1
        pic_part = _shape_bindings(output)[3]
        assert pic_part != media_part
        assert hashlib.sha256(archive.read(pic_part)).hexdigest() == receipt["image_bindings"][0]["replacement_media_sha256"]
        assert receipt["image_bindings"][0]["remaining_original_consumers"] == 1
        assert receipt["image_bindings"][0]["privacy_cleanup_pending"] is True
    rel_part = "ppt/slides/_rels/slide1.xml.rels"
    assert changed[media_part] == original[media_part]
    assert all(changed[name] == data for name, data in original.items()
               if name not in {slides[0], rel_part})
    assert changed[slides[0]] != original[slides[0]]
    assert changed[rel_part] != original[rel_part]
    assert receipt["non_target_package_parts_preserved"] is True
    assert receipt["native_visual_review_required"] is True


def test_shared_media_with_unmapped_second_slide_is_isolated(sample, tmp_path):
    source, replacement = sample
    _, _, pic_id, _, media = _shape_bindings(source)
    output = tmp_path / "isolated.pptx"
    receipt = replace_template_content(
        source, output, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        image_replacements=[{"slide": 1, "shape_id": pic_id,
                             "original_media_sha256": hashlib.sha256(media).hexdigest(),
                             "path": str(replacement)}],
    )
    with ZipFile(output) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        first = _read_xml(archive, entries, slides[0])
        second = _read_xml(archive, entries, slides[1])
        ref1 = next(r for r in _image_refs(archive, entries, slides[0], first, 1) if r["shape_id"] == pic_id)
        second_pic = _shape_bindings(source, 2)[2]
        ref2 = next(r for r in _image_refs(archive, entries, slides[1], second, 2) if r["shape_id"] == second_pic)
        assert ref1["media_part"] != ref2["media_part"]
        assert archive.read(ref2["media_part"]) == media
    assert receipt["image_bindings"][0]["copy_on_write"] is True


def test_text_hash_and_source_hash_are_required(sample, tmp_path):
    source, _ = sample
    text_id, old_text, _, _, _ = _shape_bindings(source)
    args = {"source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "text_replacements": [{"slide": 1, "shape_id": text_id,
                                   "original_text_sha256": "0" * 64, "text": "替换"}]}
    with pytest.raises(TemplateContentReplaceError, match="text SHA"):
        replace_template_content(source, tmp_path / "bad-text.pptx", **args)
    with pytest.raises(TemplateContentReplaceError, match="source_sha256"):
        replace_template_content(source, tmp_path / "bad-source.pptx", source_sha256="0" * 64,
                                 text_replacements=[{"slide": 1, "shape_id": text_id,
                                                     "original_text_sha256": text_sha256(old_text), "text": "替换"}])


def test_same_slide_shapes_sharing_one_rid_get_a_new_relationship_for_target(tmp_path):
    image = tmp_path / "old.png"
    replacement = tmp_path / "replacement.png"
    Image.new("RGB", (40, 30), "red").save(image)
    Image.new("RGB", (40, 30), "green").save(replacement)
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_picture(str(image), Inches(1), Inches(1), width=Inches(1))
    slide.shapes.add_picture(str(image), Inches(3), Inches(1), width=Inches(1))
    source = tmp_path / "same-slide-shared.pptx"
    ppt.save(source)
    with ZipFile(source) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        source_root = _read_xml(archive, entries, slides[0])
        refs = _image_refs(archive, entries, slides[0], source_root, 1)
        first_id, media_part = refs[0]["shape_id"], refs[0]["media_part"]
        media = archive.read(media_part)
    output = tmp_path / "same-slide-cow.pptx"
    replace_template_content(source, output, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                             image_replacements=[{"slide": 1, "shape_id": first_id,
                                                  "original_media_sha256": hashlib.sha256(media).hexdigest(),
                                                  "path": str(replacement)}])
    with ZipFile(output) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        root = _read_xml(archive, entries, slides[0])
        refs = _image_refs(archive, entries, slides[0], root, 1)
        assert len(refs) == 2
        assert refs[0]["media_part"] != refs[1]["media_part"]
        assert archive.read(refs[0]["media_part"]) == replacement.read_bytes()
        assert archive.read(refs[1]["media_part"]) == media
        assert refs[1]["media_part"] == media_part


def test_single_consumer_picture_replaces_media_in_place(tmp_path):
    image = tmp_path / "only.png"
    replacement = tmp_path / "replacement.png"
    Image.new("RGB", (20, 20), "black").save(image)
    Image.new("RGB", (20, 20), "white").save(replacement)
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_picture(str(image), Inches(1), Inches(1), width=Inches(1))
    source = tmp_path / "one-consumer.pptx"
    ppt.save(source)
    with ZipFile(source) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        root = _read_xml(archive, entries, slides[0])
        binding = _image_refs(archive, entries, slides[0], root, 1)[0]
        media = archive.read(binding["media_part"])
    output = tmp_path / "one-consumer-candidate.pptx"
    receipt = replace_template_content(
        source, output, source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        image_replacements=[{"slide": 1, "shape_id": binding["shape_id"],
                             "original_media_sha256": hashlib.sha256(media).hexdigest(),
                             "path": str(replacement)}],
    )
    assert receipt["image_bindings"][0]["copy_on_write"] is False
    assert receipt["image_bindings"][0]["privacy_cleanup_pending"] is False
    with ZipFile(output) as archive:
        assert archive.read(binding["media_part"]) == replacement.read_bytes()


def test_shape_image_fill_is_listed_and_replaced_without_changing_other_users(tmp_path):
    image = tmp_path / "fill.png"
    replacement = tmp_path / "replacement.png"
    Image.new("RGB", (40, 40), "orange").save(image)
    Image.new("RGB", (40, 40), "purple").save(replacement)
    ppt = Presentation()
    shape_ids = []
    for index in range(2):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        picture = slide.shapes.add_picture(str(image), Inches(1), Inches(1), width=Inches(1))
        if index == 0:
            rid = picture._element.blipFill.blip.rEmbed
            fill_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(1), Inches(1), Inches(1))
            shape_ids.append(str(fill_shape.shape_id))
            sp_pr = fill_shape._element.spPr
            for child in list(sp_pr):
                if child.tag in {qn("a:solidFill"), qn("a:noFill"), qn("a:gradFill"), qn("a:blipFill"), qn("a:pattFill")}:
                    sp_pr.remove(child)
            blip_fill = parse_xml(f'<a:blipFill {nsdecls("a", "r")}><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>')
            sp_pr.append(blip_fill)
    layout = ppt.slide_layouts[6]
    _, layout_rid = layout.part.get_or_add_image_part(str(image))
    layout_pic = CT_Picture.new_pic(99, "Layout Avatar", "", layout_rid,
                                    Inches(6), Inches(4), Inches(1), Inches(1))
    layout.shapes._spTree.insert_element_before(layout_pic, "p:extLst")
    source = tmp_path / "shape-fill.pptx"
    ppt.save(source)
    inspection = inspect_page_images(source, 1)
    fill_obj = next(obj for obj in inspection["objects"] if obj["shape_id"] == shape_ids[0])
    assert fill_obj["visual_object_type"] == "shape_image_fill"
    assert fill_obj["shape_name"]
    assert fill_obj["geometry_emu"] and fill_obj["media"][0]["media_sha256"]
    assert fill_obj["safe_for_single_image_replacement"]
    layout_obj = next(obj for obj in inspection["objects"] if obj["layer"] == "layout")
    assert layout_obj["media"][0]["media_sha256"] == fill_obj["media"][0]["media_sha256"]
    original_media = fill_obj["media"][0]
    output = tmp_path / "shape-fill-cow.pptx"
    receipt = replace_template_content(source, output, source_sha256=inspection["source_sha256"],
                                      image_replacements=[{"slide": 1, "shape_id": shape_ids[0],
                                                           "original_media_sha256": original_media["media_sha256"],
                                                           "path": str(replacement)}])
    assert receipt["image_bindings"][0]["visual_object_type"] == "shape_image_fill"
    assert receipt["image_bindings"][0]["copy_on_write"] is True
    assert receipt["image_bindings"][0]["remaining_consumer_layers"] == ["layout", "slide"]
    with ZipFile(output) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        refs = _image_refs(archive, entries, slides[0], _read_xml(archive, entries, slides[0]), 1)
        target = next(row for row in refs if row["shape_id"] == shape_ids[0])
        untouched_picture = next(row for row in refs if row["visual_object_type"] == "picture")
        assert archive.read(target["media_part"]) == replacement.read_bytes()
        assert archive.read(untouched_picture["media_part"]) == archive.read(original_media["media_part"])
    output_inspection = inspect_page_images(output, 1)
    output_layout_obj = next(obj for obj in output_inspection["objects"] if obj["layer"] == "layout")
    assert output_layout_obj["media"][0]["media_sha256"] == original_media["media_sha256"]


def test_shape_with_multiple_image_blips_is_reported_but_rejected(tmp_path):
    image = tmp_path / "multi.png"
    replacement = tmp_path / "replacement.png"
    Image.new("RGB", (24, 24), "gold").save(image)
    Image.new("RGB", (24, 24), "navy").save(replacement)
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    picture = slide.shapes.add_picture(str(image), Inches(1), Inches(1), width=Inches(1))
    rid = picture._element.blipFill.blip.rEmbed
    fill_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(1), Inches(1), Inches(1))
    blip_fill = parse_xml(f'<a:blipFill {nsdecls("a", "r")}><a:blip r:embed="{rid}"/><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>')
    sp_pr = fill_shape._element.spPr
    for child in list(sp_pr):
        if child.tag in {qn("a:solidFill"), qn("a:noFill")}:
            sp_pr.remove(child)
    sp_pr.append(blip_fill)
    source = tmp_path / "ambiguous-fill.pptx"
    ppt.save(source)
    inspection = inspect_page_images(source, 1)
    target = next(obj for obj in inspection["objects"] if obj["shape_id"] == str(fill_shape.shape_id))
    assert target["safe_for_single_image_replacement"] is False
    assert len(target["media"]) == 2
    with pytest.raises(TemplateContentReplaceError, match="exactly one"):
        replace_template_content(source, tmp_path / "rejected.pptx",
                                 source_sha256=inspection["source_sha256"],
                                 image_replacements=[{"slide": 1, "shape_id": str(fill_shape.shape_id),
                                                      "original_media_sha256": target["media"][0]["media_sha256"],
                                                      "path": str(replacement)}])
    assert not (tmp_path / "rejected.pptx").exists()
