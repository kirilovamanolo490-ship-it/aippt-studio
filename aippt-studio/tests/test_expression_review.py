import json
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from typer.testing import CliRunner
from aippt_studio.cli import app
from aippt_studio.expression_review import copy_signals, prepare_expression_review, record_expression_review, require_expression_review
from aippt_studio.models import DeckPlan, JobState, SlideSpec
from aippt_studio.storage import job_dir, read_json, save_job, write_json
from aippt_studio.workflow import WorkflowError, approve_gate, create_job

@pytest.fixture
def candidate(monkeypatch):
    monkeypatch.delenv('AIPPT_FAKE_RENDER')
    # A synthetic review fixture must never enter a real Work's learning log
    # when pytest's evidence directory happens to be inside that Work package.
    monkeypatch.setattr('aippt_studio.authoring_policy.record_quality_findings',lambda *args: {'status':'test_fixture_only'})
    def render_fixture(source, destination):
        p=Presentation(source);(destination/'slides').mkdir(exist_ok=True)
        for i,slide in enumerate(p.slides):
            # Explicit test renderer only: color changes bind this fixture to
            # actual source bytes, without claiming Office visual acceptance.
            value=sum(source.read_bytes())%255
            Image.new('RGB',(320,180),(value,i*30,240)).save(destination/f'slides/slide-{i+1}.png')
        return {'status':'test_renderer_only'}
    monkeypatch.setattr('aippt_studio.expression_review.powerpoint_roundtrip',render_fixture)
    job=create_job('表达测试',job_id='expression-fixture');job.state=JobState.FINAL_REVIEW;save_job(job)
    root=job_dir(job.id)
    plan=DeckPlan(title='测试数据',slides=[SlideSpec(id='s1',title='本项目完成三地访谈',purpose='说明范围',layout='content',points=['保留24份原始记录。']),SlideSpec(id='s2',title='下一阶段验证',purpose='承接验证问题',layout='content',points=['按相同口径复核服务覆盖情况。'])])
    write_json(root/'planning/deck-plan.json',plan)
    prs=Presentation();prs.slide_width=Inches(13.3333);prs.slide_height=Inches(7.5)
    for i,spec in enumerate(plan.slides):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);shape=slide.shapes.add_textbox(Inches(1),Inches(2),Inches(10),Inches(2));shape.text=spec.title+'\n'+spec.points[0]
        for p in shape.text_frame.paragraphs:p.runs[0].font.size=Pt(20)
        Image.new('RGB',(320,180),(245,i*30,240)).save(root/f'build/previews/{spec.id}.png')
    prs.save(root/'build/deck.pptx');prepare_expression_review(job.id,'final')
    return job,root

def request_for(root):
    p=read_json(root/'reviews/final-expression-packet.json')
    return {'packet_sha256':p['packet_sha256'],'reviewer':'current_assistant','narrative_assessment':'Unit-test payload: next validation follows completed visits, not actual visual acceptance.','slides':[{'slide_id':s['slide_id'],'preview_sha256':s['preview_sha256'],'audience_takeaway':s['internal_plan']['title'],'page_assessment':'Unit-test review fixture: claim and evidence have the specified page role.','decision':'ready','signal_dispositions':[]} for s in p['slides']]}

@pytest.mark.parametrize('text',['本项目完成24份访谈，记录按地区编号。','工艺流程：先清洗，再沉积涂层，最后检测附着力。','比较本项目与既有方法的准确率，统计口径一致。','第三阶段需要检查设备是否漏气，再启动系统。'])
def test_legitimate_business_language_not_banned(text):assert copy_signals(text)==[]

@pytest.mark.parametrize('text',['这页需要注意与前页的关联。','我会帮你调整这个PPT的排版。','制作说明：请放大标题，配图稍后补充。','请检查图文遮挡情况。'])
def test_production_copy_is_context_signal(text):assert copy_signals(text) and all(s['kind']=='context_required' for s in copy_signals(text))

@pytest.mark.parametrize('text',['我会帮你优化PPT。','以下是为您准备的汇报PPT。','这一页可以放项目背景。'])
def test_direct_deck_production_phrasing_is_signaled(text):
    expected='PRODUCTION_INSTRUCTION' if text.startswith('这一页') else 'ASSISTANT_REPLY'
    assert any(row['code']==expected for row in copy_signals(text))

@pytest.mark.parametrize('text',['培训材料引用：“我会帮你优化PPT”，说明助手口吻。','本项目将按技术流程优化PPT涂布参数。','受众请先完成设备断电，再开始检修。'])
def test_contextual_and_audience_language_is_preserved(text):
    assert copy_signals(text)==[]

def test_density_signals_explain_geometry_and_are_contextual():
    from aippt_studio.expression_review import density_signals
    objects=[
        {'object_id':'2','text':'主要结论','bounds_pt':{'left':20,'top':20,'width':500,'height':60},'font_sizes_pt':[36]},
        {'object_id':'3','text':'这是一段较长的说明文字。'*12,'bounds_pt':{'left':40,'top':120,'width':150,'height':320},'font_sizes_pt':[14]},
    ]
    signals=density_signals(objects,960,540)
    assert len(signals)==1
    signal=signals[0]
    assert signal['object_id']=='3' and signal['kind']=='context_required'
    assert signal['components']['width_ratio']==pytest.approx(150/960,abs=.0001)
    assert signal['components']['font_contrast_ratio']==pytest.approx(36/14,abs=.001)
    assert 'does not judge meaning' in signal['limitations']

def test_packet_actual_copy_and_neighbors_remain_pending(candidate):
    _,root=candidate;p=read_json(root/'reviews/final-expression-packet.json')
    assert p['status']=='awaiting_assistant_review'
    assert p['slides'][0]['objects'][0]['text'].startswith('本项目完成三地访谈')
    assert p['slides'][0]['next_page']['id']=='s2' and p['slides'][1]['previous_page']['id']=='s1'
    assert all(signal['signal_id'] for signal in p['slides'][0]['signals'])
    with pytest.raises(ValueError,match='missing'):require_expression_review(root,'final')


def test_design_context_is_visible_and_expires_with_its_scene(candidate):
    from aippt_studio.renderer import sha256
    job, root = candidate
    scene = root / 'planning/native-scene.json'
    write_json(scene, {'slides': []})
    context = {'status': 'advisory', 'form': 'business_report', 'missing_narrative_fields': ['audience_need'],
               'plan_binding': {'plan_sha256': sha256(root / 'planning/deck-plan.json'), 'brief_sha256': None},
               'scene_sha256': sha256(scene)}
    write_json(root / 'planning/design-context.json', context)
    prepare_expression_review(job.id, 'final')
    packet = read_json(root / 'reviews/final-expression-packet.json')
    assert packet['design_context_status'] == 'current_advisory'
    assert packet['context']['planning/design-context.json']['missing_narrative_fields'] == ['audience_need']
    assert 'planning/design-context.json' in packet['inputs']
    write_json(scene, {'slides': [], 'changed': True})
    prepare_expression_review(job.id, 'final')
    stale = read_json(root / 'reviews/final-expression-packet.json')
    assert stale['design_context_status'] == 'stale_or_not_recorded'
    assert stale['packet_sha256'] != packet['packet_sha256']

def test_density_candidate_is_packet_bound_and_requires_disposition(candidate):
    job,root=candidate
    deck_path=root/'build/deck.pptx';deck=Presentation(deck_path);slide=deck.slides[0]
    title=slide.shapes.add_textbox(Inches(1),Inches(.3),Inches(8),Inches(.7));title.text='核心结论'
    title.text_frame.paragraphs[0].runs[0].font.size=Pt(36)
    body=slide.shapes.add_textbox(Inches(.4),Inches(1.2),Inches(1.5),Inches(4));body.text='这是一段需要结合证据阅读的说明。'*10
    body.text_frame.paragraphs[0].runs[0].font.size=Pt(14)
    deck.save(deck_path);prepare_expression_review(job.id,'final')
    packet=read_json(root/'reviews/final-expression-packet.json')
    signal=next(row for row in packet['slides'][0]['signals'] if row['code']=='TEXT_DENSITY_CONTEXT')
    assert signal['object_id']==str(body.shape_id) and signal['signal_id']
    request=request_for(root)
    with pytest.raises(ValueError,match='Address each'):
        record_expression_review(job.id,'final',request)
    request['slides'][0]['signal_dispositions']=[{'signal_id':row['signal_id'],'decision':'legitimate','reason':'这是本页必要说明，已在实际页面尺度检查其可读性。'} for row in packet['slides'][0]['signals']]
    assert record_expression_review(job.id,'final',request)['status']=='ready'

def test_geometry_pass_cannot_replace_expression_judgment(candidate):
    job,_=candidate
    with pytest.raises(WorkflowError,match='whole-page expression review is missing'):approve_gate(job.id,'final')

def test_review_bound_to_render(candidate):
    job,root=candidate;report=record_expression_review(job.id,'final',request_for(root))
    assert require_expression_review(root,'final')==report
    packet=read_json(root/'reviews/final-expression-packet.json')
    Image.new('RGB',(320,180),'black').save(root/packet['slides'][0]['preview'])
    with pytest.raises(ValueError,match='stale'):require_expression_review(root,'final')

@pytest.mark.parametrize('relative',['build/deck.pptx','planning/deck-plan.json','planning/authoring-policy.json'])
def test_content_or_context_change_requires_new_judgment(candidate,relative):
    job,root=candidate;record_expression_review(job.id,'final',request_for(root));path=root/relative
    if relative.endswith('.pptx'):
        p=Presentation(path);p.slides[0].shapes[0].text='新的说法，需要重新判断';p.save(path)
    else:
        p=read_json(path)
        if 'deck-plan' in relative:p['slides'][0]['purpose']='另一论证意图'
        else:p['precedence']='新的当前参考'
        write_json(path,p)
    with pytest.raises(ValueError,match='stale'):require_expression_review(root,'final')

def test_missing_duplicate_and_empty_assessment_rejected(candidate):
    job,root=candidate;r=request_for(root);r['slides']=r['slides'][:1]*2
    with pytest.raises(ValueError,match='exactly once'):record_expression_review(job.id,'final',r)
    r=request_for(root);r['slides'][0]['page_assessment']='通过'
    with pytest.raises(ValueError,match='concrete observation'):record_expression_review(job.id,'final',r)

def test_copy_signal_requires_contextual_disposition(candidate):
    job,root=candidate;path=root/'build/deck.pptx';p=Presentation(path);p.slides[0].shapes[0].text='这页需要注意与前页的关联。';p.save(path)
    prepare_expression_review(job.id,'final');packet=read_json(root/'reviews/final-expression-packet.json');r=request_for(root)
    with pytest.raises(ValueError,match='Address each'):record_expression_review(job.id,'final',r)
    signal=packet['slides'][0]['signals'][0]
    r['slides'][0]['signal_dispositions']=[{'signal_id':signal['signal_id'],'decision':'revise','reason':'这是对制作者的编辑指令，不能作为当前项目汇报的结论。'}]
    with pytest.raises(ValueError,match='unresolved copy'):record_expression_review(job.id,'final',r)
    r['slides'][0]['decision']='revise';record_expression_review(job.id,'final',r)
    with pytest.raises(ValueError,match='still needs revision'):require_expression_review(root,'final')
    r['slides'][0]['decision']='ready';r['slides'][0]['signal_dispositions'][0].update(decision='legitimate',reason='另一场景的编辑培训明确引用该句作为错误示例，需保留原文。')
    assert record_expression_review(job.id,'final',r)['status']=='ready'

def test_cli_prepares_evidence_not_auto_pass(candidate):
    job,root=candidate;r=CliRunner().invoke(app,['expression-review',job.id,'--gate','final','--json'])
    assert r.exit_code==0,r.output
    assert json.loads(r.output)['status']=='awaiting_assistant_review'
    assert not (root/'reviews/final-expression-review.json').exists()

def test_already_approved_fastpath_still_requires_expression_review(candidate):
    job,_=candidate;job.approvals['final']=True;save_job(job)
    with pytest.raises(WorkflowError,match='whole-page expression review is missing'):approve_gate(job.id,'final')


def test_changed_approved_candidate_can_recover_through_formal_review(candidate,monkeypatch):
    from aippt_studio.workflow import _gate_hashes,run_job
    from aippt_studio.reconcile import audit_job
    job,root=candidate
    write_json(root/'planning/brief.json',{'test_fixture':True})
    (root/'build/deck.html').write_text('<html>test</html>')
    Image.new('RGB',(320,180),'white').save(root/'build/montage.png')
    prepare_expression_review(job.id,'final')
    record_expression_review(job.id,'final',request_for(root))
    job.approvals['final']=True;job.approval_hashes['final']=_gate_hashes(root,'final',job.route);save_job(job)
    prs=Presentation(root/'build/deck.pptx');prs.slides[0].shapes[0].left=Inches(2);prs.save(root/'build/deck.pptx')
    assert run_job(job.id)['state']=='blocked'
    prepare_expression_review(job.id,'final')
    assert record_expression_review(job.id,'final',request_for(root))['status']=='ready'
    assert audit_job(job.id).status=='consistent'
    monkeypatch.setattr('aippt_studio.workflow._read_deck_plan',lambda root: DeckPlan.model_validate(read_json(root/'planning/deck-plan.json')))
    monkeypatch.setattr('aippt_studio.workflow._validate_evidence_traceability',lambda *args: {'status':'test_fixture'})
    assert approve_gate(job.id,'final').approvals['final'] is True


def test_unrelated_blocker_cannot_be_cleared_by_expression_review(candidate):
    job,root=candidate;job.state=JobState.BLOCKED;job.resume_state=JobState.FINAL_REVIEW;job.blocked_reason='资产来源缺失';save_job(job)
    with pytest.raises(ValueError,match='corresponding candidate'):
        record_expression_review(job.id,'final',request_for(root))


def test_new_pptx_cannot_reuse_old_render(candidate):
    job,root=candidate
    old=read_json(root/'reviews/final-expression-packet.json')
    prs=Presentation(root/'build/deck.pptx');prs.slides[0].shapes[0].left=Inches(7);prs.save(root/'build/deck.pptx')
    with pytest.raises(ValueError,match='stale'):
        record_expression_review(job.id,'final',request_for(root))
    prepare_expression_review(job.id,'final')
    new=read_json(root/'reviews/final-expression-packet.json')
    assert new['slides'][0]['preview'] != old['slides'][0]['preview']
    assert new['slides'][0]['objects'][0]['bounds_pt']['left'] == 504


def test_native_render_failure_is_actionable_cli_error(monkeypatch):
    from aippt_studio.renderer import RenderError
    def fail(*args):raise RenderError('Native preview could not be rendered; inspect the Office error')
    monkeypatch.setattr('aippt_studio.expression_review.prepare_expression_review',fail)
    result=CliRunner().invoke(app,['expression-review','fixture','--gate','final','--json'])
    assert result.exit_code==2 and 'Native preview could not be rendered' in result.output


def _corner_number(root, text, *, slide_index=0, x=.3, y=.3):
    path=root/'build/deck.pptx';prs=Presentation(path)
    shape=prs.slides[slide_index].shapes.add_textbox(Inches(x),Inches(y),Inches(.6),Inches(.45))
    shape.text=text;shape.text_frame.paragraphs[0].runs[0].font.size=Pt(16)
    prs.save(path)
    return str(shape.shape_id)


@pytest.mark.parametrize('number',['01','02','03','04'])
def test_corner_ordinal_requires_context_without_automatic_bad_design_verdict(candidate,number):
    job,root=candidate;object_id=_corner_number(root,number)
    prepare_expression_review(job.id,'final')
    packet=read_json(root/'reviews/final-expression-packet.json')
    signal=next(s for s in packet['slides'][0]['signals'] if s['code']=='DECORATIVE_ORDINAL')
    assert signal['object_id']==object_id and signal['excerpt']==number
    assert signal['kind']=='context_required' and signal['location']=='top_left'
    assert signal['bounds_pt']['left']==pytest.approx(21.6)
    assert '本项目完成三地访谈' in signal['nearby_text'][0]['text']
    assert packet['status']=='awaiting_assistant_review'
    with pytest.raises(ValueError,match='Address each expression signal'):
        record_expression_review(job.id,'final',request_for(root))


@pytest.mark.parametrize('meaning',[
    '这是章节导航01，与目录及后续现场采样页逐一对应，服务于当前汇报结构。',
    '这是实验步骤01的独立标识，右侧相邻标题说明现场采样，后续步骤沿同一流程展开。',
])
def test_genuine_chapter_or_step_number_can_be_legitimate(candidate,meaning):
    job,root=candidate;_corner_number(root,'01')
    prepare_expression_review(job.id,'final')
    packet=read_json(root/'reviews/final-expression-packet.json')
    signal=packet['slides'][0]['signals'][0];request=request_for(root)
    request['slides'][0]['signal_dispositions']=[{'signal_id':signal['signal_id'],'decision':'legitimate','reason':meaning}]
    assert record_expression_review(job.id,'final',request)['status']=='ready'
    assert require_expression_review(root,'final')['slides'][0]['signal_dispositions'][0]['reason']==meaning


def test_native_page_number_keeps_its_semantic_hint_and_can_be_legitimate(candidate):
    job,root=candidate;path=root/'build/deck.pptx';prs=Presentation(path)
    shape=prs.slides[1].shapes.add_textbox(Inches(12),Inches(6.9),Inches(.5),Inches(.4))
    field=OxmlElement('a:fld');field.set('id','{16E8C133-3AAF-4E51-B0A4-DAD82535818D}');field.set('type','slidenum')
    text=OxmlElement('a:t');text.text='2';field.append(text);shape.text_frame.paragraphs[0]._p.append(field)
    prs.save(path);prepare_expression_review(job.id,'final')
    packet=read_json(root/'reviews/final-expression-packet.json')
    signal=next(s for s in packet['slides'][1]['signals'] if s['code']=='DECORATIVE_ORDINAL')
    assert signal['semantic_hints']==['ooxml_page_number','matches_deck_page_number']
    request=request_for(root)
    request['slides'][1]['signal_dispositions']=[{'signal_id':signal['signal_id'],'decision':'legitimate','reason':'原生页码字段标识当前第二页，供观众按页定位材料。'}]
    assert record_expression_review(job.id,'final',request)['status']=='ready'


def test_real_numbers_with_text_and_centered_metric_are_not_a_numeric_ban(candidate):
    job,root=candidate;path=root/'build/deck.pptx';prs=Presentation(path);slide=prs.slides[0]
    for i,text in enumerate(['步骤01：现场采样','第2页','第1章','2026','99%','24份原始记录']):
        slide.shapes.add_textbox(Inches(.3),Inches(.1+i*.15),Inches(2),Inches(.2)).text=text
    slide.shapes.add_textbox(Inches(6),Inches(3),Inches(1),Inches(1)).text='24'
    prs.save(path);prepare_expression_review(job.id,'final')
    packet=read_json(root/'reviews/final-expression-packet.json')
    assert all(s['code']!='DECORATIVE_ORDINAL' for s in packet['slides'][0]['signals'])
    assert 'master-only' in packet['ordinal_signal_coverage']


def test_unexplained_decoration_is_revised_in_the_candidate_not_just_marked_ready(candidate):
    job,root=candidate;object_id=_corner_number(root,'04')
    prepare_expression_review(job.id,'final');packet=read_json(root/'reviews/final-expression-packet.json')
    signal=packet['slides'][0]['signals'][0];request=request_for(root)
    request['slides'][0]['signal_dispositions']=[{'signal_id':signal['signal_id'],'decision':'revise','reason':'此04不对应页码、章节或步骤，只是角落装饰，当前表达不需要。'}]
    with pytest.raises(ValueError,match='unresolved copy'):
        record_expression_review(job.id,'final',request)
    request['slides'][0]['decision']='revise'
    assert record_expression_review(job.id,'final',request)['status']=='needs_revision'
    path=root/'build/deck.pptx';prs=Presentation(path)
    shape=next(s for s in prs.slides[0].shapes if str(s.shape_id)==object_id)
    shape.element.getparent().remove(shape.element);prs.save(path)
    prepare_expression_review(job.id,'final')
    revised=read_json(root/'reviews/final-expression-packet.json')
    assert not revised['slides'][0]['signals']
    assert record_expression_review(job.id,'final',request_for(root))['status']=='ready'


def test_group_transform_locates_corner_numeral_but_hidden_objects_are_ignored(candidate):
    job,root=candidate;path=root/'build/deck.pptx';prs=Presentation(path);slide=prs.slides[0]
    group=slide.shapes.add_group_shape()
    number=group.shapes.add_textbox(0,0,Inches(.6),Inches(.4));number.text='03'
    group.left,group.top=Inches(11.8),Inches(.3)
    hidden=slide.shapes.add_textbox(Inches(.2),Inches(.2),Inches(.6),Inches(.4));hidden.text='02'
    hidden.element.nvSpPr.cNvPr.set('hidden','1')
    prs.save(path);prepare_expression_review(job.id,'final')
    packet=read_json(root/'reviews/final-expression-packet.json')
    signals=[s for s in packet['slides'][0]['signals'] if s['code']=='DECORATIVE_ORDINAL']
    assert len(signals)==1 and signals[0]['object_id']==str(number.shape_id)
    assert signals[0]['location']=='top_right' and signals[0]['bounds_pt']['left']==pytest.approx(849.6)
    assert all(o['object_id']!=str(hidden.shape_id) for o in packet['slides'][0]['objects'])
