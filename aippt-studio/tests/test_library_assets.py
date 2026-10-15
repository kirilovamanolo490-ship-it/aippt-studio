from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from PIL import Image
from pptx import Presentation
import pytest

from aippt_studio.html_deck import _asset_info, write_deck_html
from aippt_studio.library_assets import LibraryAssetError, bind_library_media
from aippt_studio.models import DeckPlan, JobState, SlideSpec
from aippt_studio.storage import job_dir, load_job, read_json, read_model, save_job, write_json
from aippt_studio.workflow import create_job


def _fixture_library(tmp_path: Path, *, indexed_sha: str | None = None) -> tuple[Path, int]:
    library = tmp_path / "模板库"
    library.mkdir(parents=True)
    picture = tmp_path / "chart.png"
    Image.new("RGBA", (32, 24), (20, 110, 70, 255)).save(picture)
    source = library / "公益/diagram_asset.pptx"
    source.parent.mkdir()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_picture(str(picture), 100000, 200000, width=1000000, height=750000)
    presentation.save(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    (library / "CATALOG.jsonl").write_text(json.dumps({
        "path": "公益/diagram_asset.pptx", "asset_type": "diagram_asset",
        "media_rights": "unverified",
        # Some diagram_asset rows do not carry source_sha256; the structural
        # index remains the hash source for this path.
    }, ensure_ascii=False) + "\n", encoding="utf-8")
    (library / "DESIGN-STRUCTURE.jsonl").write_text(json.dumps({
        "schema": "aippt.template-design-index.v1", "source_path": "公益/diagram_asset.pptx",
        "source_sha256": indexed_sha or source_hash, "status": "parsed",
        "coverage": {"full_slide_coverage": True, "slides_in_deck": 1, "slides_parsed": 1},
    }, ensure_ascii=False) + "\n", encoding="utf-8")
    return source, shape.shape_id


def _job_with_plan(job_id: str = "library-media") -> str:
    job = create_job("库内素材", job_id=job_id)
    root = job_dir(job.id)
    write_json(root / "planning/deck-plan.json", DeckPlan(
        title="测试素材", slides=[SlideSpec(id="slide-01", title="项目概览", purpose="说明核心项目", layout="content")],
    ))
    return job.id


def test_library_picture_binds_to_slide_and_is_resolved_for_rendering(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, shape_id = _fixture_library(tmp_path)
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan()
    root = job_dir(job_id)
    job = load_job(job_id)
    job.state = JobState.PROTOTYPE_REVIEW
    job.approvals["prototype"] = True
    job.approval_hashes["prototype"] = {"deck-plan.json": "old"}
    save_job(job)
    preview = root / "build/prototype.html"
    preview.parent.mkdir(parents=True, exist_ok=True)
    preview.write_text("stale render", encoding="utf-8")

    result = bind_library_media(
        job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
        slide_id="slide-01", purpose="解释服务覆盖路径", selection_note="用户选中的项目路径图",
        task_reuse_authorized=True, rights_handling_note="本任务已单独确认当前稿件使用权限；目录权利状态仍按未验证保留。",
    )

    assert result["status"] == "bound"
    assert result["asset_ref"].startswith("library:lib-")
    item = result["asset"]
    copied = root / item["stored_path"]
    assert copied.is_file()
    assert hashlib.sha256(copied.read_bytes()).hexdigest() == item["media_sha256"]
    plan = read_model(root / "planning/deck-plan.json", DeckPlan)
    assert result["asset_ref"] in plan.slides[0].asset_refs
    resolved = _asset_info(root, result["asset_ref"])
    assert resolved is not None and resolved[1] == "image"
    assert resolved[0].startswith("file:")
    rendered_html = write_deck_html(job_id).read_text(encoding="utf-8")
    assert copied.as_uri() in rendered_html
    ledger = read_json(root / "assets/library/ledger.json")
    saved = ledger["assets"][0]
    assert saved["source_sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert saved["source_slide"] == 1 and saved["shape_id"] == shape_id
    assert saved["media_part"].startswith("ppt/media/")
    assert saved["rights"] == "unverified"
    assert saved["task_reuse_authorized"] is True
    assert saved["selection_note"] == "用户选中的项目路径图"
    assert "目录权利状态仍按未验证保留" in saved["rights_handling_note"]
    refreshed = load_job(job_id)
    assert refreshed.approvals["prototype"] is False
    assert "prototype" not in refreshed.approval_hashes
    assert not preview.exists()


@pytest.mark.parametrize("bad_shape", [True, 99999])
def test_library_picture_rejects_a_non_picture_shape_on_the_requested_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, bad_shape: int | bool,
) -> None:
    _source, shape_id = _fixture_library(tmp_path)
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-bad-shape")
    requested_id = bad_shape if bad_shape is True else int(bad_shape)
    with pytest.raises(LibraryAssetError, match="positive integers|picture"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=requested_id,
            slide_id="slide-01", purpose="项目说明", selection_note="目标页选择此图示",
            task_reuse_authorized=True, rights_handling_note="本任务已单独确认当前稿件使用权限。",
        )
    assert not (job_dir(job_id) / "assets/library/ledger.json").exists()
    assert not (job_dir(job_id) / "assets/library").exists()
    assert shape_id > 0


def test_library_picture_binding_requires_explicit_task_rights_disposition(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _source, shape_id = _fixture_library(tmp_path)
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-rights-not-authorized")

    with pytest.raises(LibraryAssetError, match="task_reuse_authorized=True"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", selection_note="选择这张服务图",
            rights_handling_note="仅作参考，不嵌入当前稿件。",
        )
    assert not (job_dir(job_id) / "assets/library/ledger.json").exists()


def test_library_picture_binding_requires_selection_and_rights_notes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _source, shape_id = _fixture_library(tmp_path)
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-rights-note-required")

    with pytest.raises(LibraryAssetError, match="selection_note is required"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", task_reuse_authorized=True,
            rights_handling_note="已取得本任务许可。",
        )

    with pytest.raises(LibraryAssetError, match="rights_handling_note is required"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", selection_note="选择服务图",
            task_reuse_authorized=True,
        )


def test_library_picture_rejects_a_stale_structural_index_hash(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _source, shape_id = _fixture_library(tmp_path, indexed_sha="0" * 64)
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-bad-hash")

    with pytest.raises(LibraryAssetError, match="no longer match"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", selection_note="目标页选择此图示",
            task_reuse_authorized=True, rights_handling_note="本任务已单独确认当前稿件使用权限。",
        )
    assert not (job_dir(job_id) / "assets/library/ledger.json").exists()


def test_library_picture_rejects_a_package_external_relationship(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, shape_id = _fixture_library(tmp_path)
    with zipfile.ZipFile(source) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    rels_path = "ppt/slides/_rels/slide1.xml.rels"
    rels = parts[rels_path].decode("utf-8")
    rels = rels.replace('Target="../media/image1.png"', 'Target="https://example.test/image.png" TargetMode="External"')
    parts[rels_path] = rels.encode("utf-8")
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for name, payload in parts.items():
            output.writestr(name, payload)
    # Index now claims the mutated file, ensuring the external relationship is
    # the failure reason rather than the expected-byte mismatch.
    structure_path = tmp_path / "模板库/DESIGN-STRUCTURE.jsonl"
    structure = json.loads(structure_path.read_text(encoding="utf-8"))
    structure["source_sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    structure_path.write_text(json.dumps(structure, ensure_ascii=False) + "\n", encoding="utf-8")
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-external-media")

    with pytest.raises(LibraryAssetError, match="external"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", selection_note="目标页选择此图示",
            task_reuse_authorized=True, rights_handling_note="本任务已单独确认当前稿件使用权限。",
        )


def test_library_picture_rejects_active_svg_media(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, shape_id = _fixture_library(tmp_path)
    with zipfile.ZipFile(source) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    rels_path = "ppt/slides/_rels/slide1.xml.rels"
    rels = parts[rels_path].decode("utf-8").replace("../media/image1.png", "../media/image1.svg")
    parts[rels_path] = rels.encode("utf-8")
    parts.pop("ppt/media/image1.png")
    parts["ppt/media/image1.svg"] = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for name, payload in parts.items():
            output.writestr(name, payload)
    structure_path = tmp_path / "模板库/DESIGN-STRUCTURE.jsonl"
    structure = json.loads(structure_path.read_text(encoding="utf-8"))
    structure["source_sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    structure_path.write_text(json.dumps(structure, ensure_ascii=False) + "\n", encoding="utf-8")
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-active-svg")

    with pytest.raises(LibraryAssetError, match="active SVG"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", selection_note="目标页选择此图示",
            task_reuse_authorized=True, rights_handling_note="本任务已单独确认当前稿件使用权限。",
        )


def test_library_picture_rejects_oversized_media(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, shape_id = _fixture_library(tmp_path)
    with zipfile.ZipFile(source) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    parts["ppt/media/image1.png"] = b"x" * (20 * 1024 * 1024 + 1)
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for name, payload in parts.items():
            output.writestr(name, payload)
    structure_path = tmp_path / "模板库/DESIGN-STRUCTURE.jsonl"
    structure = json.loads(structure_path.read_text(encoding="utf-8"))
    structure["source_sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    structure_path.write_text(json.dumps(structure, ensure_ascii=False) + "\n", encoding="utf-8")
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-large-media")

    with pytest.raises(LibraryAssetError, match="20 MiB"):
        bind_library_media(
            job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", selection_note="目标页选择此图示",
            task_reuse_authorized=True, rights_handling_note="本任务已单独确认当前稿件使用权限。",
        )


def test_library_picture_prefers_office_embedded_svg_over_png_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, shape_id = _fixture_library(tmp_path)
    with zipfile.ZipFile(source) as original:
        parts = {name: original.read(name) for name in original.namelist()}

    ns_p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns_svg = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
    slide_path = "ppt/slides/slide1.xml"
    slide = ElementTree.fromstring(parts[slide_path])
    picture = next(node for node in slide.iter(f"{{{ns_p}}}pic")
                   if node.find(f"{{{ns_p}}}nvPicPr/{{{ns_p}}}cNvPr").get("id") == str(shape_id))
    blip = picture.find(f"{{{ns_p}}}blipFill/{{{ns_a}}}blip")
    ext_list = ElementTree.SubElement(blip, f"{{{ns_a}}}extLst")
    extension = ElementTree.SubElement(
        ext_list, f"{{{ns_a}}}ext", {"uri": "{96AEC541-7D9C-4A6A-85BD-1B6DE7D36B8A}"},
    )
    ElementTree.SubElement(extension, f"{{{ns_svg}}}svgBlip", {f"{{{ns_r}}}embed": "rIdSVG"})
    parts[slide_path] = ElementTree.tostring(slide, encoding="utf-8", xml_declaration=True)

    rels_path = "ppt/slides/_rels/slide1.xml.rels"
    rels = ElementTree.fromstring(parts[rels_path])
    ElementTree.SubElement(rels, f"{{{ns_rel}}}Relationship", {
        "Id": "rIdSVG", "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
        "Target": "../media/image-vector.svg",
    })
    parts[rels_path] = ElementTree.tostring(rels, encoding="utf-8", xml_declaration=True)
    parts["ppt/media/image-vector.svg"] = (
        b'<svg xmlns="http://www.w3.org/2000/svg" width="32" height="24">'
        b'<path d="M1 1h30v22H1z" fill="#146e46"/></svg>'
    )
    content_types = ElementTree.fromstring(parts["[Content_Types].xml"])
    if not any(node.get("Extension") == "svg" for node in content_types):
        ElementTree.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default", {
            "Extension": "svg", "ContentType": "image/svg+xml",
        })
    parts["[Content_Types].xml"] = ElementTree.tostring(content_types, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for name, payload in parts.items():
            output.writestr(name, payload)

    structure_path = tmp_path / "模板库/DESIGN-STRUCTURE.jsonl"
    structure = json.loads(structure_path.read_text(encoding="utf-8"))
    structure["source_sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    structure_path.write_text(json.dumps(structure, ensure_ascii=False) + "\n", encoding="utf-8")
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-office-svg")

    result = bind_library_media(
        job_id, "公益/diagram_asset.pptx", source_slide=1, shape_id=shape_id,
        slide_id="slide-01", purpose="复用可编辑图示的矢量源素材",
        selection_note="选中该图示用于解释服务流程", task_reuse_authorized=True,
        rights_handling_note="本任务已单独确认当前稿件使用权限。",
    )

    assert result["asset_ref"].startswith("library:lib-")
    assert result["asset"]["media_type"] == "image/svg+xml"
    assert result["asset"]["media_part"] == "ppt/media/image-vector.svg"
    assert (job_dir(job_id) / result["asset"]["stored_path"]).suffix == ".svg"
    resolved = _asset_info(job_dir(job_id), result["asset_ref"])
    assert resolved is not None and resolved[1] == "image"
    assert "PHN2Zy" in resolved[0]


def test_library_picture_resolves_old_path_alias_with_exact_source_hash(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, shape_id = _fixture_library(tmp_path)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    (tmp_path / "模板库/RENAME-ALIASES.jsonl").write_text(json.dumps({
        "schema": "aippt.library-path-alias.v1", "old_path": "公益/01.pptx",
        "current_path": "公益/diagram_asset.pptx", "source_sha256": digest,
    }, ensure_ascii=False) + "\n", encoding="utf-8")
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-legacy-path")

    result = bind_library_media(
        job_id, "公益/01.pptx", source_slide=1, shape_id=shape_id,
        slide_id="slide-01", purpose="项目说明", selection_note="选中该服务图",
        task_reuse_authorized=True, rights_handling_note="已确认本任务使用许可。",
    )
    assert result["asset"]["source_path"] == "公益/diagram_asset.pptx"
    assert result["asset"]["requested_source_path"] == "公益/01.pptx"
    assert result["asset"]["source_sha256"] == digest


def test_library_picture_rejects_stale_source_hash_in_old_path_alias(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _source, shape_id = _fixture_library(tmp_path)
    (tmp_path / "模板库/RENAME-ALIASES.jsonl").write_text(json.dumps({
        "schema": "aippt.library-path-alias.v1", "old_path": "公益/01.pptx",
        "current_path": "公益/diagram_asset.pptx", "source_sha256": "0" * 64,
    }, ensure_ascii=False) + "\n", encoding="utf-8")
    monkeypatch.setattr("aippt_studio.library_assets.WORKSPACE_ROOT", tmp_path)
    job_id = _job_with_plan("library-legacy-path-stale")

    with pytest.raises(LibraryAssetError, match="alias hash"):
        bind_library_media(
            job_id, "公益/01.pptx", source_slide=1, shape_id=shape_id,
            slide_id="slide-01", purpose="项目说明", selection_note="选中该服务图",
            task_reuse_authorized=True, rights_handling_note="已确认本任务使用许可。",
        )
