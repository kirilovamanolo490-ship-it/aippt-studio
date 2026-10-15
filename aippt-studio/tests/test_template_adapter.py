from __future__ import annotations

import zipfile
import hashlib
import json
import importlib
from pathlib import Path

from pptx import Presentation
import pytest

from adapter_test_helpers import run_adapter, write_contract_job
from aippt_studio.mastering import audit_master_layout_ids
from aippt_studio.notes import audit_pptx_notes


def test_template_adapter_clones_template_and_replaces_tags(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "{{title}}"
    body = slide.shapes.add_textbox(1000000, 2000000, 6000000, 2000000)
    body.text = "{{body}}"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "{{title}}"
    second_body = second.shapes.add_textbox(1000000, 2000000, 6000000, 2000000)
    second_body.text = "{{body}}"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=2)
    strict_request = json.loads(request.read_text(encoding="utf-8"))
    strict_request["template_application_required"] = True
    request.write_text(json.dumps(strict_request), encoding="utf-8")

    payload = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py",
        request,
        result,
    )

    assert payload["status"] == "succeeded", payload
    assert payload["receipt"]["engine"] == "pptx-automizer"
    assert payload["receipt"]["version"] == "0.8.2"
    provenance = payload["receipt"]["provenance"]
    assert provenance["execution_kind"] == "managed_package_runtime"
    assert provenance["runtime"]["package"] == "pptx-automizer"
    assert provenance["runtime"]["version"] == "0.8.2"
    assert len(provenance["runtime"]["lock_sha256"]) == 64
    assert provenance["bridge"]["version"] == "0.9.1"
    output = root / "build" / "adapter-staging" / "deck.pptx"
    assert output.is_file()
    with zipfile.ZipFile(output) as archive:
        text = "".join(archive.read(name).decode("utf-8", errors="ignore") for name in archive.namelist() if name.startswith("ppt/slides/slide"))
        output_themes = {name: archive.read(name) for name in archive.namelist() if name.startswith("ppt/theme/theme")}
    with zipfile.ZipFile(template) as archive:
        source_themes = {name: archive.read(name) for name in archive.namelist() if name.startswith("ppt/theme/theme")}
    assert "Fixture title 1" in text
    assert "Fixture title 2" in text
    assert all(output_themes.get(name) == data for name, data in source_themes.items())
    master_audit = audit_master_layout_ids(output)
    assert master_audit["master_count"] >= 1
    assert master_audit["layout_count"] >= 1
    assert audit_pptx_notes(output) == []
    assert payload["metadata"]["template_application"]["verified_slide_count"] == 2


def test_required_template_application_rejects_a_noop_source(tmp_path: Path) -> None:
    template = tmp_path / "unprepared-service-proposal.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "旧项目服务方案"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    strict_request = json.loads(request.read_text(encoding="utf-8"))
    strict_request["template_application_required"] = True
    request.write_text(json.dumps(strict_request), encoding="utf-8")
    payload = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py",
        request,
        result,
    )
    assert payload["status"] == "blocked"
    assert payload["failure"]["category"] == "template_application_noop"


def test_required_template_application_edits_source_bound_text_shapes(tmp_path: Path) -> None:
    template = tmp_path / "contest-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "原竞赛项目标题"
    body = slide.shapes.add_textbox(1000000, 2000000, 6000000, 2000000)
    body.text = "原项目做法与成效"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    strict_request = json.loads(request.read_text(encoding="utf-8"))
    strict_request["template_application_required"] = True
    request.write_text(json.dumps(strict_request), encoding="utf-8")
    plan_path = root / "planning/deck-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    mapping = {
        "schema_version": 2,
        "template_source_id": "source-1",
        "template_source_sha256": hashlib.sha256((root / "sources" / template.name).read_bytes()).hexdigest(),
        "template_source_slide_count": 1,
        "deck_plan_sha256": hashlib.sha256(plan_path.read_bytes()).hexdigest(),
        "allow_animations": False,
        "slides": [{
            "deck_slide_id": plan["slides"][0]["id"], "source_slide": 1,
            "replacements": {"title": plan["slides"][0]["title"], "body": "\n".join(plan["slides"][0]["points"])},
            "shape_targets": {
                "title": {"shape_name": slide.shapes.title.name, "source_text": "原竞赛项目标题"},
                "body": {"shape_name": body.name, "source_text": "原项目做法与成效"},
            },
        }],
    }
    mapping_path = root / "planning/template-map.json"
    mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    payload = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py",
        request, result,
    )
    assert payload["status"] == "succeeded", payload
    proof = payload["metadata"]["template_application"]
    assert proof["status"] == "verified"
    assert proof["pages"][0]["replacement_methods"] == {"body": "source_shape", "title": "source_shape"}
    built = Presentation(root / "build/adapter-staging/deck.pptx")
    assert len(built.slides) == 1
    visible = "\n".join(shape.text for shape in built.slides[0].shapes if shape.has_text_frame)
    assert plan["slides"][0]["title"] in visible
    assert plan["slides"][0]["points"][0] in visible
    assert "原竞赛项目标题" not in visible


def test_required_template_application_rejects_unreviewed_old_project_copy(tmp_path: Path) -> None:
    template = tmp_path / "contest-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "{{title}}"
    slide.shapes.add_textbox(1000000, 2000000, 6000000, 2000000).text = "{{body}}"
    slide.shapes.add_textbox(1000000, 5000000, 6000000, 500000).text = "原项目已服务200户"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    strict_request = json.loads(request.read_text(encoding="utf-8"))
    strict_request["template_application_required"] = True
    request.write_text(json.dumps(strict_request), encoding="utf-8")
    payload = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py",
        request, result,
    )
    assert payload["status"] == "blocked"
    assert payload["failure"]["category"] == "template_application_noop"
    assert "unreviewed source text" in payload["failure"]["message"]


def test_application_verifier_checks_mapped_shape_not_duplicate_page_text(tmp_path: Path, monkeypatch) -> None:
    script = Path(__file__).resolve().parents[1] / "scripts/adapters"
    monkeypatch.syspath_prepend(str(script))
    adapter = importlib.import_module("template_adapter")
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "旧标题"
    slide.shapes.add_textbox(1000000, 2000000, 6000000, 600000).text = "新标题"
    presentation.save(template)
    output = tmp_path / "wrong-output.pptx"
    output.write_bytes(template.read_bytes())
    mapping = {"slides": [{"replacements": {"title": "新标题"}}]}
    proof = [{"source_slide": 1, "replacement_tags": ["title"],
              "field_shapes": {"title": [slide.shapes.title.name]}, "preserved_text_shapes": []}]
    error = adapter._verify_template_application(template, output, mapping, proof)
    assert error is not None and "mapped title replacement" in error


def test_application_verifier_requires_active_theme_chain(tmp_path: Path, monkeypatch) -> None:
    script = Path(__file__).resolve().parents[1] / "scripts/adapters"
    monkeypatch.syspath_prepend(str(script))
    adapter = importlib.import_module("template_adapter")
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "{{title}}"
    presentation.save(template)
    with zipfile.ZipFile(template) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    original_theme = entries["ppt/theme/theme1.xml"]
    entries["ppt/theme/theme1.xml"] = original_theme.replace(b"</a:theme>", b"<!--different active theme--></a:theme>")
    assert entries["ppt/theme/theme1.xml"] != original_theme
    entries["ppt/theme/theme2.xml"] = original_theme  # present, but unreferenced
    output = tmp_path / "wrong-theme.pptx"
    with zipfile.ZipFile(output, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    mapping = {"slides": [{"replacements": {"title": "新标题"}}]}
    proof = [{"source_slide": 1, "replacement_tags": ["title"],
              "field_shapes": {"title": [slide.shapes.title.name]}, "preserved_text_shapes": []}]
    error = adapter._verify_template_application(template, output, mapping, proof)
    assert error is not None and "layout/master/theme" in error


def test_template_adapter_blocks_unreviewed_animation(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[5])
    presentation.save(template)
    with zipfile.ZipFile(template) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["ppt/slides/slide1.xml"] = entries["ppt/slides/slide1.xml"].replace(b"</p:sld>", b"<p:timing/></p:sld>")
    rewritten = template.with_name("template-with-timing.pptx")
    with zipfile.ZipFile(rewritten, "w") as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    rewritten.replace(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    payload = run_adapter(Path(__import__("sys").executable), Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py", request, result)
    assert payload["status"] == "blocked"
    assert payload["failure"]["category"] == "unsupported_animation"


def test_template_adapter_rejects_zip_traversal_entries(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[5])
    presentation.save(template)
    with zipfile.ZipFile(template, "a") as archive:
        archive.writestr("../escape.txt", "blocked")
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template)
    payload = run_adapter(Path(__import__("sys").executable), Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py", request, result)
    assert payload["status"] == "blocked"
    assert payload["failure"]["category"] == "unsafe_template_archive"


@pytest.mark.parametrize("subtitle", [None, "三地访谈，保留24份原始记录"])
def test_template_default_does_not_publish_internal_purpose_or_remove_source_numbering(tmp_path: Path, subtitle):
    template = tmp_path / "source-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, text in enumerate(("{{title}}", "{{subtitle}}", "{{body}}", "01", "客户原有章节导航")):
        box = slide.shapes.add_textbox(500000, 500000 + index*800000, 6500000, 600000)
        box.text = text
    presentation.save(template)
    original_hash = hashlib.sha256(template.read_bytes()).hexdigest()
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    plan_path = root / "planning/deck-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    plan["slides"][0]["purpose"] = "本页需要补充与前页的关联，并调整字号和配图"
    if subtitle is not None:
        plan["slides"][0]["subtitle"] = subtitle
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    payload = run_adapter(Path(__import__("sys").executable), Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py", request, result)
    assert payload["status"] == "succeeded", payload
    actual = Presentation(root / "build/adapter-staging/deck.pptx")
    texts = [shape.text for shape in actual.slides[0].shapes if shape.has_text_frame]
    assert "01" in texts and "客户原有章节导航" in texts
    assert not any("本页需要" in text or "{{subtitle}}" in text for text in texts)
    if subtitle is not None:
        assert subtitle in texts
    mapping = json.loads((root / "planning/template-map.json").read_text(encoding="utf-8"))
    assert mapping["slides"][0]["replacements"]["subtitle"] == (subtitle or "")
    assert hashlib.sha256(template.read_bytes()).hexdigest() == original_hash


def test_generated_template_mapping_does_not_repeat_an_exact_title_in_body(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(500000, 500000, 6500000, 600000).text = "{{title}}"
    slide.shapes.add_textbox(500000, 1500000, 6500000, 1800000).text = "{{body}}"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    plan_path = root / "planning/deck-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    repeated = "现场验收已完成，关键记录可以复核。"
    support = "后续安排已明确到责任单位和时间节点。"
    plan["slides"][0]["title"] = repeated
    plan["slides"][0]["points"] = [repeated, support]
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")

    payload = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py",
        request,
        result,
    )

    assert payload["status"] == "succeeded", payload
    mapping = json.loads((root / "planning/template-map.json").read_text(encoding="utf-8"))
    assert mapping["generated_content_policy"] == "title_body_deduplicated_v1"
    assert mapping["slides"][0]["replacements"]["body"] == support
    actual = Presentation(root / "build/adapter-staging/deck.pptx")
    visible_text = [shape.text for shape in actual.slides[0].shapes if shape.has_text_frame]
    assert visible_text.count(repeated) == 1
    assert support in visible_text


def test_valid_legacy_generated_mapping_is_refreshed_but_stale_one_remains_blocked(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(500000, 500000, 6500000, 600000).text = "{{title}}"
    slide.shapes.add_textbox(500000, 1500000, 6500000, 1800000).text = "{{body}}"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    plan_path = root / "planning/deck-plan.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    repeated = "本次调研形成可复核的结论。"
    plan["slides"][0]["title"] = repeated
    plan["slides"][0]["points"] = [repeated]
    plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")

    first = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py",
        request,
        result,
    )
    assert first["status"] == "succeeded", first
    mapping_path = root / "planning/template-map.json"
    legacy_mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    legacy_mapping.pop("generated_content_policy")
    legacy_mapping["slides"][0]["replacements"]["body"] = repeated
    mapping_path.write_text(json.dumps(legacy_mapping, ensure_ascii=False), encoding="utf-8")

    refreshed = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py",
        request,
        result,
    )
    assert refreshed["status"] == "succeeded", refreshed
    refreshed_mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    assert refreshed_mapping["generated_content_policy"] == "title_body_deduplicated_v1"
    assert refreshed_mapping["slides"][0]["replacements"]["body"] == ""

    changed_plan = json.loads(plan_path.read_text(encoding="utf-8"))
    changed_plan["slides"][0]["title"] = "修改后的标题不能自动重绑旧模板映射"
    plan_path.write_text(json.dumps(changed_plan, ensure_ascii=False), encoding="utf-8")
    stale_mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
    stale_mapping.pop("generated_content_policy")
    mapping_path.write_text(json.dumps(stale_mapping, ensure_ascii=False), encoding="utf-8")

    stale = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py",
        request,
        result,
    )
    assert stale["status"] == "blocked"
    assert stale["failure"]["category"] == "template_map_binding"


def test_explicit_template_mapping_remains_source_authoritative(tmp_path: Path):
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(500000, 500000, 6500000, 600000).text = "{{subtitle}}"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    stored = root / "sources" / template.name
    mapping = {
        "schema_version": 2,
        "template_source_id": "source-1",
        "template_source_sha256": hashlib.sha256(stored.read_bytes()).hexdigest(),
        "template_source_slide_count": 1,
        "deck_plan_sha256": hashlib.sha256((root / "planning" / "deck-plan.json").read_bytes()).hexdigest(),
        "slides": [
            {
                "deck_slide_id": "slide-01",
                "source_slide": 1,
                "replacements": {"subtitle": "01 现场采样"},
            }
        ],
    }
    target = root / "planning/template-map.json"
    target.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
    original = target.read_bytes()
    payload = run_adapter(Path(__import__("sys").executable), Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py", request, result)
    assert payload["status"] == "succeeded", payload
    actual = Presentation(root / "build/adapter-staging/deck.pptx")
    assert any(shape.has_text_frame and shape.text == "01 现场采样" for shape in actual.slides[0].shapes)
    assert target.read_bytes() == original


def test_template_adapter_requires_a_bound_page_map_before_reusing_a_layout(tmp_path: Path) -> None:
    template = tmp_path / "single-layout-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(500000, 500000, 6500000, 600000).text = "{{title}}"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=2)

    payload = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py",
        request,
        result,
    )

    assert payload["status"] == "blocked"
    assert payload["failure"]["category"] == "template_map_required"
    assert not (root / "planning" / "template-map.json").exists()


def test_template_adapter_rejects_stale_or_ambiguous_template_binding(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    stored = root / "sources" / template.name
    mapping_path = root / "planning" / "template-map.json"
    mapping_path.write_text(
        json.dumps(
            {
                "schema_version": 2,
                "template_source_id": "source-1",
                "template_source_sha256": "0" * 64,
                "template_source_slide_count": 1,
                "deck_plan_sha256": hashlib.sha256((root / "planning" / "deck-plan.json").read_bytes()).hexdigest(),
                "slides": [{"deck_slide_id": "slide-01", "source_slide": 1, "replacements": {}}],
            }
        ),
        encoding="utf-8",
    )
    stale = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py",
        request,
        result,
    )
    assert stale["status"] == "blocked"
    assert stale["failure"]["category"] == "template_map_binding"

    second = root / "sources" / "other-template.pptx"
    second.write_bytes(stored.read_bytes())
    (root / "normalized" / "sources" / "source-2.json").write_text(
        json.dumps(
            {
                "id": "source-2",
                "original": "other-template.pptx",
                "stored_path": "sources/other-template.pptx",
                "kind": "powerpoint",
                "sha256": hashlib.sha256(second.read_bytes()).hexdigest(),
                "size_bytes": second.stat().st_size,
                "metadata": {},
            }
        ),
        encoding="utf-8",
    )
    request_payload = json.loads(request.read_text(encoding="utf-8"))
    request_payload["input_roles"]["source-2"] = "edit_target"
    request.write_text(json.dumps(request_payload), encoding="utf-8")
    ambiguous = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py",
        request,
        result,
    )
    assert ambiguous["status"] == "blocked"
    assert ambiguous["failure"]["category"] == "template_source_ambiguous"


def test_template_adapter_rejects_map_bound_to_an_old_plan_even_when_slide_ids_match(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(500000, 500000, 6500000, 600000).text = "{{title}}"
    presentation.save(template)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="template", role="template", source=template, slides=1)
    first = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py",
        request,
        result,
    )
    assert first["status"] == "succeeded", first
    plan_path = root / "planning" / "deck-plan.json"
    changed = json.loads(plan_path.read_text(encoding="utf-8"))
    changed["slides"][0]["title"] = "新计划的正式标题"
    plan_path.write_text(json.dumps(changed, ensure_ascii=False), encoding="utf-8")

    second = run_adapter(
        Path(__import__("sys").executable),
        Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "template_adapter.py",
        request,
        result,
    )
    assert second["status"] == "blocked"
    assert second["failure"]["category"] == "template_map_binding"
