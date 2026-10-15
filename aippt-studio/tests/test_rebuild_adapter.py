from __future__ import annotations

from pathlib import Path
import json

import pytest

from PIL import Image
from pptx import Presentation

from adapter_test_helpers import run_adapter, write_contract_job


def test_rebuild_adapter_promotes_finalized_editppt_run(tmp_path: Path) -> None:
    image = tmp_path / "scan.png"
    Image.new("RGB", (320, 180), "white").save(image)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="rebuild", role="flattened_input", source=image)
    run_dir = root / "workspace" / "editppt"
    run_dir.mkdir(parents=True)
    Presentation().save(run_dir / "final.pptx")
    (run_dir / "final-manifest.json").write_text('{"pages": 1}', encoding="utf-8")

    payload = run_adapter(Path(__import__("sys").executable), Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "rebuild_adapter.py", request, result)

    assert payload["status"] == "succeeded"
    assert payload["editability"] == "reconstructed_editable"
    assert (root / "build" / "adapter-staging" / "deck.pptx").is_file()


def test_rebuild_adapter_blocks_when_no_reconstructable_source_is_designated(tmp_path: Path) -> None:
    notes = tmp_path / "notes.txt"
    notes.write_text("content only", encoding="utf-8")
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="rebuild", role="asset", source=notes)
    payload = run_adapter(Path(__import__("sys").executable), Path(__file__).resolve().parents[1] / "scripts" / "adapters" / "rebuild_adapter.py", request, result)
    assert payload["status"] == "blocked"
    assert payload["failure"]["category"] == "input_role"


def test_latest_rebuild_prepare_resumes_without_resetting_state(tmp_path: Path) -> None:
    image = tmp_path / "scan.png"
    Image.new("RGB", (320, 180), "white").save(image)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="rebuild", role="flattened_input", source=image)
    adapter = Path(__file__).resolve().parents[1] / "scripts/adapters/rebuild_adapter.py"
    first = run_adapter(Path(__import__("sys").executable), adapter, request, result)
    assert first["failure"]["category"] == "reconstruction_pending"
    state = root / "workspace/editppt/deck_manifest.json"
    before = state.read_bytes()
    assert json.loads(before)["image_backend"]["backend_id"] == "builtin-imagegen"
    second = run_adapter(Path(__import__("sys").executable), adapter, request, result)
    assert second["failure"]["category"] == "reconstruction_pending"
    assert state.read_bytes() == before


@pytest.mark.parametrize("complete,passed,slides,expected", [(True, True, 1, "succeeded"), (False, True, 1, "blocked"), (True, False, 1, "blocked"), (True, True, 2, "blocked")])
def test_latest_finalization_requires_complete_validated_matching_pages(tmp_path: Path, complete, passed, slides, expected) -> None:
    image = tmp_path / "scan.png"
    Image.new("RGB", (320, 180), "white").save(image)
    root = tmp_path / "job"
    request, result = write_contract_job(root, route="rebuild", role="flattened_input", source=image)
    run_dir = root / "workspace/editppt"
    final = run_dir / "final"
    final.mkdir(parents=True)
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(final / "scan_edited.pptx")
    (run_dir / "deck_manifest.json").write_text(json.dumps({"output": "final/scan_edited.pptx", "page_count": 1}))
    (run_dir / "run_state.json").write_text(json.dumps({"status": "complete" if complete else "deck_built"}))
    (final / "run_summary.json").write_text(json.dumps({"status": "complete", "page_count": 1}))
    (final / "validation.json").write_text(json.dumps({"passed": passed, "slides": slides}))
    payload = run_adapter(Path(__import__("sys").executable), Path(__file__).resolve().parents[1] / "scripts/adapters/rebuild_adapter.py", request, result)
    assert payload["status"] == expected, payload
    assert (root / "build/adapter-staging/deck.pptx").is_file() == (expected == "succeeded")
