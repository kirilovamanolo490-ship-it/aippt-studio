from __future__ import annotations

import hashlib
from copy import deepcopy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
import pytest

from aippt_studio.presentation_quality import audit_presentation


def _deck():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333333), Inches(7.5)
    return presentation, presentation.slides.add_slide(presentation.slide_layouts[6])


def _text(slide, text="正文示例", size=24, *, x=1, y=2, w=5, h=1):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size) if size is not None else None
    return shape


def _image(tmp_path: Path, *, alpha=False):
    path = tmp_path / ("transparent.png" if alpha else "opaque.png")
    Image.new("RGBA" if alpha else "RGB", (400, 240), (255, 255, 255, 0) if alpha else (32, 50, 70)).save(path)
    return path


def _audit(presentation, tmp_path: Path, policy=None):
    source = tmp_path / "fixture.pptx"
    presentation.save(source)
    before = source.read_bytes()
    result = audit_presentation(source, policy)
    assert source.read_bytes() == before
    assert result["source_modified"] is False
    return result


def _issues(result, code):
    return [item for item in result["issues"] if item["code"] == code and item.get("disposition") == "open"]


def test_normal_background_and_large_type_pass_without_visual_quality_claim(tmp_path: Path):
    deck, slide = _deck()
    slide.shapes.add_picture(str(_image(tmp_path)), 0, 0, width=deck.slide_width, height=deck.slide_height)
    title = _text(slide, "客户问题与解决方案", 32, x=1, y=.4, w=11, h=.7)
    _text(slide, "清晰说明当前业务结论", 24)
    _text(slide, "来源：当前客户提供的业务材料", 16, y=6.6, w=8, h=.4)
    _text(slide, "1", 16, x=11.5, y=6.8, w=.5, h=.3)
    result = _audit(deck, tmp_path, {"object_roles": [{"slide": 1, "object_id": str(title.shape_id), "role": "title", "reason": "主标题"}]})
    assert result["status"] == "passed", result["issues"]
    assert result["visual_quality_verified"] is False
    assert result["slides"][0]["objects"][0]["background_container"] is True
    assert "not rendered" in result["coverage"]["text_wrapping"]


def test_selected_slides_are_explicitly_scoped_and_do_not_claim_full_deck_pass(tmp_path: Path):
    deck, first = _deck()
    _text(first, "第一页")
    second = deck.slides.add_slide(deck.slide_layouts[6])
    _text(second, "第二页")
    source = tmp_path / "scoped.pptx"
    deck.save(source)

    result = audit_presentation(source, selected_slides=[2])

    assert result["status"] == "passed_scoped"
    assert [slide["slide"] for slide in result["slides"]] == [2]
    assert result["summary"]["total_deck_slides"] == 2
    assert result["scope"] == {"kind": "selected_slides", "selected_slides": [2], "full_deck_verified": False,
                               "policy_records_not_evaluated": {"object_roles": 0, "exceptions": 0}}


def test_static_layout_text_overlap_is_reported_for_rendered_slide(tmp_path: Path):
    deck, slide = _deck()
    inherited = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    inherited.text = "版式固定文字"
    deck.slide_layouts[6]._element.spTree.insert_element_before(deepcopy(inherited._element), "p:extLst")
    slide.shapes._spTree.remove(inherited._element)
    inherited = deck.slide_layouts[6].shapes[-1]
    _text(slide, "页面正文", 24, x=1, y=2, w=5, h=1)

    result = _audit(deck, tmp_path)

    issue = _issues(result, "INHERITED_TEXT_OVERLAP")[0]
    assert issue["severity"] == "review"
    assert issue["details"]["other_object_id"] == f"layout:{inherited.shape_id}"
    assert issue["details"]["certainty"] == "text_container_geometry_only"
    assert result["status"] == "needs_review"


def test_scoped_audit_ignores_policy_records_for_untouched_pages(tmp_path: Path):
    deck, first = _deck()
    first_shape = _text(first, "第一页")
    second = deck.slides.add_slide(deck.slide_layouts[6])
    _text(second, "第二页")
    source = tmp_path / "scoped-policy.pptx"
    deck.save(source)
    policy = {"object_roles": [{"slide": 1, "object_id": str(first_shape.shape_id), "role": "body",
                                "reason": "Whole-deck policy for the untouched first page"}]}
    result = audit_presentation(source, policy, selected_slides=[2])
    assert result["status"] == "passed_scoped"
    assert result["scope"]["policy_records_not_evaluated"]["object_roles"] == 1
    assert result["policy"]["object_roles"] == []


@pytest.mark.parametrize("selected", [[], [1, 1], [True], [0]])
def test_selected_slides_reject_empty_duplicate_or_invalid_pages(tmp_path: Path, selected):
    deck, slide = _deck()
    _text(slide)
    source = tmp_path / "scoped-invalid.pptx"
    deck.save(source)
    with pytest.raises(ValueError, match="selected_slides"):
        audit_presentation(source, selected_slides=selected)


def test_selected_slides_reports_nonexistent_page_as_audit_error(tmp_path: Path):
    deck, slide = _deck()
    _text(slide)
    source = tmp_path / "scoped-invalid-page.pptx"
    deck.save(source)
    result = audit_presentation(source, selected_slides=[2])
    assert result["status"] == "errors"
    assert result["scope"]["full_deck_verified"] is False
    assert "existing slides" in result["errors"][0]["message"]


def test_explicit_project_font_constraint_is_blocker_with_exact_object_and_role(tmp_path: Path):
    deck, slide = _deck()
    shape = _text(slide, size=14)
    result = _audit(deck, tmp_path, {"typography_mode": "explicit", "minimum_font_sizes": {"body": 24}})
    issue = _issues(result, "SMALL_FONT")[0]
    assert result["status"] == "needs_review"
    assert issue["severity"] == "blocker"
    assert issue["object_id"] == str(shape.shape_id)
    assert issue["details"]["role"] == "body"
    assert issue["details"]["minimum_font_pt"] == 24
    assert issue["details"]["actual_min_font_pt"] == 14
    assert issue["details"]["text_samples"][0]["font_source"] == "run.rPr"
    assert issue["details"]["policy_basis"] == "explicit_project_constraint"


def test_small_annotation_is_advisory_and_does_not_claim_visual_success(tmp_path: Path):
    deck, slide = _deck()
    _text(slide, "业务结论", 36, y=.5, w=9)
    _text(slide, "按当前样本口径统计", 11, y=5.5, w=4, h=.5)
    result = _audit(deck, tmp_path)
    hint = _issues(result, "SMALL_FONT")[0]
    assert result["status"] == "passed"
    assert result["blocking_issues"] == []
    assert result["summary"]["open_issues"] == 0
    assert result["summary"]["advisories"] == 1
    assert result["advisory_issues"] == [hint]
    assert hint["severity"] == "advisory"
    assert hint["details"]["actual_min_font_pt"] == 11
    assert hint["details"]["policy_basis"] == "diagnostic_reference"
    assert "minimum_font_pt" not in hint["details"]
    assert result["typography_mode"] == "advisory"
    assert result["visual_quality_verified"] is False
    assert "actual slide review" in result["coverage"]["typography_judgment"]


@pytest.mark.parametrize("policy", [
    {"minimum_font_sizes": {"body": 30}},
    {"typography_mode": "advisory", "minimum_font_sizes": {"body": 30}},
    {"typography_mode": "explicit"},
])
def test_reference_defaults_or_mode_alone_cannot_create_a_font_constraint(tmp_path: Path, policy):
    deck, slide = _deck()
    _text(slide, size=12)
    result = _audit(deck, tmp_path, policy)
    assert result["status"] == "passed"
    assert result["blocking_issues"] == []
    assert _issues(result, "SMALL_FONT")[0]["severity"] == "advisory"


def test_explicit_constraints_only_enforce_supplied_roles(tmp_path: Path):
    deck, slide = _deck()
    body = _text(slide, size=14)
    caption = _text(slide, "数据标注", 10, y=4)
    policy = {"typography_mode": "explicit", "minimum_font_sizes": {"body": 18},
              "object_roles": [{"slide": 1, "object_id": str(caption.shape_id), "role": "caption", "reason": "本页辅助数据标注"}]}
    result = _audit(deck, tmp_path, policy)
    assert result["status"] == "needs_review"
    assert [issue["object_id"] for issue in result["blocking_issues"]] == [str(body.shape_id)]
    assert result["blocking_issues"][0]["details"]["minimum_font_pt"] == 18
    assert [issue["object_id"] for issue in result["advisory_issues"]] == [str(caption.shape_id)]
    assert result["policy"]["minimum_font_sizes"] == {"body": 18}


def test_explicit_project_minimum_can_be_below_diagnostic_reference(tmp_path: Path):
    deck, slide = _deck()
    _text(slide, size=12)
    result = _audit(deck, tmp_path, {"typography_mode": "explicit", "minimum_font_sizes": {"body": 10}})
    assert result["status"] == "passed"
    assert _issues(result, "SMALL_FONT") == []
    below = _audit(deck, tmp_path, {"typography_mode": "explicit", "minimum_font_sizes": {"body": 13}})
    assert below["status"] == "needs_review"
    assert below["blocking_issues"][0]["details"]["minimum_font_pt"] == 13


def test_small_project_exception_has_no_implicit_sixteen_point_floor(tmp_path: Path):
    deck, slide = _deck()
    shape = _text(slide, "限定样本标注", 10)
    policy = {"typography_mode": "explicit", "minimum_font_sizes": {"body": 14},
              "exceptions": [{"slide": 1, "object_id": str(shape.shape_id), "code": "SMALL_FONT", "min_font_pt": 10,
                              "reason": "本页已确认的独立标注", "evidence": "当前对象阅读尺度确认"}]}
    result = _audit(deck, tmp_path, policy)
    assert result["status"] == "passed"
    assert result["exceptions"]["applied"][0]["min_font_pt"] == 10
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    changed = _audit(deck, tmp_path, policy)
    assert changed["status"] == "needs_review"
    assert changed["exceptions"]["applied"] == []


def test_title_inherits_layout_size_and_body_inherits_master_size(tmp_path: Path):
    deck = Presentation()
    layout = deck.slide_layouts[1]
    layout.placeholders[0].text_frame.paragraphs[0].font.size = Pt(26)
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = "继承标题"
    body = slide.placeholders[1]
    body.text = "继承正文"
    ppr = deck.slide_master._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txStyles/{http://schemas.openxmlformats.org/presentationml/2006/main}bodyStyle/{http://schemas.openxmlformats.org/drawingml/2006/main}lvl1pPr")
    ppr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr").set("sz", "2000")
    result = _audit(deck, tmp_path)
    issues = _issues(result, "SMALL_FONT")
    title_issue = next(item for item in issues if item["object_id"] == str(slide.shapes.title.shape_id))
    body_issue = next(item for item in issues if item["object_id"] == str(body.shape_id))
    assert title_issue["details"]["actual_min_font_pt"] == 26
    assert title_issue["details"]["text_samples"][0]["font_source"].startswith("layout_placeholder")
    assert body_issue["details"]["actual_min_font_pt"] == 20
    assert body_issue["details"]["text_samples"][0]["font_source"].startswith("master.bodyStyle")


def test_paragraph_default_and_normal_autofit_are_measured(tmp_path: Path):
    deck, slide = _deck()
    shape = _text(slide, size=None)
    shape.text_frame.paragraphs[0].font.size = Pt(28)
    body = shape.text_frame._txBody.bodyPr
    for child in list(body):
        if child.tag.endswith(("spAutoFit", "normAutofit")):
            body.remove(child)
    norm = OxmlElement("a:normAutofit")
    norm.set("fontScale", "50000")
    body.append(norm)
    result = _audit(deck, tmp_path)
    sample = _issues(result, "SMALL_FONT")[0]["details"]["text_samples"][0]
    assert sample["declared_font_pt"] == 28
    assert sample["font_pt"] == 14
    assert sample["font_source"] == "paragraph.defRPr"


def test_table_cells_and_chart_labels_are_not_omitted(tmp_path: Path):
    deck, slide = _deck()
    table_shape = slide.shapes.add_table(2, 2, Inches(.5), Inches(2), Inches(5), Inches(2))
    for row in table_shape.table.rows:
        for cell in row.cells:
            cell.text = "可读表格"
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    table_shape.table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("示例", [1, 2])
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(2), Inches(5), Inches(3), data)
    chart_shape.chart.font.size = Pt(18)
    chart_shape.chart.category_axis.tick_labels.font.size = Pt(14)
    result = _audit(deck, tmp_path)
    table_issue = next(item for item in _issues(result, "SMALL_FONT") if item["object_id"] == str(table_shape.shape_id))
    chart_issue = next(item for item in _issues(result, "SMALL_FONT") if item["object_id"] == str(chart_shape.shape_id))
    assert table_issue["details"]["text_samples"][0]["text_target"] == "cell:2:2"
    assert chart_issue["details"]["reference_font_pt"] == 18
    assert chart_issue["details"]["actual_min_font_pt"] == 14


def test_nested_group_coordinates_scaling_and_out_of_bounds(tmp_path: Path):
    deck, slide = _deck()
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    text = _text(inner, "分组文字", 24, x=0, y=0, w=2, h=1)
    # python-pptx updates ext/chExt while adding children; explicitly set the
    # actual group transform to scale the text to 12pt and move it off-page.
    xfrm = outer._element.grpSpPr.xfrm
    xfrm.off.x, xfrm.off.y = Inches(13), Inches(1)
    xfrm.ext.cx, xfrm.ext.cy = Inches(1), Inches(.5)
    xfrm.chOff.x, xfrm.chOff.y = 0, 0
    xfrm.chExt.cx, xfrm.chExt.cy = Inches(2), Inches(1)
    result = _audit(deck, tmp_path)
    row = result["slides"][0]["objects"][0]
    assert row["object_id"] == str(text.shape_id)
    assert len(row["group_ids"]) == 2
    assert row["bounds_pt"][0] == pytest.approx(936)
    assert _issues(result, "OUT_OF_BOUNDS")
    assert _issues(result, "SMALL_FONT")[0]["details"]["actual_min_font_pt"] == 12


def test_opaque_foreground_image_covering_text_is_hard_occlusion(tmp_path: Path):
    deck, slide = _deck()
    text = _text(slide, size=48, x=1, y=2, w=4, h=1)
    image = slide.shapes.add_picture(str(_image(tmp_path)), Inches(.9), Inches(1.9), Inches(4.2), Inches(1.2))
    result = _audit(deck, tmp_path)
    issue = _issues(result, "TEXT_OCCLUDED")[0]
    assert issue["object_id"] == str(text.shape_id)
    assert issue["details"]["other_object_id"] == str(image.shape_id)
    assert issue["severity"] == "blocker"
    assert result["status"] == "needs_review"
    assert _issues(result, "SMALL_FONT") == []
    assert result["typography_mode"] == "advisory"


def test_text_box_with_insets_consuming_usable_area_needs_review(tmp_path: Path):
    deck, slide = _deck()
    collapsed = _text(slide, "这段文字没有可用的内容区域", 24, x=1, y=2, w=5, h=1)
    collapsed.text_frame.margin_left = Inches(3)
    collapsed.text_frame.margin_right = Inches(3)
    result = _audit(deck, tmp_path)
    issue = _issues(result, "TEXT_CONTENT_AREA_COLLAPSED")[0]
    assert result["status"] == "needs_review"
    assert issue["severity"] == "review"
    assert issue["object_id"] == str(collapsed.shape_id)
    assert issue["details"]["certainty"] == "insets_leave_no_text_area"
    assert issue["details"]["text_preview"] == "这段文字没有可用的内容区域"


def test_text_box_with_positive_usable_area_does_not_trigger_collapsed_signal(tmp_path: Path):
    deck, slide = _deck()
    _text(slide, "正常文字区域", 24, x=1, y=2, w=5, h=1)
    result = _audit(deck, tmp_path)
    assert _issues(result, "TEXT_CONTENT_AREA_COLLAPSED") == []
    assert result["status"] == "passed"


def test_transparent_foreground_image_is_review_not_proven_occlusion(tmp_path: Path):
    deck, slide = _deck()
    _text(slide, x=1, y=2, w=4, h=1)
    slide.shapes.add_picture(str(_image(tmp_path, alpha=True)), Inches(.9), Inches(1.9), Inches(4.2), Inches(1.2))
    result = _audit(deck, tmp_path)
    assert _issues(result, "TEXT_OCCLUDED") == []
    assert _issues(result, "IMAGE_TEXT_OVERLAP")[0]["severity"] == "review"


def test_text_collision_is_geometry_review_and_pair_exception_is_traceable(tmp_path: Path):
    deck, slide = _deck()
    first = _text(slide, "有意叠字底层", 28)
    second = _text(slide, "有意叠字上层", 28, x=1.1, y=2.1)
    result = _audit(deck, tmp_path)
    assert _issues(result, "TEXT_OVERLAP")[0]["severity"] == "review"
    policy = {"exceptions": [{"slide": 1, "object_id": str(first.shape_id), "other_object_id": str(second.shape_id),
                               "code": "TEXT_OVERLAP", "max_overlap_ratio": 1.0,
                               "reason": "用户选择的叠字标题效果", "evidence": "当前对象对照样稿确认"}]}
    accepted = _audit(deck, tmp_path, policy)
    assert accepted["status"] == "passed"
    assert accepted["issues"][0]["disposition"] == "accepted_object_exception"
    assert accepted["exceptions"]["applied"][0]["evidence"] == "当前对象对照样稿确认"
    _text(slide, "第三个意外相交文字", 24, x=1.2, y=2.2)
    added = _audit(deck, tmp_path, policy)
    assert added["status"] == "needs_review"
    assert len(_issues(added, "TEXT_OVERLAP")) == 2


def test_font_exception_is_bounded_to_one_object_and_keeps_original_issue(tmp_path: Path):
    deck, slide = _deck()
    first = _text(slide, size=20)
    second = _text(slide, size=18, y=4)
    policy = {"typography_mode": "explicit", "minimum_font_sizes": {"body": 24},
              "exceptions": [{"slide": 1, "object_id": str(first.shape_id), "code": "SMALL_FONT", "min_font_pt": 20,
                               "reason": "接受当前单一参考对象字号", "evidence": "用户本轮文字确认"}]}
    result = _audit(deck, tmp_path, policy)
    assert len(result["exceptions"]["applied"]) == 1
    assert [item["object_id"] for item in _issues(result, "SMALL_FONT")] == [str(second.shape_id)]
    first.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    result = _audit(deck, tmp_path, policy)
    assert result["exceptions"]["applied"] == []
    assert len(result["exceptions"]["unused"]) == 1


@pytest.mark.parametrize("policy", [
    {"skip_all": True}, {"minimum_font_sizes": {"body": 0}},
    {"minimum_font_sizes": {"body": float("nan")}}, {"minimum_font_sizes": {"body": True}},
    {"minimum_font_sizes": {"unknown-role": 18}}, {"typography_mode": "auto"},
    {"exceptions": [{"slide": 1, "object_id": "*", "code": "SMALL_FONT", "reason": "all", "evidence": "all", "min_font_pt": 16}]},
    {"exceptions": [{"slide": 1, "object_id": "2", "code": "SMALL_FONT", "reason": "tiny", "evidence": "current", "min_font_pt": 0}]},
])
def test_global_or_unbounded_waivers_are_rejected(tmp_path: Path, policy):
    deck, slide = _deck()
    _text(slide)
    assert _audit(deck, tmp_path, policy)["status"] == "errors"


def test_label_caption_roles_are_explicit_and_body_cannot_masquerade_as_source(tmp_path: Path):
    deck, slide = _deck()
    label = _text(slide, "年度标签", 18)
    caption = _text(slide, "图像说明", 18, y=4)
    roles = [{"slide": 1, "object_id": str(shape.shape_id), "role": role, "reason": "本页语义标注"}
             for shape, role in [(label, "label"), (caption, "caption")]]
    assert _audit(deck, tmp_path, {"object_roles": roles})["status"] == "passed"
    roles[0]["role"] = "source"
    result = _audit(deck, tmp_path, {"object_roles": roles})
    assert result["status"] == "errors"
    assert "without source/reference semantics" in result["errors"][0]["message"]


def test_normal_text_box_near_top_is_explainable_title_candidate_not_hard_title(tmp_path: Path):
    deck, slide = _deck()
    label = _text(slide, "这个短句可能是标题也可能是辅助标注", 24, y=.2, w=8)
    result = _audit(deck, tmp_path)
    assert _issues(result, "TITLE_ROLE_UNCONFIRMED") == []
    assert _issues(result, "TITLE_ROLE_HINT")[0]["severity"] == "advisory"
    assert result["status"] == "passed"
    assert result["summary"]["open_issues"] == 0
    assert result["visual_quality_verified"] is False
    assert _issues(result, "SMALL_FONT") == []
    policy = {"object_roles": [{"slide": 1, "object_id": str(label.shape_id), "role": "label", "reason": "本页辅助标注"}]}
    assert _audit(deck, tmp_path, policy)["status"] == "passed"


def test_legacy_mode_is_read_only_advisory_and_corrupt_file_reports_error(tmp_path: Path):
    deck, slide = _deck()
    _text(slide, size=12)
    result = _audit(deck, tmp_path, {"mode": "legacy_audit"})
    assert result["status"] == "passed"
    assert result["blocking_issues"] == []
    assert result["summary"]["enforcement"] == "legacy_read_only_advisory"
    source = tmp_path / "broken.pptx"
    source.write_bytes(b"not a PPTX")
    assert audit_presentation(source)["status"] == "errors"


def test_hidden_child_does_not_hide_visible_small_text_in_same_group(tmp_path: Path):
    deck, slide = _deck()
    group = slide.shapes.add_group_shape()
    hidden = _text(group, "隐藏的辅助对象", 12)
    hidden._element.nvSpPr.cNvPr.set("hidden", "1")
    visible = _text(group, "仍可见的小字", 12, y=4)
    result = _audit(deck, tmp_path)
    assert result["slides"][0]["hidden_object_ids"] == [str(hidden.shape_id)]
    assert [item["object_id"] for item in _issues(result, "SMALL_FONT")] == [str(visible.shape_id)]


def test_source_word_in_body_does_not_grant_source_font_minimum(tmp_path: Path):
    deck, slide = _deck()
    _text(slide, "来源多元促进客户协同", 16)
    result = _audit(deck, tmp_path)
    assert _issues(result, "SMALL_FONT")[0]["details"]["role"] == "body"
    assert _issues(result, "SMALL_FONT")[0]["details"]["reference_font_pt"] == 24


def test_foreground_picture_covers_table_and_opaque_text_box_covers_text(tmp_path: Path):
    deck, slide = _deck()
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(4), Inches(1))
    table.table.cell(0, 0).text = "关键结论"
    table.table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    slide.shapes.add_picture(str(_image(tmp_path)), Inches(.9), Inches(1.9), Inches(4.2), Inches(1.2))
    first = _text(slide, "被后景文本框填充遮住", 24, y=4)
    cover = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.9), Inches(3.9), Inches(5.2), Inches(1.2))
    cover.fill.solid()
    cover.fill.fore_color.rgb = RGBColor(255, 255, 255)
    cover.text = "上层文字"
    cover.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    result = _audit(deck, tmp_path)
    assert {item["object_id"] for item in _issues(result, "TEXT_OCCLUDED")} == {str(table.shape_id), str(first.shape_id)}


def test_bounded_overflow_exception_does_not_hide_larger_future_overflow(tmp_path: Path):
    deck, slide = _deck()
    text = _text(slide, x=10, w=3.5)
    policy = {"exceptions": [{"slide": 1, "object_id": str(text.shape_id), "code": "OUT_OF_BOUNDS", "max_overflow_pt": 18,
                               "reason": "当前单对象边缘延伸样稿已审阅", "evidence": "本轮参考对象确认"}]}
    result = _audit(deck, tmp_path, policy)
    assert result["status"] == "passed"
    text.width = Inches(4)
    result = _audit(deck, tmp_path, policy)
    assert _issues(result, "OUT_OF_BOUNDS")


def test_disabled_chart_label_format_does_not_create_false_small_text_blocker(tmp_path: Path):
    deck, slide = _deck()
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("示例", [1, 2])
    shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4), data)
    shape.chart.font.size = Pt(18)
    shape.chart.plots[0].has_data_labels = True
    labels = shape.chart.plots[0].data_labels
    labels.font.size = Pt(12)
    labels.show_value = False
    labels.show_category_name = False
    labels.show_series_name = False
    labels.show_percentage = False
    result = _audit(deck, tmp_path)
    assert result["status"] == "passed", result["issues"]


def test_unknown_chart_font_is_advisory_unless_a_project_constraint_needs_verification(tmp_path: Path):
    deck, slide = _deck()
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("示例", [1, 2])
    shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4), data)
    for node in shape.chart._chartSpace.iter():
        node.attrib.pop("sz", None)
    advisory = _audit(deck, tmp_path)
    assert advisory["status"] == "passed"
    assert _issues(advisory, "FONT_UNRESOLVED")[0]["severity"] == "advisory"
    explicit = _audit(deck, tmp_path, {"typography_mode": "explicit", "minimum_font_sizes": {"label": 12}})
    assert explicit["status"] == "needs_review"
    assert _issues(explicit, "FONT_UNRESOLVED")[0]["severity"] == "review"


def test_legacy_explicit_diagnostics_and_geometry_remain_read_only(tmp_path: Path):
    deck, slide = _deck()
    _text(slide, size=12, x=12, w=3)
    result = _audit(deck, tmp_path, {"mode": "legacy_audit", "typography_mode": "explicit", "minimum_font_sizes": {"body": 14}})
    assert result["status"] == "needs_review"
    assert result["blocking_issues"] == []
    assert _issues(result, "OUT_OF_BOUNDS")[0]["severity"] == "blocker"
    assert _issues(result, "SMALL_FONT")[0]["severity"] == "blocker"
