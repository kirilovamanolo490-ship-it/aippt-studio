from pathlib import Path
import json

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from typer.testing import CliRunner

from aippt_studio.authoring_policy import authoring_defaults, check_job_presentation
from aippt_studio.cli import app
from aippt_studio.context_intelligence import recall_context
from aippt_studio.models import JobState
from aippt_studio.renderer import RenderError
from aippt_studio.storage import job_dir, load_job, save_job
from aippt_studio.workflow import create_job, status_payload


def tiny_pptx(path: Path) -> Path:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    shape.text = "正文不能因为内容多就自动缩成小字"
    for run in shape.text_frame.paragraphs[0].runs:
        run.font.size = Pt(12)
    presentation.save(path)
    return path


def test_new_job_keeps_persistent_user_defaults():
    job = create_job("新的无关主题", job_id="default-policy")
    assert load_job(job.id).authoring_policy_id == "presentation-expression-v3"
    policy = json.loads((job_dir(job.id) / "planning/authoring-policy.json").read_text())
    assert policy["minimum_font_sizes"] == {}
    assert policy["typography_mode"] == "advisory"
    assert policy["expression_review_required"] is True
    assert policy["rules"][0]["id"] == "audience_message"
    actions=' '.join(rule['action'] for rule in policy['rules'])
    assert 'concise' in actions and 'evidence detail' in actions and 'speaker notes' in actions
    assert 'native picture geometry' in actions and 'template' in actions


def test_unrelated_query_still_receives_explicit_preferences(tmp_path, monkeypatch):
    monkeypatch.setattr("aippt_studio.context_intelligence.recommend_upstream_tools", lambda _: [])
    context = recall_context("新的无关主题", tmp_path)
    assert context["completed_lessons"] == []
    assert context["user_authoring_defaults"]["expression_review_required"] is True
    assert "user_authoring_defaults" in context["precedence"]


def test_font_reference_is_advisory_and_report_survives(tmp_path, monkeypatch):
    monkeypatch.delenv("AIPPT_FAKE_RENDER")
    job = create_job("文字检查", job_id="quality-block")
    root = job_dir(job.id)
    check_job_presentation(root, tiny_pptx(tmp_path / "tiny.pptx"), "test")
    report = json.loads((root / "qa/presentation-quality-test.json").read_text())
    assert report["status"] == "passed"
    assert report["enforced"] is True
    assert any(issue["code"] == "SMALL_FONT" for issue in report["issues"])


def test_removing_new_job_policy_does_not_disable_quality_gate(tmp_path, monkeypatch):
    monkeypatch.delenv("AIPPT_FAKE_RENDER")
    job = create_job("不能移除门禁", job_id="quality-missing")
    root = job_dir(job.id)
    (root / "planning/authoring-policy.json").unlink()
    with pytest.raises(RenderError, match="policy is missing"):
        check_job_presentation(root, tiny_pptx(tmp_path / "tiny.pptx"), "test")


def test_legacy_audit_reports_problems_without_rewriting_source(tmp_path, monkeypatch):
    monkeypatch.delenv("AIPPT_FAKE_RENDER")
    root = tmp_path / "legacy"
    source = tiny_pptx(tmp_path / "source.pptx")
    before = source.read_bytes()
    report = check_job_presentation(root, source, "legacy", enforce=False)
    assert report["status"] == "passed"
    assert report["enforced"] is False
    assert source.read_bytes() == before


def test_fake_render_never_claims_readability_pass(tmp_path):
    result = check_job_presentation(tmp_path, tmp_path / "fake.pptx", "fake")
    assert result["status"] == "skipped-test-render"


def test_quality_cli_keeps_font_diagnostic_without_design_verdict(tmp_path):
    source = tiny_pptx(tmp_path / "tiny.pptx")
    output = tmp_path / "report.json"
    result = CliRunner().invoke(app, ["quality-check", str(source), "--output", str(output), "--json"])
    assert result.exit_code == 0, result.output
    assert json.loads(output.read_text())["status"] == "passed"
    assert json.loads(output.read_text())["visual_quality_verified"] is False


def test_defaults_are_not_mutated_by_one_task():
    first = authoring_defaults()
    first["minimum_font_sizes"]["body"] = 10
    first["rules"].clear()
    assert authoring_defaults()["minimum_font_sizes"] == {}
    assert authoring_defaults()["rules"]


def test_observed_failures_are_deduplicated_with_immutable_evidence(tmp_path):
    from aippt_studio.authoring_policy import record_quality_findings
    from aippt_studio.learning import load_events
    (tmp_path / "work.json").write_text('{"status":"in_progress"}')
    root = tmp_path / "workspace/studio-jobs/demo"
    root.mkdir(parents=True)
    bad = {"status": "needs_review", "issues": [{"code": "SMALL_FONT", "slide": 1, "object_id": "2"}]}
    first = record_quality_findings(root, {**bad, "generated_at": "first-run", "source": "temporary-1.pptx"})
    second = record_quality_findings(root, {**bad, "generated_at": "second-run", "source": "temporary-2.pptx"})
    assert first["event_id"] == second["event_id"]
    assert len(load_events(tmp_path)) == 1
    assert load_events(tmp_path)[0]["reusable"] is False
    assert load_events(tmp_path)[0]["verification"] == "failed"
    assert json.loads((tmp_path / first["evidence"]).read_text()) == bad


def test_actual_pptx_requires_expression_review_through_normal_workflow(tmp_path, monkeypatch):
    from aippt_studio.ingest import ingest_sources
    from aippt_studio.workflow import run_job, approve_gate
    job = create_job("可读性门禁流程", job_id="quality-approval")
    source = tmp_path / "input.md"
    source.write_text("\n".join(f"第{i}条测试材料包含独立的信息，用来验证流程中的内容保留和页面检查。" for i in range(1, 8)))
    ingest_sources(job.id, [str(source)])
    run_job(job.id)
    approve_gate(job.id, "brief")
    run_job(job.id)
    approve_gate(job.id, "prototype")
    run_job(job.id)
    root = job_dir(job.id)
    tiny_pptx(root / "build/deck.pptx")
    monkeypatch.delenv("AIPPT_FAKE_RENDER")
    from aippt_studio.workflow import WorkflowError
    with pytest.raises(WorkflowError, match="whole-page expression review is missing"):
        approve_gate(job.id, "final")
    assert load_job(job.id).approvals["final"] is False
    assert not (root / "reviews/final-expression-review.json").is_file()


def test_native_prototype_order_matches_object_policy_numbering(tmp_path, monkeypatch):
    from aippt_studio.models import DeckPlan, SlideSpec
    from aippt_studio.native_scene import render_scene
    from aippt_studio.renderer import sha256
    from aippt_studio.storage import write_json
    job = create_job("样稿次序", job_id="quality-prototype-order")
    root = job_dir(job.id)
    plan = DeckPlan(title="样稿次序", prototype_slide_ids=["slide-03", "slide-01"],
        slides=[SlideSpec(id=f"slide-{i:02}", title=f"标题{i}", purpose="示例", layout="content") for i in range(1,4)])
    write_json(root / "planning/deck-plan.json", plan)
    write_json(root / "planning/native-scene.json", {
        "deck_plan_sha256": sha256(root / "planning/deck-plan.json"),
        "slides": [{"id": s.id, "elements": [{"type":"text", "text":s.title}]} for s in plan.slides]})
    policy = authoring_defaults()
    policy["object_roles"] = [{"slide": 3, "object_id": "2", "role": "title", "reason": "第三页标题"}]
    write_json(root / "planning/authoring-policy.json", policy)
    captured = {}
    def capture(provider, operation, scene, output):
        captured.update(scene)
        raise ValueError("stop before native rendering")
    monkeypatch.setattr("aippt_studio.native_scene.invoke", capture)
    with pytest.raises(RenderError, match="stop before native"):
        render_scene(root, tmp_path / "candidate.pptx", tmp_path / "previews", prototype=True)
    assert [page["id"] for page in captured["slides"]] == ["slide-03", "slide-01"]
    assert captured["quality_policy"]["object_roles"][0]["slide"] == 1


def test_policy_update_recovers_visual_review_without_reapproving_brief(tmp_path):
    from aippt_studio.authoring_policy import update_job_policy
    from aippt_studio.ingest import ingest_sources
    from aippt_studio.workflow import run_job, approve_gate
    from aippt_studio.reconcile import audit_job
    job = create_job("规则修正", job_id="quality-policy-recovery")
    source = tmp_path / "input.md"
    source.write_text("\n".join(f"第{i}条材料用于验证规则修正后工作流程能够继续并保持来源一致。" for i in range(1, 8)))
    ingest_sources(job.id, [str(source)])
    run_job(job.id); approve_gate(job.id, "brief"); run_job(job.id)
    approve_gate(job.id, "prototype"); run_job(job.id)
    changed = update_job_policy(job.id, {"object_roles": [{"slide": 1, "object_id": "2", "role": "caption", "reason": "明确图注"}]})
    assert changed["state"] == "prototype_review"
    assert load_job(job.id).approvals["brief"] is True
    assert load_job(job.id).approvals["prototype"] is False
    assert audit_job(job.id).status == "consistent"
    assert run_job(job.id)["state"] == "prototype_review"
    approve_gate(job.id, "prototype")
    assert run_job(job.id)["state"] == "final_review"
    before = (job_dir(job.id) / "planning/authoring-policy.json").read_bytes()
    with pytest.raises(ValueError, match="font"):
        update_job_policy(job.id, {"minimum_font_sizes": {"body": 0}})
    assert (job_dir(job.id) / "planning/authoring-policy.json").read_bytes() == before


def test_policy_change_archives_live_candidate_delivery(tmp_path):
    from aippt_studio.authoring_policy import update_job_policy

    job = create_job("候选交付规则修正", job_id="quality-policy-delivery-invalidation")
    root = job_dir(job.id)
    candidate = root / "deliverables" / "candidate_editable.pptx"
    candidate.write_bytes(b"candidate before policy change")
    (root / "deliverables" / "build-manifest.json").write_text("{}\n", encoding="utf-8")
    job.state = JobState.FINAL_REVIEW
    save_job(job)

    changed = update_job_policy(job.id, {"minimum_font_sizes": {"body": 28}})

    assert changed["state"] == JobState.PROTOTYPE_REVIEW.value
    assert not candidate.exists()
    archived = list((root / "history" / "superseded-deliverables").rglob("candidate_editable.pptx"))
    assert len(archived) == 1
    assert archived[0].read_bytes() == b"candidate before policy change"
    receipts = list((root / "reviews" / "delivery-invalidations").glob("*-deliverables.json"))
    assert len(receipts) == 1
    receipt = json.loads(receipts[0].read_text(encoding="utf-8"))
    assert receipt["reason"] == "authoring policy changed"
    assert receipt["archive_path"] == archived[0].parent.relative_to(root).as_posix()
    status = status_payload(load_job(job.id))
    assert status["artifacts"]["delivery"] is None
    assert status["artifacts"]["candidate_delivery"] is None


def test_early_native_font_failure_records_the_real_report(tmp_path, monkeypatch):
    from aippt_studio.models import DeckPlan, SlideSpec
    from aippt_studio.native_scene import render_scene
    from aippt_studio.renderer import sha256
    from aippt_studio.storage import write_json
    job = create_job("原生早期失败", job_id="native-quality-record")
    root = job_dir(job.id)
    plan = DeckPlan(title="样例", slides=[SlideSpec(id="s1", title="标题", purpose="样例", layout="content")])
    write_json(root / "planning/deck-plan.json", plan)
    write_json(root / "planning/native-scene.json", {"deck_plan_sha256": sha256(root / "planning/deck-plan.json"), "slides": [{"id": "s1", "elements": [{"type": "text", "text": "标题"}]}]})
    def invoke(provider, operation, scene, output):
        write_json(output / "native-qa.json", {"status": "passed"})
        write_json(output / "presentation-quality.json", {"status":"needs_review", "issues":[{"code":"SMALL_FONT"}]})
        return {"result": {"status": "review_required"}}
    recorded = []
    monkeypatch.setattr("aippt_studio.native_scene.invoke", invoke)
    monkeypatch.setattr("aippt_studio.authoring_policy.record_quality_findings", lambda root, report: recorded.append(report))
    with pytest.raises(RenderError, match="presentation-quality.json"):
        render_scene(root, tmp_path / "candidate.pptx", tmp_path / "previews")
    assert recorded[0]["issues"][0]["code"] == "SMALL_FONT"


@pytest.mark.parametrize("state", ["delivered", "failed"])
def test_policy_update_rejects_terminal_job_before_any_write(state):
    from aippt_studio.authoring_policy import update_job_policy
    from aippt_studio.models import JobState
    from aippt_studio.storage import save_job
    job = create_job("终态保护", job_id=f"terminal-policy-{state}")
    job.state = JobState(state)
    save_job(job)
    root = job_dir(job.id)
    before = {p.relative_to(root): p.read_bytes() for p in root.rglob("*") if p.is_file()}
    with pytest.raises(ValueError, match="Terminal job"):
        update_job_policy(job.id, {"minimum_font_sizes": {"body": 28}})
    after = {p.relative_to(root): p.read_bytes() for p in root.rglob("*") if p.is_file()}
    assert after == before


def test_standalone_policy_accepts_job_wrapper_but_not_unknown_skip():
    from aippt_studio.authoring_policy import standalone_audit_policy
    from aippt_studio.presentation_quality import _policy
    policy = authoring_defaults()
    rules = standalone_audit_policy(policy)
    assert set(rules) == {"mode", "typography_mode", "minimum_font_sizes", "object_roles", "exceptions"}
    assert _policy(rules)["typography_mode"] == "advisory"
    with pytest.raises(ValueError, match="unknown fields"):
        standalone_audit_policy({**policy, "skip_quality": True})
