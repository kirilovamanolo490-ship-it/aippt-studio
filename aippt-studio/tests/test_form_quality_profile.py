import json

from pptx import Presentation
from pptx.util import Inches

from aippt_studio.form_quality_profile import audit_form_quality, select_work_form


def _deck(path, rows):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for texts in rows:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for text, x, y in texts:
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(1), Inches(.4))
            shape.text = text
    presentation.save(path)
    return path


def test_contextual_signals_are_advisories_and_keep_meaningful_numerals(tmp_path):
    path = _deck(tmp_path / "signals.pptx", [[
        ("01", .2, .2), ("24", 5, 3), ("这页需要注意与前页的关联。", 2, 2),
        ("本项目完成24份访谈，记录按地区编号。", 2, 3),
    ]])
    report = audit_form_quality(path, "business_report")

    assert report["status"] == "advisory_review"
    assert report["source_sha256"]
    assert len(report["review_lens"]) == 2
    assert report["summary"]["decorative_ordinal_candidates"] == 1
    assert report["summary"]["production_copy_candidates"] == 1
    assert all(row["severity"] == "advisory" and row["disposition"] == "context_required"
               for row in report["signals"])
    assert not any("24份访谈" in row["excerpt"] for row in report["signals"])
    assert report["summary"]["automatic_removals"] == 0
    assert report["summary"]["automatic_visual_score"] is None


def test_form_changes_review_questions_not_font_rules_or_signal_disposition(tmp_path):
    path = _deck(tmp_path / "forms.pptx", [[("机制与适用边界", 1, 1)]])
    technical = audit_form_quality(path, "technical_explanation")
    teaching = audit_form_quality(path, "teaching")

    assert technical["review_lens"] != teaching["review_lens"]
    assert technical["signals"] == teaching["signals"] == []
    assert "font" not in str(technical).lower()


def test_unknown_form_is_rejected(tmp_path):
    path = _deck(tmp_path / "unknown.pptx", [[("页", 1, 1)]])
    try:
        audit_form_quality(path, "invented_form")
    except ValueError as error:
        assert "unsupported presentation form" in str(error)
    else:
        raise AssertionError("unknown forms must not silently use a guessed profile")


def test_work_form_uses_brief_consensus_and_surfaces_conflicts(tmp_path):
    root = tmp_path / "work" / "1"
    for job in ("first", "second"):
        path = root / "workspace" / "studio-jobs" / job / "planning" / "brief.json"
        path.parent.mkdir(parents=True)
        path.write_text(json.dumps({"presentation_intent": {"presentation_form": "teaching"}}))
    selected = select_work_form(root)
    assert selected["form"] == "teaching"
    assert selected["source"] == "unanimous_work_briefs_unbound_to_output"

    path.write_text(json.dumps({"presentation_intent": {"presentation_form": "government_briefing"}}))
    ambiguous = select_work_form(root)
    assert ambiguous["form"] == "general"
    assert ambiguous["source"] == "conflicting_work_briefs"
    assert select_work_form(root, "competition_defense")["form"] == "competition_defense"

def test_form_audit_includes_production_and_density_context_signals(tmp_path):
    path=Presentation(); slide=path.slides.add_slide(path.slide_layouts[6])
    from pptx.util import Pt
    title=slide.shapes.add_textbox(Inches(.3),Inches(.2),Inches(8),Inches(.6));title.text='年度服务改进';title.text_frame.paragraphs[0].runs[0].font.size=Pt(36)
    body=slide.shapes.add_textbox(Inches(.3),Inches(1.1),Inches(1.7),Inches(4));body.text='这一页可以放项目背景。\n'+'说明内容需要结合资料进行解释。\n'*6
    deck_copy=slide.shapes.add_textbox(Inches(3),Inches(1),Inches(5),Inches(.5));deck_copy.text='我会帮你优化PPT。'
    for paragraph in body.text_frame.paragraphs:
        for run in paragraph.runs: run.font.size=Pt(14)
    source=tmp_path/'context.pptx';path.save(source)
    report=audit_form_quality(source,'business_report')
    assert {'ASSISTANT_REPLY','TEXT_DENSITY_CONTEXT'} <= {row['code'] for row in report['signals']}
    signal=next(row for row in report['signals'] if row['code']=='TEXT_DENSITY_CONTEXT')
    assert signal['disposition']=='context_required'
    assert signal['components'] and signal['limitations']
