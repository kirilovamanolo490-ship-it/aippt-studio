"""Remaining conversion/asset/video entrypoints honor the selected host runtime."""
import importlib.util
import json
import shutil
import struct
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from adapter_test_helpers import write_contract_job
from aippt_studio import ingest, mac_video, runtime, support


def executable(path, content=b"#!/bin/sh\nexit 0\n"):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    path.chmod(0o755)
    return path


def test_markitdown_ignores_foreign_path_and_sibling_launcher(tmp_path, monkeypatch):
    python = executable(tmp_path / "selected Python" / "python")
    executable(python.with_name("markitdown.exe"), b"MZ" + bytes(100))
    foreign = executable(tmp_path / "legacy" / "markitdown", b"MZ" + bytes(100))
    monkeypatch.setenv("PATH", str(foreign.parent))
    monkeypatch.setattr(ingest.sys, "executable", str(python))
    source = tmp_path / "原始 文件.pdf"
    source.write_bytes(b"fixture")
    calls = []
    def run(command, **kwargs):
        calls.append(command)
        return subprocess.CompletedProcess(command, 0, "真实转换结果", "")
    monkeypatch.setattr(ingest.subprocess, "run", run)
    assert ingest._markitdown(source) == "真实转换结果"
    assert calls == [[str(python), "-m", "markitdown", str(source)]]


def test_markitdown_start_failure_is_explicit_without_xml_fallback(tmp_path, monkeypatch):
    source = tmp_path / "source.pptx"
    Presentation().save(source)
    def fail(*args, **kwargs):
        raise OSError("Exec format error")
    monkeypatch.setattr(ingest.subprocess, "run", fail)
    monkeypatch.setattr(ingest, "_xml_text_from_zip", lambda *a: pytest.fail("conversion failure was hidden"))
    with pytest.raises(ValueError, match="MarkItDown could not start.*Exec format error") as result:
        ingest.extract_text(source, "powerpoint")
    assert isinstance(result.value.__cause__, OSError)


def support_fixture(tmp_path, monkeypatch):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    monkeypatch.setenv("AIPPT_JOBS_DIR", str(workspace / "jobs"))
    source = workspace / "数据.json"
    source.write_text("{}")
    runtime_root = tmp_path / ".tools/extension-runtime"
    package = runtime_root / "node_modules/echarts"
    package.mkdir(parents=True)
    (package / "package.json").write_text('{"version":"fixture"}')
    (runtime_root / "package-lock.json").write_text("{}")
    monkeypatch.setattr(support, "PROJECT_ROOT", tmp_path)
    return source, workspace / "chart.svg"


def test_support_uses_aippt_node_when_path_empty(tmp_path, monkeypatch):
    source, target = support_fixture(tmp_path, monkeypatch)
    selected = executable(tmp_path / "selected runtime" / "node")
    monkeypatch.setenv("AIPPT_NODE", str(selected))
    monkeypatch.setenv("PATH", "")
    monkeypatch.setattr(support, "probe_node_runtime", lambda: {"status": "passed", "path": str(selected.resolve())})
    calls = []
    def render(command, **kwargs):
        calls.append(command)
        Path(command[-1]).write_text('<svg xmlns="http://www.w3.org/2000/svg"/>')
        return subprocess.CompletedProcess(command, 0, "", "")
    monkeypatch.setattr(support.subprocess, "run", render)
    result = support.render_support("chart", source, target)
    assert result["status"] == "succeeded"
    assert calls[0][0] == str(selected.resolve())


def test_support_rejects_foreign_node_before_render(tmp_path, monkeypatch):
    source, target = support_fixture(tmp_path, monkeypatch)
    selected = executable(tmp_path / "windows-node", b"MZ" + bytes(100))
    monkeypatch.setattr(runtime.sys, "platform", "darwin")
    monkeypatch.setenv("AIPPT_NODE", str(selected))
    monkeypatch.setattr(support, "probe_node_runtime", lambda: {"status": "failed", "reason": "foreign native executable"})
    monkeypatch.setattr(support.subprocess, "run", lambda *a, **kw: pytest.fail("foreign Node executed"))
    with pytest.raises(ValueError, match="host-native Node.js"):
        support.render_support("chart", source, target)
    assert not target.exists()


def test_template_adapter_uses_aippt_node_when_path_empty(tmp_path, monkeypatch):
    script = Path(__file__).resolve().parents[1] / "scripts/adapters/template_adapter.py"
    monkeypatch.syspath_prepend(str(script.parent))
    spec = importlib.util.spec_from_file_location("template_platform_fixture", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    template = tmp_path / "template.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(template)
    request, result = write_contract_job(tmp_path / "job", route="template", role="template", source=template, slides=1)
    monkeypatch.setattr(module, "parse_contract_args", lambda: SimpleNamespace(request=str(request), result=str(result)))
    selected = executable(tmp_path / "selected runtime" / "node")
    monkeypatch.setenv("AIPPT_NODE", str(selected))
    monkeypatch.setenv("PATH", "")
    monkeypatch.setattr(module, "probe_node_runtime", lambda: {"status": "passed", "path": str(selected.resolve())})
    calls = []
    def render(command, **kwargs):
        calls.append(command)
        shutil.copy2(template, command[-1])
        return subprocess.CompletedProcess(command, 0, "", "")
    monkeypatch.setattr(module.subprocess, "run", render)
    assert module.main() == 0
    assert json.loads(result.read_text())["status"] == "succeeded"
    assert calls[0][0] == str(selected.resolve())


@pytest.mark.parametrize("header,valid", [
    (b"MZ" + bytes(100), False),
    (b"\xcf\xfa\xed\xfe" + struct.pack("<I", 0x1000007), False),
    (b"\xcf\xfa\xed\xfe" + struct.pack("<I", 0x100000C), True),
])
def test_ffmpeg_override_requires_host_binary(tmp_path, monkeypatch, header, valid):
    monkeypatch.setattr(runtime.sys, "platform", "darwin")
    monkeypatch.setattr(runtime.platform, "machine", lambda: "arm64")
    selected = executable(tmp_path / "ffmpeg", header)
    monkeypatch.setenv("AIPPT_FFMPEG", str(selected))
    assert mac_video.media_executable() == (selected.resolve() if valid else None)


def test_ffmpeg_skips_incompatible_bundled_candidates(tmp_path, monkeypatch):
    monkeypatch.setattr(runtime.sys, "platform", "darwin")
    monkeypatch.setattr(runtime.platform, "machine", lambda: "arm64")
    monkeypatch.setattr(mac_video, "PROJECT_ROOT", tmp_path)
    monkeypatch.delenv("AIPPT_FFMPEG", raising=False)
    monkeypatch.setenv("PATH", "")
    folder = tmp_path / ".tools/media-runtime/imageio_ffmpeg/binaries"
    executable(folder / "ffmpeg-a-old-windows", b"MZ" + bytes(100))
    selected = executable(folder / "ffmpeg-b-mac", b"\xcf\xfa\xed\xfe" + struct.pack("<I", 0x100000C))
    assert mac_video.media_executable() == selected.resolve()


@pytest.mark.skipif(sys.platform == "win32", reason="host-native browser rejection is exercised on Unix hosts")
def test_node_browser_rejects_a_windows_chromium_override_before_puppeteer_starts(tmp_path):
    node = runtime.resolve_executable("node")
    if node is None:
        pytest.skip("requires a host Node runtime")
    foreign = executable(tmp_path / "chrome.exe", b"MZ" + bytes(128))
    env = dict(__import__("os").environ, AIPPT_CHROMIUM=str(foreign))
    command = (
        "import {chromiumExecutable} from './node/browser.mjs';"
        "chromiumExecutable();"
    )
    result = subprocess.run(
        [str(node), "--input-type=module", "-e", command],
        cwd=Path(__file__).resolve().parents[1],
        env=env,
        text=True,
        capture_output=True,
        timeout=20,
        check=False,
    )
    assert result.returncode != 0
    assert "compatible host Chromium" in (result.stderr + result.stdout)
