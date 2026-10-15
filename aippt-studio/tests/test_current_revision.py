from __future__ import annotations

import hashlib
import json
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from typer.testing import CliRunner

from aippt_studio import cli
from aippt_studio.fast_revision import revise_pptx
from aippt_studio.pptx_edit import inspect_pptx


def _deck(path: Path) -> Path:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "before"
    deck.save(path)
    return path


def _prepare(work: Path, *, source_area: str = "output") -> tuple[Path, Path]:
    workspace = work / "workspace"
    output = work / "output"
    source_root = work / source_area
    for path in (workspace, output, work / "input" / "source"):
        path.mkdir(parents=True, exist_ok=True)
    source = _deck(source_root / ("source.pptx" if source_area != "output" else "delivery.pptx"))
    inspected = inspect_pptx(source, slides=[1])
    shape_id = inspected["slides"][0]["objects"][0]["shape_id"]
    request = {"source": str(source), "source_sha256": inspected["source_sha256"],
               "edits": [{"action": "replace_text", "slide": 1, "shape_id": shape_id,
                          "expected_text": "before", "replacement_text": "after"}]}
    revision = workspace / "revisions" / "wording-01"
    revised = revise_pptx(request, revision, preview_backend="none")
    assert revised["revision_qa_status"] == "passed_scoped"
    return source, revision


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_revision_publish_preserves_path_and_creates_hash_named_backup(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    original = source.read_bytes()
    candidate = revision / "candidate.pptx"
    candidate_bytes = candidate.read_bytes()
    artifacts = {name: (revision / name).read_bytes()
                 for name in ("receipt.json", "fast-revision.json", "revision-qa.json")}

    result = CliRunner().invoke(cli.app, ["revision-publish", "--work-dir", str(work),
                                           "--revision", str(revision), "--json"])

    assert result.exit_code == 0, result.output
    payload = json.loads(result.output)
    backup = Path(payload["backup"])
    assert payload["status"] == "published"
    assert payload["source"] == str(source)
    assert payload["current_sha256"] == _sha(source) == _sha(candidate)
    assert backup == work / "workspace" / "history" / "revisions" / f"{_sha_bytes(original)}.pptx"
    assert backup.read_bytes() == original
    assert payload["native_visual_verified"] is False
    assert payload["visual_review_required"] is True
    assert source.name == "delivery.pptx"
    assert candidate.read_bytes() == candidate_bytes
    assert {name: (revision / name).read_bytes() for name in artifacts} == artifacts
    assert (revision / "publish-receipt.json").is_file()


def _sha_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def test_revision_publish_rejects_source_drift(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    source.write_bytes(source.read_bytes() + b"drift")
    changed = source.read_bytes()
    result = CliRunner().invoke(cli.app, ["revision-publish", "--work-dir", str(work),
                                           "--revision", str(revision), "--json"])
    assert result.exit_code == 2
    assert "source changed" in result.output
    assert source.read_bytes() == changed


def test_revision_publish_rejects_candidate_drift(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    before = source.read_bytes()
    (revision / "candidate.pptx").write_bytes(b"altered candidate")
    with pytest.raises(ValueError, match="candidate SHA-256 changed"):
        from aippt_studio.current_revision import publish_revision
        publish_revision(work, revision)
    assert source.read_bytes() == before


def test_revision_publish_rejects_input_source(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work, source_area="input/source")
    before = source.read_bytes()
    with pytest.raises(ValueError, match="never input"):
        from aippt_studio.current_revision import publish_revision
        publish_revision(work, revision)
    assert source.read_bytes() == before


def test_revision_publish_rejects_nonpassing_scoped_qa(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    before = source.read_bytes()
    qa_path = revision / "revision-qa.json"
    qa = json.loads(qa_path.read_text())
    qa["status"] = "needs_review"
    qa_path.write_text(json.dumps(qa), encoding="utf-8")
    with pytest.raises(ValueError, match="passed_scoped"):
        from aippt_studio.current_revision import publish_revision
        publish_revision(work, revision)
    assert source.read_bytes() == before


def test_revision_publish_is_idempotent_for_same_candidate(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    from aippt_studio.current_revision import publish_revision
    first = publish_revision(work, revision)
    after_first = source.read_bytes()
    second = publish_revision(work, revision)
    assert second["idempotent"] is True
    assert second["current_sha256"] == first["current_sha256"]
    assert source.read_bytes() == after_first
    assert len(list((work / "workspace/history/revisions").glob("*.pptx"))) == 1


def test_revision_publish_excludes_concurrent_candidates_for_same_source(tmp_path: Path, monkeypatch):
    import aippt_studio.current_revision as current_revision

    work = tmp_path / "work" / "52"
    source, first_revision = _prepare(work)
    inspected = inspect_pptx(source, slides=[1])
    shape_id = inspected["slides"][0]["objects"][0]["shape_id"]
    second_request = {
        "source": str(source),
        "source_sha256": inspected["source_sha256"],
        "edits": [{"action": "replace_text", "slide": 1, "shape_id": shape_id,
                   "expected_text": "before", "replacement_text": "alternate"}],
    }
    second_revision = work / "workspace" / "revisions" / "alternate-01"
    revise_pptx(second_request, second_revision, preview_backend="none")
    first_candidate_sha = _sha(first_revision / "candidate.pptx")

    lock_held = threading.Event()
    allow_first_to_continue = threading.Event()
    original_sha = current_revision._sha
    paused = False

    def pause_first_after_lock(path: Path) -> str:
        nonlocal paused
        if path == source and threading.current_thread().name.startswith("publisher-one") and not paused:
            paused = True
            lock_held.set()
            assert allow_first_to_continue.wait(timeout=5)
        return original_sha(path)

    monkeypatch.setattr(current_revision, "_sha", pause_first_after_lock)
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix="publisher-one") as pool:
        first = pool.submit(current_revision.publish_revision, work, first_revision)
        assert lock_held.wait(timeout=5)
        with pytest.raises(ValueError, match="already in progress"):
            current_revision.publish_revision(work, second_revision)
        allow_first_to_continue.set()
        result = first.result(timeout=5)

    assert result["candidate_sha256"] == first_candidate_sha
    assert _sha(source) == first_candidate_sha
    assert not list((work / "workspace" / "history" / "revisions").glob(".publish-*.lock"))


def test_revision_publish_cleans_lock_after_publish_exception(tmp_path: Path, monkeypatch):
    import aippt_studio.current_revision as current_revision

    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    original = source.read_bytes()
    original_copy2 = current_revision.shutil.copy2
    candidate = revision / "candidate.pptx"

    def fail_candidate_copy(src, dst, *args, **kwargs):
        if Path(src) == candidate:
            raise OSError("controlled copy failure")
        return original_copy2(src, dst, *args, **kwargs)

    monkeypatch.setattr(current_revision.shutil, "copy2", fail_candidate_copy)
    with pytest.raises(OSError, match="controlled copy failure"):
        current_revision.publish_revision(work, revision)
    assert source.read_bytes() == original
    assert not list((work / "workspace" / "history" / "revisions").glob(".publish-*.lock"))


def test_revision_publish_rejects_symlinked_revision_path(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    other = tmp_path / "other"
    other.mkdir()
    link = work / "workspace" / "revisions" / "linked"
    link.symlink_to(revision, target_is_directory=True)
    with pytest.raises(ValueError, match="symlink"):
        from aippt_studio.current_revision import publish_revision
        publish_revision(work, link)


def test_revision_publish_rejects_symlinked_source(tmp_path: Path):
    work = tmp_path / "work" / "52"
    source, revision = _prepare(work)
    outside = tmp_path / "outside.pptx"
    outside.write_bytes(source.read_bytes())
    source.unlink()
    source.symlink_to(outside)
    patch_path = revision / "receipt.json"
    patch = json.loads(patch_path.read_text())
    patch["source"] = str(source)
    patch_path.write_text(json.dumps(patch), encoding="utf-8")

    with pytest.raises(ValueError, match="symlink"):
        from aippt_studio.current_revision import publish_revision
        publish_revision(work, revision)
    assert outside.read_bytes()
