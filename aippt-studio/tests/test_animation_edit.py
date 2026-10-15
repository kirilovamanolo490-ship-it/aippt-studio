from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches
from lxml import etree as ET

from aippt_studio import animation_edit as ae


def deck(tmp_path):
    path = tmp_path / "source.pptx"
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    for text in ["Revenue", "Revenue"]:
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = text
    ppt.save(path)
    with ZipFile(path) as z:
        parts = {i: z.read(i) for i in z.namelist()}
    root = ae._xml(parts["ppt/slides/slide1.xml"])
    timing = ET.SubElement(root, f"{{{ae.NS['p']}}}timing")
    # Two different objects have interleaved effects with explicit triggers,
    # delays and repeat/autoReverse metadata. Only durations may change.
    for effect_id, spid, duration in [("10", "2", "800"), ("20", "3", "500"), ("30", "2", "400")]:
        timing.append(ET.fromstring(f'''<p:par xmlns:p="{ae.NS['p']}"><p:cTn id="{effect_id}" presetClass="entr" presetID="10" nodeType="clickEffect" repeatCount="2000" autoRev="1"><p:stCondLst><p:cond delay="350"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="{int(effect_id)+1}" dur="1"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:set><p:animEffect filter="fade"><p:cBhvr><p:cTn id="{int(effect_id)+2}" dur="{duration}"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>'''))
    parts["ppt/slides/slide1.xml"] = ET.tostring(root)
    with ZipFile(path, "w") as z:
        for key, value in parts.items():
            z.writestr(key, value)
    return path


def request(path):
    return {"source": str(path), "source_sha256": ae._sha(path), "slide": 1,
            "shape_id": "2", "animation_id": "10", "delta_seconds": 2}


def test_duration_preserves_interleaved_chain_and_other_bytes(tmp_path, monkeypatch):
    path = deck(tmp_path)
    before = ae._package(path)
    output = tmp_path / "call"
    output.mkdir()
    monkeypatch.setattr("aippt_studio.renderer.officecli_validate", lambda p: {"status": "passed"})
    result = ae.run_animation_operation("animation-duration", request(path), output)
    after = ae._package(output / "candidate.pptx")
    expected = dict(before)
    expected["ppt/slides/slide1.xml"] = expected["ppt/slides/slide1.xml"].replace(b'dur="800"', b'dur="2800"')
    assert after == expected
    assert ae._package(path) == before
    assert result["before_ms"] == 800 and result["after_ms"] == 2800
    assert result["all_other_xml_bytes_preserved"] is True
    assert result["playback_verified"] is False
    assert [a["duration_ms"] for a in ae.inspect_animations(output / "candidate.pptx")["slides"][0]["animations"]] == [2800, 500, 400]


@pytest.mark.parametrize("change,match", [
    ({"source_sha256": "stale"}, "source changed"),
    ({"shape_id": "999"}, "ambiguous"),
    ({"animation_id": "20"}, "animation is missing"),
    ({"delta_seconds": -2}, "positive"),
    ({"duration_seconds": 3}, "exactly one"),
    ({"delta_seconds": 0.0001}, "whole milliseconds"),
    ({"delta_seconds": float("nan")}, "whole milliseconds"),
    ({"delay": 2}, "unknown animation fields"),
])
def test_rejected_edits_never_publish_candidate(tmp_path, change, match):
    path = deck(tmp_path)
    args = request(path) | change
    output = tmp_path / "call"
    output.mkdir()
    with pytest.raises(ValueError, match=match):
        ae.run_animation_operation("animation-duration", args, output)
    assert not (output / "candidate.pptx").exists()
    assert ae._sha(path) == request(path)["source_sha256"]


def test_ambiguous_text_requires_stable_id(tmp_path):
    path = deck(tmp_path)
    with pytest.raises(ValueError, match="ambiguous"):
        ae._select(ae.inspect_animations(path), {"slide": 1, "shape_text": "Revenue"})


def test_complex_durations_are_visible_but_not_editable(tmp_path):
    path = deck(tmp_path)
    report = ae.inspect_animations(path)
    assert report["slides"][0]["animations"][0]["duration_editable"]
    root = ae._xml(ae._package(path)["ppt/slides/slide1.xml"])
    effect = root.xpath(".//p:cTn[@id='10']", namespaces=ae.NS)[0]
    behavior = ET.fromstring(f'<p:anim xmlns:p="{ae.NS["p"]}"><p:cBhvr><p:cTn id="99" dur="300"/></p:cBhvr></p:anim>')
    effect.find("p:childTnLst", ae.NS).append(behavior)
    value, nodes, reason = ae._duration(effect)
    assert value is None and reason


def test_validation_failure_does_not_publish(tmp_path, monkeypatch):
    path = deck(tmp_path)
    output = tmp_path / "call"
    output.mkdir()
    def fail(p):
        raise RuntimeError("invalid candidate")
    monkeypatch.setattr("aippt_studio.renderer.officecli_validate", fail)
    with pytest.raises(RuntimeError, match="invalid candidate"):
        ae.run_animation_operation("animation-duration", request(path), output)
    assert not (output / "candidate.pptx").exists()


def test_add_refuses_to_replace_existing_chain(tmp_path):
    path = deck(tmp_path)
    output = tmp_path / "call"
    output.mkdir()
    args = request(path)
    args.pop("animation_id")
    args.pop("delta_seconds")
    with pytest.raises(ValueError, match="without existing effects"):
        ae.run_animation_operation("animation-add", args, output)


def test_slide_order_follows_presentation_relationships(tmp_path):
    path = deck(tmp_path)
    parts = ae._package(path)
    pres = ae._xml(parts["ppt/presentation.xml"])
    rid = pres.find("p:sldIdLst/p:sldId", ae.NS).get(f"{{{ae.NS['r']}}}id")
    rels = ae._xml(parts["ppt/_rels/presentation.xml.rels"])
    next(r for r in rels if r.get("Id") == rid).set("Target", "slides/slide7.xml")
    parts["ppt/_rels/presentation.xml.rels"] = ET.tostring(rels)
    parts["ppt/slides/slide7.xml"] = parts.pop("ppt/slides/slide1.xml")
    assert ae._slides(parts) == ["ppt/slides/slide7.xml"]
