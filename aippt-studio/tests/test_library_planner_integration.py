from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from aippt_studio.library_reuse_gate import library_reuse_block_reason, record_no_suitable_library_candidate
from aippt_studio.models import InputRole, JobState, ProductionRoute
from aippt_studio.planner import create_brief, create_deck_plan
from aippt_studio.storage import job_dir, load_job, save_job, write_json
from aippt_studio.workflow import (WorkflowError, approve_gate, bind_catalog_template,
                                   configure_route, confirm_catalog_template, initialize_work_intake, run_job)


def test_new_work_plan_emits_source_page_reuse_candidates(tmp_path: Path, monkeypatch) -> None:
    from aippt_studio import config, workflow

    monkeypatch.setattr(config, "WORKSPACE_ROOT", tmp_path)
    monkeypatch.setattr(workflow, "WORKSPACE_ROOT", tmp_path)
    work = tmp_path / "work" / "88"
    (work / "input" / "source").mkdir(parents=True)
    (work / "input" / "reference").mkdir()
    (work / "workspace").mkdir()
    (work / "work.json").write_text(json.dumps({"work_id": "88", "name": "社区创新项目"}), encoding="utf-8")
    (work / "input" / "request.md").write_text(
        "新建一套竞赛答辩PPT，面向社工创新设计大赛评委，展示社区养老服务问题、协作机制和成效证据。",
        encoding="utf-8",
    )
    (work / "input" / "source" / "facts.md").write_text(
        "社区老人提出维修报事和助餐协助需求。\n社工联合物业建立入户走访和服务转介机制。\n"
        "项目记录了走访覆盖、服务参与和回访评价。\n"
        "团队把需求登记、资源协调、上门服务与后续回访连接成连续流程，所有数据均以真实服务台账为依据。\n"
        "项目准备向评委说明需求证据、创新做法、执行分工、实施风险及评估方式。",
        encoding="utf-8",
    )
    monkeypatch.setenv("AIPPT_WORK_DIR", str(work))
    monkeypatch.setenv("AIPPT_JOBS_DIR", str(work / "workspace" / "studio-jobs"))

    library = tmp_path / "模板库"
    source = library / "竞赛与答辩" / "社工创新项目.pptx"
    source.parent.mkdir(parents=True)
    presentation = Presentation()
    media = tmp_path / "embedded-picture.png"
    Image.new("RGB", (24, 24), (34, 93, 63)).save(media)
    for title in ("社区养老创新项目", "需求与问题", "社工物业协作机制"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(500000, 500000, 6000000, 800000).text = title
        if title == "社工物业协作机制":
            slide.shapes.add_picture(str(media), 900000, 1800000, 500000, 500000)
    presentation.save(source)
    catalog = {"path": source.relative_to(library).as_posix(), "asset_type": "template",
               "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
               "occasion_tags": ["竞赛与答辩", "竞赛"], "media_rights": "unverified"}
    (library / "CATALOG.jsonl").write_text(json.dumps(catalog, ensure_ascii=False) + "\n", encoding="utf-8")

    initialize_work_intake(work, job_id="library-planner-integration")
    create_brief("library-planner-integration")
    plan = create_deck_plan("library-planner-integration")
    packet_path = job_dir("library-planner-integration") / "planning" / "library-reuse-candidates.json"
    packet = json.loads(packet_path.read_text(encoding="utf-8"))
    assert packet["status"] == "ready"
    assert packet["candidates"][0]["path"] == catalog["path"]
    assert packet["candidates"][0]["target_pages"]
    assert packet["asset_candidates"]
    assert plan.planning_strategy["library_reuse"]["asset_candidate_paths"] == [catalog["path"]]
    assert packet["candidates"][0]["evidence"]["sequence_matches"]
    assert plan.planning_strategy["library_reuse"]["application"] == "candidate_selection_required"
    assert "社区养老创新项目" not in " ".join(slide.title for slide in plan.slides)

    root = job_dir("library-planner-integration")
    job = load_job("library-planner-integration")
    assert "尚未实际套用" in (library_reuse_block_reason(job, root, plan) or "")
    approve_gate(job.id, "brief")
    blocked = run_job(job.id)
    assert blocked["state"] == "blocked"
    assert "尚未实际套用" in blocked["blocked_reason"]
    reviewed = {"decision": "no_suitable_candidate",
                "reason": "来源页只有标题文本，没有足以承载本次需求证据与协作机制的结构。",
                "reviewed_pages": [{"source_path": catalog["path"], "source_slide": 1,
                                    "finding": "封面只有一处标题，没有需要的论点和证据区。"}]}
    receipt = record_no_suitable_library_candidate(job.id, reviewed)
    assert receipt["status"] == "recorded_for_current_plan"
    assert library_reuse_block_reason(job, root, plan) is None
    plan.slides[0].title += "（修订）"
    write_json(root / "planning/deck-plan.json", plan)
    assert "页内容已变化" in (library_reuse_block_reason(job, root, plan) or "")

    # A plain "new PPT" request reaches retrieval after planning. It must be
    # able to promote a reviewed current whole-deck candidate to the actual
    # template route, including when the prototype gate has blocked the job.
    template_work = tmp_path / "work" / "89"
    (template_work / "input" / "source").mkdir(parents=True)
    (template_work / "input" / "reference").mkdir()
    (template_work / "workspace").mkdir()
    (template_work / "work.json").write_text(json.dumps({"work_id": "89", "name": "社区创新项目模板复用"}), encoding="utf-8")
    (template_work / "input" / "request.md").write_text(
        "新建一套竞赛答辩PPT，面向社工创新设计大赛评委，展示社区养老服务问题、协作机制和成效证据。",
        encoding="utf-8",
    )
    (template_work / "input" / "source" / "facts.md").write_text(
        (work / "input" / "source" / "facts.md").read_text(encoding="utf-8"), encoding="utf-8",
    )
    monkeypatch.setenv("AIPPT_WORK_DIR", str(template_work))
    monkeypatch.setenv("AIPPT_JOBS_DIR", str(template_work / "workspace" / "studio-jobs"))
    initialize_work_intake(template_work, job_id="library-template-promotion")
    create_brief("library-template-promotion")
    create_deck_plan("library-template-promotion")
    approve_gate("library-template-promotion", "brief")
    assert run_job("library-template-promotion")["state"] == "blocked"
    promotion_root = job_dir("library-template-promotion")
    candidates_path = promotion_root / "planning" / "library-reuse-candidates.json"
    candidates = json.loads(candidates_path.read_text(encoding="utf-8"))
    current_signature = candidates["plan_message_signature"]
    candidates["plan_message_signature"] = "0" * 64
    write_json(candidates_path, candidates)
    with pytest.raises(WorkflowError, match="current source-bound page-level candidate"):
        bind_catalog_template("library-template-promotion", catalog["path"])
    candidates["plan_message_signature"] = current_signature
    write_json(candidates_path, candidates)
    with pytest.raises(WorkflowError, match="confirmation"):
        bind_catalog_template("library-template-promotion", catalog["path"])
    assert not (template_work / "input" / "reference" / f"{source.stem}-{catalog['source_sha256'][:12]}.pptx").exists()
    from aippt_studio.ingest import ingest_sources
    intake_job = load_job("library-template-promotion")
    intake_job.state = JobState.INTAKE
    intake_job.resume_state = None
    save_job(intake_job)
    imported_catalog_source = ingest_sources("library-template-promotion", [str(source)])[0]
    with pytest.raises(WorkflowError, match="confirmation"):
        configure_route("library-template-promotion", roles={imported_catalog_source.id: InputRole.TEMPLATE})
    # Legacy catalog entries may leave the source hash in the complete OOXML
    # structure index; route checks must resolve the verified current bytes.
    catalog_file = library / "CATALOG.jsonl"
    original_catalog = catalog_file.read_bytes()
    legacy_row = dict(catalog)
    legacy_row.pop("source_sha256")
    catalog_file.write_text(json.dumps(legacy_row, ensure_ascii=False) + "\n", encoding="utf-8")
    structure = {"schema": "aippt.template-design-index.v1", "source_path": catalog["path"],
                 "source_sha256": catalog["source_sha256"], "status": "parsed",
                 "coverage": {"full_slide_coverage": True}}
    (library / "DESIGN-STRUCTURE.jsonl").write_text(json.dumps(structure) + "\n", encoding="utf-8")
    with pytest.raises(WorkflowError, match="confirmation"):
        configure_route("library-template-promotion", roles={imported_catalog_source.id: InputRole.TEMPLATE})
    catalog_file.write_bytes(original_catalog)
    previews = []
    for page in (1, 2):
        preview = template_work / "workspace" / f"template-page-{page}.png"
        Image.new("RGB", (24, 24), (34, 93, 63)).save(preview)
        render_receipt = preview.with_suffix(".preview.json")
        rendered = {
            "schema_version": 1, "status": "rendered_advisory",
            "source_sha256": catalog["source_sha256"], "slide_number": page,
            "approximate_renderer": True,
            "limitations": ["Approximate renderer; not native Office rendering or acceptance."],
            "output_png": str(preview),
            "output_sha256": hashlib.sha256(preview.read_bytes()).hexdigest(),
        }
        render_receipt.write_text(json.dumps(rendered, ensure_ascii=False), encoding="utf-8")
        previews.append({"source_slide": page, "work_local_path": str(preview),
                         "sha256": hashlib.sha256(preview.read_bytes()).hexdigest(),
                         "render_receipt_path": str(render_receipt),
                         "render_receipt_sha256": hashlib.sha256(render_receipt.read_bytes()).hexdigest(),
                         "observation": f"来源第{page}页的实际构图观察记录。"})
    no_render_receipt = [dict(item) for item in previews]
    no_render_receipt[0].pop("render_receipt_path")
    no_render_receipt[0].pop("render_receipt_sha256")
    with pytest.raises(WorkflowError, match="preview-slide .preview.json"):
        confirm_catalog_template("library-template-promotion", {
            "selected_catalog_path": catalog["path"], "source_sha256": catalog["source_sha256"],
            "source_previews": no_render_receipt,
            "user_confirmation_text": "我确认使用这份模板，保留构图主题并替换本项目文字与图片。",
            "adaptation_scope": {
                "preserve_composition_theme": "保留构图和竞赛主题。",
                "text_image_replacement": "替换旧文字图片。",
                "allowed_layout_adjustments": "允许小幅微调布局。",
            },
        })
    first_receipt = Path(previews[0]["render_receipt_path"])
    render_bytes = first_receipt.read_bytes()
    render_data = json.loads(render_bytes)
    render_data["source_sha256"] = "0" * 64
    first_receipt.write_text(json.dumps(render_data), encoding="utf-8")
    stale_preview_receipt = dict(previews[0])
    stale_preview_receipt["render_receipt_sha256"] = hashlib.sha256(first_receipt.read_bytes()).hexdigest()
    with pytest.raises(WorkflowError, match="does not bind an advisory render"):
        confirm_catalog_template("library-template-promotion", {
            "selected_catalog_path": catalog["path"], "source_sha256": catalog["source_sha256"],
            "source_previews": [stale_preview_receipt, previews[1]],
            "user_confirmation_text": "我确认使用这份模板，保留构图主题并替换本项目文字与图片。",
            "adaptation_scope": {
                "preserve_composition_theme": "保留构图和竞赛主题。",
                "text_image_replacement": "替换旧文字图片。",
                "allowed_layout_adjustments": "允许小幅微调布局。",
            },
        })
    first_receipt.write_bytes(render_bytes)
    confirmation = confirm_catalog_template("library-template-promotion", {
        "selected_catalog_path": catalog["path"], "source_sha256": catalog["source_sha256"],
        "source_previews": previews,
        "user_confirmation_text": "我确认使用这份模板，保留构图主题并替换本项目文字与图片。",
        "adaptation_scope": {
            "preserve_composition_theme": "保留现有版式构图和竞赛主题语言。",
            "text_image_replacement": "替换所有旧项目文字与图片为本项目内容。",
            "allowed_layout_adjustments": "允许为新文案长度做小幅间距与字号调整。",
        },
    })
    assert confirmation["status"] == "confirmation_recorded"
    receipt_path = promotion_root / "planning/template-selection-confirmation.json"
    receipt_bytes = receipt_path.read_bytes()
    current_brief_path = promotion_root / "planning/brief.json"
    current_brief = current_brief_path.read_bytes()
    current_brief_path.unlink()
    with pytest.raises(WorkflowError, match="brief.json changed or is missing"):
        configure_route("library-template-promotion", route=ProductionRoute.TEMPLATE,
                        roles={imported_catalog_source.id: InputRole.TEMPLATE})
    current_brief_path.write_bytes(current_brief)
    current_plan_path = promotion_root / "planning/deck-plan.json"
    current_plan = current_plan_path.read_bytes()
    current_plan_path.unlink()
    with pytest.raises(WorkflowError, match="deck-plan.json changed or is missing"):
        configure_route("library-template-promotion", route=ProductionRoute.TEMPLATE,
                        roles={imported_catalog_source.id: InputRole.TEMPLATE})
    current_plan_path.write_bytes(current_plan)
    candidate_bytes = candidates_path.read_bytes()
    changed_candidates = json.loads(candidate_bytes)
    changed_candidates["review_probe"] = True
    write_json(candidates_path, changed_candidates)
    with pytest.raises(WorkflowError, match="candidate packet changed"):
        bind_catalog_template("library-template-promotion", catalog["path"])
    candidates_path.write_bytes(candidate_bytes)
    receipt_data = json.loads(receipt_bytes)
    receipt_data["selected_catalog_path"] = "different-template.pptx"
    write_json(receipt_path, receipt_data)
    with pytest.raises(WorkflowError, match="stale or belongs"):
        bind_catalog_template("library-template-promotion", catalog["path"])
    receipt_path.write_bytes(receipt_bytes)
    bound_job = load_job("library-template-promotion")
    request_snapshot = promotion_root / bound_job.work_request.snapshot_path
    original_request = request_snapshot.read_bytes()
    request_snapshot.write_bytes(original_request + " 修订".encode("utf-8"))
    with pytest.raises(WorkflowError):
        bind_catalog_template("library-template-promotion", catalog["path"])
    request_snapshot.write_bytes(original_request)
    preview = Path(json.loads(receipt_bytes)["source_previews"][0]["work_local_path"])
    preview_bytes = preview.read_bytes()
    preview.write_bytes(preview_bytes + b"changed")
    with pytest.raises(WorkflowError, match="preview is missing or changed"):
        bind_catalog_template("library-template-promotion", catalog["path"])
    preview.write_bytes(preview_bytes)
    bound = bind_catalog_template("library-template-promotion", catalog["path"])
    assert bound["route"] == "template"
    assert bound["source_sha256"] == catalog["source_sha256"]
    promoted = load_job("library-template-promotion")
    assert promoted.route.value == "template"
    assert (promotion_root / "planning/template-selection-adoption.json").is_file()
    # Route reconfiguration after adoption remains valid even though binding
    # invalidated/deleted the pre-adoption plan artifacts.
    assert configure_route("library-template-promotion", route=ProductionRoute.TEMPLATE).route == ProductionRoute.TEMPLATE
    adopted_catalog = json.loads(catalog_file.read_text(encoding="utf-8"))
    adopted_catalog["source_sha256"] = "0" * 64
    catalog_file.write_text(json.dumps(adopted_catalog, ensure_ascii=False) + "\n", encoding="utf-8")
    with pytest.raises(WorkflowError, match="missing or changed in the current library"):
        configure_route("library-template-promotion", route=ProductionRoute.TEMPLATE)
    catalog_file.write_bytes(original_catalog)
    assert sum(role.value == "template" for role in promoted.input_roles.values()) == 1
    assert not (promotion_root / "planning" / "deck-plan.json").exists()
