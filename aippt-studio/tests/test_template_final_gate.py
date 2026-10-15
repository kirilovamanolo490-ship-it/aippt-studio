from __future__ import annotations

import json
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from adapter_test_helpers import write_contract_job
from aippt_studio.models import DeckPlan, ExternalEngineManifest, InputRole, ProductionRoute
from aippt_studio.renderer import sha256
from aippt_studio.storage import read_model, write_json
from aippt_studio.template_verification import _template_application_inputs
from aippt_studio.workflow import WorkflowError, _validate_external_output


def test_external_template_manifest_cannot_self_certify_an_unchanged_deck(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    template = tmp_path / "competition-template.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "{{title}}"
    slide.shapes.add_textbox(1000000, 2000000, 6000000, 2000000).text = "{{body}}"
    presentation.save(template)
    root = tmp_path / "job"
    write_contract_job(root, route="template", role="template", source=template, slides=1)
    plan_path = root / "planning/deck-plan.json"
    plan = read_model(plan_path, DeckPlan)
    mapping = {
        "schema_version": 2,
        "template_source_id": "source-1",
        "template_source_sha256": sha256(root / "sources" / template.name),
        "template_source_slide_count": 1,
        "deck_plan_sha256": sha256(plan_path),
        "allow_animations": False,
        "slides": [{
            "deck_slide_id": plan.slides[0].id,
            "source_slide": 1,
            "replacements": {"title": plan.slides[0].title, "body": "\n".join(plan.slides[0].points)},
        }],
    }
    mapping_path = root / "planning/template-map.json"
    mapping_path.write_text(json.dumps(mapping), encoding="utf-8")
    pages, error = _template_application_inputs(root / "sources" / template.name, mapping, plan.model_dump(mode="json"))
    assert error is None
    deck = root / "build/deck.pptx"
    shutil.copy2(template, deck)  # Claim success while leaving original placeholders unchanged.
    preview = root / "build/previews" / f"{plan.slides[0].id}.png"
    preview.parent.mkdir()
    Image.new("RGB", (640, 360), "#12314a").save(preview)
    job = SimpleNamespace(
        id="fixture", route=ProductionRoute.TEMPLATE,
        work_request=object(), input_roles={"source-1": InputRole.TEMPLATE},
    )
    monkeypatch.setattr(
        "aippt_studio.workflow.read_bound_work_request",
        lambda _job_id: ("根据模板库里面合适的模板进行套用。", None),
    )
    write_json(root / "build/external-engine-manifest.json", ExternalEngineManifest(
        job_id=job.id, route=job.route, engine="external", version="1",
        editable_class="native_editable", deck_plan_sha256=sha256(plan_path),
        pptx_sha256=sha256(deck), preview_hashes={plan.slides[0].id: sha256(preview)},
        build_source="planning/template-map.json", build_source_sha256=sha256(mapping_path),
        template_application={
            "status": "verified", "source_sha256": mapping["template_source_sha256"],
            "page_map_sha256": sha256(mapping_path), "verified_slide_count": 1,
            "pages": pages,
        },
    ))

    with pytest.raises(WorkflowError, match="mapped .* replacement"):
        _validate_external_output(job, root, plan)
