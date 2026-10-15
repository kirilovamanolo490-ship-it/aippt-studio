import hashlib
import json
from pathlib import Path

import pytest
from pptx import Presentation

from aippt_studio.tool_evidence import collect_tool_evidence
from test_work_closure import POWERSHELL, _make_work, _run_script


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def batch_fixture(tmp_path):
    work = tmp_path / "work/1"
    job = work / "workspace/studio-jobs/example"
    (job / "edits").mkdir(parents=True)
    (job / "sources").mkdir()
    (job / "build").mkdir()
    (work / "output").mkdir(exist_ok=True)
    source, output = job / "sources/source.pptx", job / "build/edited.pptx"
    source.write_bytes(b"source fixture")
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(output)
    (work / "output/delivery.pptx").write_bytes(output.read_bytes())
    batch = {"schema_version": 1, "id": "revision", "job_id": "example",
        "input_path": "sources/source.pptx", "output_path": "build/edited.pptx",
        "input_sha256": digest(source), "output_sha256": digest(output),
        "operations": [{"id": "one", "action": "set", "target": "/slide[1]/shape[1]", "properties": {"text": "改后"}}],
        "results": [{"id": "one", "response": "Updated shape"}],
        "validation": {"status": "passed"}, "created_at": "2026-09-15T01:00:00Z", "applied_at": "2026-09-15T01:00:01Z"}
    path = job / "edits/revision.json"
    path.write_text(json.dumps(batch), encoding="utf-8")
    return work, job, path, source, output, batch


def test_edit_batch_has_separate_counts_and_exact_output_identity(tmp_path):
    work, job, path, source, output, batch = batch_fixture(tmp_path)
    result = collect_tool_evidence(work)
    assert result["counts"]["recorded_calls"] == 0
    assert result["counts"]["officecli_edit_batches"] == 1
    assert result["counts"]["officecli_recorded_applied_batches"] == 1
    row = result["officecli_edit_batches"][0]
    assert row["execution_status"] == "recorded_applied"
    assert row["input_identity"]["current_status"] == "matches_recorded"
    assert row["output_identity"]["current_status"] == "matches_recorded"
    assert row["exact_output_copies"][0]["outputs"] == ["output/delivery.pptx"]
    assert row["final_adoption"] == "not_established"
    assert row["runtime_version"] == "not_recorded"
    assert {x["path"] for x in result["records"]} >= {p.relative_to(work).as_posix() for p in [path, source, output]}


def test_revised_input_does_not_replace_historical_source_identity(tmp_path):
    work, job, path, source, output, batch = batch_fixture(tmp_path)
    source.write_bytes(b"user revised input later")
    row = collect_tool_evidence(work)["officecli_edit_batches"][0]
    assert row["execution_status"] == "recorded_applied"
    assert row["input_identity"]["recorded_sha256"] == batch["input_sha256"]
    assert row["input_identity"]["current_status"] == "differs_from_recorded"
    assert row["input_identity"]["current"]["sha256"] == digest(source)


def test_missing_output_retains_execution_but_no_current_copy_claim(tmp_path):
    work, job, path, source, output, batch = batch_fixture(tmp_path)
    output.unlink()
    row = collect_tool_evidence(work)["officecli_edit_batches"][0]
    assert row["execution_status"] == "recorded_applied"
    assert row["output_identity"]["current_status"] == "missing"
    assert row["exact_output_copies"] == []


@pytest.mark.parametrize("absolute", [True, False])
def test_external_source_not_read_or_preserved(tmp_path, absolute):
    work, job, path, source, output, batch = batch_fixture(tmp_path)
    outside = tmp_path / "outside.pptx"
    outside.write_bytes(b"outside")
    if absolute:
        batch["input_path"] = str(outside)
    else:
        source.unlink()
        source.symlink_to(outside)
    path.write_text(json.dumps(batch))
    result = collect_tool_evidence(work)
    assert result["officecli_edit_batches"][0]["input_identity"]["current_status"] == "outside_scope"
    assert all(x["sha256"] != digest(outside) for x in result["records"])


def test_corrupt_batch_isolated_and_only_canonical_edit_directory_recognized(tmp_path):
    work, job, path, source, output, batch = batch_fixture(tmp_path)
    (job / "edits/bad.json").write_text('{"schema_version": 500}')
    (work / "workspace/elsewhere.json").write_text(json.dumps(batch))
    result = collect_tool_evidence(work)
    assert len(result["officecli_edit_batches"]) == 1
    assert len(result["errors"]) == 1
    assert result["errors"][0]["kind"] == "officecli_edit_batch"


def test_failed_batch_with_output_is_not_counted_as_applied(tmp_path):
    work, job, path, source, output, batch = batch_fixture(tmp_path)
    batch.update(applied_at=None, validation={"status": "failed", "error": "invalid operation"})
    path.write_text(json.dumps(batch))
    result = collect_tool_evidence(work)
    assert result["counts"]["officecli_recorded_applied_batches"] == 0
    assert result["officecli_edit_batches"][0]["execution_status"] == "recorded_failed"


@pytest.mark.skipif(POWERSHELL is None, reason="PowerShell is not installed")
def test_only_edit_batches_write_summary_through_real_prepare_entrypoint(tmp_path):
    _make_work(tmp_path)
    work, job, path, source, output, batch = batch_fixture(tmp_path)
    run = _run_script("prepare-work-close.ps1", tmp_path, "-WorkId", "1")
    assert run.returncode == 0, run.stderr
    summary_path = work / "workspace/qa/tool-execution-summary.json"
    assert summary_path.is_file()
    summary = json.loads(summary_path.read_text())
    assert summary["counts"]["recorded_calls"] == 0
    assert summary["counts"]["officecli_edit_batches"] == 1
    manifest = json.loads((work / "workspace/qa/final-output-manifest.json").read_text(encoding="utf-8-sig"))
    collected = {item["path"] for item in manifest["adapter_evidence"]}
    assert path.relative_to(work).as_posix() in collected
    assert source.relative_to(work).as_posix() in collected
    assert output.relative_to(work).as_posix() in collected
