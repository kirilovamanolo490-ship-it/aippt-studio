from __future__ import annotations

import json
import hashlib
import shutil
from pathlib import Path

import pytest
from pptx import Presentation
from PIL import Image

from aippt_studio.html_deck import write_deck_html
from aippt_studio.ingest import ingest_sources
from aippt_studio.models import (
    ChangeMode,
    DeckPlan,
    InputRole,
    JobBrief,
    PresentationForm,
    SourceRecord,
    VisualSystemMode,
)
from aippt_studio.planner import create_brief, create_deck_plan
from aippt_studio.presentation_intent import PresentationIntentError, WorkBindingEnvironmentError, requests_bounded_revision, requests_template_adaptation, select_presentation_intent, source_document_identity
from aippt_studio.participation import _style_selection
from aippt_studio.reconcile import audit_job
from aippt_studio.storage import job_dir, load_job, read_model, write_json
from aippt_studio.workflow import (WorkflowError, _template_application_proof_reason, approve_gate,
                                   bind_catalog_template, confirm_catalog_template,
                                   configure_presentation_intent, configure_route, create_job,
                                   initialize_work_intake, run_job)


def _work(tmp_path: Path, request: str, *, name: str = "需求分类回归", work_id: str = "88") -> Path:
    work = tmp_path / "work" / work_id
    (work / "input" / "source").mkdir(parents=True)
    (work / "input" / "reference").mkdir()
    (work / "workspace").mkdir()
    (work / "work.json").write_text(json.dumps({"work_id": work_id, "name": name}, ensure_ascii=False), encoding="utf-8")
    (work / "input" / "request.md").write_text(request, encoding="utf-8")
    return work


def _content(path: Path, count: int = 10) -> Path:
    source = path / "content.md"
    source.write_text(
        "\n".join(f"培训事实{i}：现场人员完成第{i}项真实护理核对，并保留对应记录。" for i in range(1, count + 1)),
        encoding="utf-8",
    )
    return source


def _activate_work(monkeypatch, work: Path) -> None:
    monkeypatch.setenv("AIPPT_WORK_DIR", str(work))
    monkeypatch.setenv("AIPPT_JOBS_DIR", str(work / "workspace" / "studio-jobs"))


def _pptx(path: Path, slides: int = 1) -> Path:
    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(400000, 400000, 7000000, 700000).text = f"原稿第 {index + 1} 页"
    presentation.save(path)
    return path


def test_local_patch_keeps_scope_and_visual_system_separate_from_route() -> None:
    intent = select_presentation_intent(
        JobBrief(goal="铁路两页优化"),
        {"deck": InputRole.EDIT_TARGET},
        request_text="仅改第4、5页，其他页保持不动；保留原主题、Logo、图片和GIF。",
    )
    assert intent.change_mode == ChangeMode.LOCAL_PATCH
    assert intent.visual_system == VisualSystemMode.PRESERVE_EXISTING
    assert intent.preservation.target_slide_numbers == [4, 5]
    assert intent.preservation.preserve_non_target_slides is True
    assert intent.preservation.preserve_media is True
    assert intent.preservation.allow_auto_outline is False


def test_forms_change_expression_rules_without_selecting_a_fixed_palette() -> None:
    government = select_presentation_intent(
        JobBrief(goal="防汛工作"), {}, request_text="为政府领导进行正式防汛工作汇报，保持口径严谨。"
    )
    competition = select_presentation_intent(
        JobBrief(goal="创新项目"), {}, request_text="这是竞赛答辩，围绕创新、证据和可行性呈现。"
    )
    premium = select_presentation_intent(
        JobBrief(goal="作品案例"), {}, request_text="制作高级华丽的品牌作品集，展示真实项目截图。"
    )
    assert government.presentation_form == PresentationForm.GOVERNMENT_BRIEFING
    assert competition.presentation_form == PresentationForm.COMPETITION_DEFENSE
    assert premium.presentation_form == PresentationForm.PREMIUM_BRAND_PORTFOLIO
    assert "蓝" not in government.reason and "金" not in premium.reason


def test_academic_defense_is_not_classified_as_competition() -> None:
    academic = select_presentation_intent(
        JobBrief(goal="研究展示"), {}, request_text="这是毕业论文答辩，请说明研究方法与证据。"
    )
    assert academic.presentation_form == PresentationForm.ACADEMIC_DEFENSE


def test_negative_style_request_does_not_select_that_form() -> None:
    intent = select_presentation_intent(
        JobBrief(goal="业务说明"), {}, request_text="这是正式业务说明，不要华丽品牌风，也不要庆典式表达。"
    )
    assert intent.presentation_form == PresentationForm.GENERAL


def test_government_and_portfolio_forms_do_not_fall_back_to_consulting_reference() -> None:
    government = JobBrief(goal="防汛工作", audience_profile="leadership")
    government.presentation_intent = select_presentation_intent(
        government, {}, request_text="向政府领导进行正式防汛工作汇报。"
    )
    portfolio = JobBrief(goal="数据分析案例", audience_profile="leadership")
    portfolio.presentation_intent = select_presentation_intent(
        portfolio, {}, request_text="制作正式华丽的高端品牌作品集。"
    )
    assert _style_selection(government) == (None, "presentation_intent:government_briefing")
    assert _style_selection(portfolio) == (None, "presentation_intent:premium_brand_portfolio")


def test_bound_work_request_guides_brief_without_becoming_slide_copy(tmp_path: Path, monkeypatch) -> None:
    work = _work(
        tmp_path,
        "需要制作10页护理培训PPT，面向护理员和家属，采用暖而正式的培训表达；不要增加目录页。",
        name="老年人跌倒预案",
    )
    _activate_work(monkeypatch, work)
    job = create_job("老年人跌倒预案", job_id="bound-work-request")
    ingest_sources(job.id, [str(_content(tmp_path))])

    brief = create_brief(job.id)
    plan = create_deck_plan(job.id)

    assert load_job(job.id).work_request is not None
    assert brief.presentation_intent.presentation_form == PresentationForm.TEACHING
    assert brief.presentation_intent.preservation.allow_auto_outline is False
    assert all(slide.layout != "summary" for slide in plan.slides)
    visible = "\n".join(piece for slide in plan.slides for piece in [slide.title, *slide.points, slide.speaker_notes])
    assert "不要增加目录页" not in visible
    assert "护理培训PPT" not in visible
    root = job_dir(job.id)
    snapshot = root / load_job(job.id).work_request.snapshot_path
    assert snapshot.is_file()
    decision = json.loads((root / "planning" / "intake-decision.json").read_text(encoding="utf-8"))
    assert decision["work_request"]["source_relative_path"] == "input/request.md"


def test_changed_work_request_invalidates_live_plan_and_preserves_old_snapshot(tmp_path: Path, monkeypatch) -> None:
    work = _work(tmp_path, "需要制作10页正式经营汇报，不要目录页。", name="经营复盘")
    _activate_work(monkeypatch, work)
    job = create_job("经营复盘", job_id="request-drift")
    ingest_sources(job.id, [str(_content(tmp_path))])
    assert run_job(job.id)["state"] == "brief_review"
    root = job_dir(job.id)
    first_snapshot = load_job(job.id).work_request.snapshot_path
    create_deck_plan(job.id)
    assert (root / "planning" / "deck-plan.json").is_file()

    (work / "input" / "request.md").write_text("需要制作10页政府正式防汛汇报，不能自动增加目录页。", encoding="utf-8")
    result = run_job(job.id)

    assert result["state"] == "brief_review"
    assert result["presentation_intent"]["presentation_form"] == PresentationForm.GOVERNMENT_BRIEFING.value
    assert not (root / "planning" / "deck-plan.json").exists()
    assert (root / first_snapshot).is_file()
    assert (root / load_job(job.id).work_request.snapshot_path).is_file()


def test_intent_override_reopens_brief_and_html_receives_form_hook(tmp_path: Path) -> None:
    job = create_job("品牌案例", job_id="intent-override")
    ingest_sources(job.id, [str(_content(tmp_path))])
    create_brief(job.id)
    create_deck_plan(job.id)
    root = job_dir(job.id)
    assert (root / "planning" / "deck-plan.json").is_file()

    configured = configure_presentation_intent(
        job.id,
        presentation_form=PresentationForm.PREMIUM_BRAND_PORTFOLIO,
        reason="本轮明确要求作品集展示。",
    )
    assert configured.presentation_intent.presentation_form == PresentationForm.PREMIUM_BRAND_PORTFOLIO
    assert load_job(job.id).state.value == "brief_review"
    assert not (root / "planning" / "deck-plan.json").exists()

    plan = create_deck_plan(job.id)
    assert plan.planning_strategy["presentation_form"] == PresentationForm.PREMIUM_BRAND_PORTFOLIO.value
    html = write_deck_html(job.id).read_text(encoding="utf-8")
    assert 'data-presentation-form="premium_brand_portfolio"' in html
    assert "presentation-form-premium_brand_portfolio" in html


def test_create_new_applies_selected_numeric_style_profile_to_html(tmp_path: Path) -> None:
    job = create_job("竞赛答辩样式", job_id="competition-style-profile")
    ingest_sources(job.id, [str(_content(tmp_path))])
    create_brief(job.id)
    configure_presentation_intent(
        job.id,
        change_mode=ChangeMode.NEW_COMPOSITION,
        presentation_form=PresentationForm.COMPETITION_DEFENSE,
        visual_system=VisualSystemMode.CREATE_NEW,
        reason="本轮明确新建竞赛答辩风格。",
    )
    plan = create_deck_plan(job.id)
    profile = plan.planning_strategy["visual_style_profile"]
    assert profile["profile_id"] == "competition-argument-v1"
    assert profile["application"] == "baseline_new_composition"
    html = write_deck_html(job.id).read_text(encoding="utf-8")
    assert "--page-title-size:68px" in html
    assert "--page-body-size:46px" in html
    assert "--page-image-left:0.92fr" in html
    assert '--page-heading-font:"Heiti SC","Avenir Next","Hiragino Sans GB"' in html
    plan.font_family = "Aptos"
    write_json(job_dir(job.id) / "planning/deck-plan.json", plan)
    explicit_font_html = write_deck_html(job.id).read_text(encoding="utf-8")
    assert "font-family:\"Aptos\"" in explicit_font_html
    assert "--page-heading-font:" not in explicit_font_html


def test_reference_led_form_keeps_style_profile_advisory(tmp_path: Path) -> None:
    job = create_job("保留参考风格", job_id="reference-style-profile")
    ingest_sources(job.id, [str(_content(tmp_path))])
    create_brief(job.id)
    configure_presentation_intent(
        job.id,
        change_mode=ChangeMode.NEW_COMPOSITION,
        presentation_form=PresentationForm.COMPETITION_DEFENSE,
        visual_system=VisualSystemMode.DERIVE_FROM_REFERENCE,
        reason="本轮需要遵循已经提供的视觉参考。",
    )
    plan = create_deck_plan(job.id)
    assert plan.planning_strategy["visual_style_profile"]["application"] == "advisory_only"
    html = write_deck_html(job.id).read_text(encoding="utf-8")
    assert "--page-title-size:68px" not in html
    assert "--page-heading-font:" not in html
    assert 'data-composition-baseline="new_composition"' not in html


def test_page_outline_is_not_a_local_patch_but_page_range_edit_is() -> None:
    outline = select_presentation_intent(
        JobBrief(goal="新建培训"),
        {},
        request_text="请新建一份 6 页培训演示：第1页说明问题，第2页讲方法，第3页给示例，第4页讲练习，第5页讲边界，第6页收束。",
    )
    edit = select_presentation_intent(
        JobBrief(goal="局部修改"),
        {"deck": InputRole.EDIT_TARGET},
        request_text="请修改第4至6页的表达，其余页面保持不动。",
    )
    assert requests_bounded_revision("请修改第4至6页的表达。") is True
    assert requests_bounded_revision("第4页改、其余不动。") is True
    assert requests_bounded_revision("请新建6页：第4页讲练习。") is False
    assert outline.change_mode == ChangeMode.NEW_COMPOSITION
    assert outline.preservation.target_slide_numbers == []
    assert edit.change_mode == ChangeMode.LOCAL_PATCH
    assert edit.preservation.target_slide_numbers == [4, 5, 6]


def test_bare_negative_instructions_do_not_become_template_or_style_requests() -> None:
    no_template = select_presentation_intent(
        JobBrief(goal="重新设计"),
        {"deck": InputRole.EDIT_TARGET},
        request_text="不使用模板，请完全重新设计内容结构。",
    )
    no_competition = select_presentation_intent(
        JobBrief(goal="正式汇报"),
        {},
        request_text="不竞赛风格，做正式经营汇报，不需要目录。",
    )
    assert requests_template_adaptation("不使用模板，请重新组织内容。") is False
    assert no_template.change_mode == ChangeMode.NEW_COMPOSITION
    assert no_competition.presentation_form == PresentationForm.BUSINESS_REPORT
    assert no_competition.preservation.allow_auto_outline is False


def test_library_template_application_wording_is_an_execution_requirement() -> None:
    request = "这是大学生社工项目创新设计竞赛，请根据模板库里面合适的模板进行套用并生成25页PPT。"
    assert requests_template_adaptation(request) is True
    assert requests_template_adaptation("不要套用模板，请根据内容重新设计。") is False
    intent = select_presentation_intent(JobBrief(goal="社工项目"), {"ref": InputRole.TEMPLATE}, request_text=request)
    assert intent.presentation_form == PresentationForm.COMPETITION_DEFENSE
    assert intent.change_mode == ChangeMode.TEMPLATE_RELAYOUT


def test_contest_application_document_identity_supplies_missing_occasion(tmp_path: Path) -> None:
    request = "根据模板库里面合适的模板进行套用并生成一套大概20到25页的PPT。"
    identity = [("source:entry:filename", "改-四川省大学生社工项目创新设计大赛申报书-星芒熠熠聚睦邻.docx")]
    inferred = select_presentation_intent(
        JobBrief(goal="社工项目"), {"entry": InputRole.CONTENT},
        request_text=request, source_identity=identity,
    )
    assert inferred.presentation_form == PresentationForm.COMPETITION_DEFENSE
    assert "source:entry:filename:presentation_form" in inferred.evidence
    explicit = select_presentation_intent(
        JobBrief(goal="社工项目"), {"entry": InputRole.CONTENT},
        request_text=request + "本轮给董事会做经营汇报。", source_identity=identity,
    )
    assert explicit.presentation_form == PresentationForm.BUSINESS_REPORT
    normalized = tmp_path / "normalized" / "entry.md"
    normalized.parent.mkdir()
    normalized.write_text("四川省大学生社工项目创新设计大赛申报书\n项目宗旨与服务计划\n竞赛材料中的普通业务段落", encoding="utf-8")
    record = SourceRecord(id="entry", original="generic.docx", stored_path="sources/generic.docx",
                          normalized_path="normalized/entry.md", kind="word", sha256="a" * 64, size_bytes=1)
    title_identity = source_document_identity(tmp_path, [record], {"entry": InputRole.CONTENT})
    from_title = select_presentation_intent(
        JobBrief(goal="社工项目"), {"entry": InputRole.CONTENT},
        request_text=request, source_identity=title_identity,
    )
    assert from_title.presentation_form == PresentationForm.COMPETITION_DEFENSE


def test_competition_template_work_rejects_service_proposal_and_native_downgrade(tmp_path: Path, monkeypatch) -> None:
    from aippt_studio import workflow

    library = tmp_path / "模板库"
    service_dir = library / "工作汇报"
    service_dir.mkdir(parents=True)
    service = _pptx(service_dir / "4055.pptx")
    contest_dir = library / "竞赛与答辩"
    contest_dir.mkdir(parents=True)
    contest = _pptx(contest_dir / "4215.pptx", slides=2)
    rows = [
        {"path": "工作汇报/4055.pptx", "asset_type": "template", "category": "工作汇报",
         "occasion_tags": ["工作汇报", "公共服务项目方案"],
         "source_sha256": hashlib.sha256(service.read_bytes()).hexdigest()},
        {"path": "竞赛与答辩/4215.pptx", "asset_type": "template", "category": "竞赛与答辩",
         "occasion_tags": ["竞赛", "比赛"],
         "source_sha256": hashlib.sha256(contest.read_bytes()).hexdigest()},
    ]
    (library / "CATALOG.jsonl").write_text(
        "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n", encoding="utf-8"
    )
    monkeypatch.setattr(workflow, "WORKSPACE_ROOT", tmp_path)
    request = "四川省大学生社工项目创新设计竞赛，根据模板库里面合适的模板进行套用并生成25页PPT。"
    work = _work(tmp_path, request, work_id="94")
    shutil.copy2(_content(tmp_path), work / "input" / "source" / "facts.md")
    shutil.copy2(service, work / "input" / "reference" / service.name)
    _activate_work(monkeypatch, work)
    payload = initialize_work_intake(work, job_id="competition-template-mismatch")
    job = load_job("competition-template-mismatch")
    assert payload["work_intake"]["role_inference"]["status"] == "library_template_confirmation_required"
    assert [item["path"] for item in payload["work_intake"]["template_selection"]["candidates"]] == [
        "竞赛与答辩/4215.pptx"
    ]
    assert job.route.value != "template"
    assert payload["presentation_intent"]["presentation_form"] == "competition_defense"
    with pytest.raises(WorkflowError, match="套用模板"):
        approve_gate(job.id, "brief")
    with pytest.raises(WorkflowError, match="套用模板"):
        configure_route(job.id, route=workflow.ProductionRoute.NATIVE, reason="只作为参考")
    with pytest.raises(WorkflowError, match="套用模板"):
        configure_presentation_intent(job.id, change_mode=ChangeMode.NEW_COMPOSITION, reason="改成新稿")
    with pytest.raises(WorkflowError, match="does not match competition_defense"):
        bind_catalog_template(job.id, "工作汇报/4055.pptx")
    from typer.testing import CliRunner
    from aippt_studio.cli import app

    review_result = CliRunner().invoke(
        app, ["library-template-review", job.id, "--path", "竞赛与答辩/4215.pptx", "--json"]
    )
    assert review_result.exit_code == 0, review_result.output
    review = json.loads(review_result.output)
    assert review["candidate_packet_kind"] == "brief_template_selection"
    source_previews = []
    for page in (1, 2):
        image = work / "workspace" / f"contest-page-{page}.png"
        Image.new("RGB", (320, 200), (24 + page, 76, 130)).save(image)
        render_receipt = image.with_suffix(".preview.json")
        receipt_data = {
            "schema_version": 1, "status": "rendered_advisory",
            "source_sha256": rows[1]["source_sha256"], "slide_number": page,
            "approximate_renderer": True,
            "limitations": ["Approximate render only; not Office native review."],
            "output_png": str(image),
            "output_sha256": hashlib.sha256(image.read_bytes()).hexdigest(),
        }
        render_receipt.write_text(json.dumps(receipt_data), encoding="utf-8")
        source_previews.append({
            "source_slide": page, "work_local_path": str(image),
            "sha256": hashlib.sha256(image.read_bytes()).hexdigest(),
            "render_receipt_path": str(render_receipt),
            "render_receipt_sha256": hashlib.sha256(render_receipt.read_bytes()).hexdigest(),
            "observation": f"竞赛模板第{page}页的具体结构观察。",
        })
    confirmation_request = {
        "selected_catalog_path": "竞赛与答辩/4215.pptx",
        "source_sha256": rows[1]["source_sha256"],
        "source_previews": source_previews,
        "user_confirmation_text": "我选用展示的竞赛与答辩模板，按当前项目内容替换文图。",
        "adaptation_scope": {
            "preserve_composition_theme": "保留竞赛答辩的主题语言和核心构图。",
            "text_image_replacement": "替换原模板中的旧文案和图片。",
            "allowed_layout_adjustments": "仅允许适配新文案的局部间距微调。",
        },
    }
    request_path = work / "workspace" / "template-confirmation.json"
    request_path.write_text(json.dumps(confirmation_request, ensure_ascii=False), encoding="utf-8")
    confirmed = CliRunner().invoke(
        app, ["library-template-confirm", job.id, "--request", str(request_path), "--json"]
    )
    assert confirmed.exit_code == 0, confirmed.output
    assert json.loads(confirmed.output)["status"] == "confirmation_recorded"
    bound = CliRunner().invoke(
        app, ["library-template", job.id, "--path", "竞赛与答辩/4215.pptx", "--json"]
    )
    assert bound.exit_code == 0, bound.output
    binding = json.loads(bound.output)
    assert binding["status"] == "template_bound_for_application"
    assert binding["source_sha256"] == rows[1]["source_sha256"]
    rebound = load_job(job.id)
    assert rebound.input_roles[binding["source_id"]] == InputRole.TEMPLATE
    assert list(rebound.input_roles.values()).count(InputRole.TEMPLATE) == 1
    assert InputRole.VISUAL_REFERENCE in rebound.input_roles.values()
    assert approve_gate(job.id, "brief").state.value == "full_build"

    suitable = _work(tmp_path, request, work_id="95")
    shutil.copy2(_content(tmp_path, count=30), suitable / "input" / "source" / "facts.md")
    shutil.copy2(contest, suitable / "input" / "reference" / contest.name)
    _activate_work(monkeypatch, suitable)
    initialize_work_intake(suitable, job_id="competition-template-compatible")
    suitable_previews = []
    for page in (1, 2):
        image = suitable / "workspace" / f"compatible-page-{page}.png"
        Image.new("RGB", (320, 200), (25 + page, 77, 131)).save(image)
        render_receipt = image.with_suffix(".preview.json")
        render_receipt.write_text(json.dumps({
            "schema_version": 1, "status": "rendered_advisory",
            "source_sha256": rows[1]["source_sha256"], "slide_number": page,
            "approximate_renderer": True,
            "limitations": ["Approximate render only; not Office native review."],
            "output_png": str(image),
            "output_sha256": hashlib.sha256(image.read_bytes()).hexdigest(),
        }), encoding="utf-8")
        suitable_previews.append({
            "source_slide": page, "work_local_path": str(image),
            "sha256": hashlib.sha256(image.read_bytes()).hexdigest(),
            "render_receipt_path": str(render_receipt),
            "render_receipt_sha256": hashlib.sha256(render_receipt.read_bytes()).hexdigest(),
            "observation": f"竞赛模板第{page}页适配观察。",
        })
    confirm_catalog_template("competition-template-compatible", {
        "selected_catalog_path": "竞赛与答辩/4215.pptx",
        "source_sha256": rows[1]["source_sha256"],
        "source_previews": suitable_previews,
        "user_confirmation_text": "我选用展示的竞赛与答辩模板，按当前项目内容替换文图。",
        "adaptation_scope": {
            "preserve_composition_theme": "保留竞赛答辩的主题语言和核心构图。",
            "text_image_replacement": "替换原模板中的旧文案和图片。",
            "allowed_layout_adjustments": "仅允许适配新文案的局部间距微调。",
        },
    })
    bind_catalog_template("competition-template-compatible", "竞赛与答辩/4215.pptx")
    accepted = approve_gate("competition-template-compatible", "brief")
    assert accepted.state.value == "full_build"
    accepted_root = job_dir(accepted.id)
    create_deck_plan(accepted.id)
    accepted_plan = read_model(accepted_root / "planning/deck-plan.json", DeckPlan)
    assert "source-bound" in _template_application_proof_reason(
        accepted, accepted_root, accepted_plan, {}, "a" * 64,
    )


def test_unclassified_override_and_scoped_contract_cannot_generate_a_generic_plan(tmp_path: Path) -> None:
    job = create_job("未分类安全门", job_id="unclassified-safety")
    ingest_sources(job.id, [str(_content(tmp_path))])
    create_brief(job.id)
    configure_presentation_intent(
        job.id,
        change_mode=ChangeMode.UNCLASSIFIED,
        reason="尚未确认这轮究竟是局部修改还是新建。",
    )
    with pytest.raises(ValueError, match="仍未分类"):
        create_deck_plan(job.id)
    with pytest.raises(WorkflowError, match="仍未分类"):
        approve_gate(job.id, "brief")

    scoped = select_presentation_intent(
        JobBrief(goal="范围编辑"),
        {"deck": InputRole.EDIT_TARGET},
        request_text="仅改第4页，其他页不动。",
    )
    with pytest.raises(PresentationIntentError, match="scoped"):
        from aippt_studio.presentation_intent import apply_intent_overrides

        apply_intent_overrides(scoped, change_mode=ChangeMode.NEW_COMPOSITION, reason="错误的覆盖")


def test_work_intake_uses_real_roles_and_never_leaks_reference_text(tmp_path: Path, monkeypatch) -> None:
    local_work = _work(tmp_path, "仅改第4页，其他页不动；保留原稿主题。", name="局部原稿", work_id="91")
    _pptx(local_work / "input" / "source" / "original.pptx", slides=4)
    _activate_work(monkeypatch, local_work)
    local_payload = initialize_work_intake(local_work, job_id="work-local-intent")
    local_job = load_job("work-local-intent")
    assert local_payload["work_intake"]["role_inference"]["status"] == "inferred_unique_edit_target"
    assert list(local_job.input_roles.values()).count(InputRole.EDIT_TARGET) == 1
    assert local_job.route.value == "template"
    assert local_payload["presentation_intent"]["change_mode"] == ChangeMode.LOCAL_PATCH.value
    with pytest.raises(WorkflowError, match="局部修订"):
        approve_gate(local_job.id, "brief")

    content_work = _work(tmp_path, "新建 6 页正式说明，不要目录。", name="参考隔离", work_id="92")
    (content_work / "input" / "source" / "facts.md").write_text(
        "\n".join(
            [
                "真实事实：本季度完成三项客户现场验证，并形成可复核记录。",
                "真实事实：下一阶段需要完成两项整改并复盘责任分工。",
                "真实事实：现场抽检显示关键节点的记录完整率持续提升。",
                "真实事实：后续安排将以客户反馈和复核结果为依据。",
            ]
        ),
        encoding="utf-8",
    )
    (content_work / "input" / "reference" / "style.md").write_text(
        "INTERNAL_STYLE_TOKEN 不得作为页内文案；这是参考模板的内部占位说明。",
        encoding="utf-8",
    )
    _activate_work(monkeypatch, content_work)
    reference_payload = initialize_work_intake(content_work, job_id="work-reference-intent")
    reference_job = load_job("work-reference-intent")
    brief = create_brief(reference_job.id)
    plan = create_deck_plan(reference_job.id)
    assert reference_payload["work_intake"]["imported_reference_count"] == 1
    assert InputRole.VISUAL_REFERENCE in reference_job.input_roles.values()
    assert "INTERNAL_STYLE_TOKEN" not in "\n".join(brief.must_include)
    visible = "\n".join(piece for slide in plan.slides for piece in [slide.title, *slide.points, slide.speaker_notes])
    assert "INTERNAL_STYLE_TOKEN" not in visible


def test_work_intake_exposes_multiple_edit_targets_without_guessing(tmp_path: Path, monkeypatch) -> None:
    work = _work(tmp_path, "仅改第4页，其他页不动。", name="多原稿", work_id="93")
    _pptx(work / "input" / "source" / "a.pptx", slides=4)
    _pptx(work / "input" / "source" / "b.pptx", slides=4)
    _activate_work(monkeypatch, work)
    payload = initialize_work_intake(work, job_id="work-ambiguous-intent")
    job = load_job("work-ambiguous-intent")
    assert payload["work_intake"]["role_inference"]["status"] == "needs_edit_target_designation"
    assert all(role != InputRole.EDIT_TARGET for role in job.input_roles.values())
    assert payload["presentation_intent"]["change_mode"] == ChangeMode.UNCLASSIFIED.value
    with pytest.raises(WorkflowError, match="仍未分类"):
        approve_gate(job.id, "brief")


def test_wrong_work_wrapper_is_rejected_without_rebinding_or_mutating_job(tmp_path: Path, monkeypatch) -> None:
    first = _work(tmp_path, "新建6页正式汇报。", name="第一工作区", work_id="94")
    _activate_work(monkeypatch, first)
    job = create_job("第一工作区", job_id="work-wrapper-safe")
    ingest_sources(job.id, [str(_content(tmp_path))])
    assert run_job(job.id)["state"] == "brief_review"
    before = (job_dir(job.id) / "job.json").read_bytes()

    second = _work(tmp_path, "新建6页正式汇报。", name="第二工作区", work_id="95")
    monkeypatch.setenv("AIPPT_WORK_DIR", str(second))
    # Keep the first Work's job root deliberately: this reproduces a stale
    # wrapper without letting the call attach a second Work's request.
    with pytest.raises(WorkflowError, match="wrapper is inconsistent"):
        run_job(job.id)
    assert (job_dir(job.id) / "job.json").read_bytes() == before


def test_work_intake_refuses_copied_or_mismatched_work_metadata(tmp_path: Path, monkeypatch) -> None:
    work = _work(tmp_path, "新建6页正式汇报。", work_id="96")
    (work / "work.json").write_text(json.dumps({"work_id": "95", "name": "错误复制的 Work"}, ensure_ascii=False), encoding="utf-8")
    _activate_work(monkeypatch, work)
    with pytest.raises(WorkBindingEnvironmentError, match="physical work"):
        create_job("mismatched-work-identity")
    with pytest.raises(WorkflowError, match="identity is inconsistent"):
        initialize_work_intake(work, job_id="mismatched-work-intake")


def test_visual_tokens_follow_explicit_direction_and_are_rendered(tmp_path: Path) -> None:
    blue_source = tmp_path / "blue.md"
    blue_source.write_text(
        "风格要求：深蓝正式。\n" + "\n".join(
            f"真实事实：第{i}项客户现场验证已完成，并保留可复核记录。" for i in range(1, 11)
        ),
        encoding="utf-8",
    )
    red_source = tmp_path / "red.md"
    red_source.write_text(
        "风格要求：红金庄重。\n" + "\n".join(
            f"真实事实：第{i}项客户现场验证已完成，并保留可复核记录。" for i in range(1, 11)
        ),
        encoding="utf-8",
    )
    blue_job = create_job("蓝色方向", job_id="visual-blue")
    red_job = create_job("红色方向", job_id="visual-red")
    ingest_sources(blue_job.id, [str(blue_source)])
    ingest_sources(red_job.id, [str(red_source)])
    blue_plan = create_deck_plan(blue_job.id) if create_brief(blue_job.id) else None
    red_plan = create_deck_plan(red_job.id) if create_brief(red_job.id) else None
    assert blue_plan is not None and red_plan is not None
    blue_visual = blue_plan.planning_strategy["visual_system"]
    red_visual = red_plan.planning_strategy["visual_system"]
    assert blue_visual["source"] == "explicit_visual_direction"
    assert red_visual["source"] == "explicit_visual_direction"
    assert blue_visual["tokens"]["cover"] != red_visual["tokens"]["cover"]
    blue_html = write_deck_html(blue_job.id).read_text(encoding="utf-8")
    red_html = write_deck_html(red_job.id).read_text(encoding="utf-8")
    assert blue_visual["tokens"]["cover"] in blue_html
    assert red_visual["tokens"]["cover"] in red_html
    assert blue_html != red_html


def test_intent_change_supersedes_live_delivery_without_destroying_its_evidence(tmp_path: Path) -> None:
    job = create_job("旧稿失效", job_id="delivery-invalidation")
    ingest_sources(job.id, [str(_content(tmp_path))])
    create_brief(job.id)
    root = job_dir(job.id)
    delivery = root / "deliverables"
    (delivery / "candidate_editable.pptx").write_bytes(b"old candidate bytes")
    (delivery / "build-manifest.json").write_text('{"status":"candidate"}', encoding="utf-8")

    configure_presentation_intent(
        job.id,
        presentation_form=PresentationForm.COMPETITION_DEFENSE,
        reason="本轮已确认是竞赛答辩。",
    )

    assert not any(path.is_file() for path in delivery.rglob("*"))
    archived = list((root / "history" / "superseded-deliverables").rglob("candidate_editable.pptx"))
    assert len(archived) == 1
    status = run_job(job.id)
    assert status["artifacts"]["delivery"] is None
    assert audit_job(job.id).status == "consistent"
