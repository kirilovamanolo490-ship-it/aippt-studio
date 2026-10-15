from __future__ import annotations

import hashlib
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from aippt_studio.style_hybrid_workflow import (
    StyleHybridWorkflowError, _validate_image_intents, build_hybrid_candidate, plan_hybrid_adaptation,
)
from aippt_studio.style_library import inspect_style_deck


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows), encoding="utf-8")


def _deck(path: Path, titles: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    for title in titles:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(.7)).text = title
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(.7)).text = "证据与方法占位"
    presentation.save(path)


def _receipt(tmp_path: Path, source: Path, slide: int, label: str) -> dict:
    image = tmp_path / f"{label}.png"
    Image.new("RGB", (20, 12), "blue").save(image)
    receipt = tmp_path / f"{label}.preview.json"
    receipt.write_text(json.dumps({"status": "rendered_advisory", "source_sha256": _sha(source),
                                   "slide_number": slide, "output_png": str(image),
                                   "output_sha256": _sha(image), "approximate_renderer": True}), encoding="utf-8")
    return {"slide": slide, "finding": "原稿页面的标题、内容区域和阅读顺序已逐项观察。",
            "preview_receipt": str(receipt)}


def _library(tmp_path: Path) -> tuple[Path, str, str, Path, Path]:
    library = tmp_path / "模板库"
    primary_rel = "创赛/静态模板/主模板.pptx"
    donor_rel = "创赛/静态模板/候选版式.pptx"
    primary, donor = library / primary_rel, library / donor_rel
    _deck(primary, ["项目封面", "实践成果"])
    _deck(donor, ["实践成果替代版式"])
    analysis = library / "创赛/分析索引"
    catalog, pages, curation = [], [], []
    for source, rel, titles in ((primary, primary_rel, ["项目封面", "实践成果"]),
                                (donor, donor_rel, ["实践成果替代版式"])):
        digest = _sha(source)
        detail_ref = f"创赛/分析索引/逐份结构/{digest}.json"
        detail = library / detail_ref
        detail.parent.mkdir(parents=True, exist_ok=True)
        detail.write_text(json.dumps({"source_sha256": digest,
                                      "presentation_size": {"width_emu": 12192000,
                                                            "height_emu": 6858000}}), encoding="utf-8")
        catalog.append({"path": rel, "source_sha256": digest, "status": "parsed",
                        "source_format": "pptx", "slide_count": len(titles), "detail_ref": detail_ref,
                        "style_profile": {"theme_palette": {"accent1": "4472C4"},
                                          "page_aspect_ratio": 16/9}})
        for number, title in enumerate(titles, 1):
            pages.append({"source_sha256": digest, "slide": number, "title": title,
                          "text_excerpt": title + " 项目实践成果证据", "role_evidence": [
                              {"role": "cover" if number == 1 and source == primary else "evidence_results"}],
                          "composition": {"text_object_count": 2, "object_counts": {"text_shape": 2}},
                          "layout_pattern": {"kind": "two_column_evidence"}})
        curation.append({"path": "模板库/" + rel, "sha256": digest,
                         "material_type": "模板/框架（占位仍显著）", "category": "创新创业"})
    _jsonl(analysis / "STYLE-CATALOG.jsonl", catalog)
    _jsonl(analysis / "STYLE-PAGES.jsonl", pages)
    _jsonl(analysis / "creativecup-curation.jsonl", curation)
    return library, primary_rel, donor_rel, primary, donor


def test_build_whole_template_with_hash_bound_text_edit(tmp_path: Path) -> None:
    library, primary_rel, _, primary, _ = _library(tmp_path)
    request = {"primary_source_path": primary_rel, "primary_source_sha256": _sha(primary),
               "reviewed_primary_pages": [_receipt(tmp_path, primary, 1, "primary-cover"),
                                          _receipt(tmp_path, primary, 2, "primary-content")],
               "selection_note": "完整保留原稿封面与正文的静态几何框架并替换当前项目文字。",
               "source_content_replacement_plan": "逐页清理原项目名称、示例事实与占位信息，本轮先验证封面目标。",
               "rights_handling_note": "本地测试稿不进入交付，真实图片使用前逐张核实来源和授权。",
               "text_replacements": [{"slide": 1, "shape_id": "2",
                                      "original_text_sha256": hashlib.sha256("项目封面".encode()).hexdigest(),
                                      "new_text": "新的竞赛项目封面"}],
               "geometry_edits": [{"action": "set", "slide": 2, "shape_id": "2",
                                   "geometry": {"unit": "mm", "x": 30}}]}
    output = tmp_path / "whole.pptx"
    with pytest.raises(StyleHybridWorkflowError, match="differs from requested page count"):
        build_hybrid_candidate(library, {**request, "project": {"page_range": [3, 3]}},
                               tmp_path / "wrong-count.pptx")
    assert not (tmp_path / "wrong-count.pptx").exists()
    result = build_hybrid_candidate(library, request, output)
    assert result["slide_count"] == 2
    assert result["page_overrides"] == []
    assert result["content_replacement_receipt"]["text_replacements"] == 1
    assert result["geometry_adjustment_receipt"]["status"] == "candidate"
    assert "新的竞赛项目封面" in Presentation(output).slides[0].shapes[0].text
    assert abs(Presentation(output).slides[1].shapes[0].left / 36000 - 30) < .01
    assert result["visual_review_pages"] == [1, 2]


def test_plan_and_build_one_local_layout_swap(tmp_path: Path) -> None:
    library, primary_rel, donor_rel, primary, donor = _library(tmp_path)
    plan = plan_hybrid_adaptation(library, {
        "project": {"occasion": "创新创业竞赛", "topic": "项目实践成果", "page_range": [2, 2]},
        "primary_source_path": primary_rel, "primary_source_sha256": _sha(primary),
        "planned_pages": [{"message": "项目主张", "role": "cover"},
                           {"message": "项目成果证据", "role": "evidence_results"}],
        "pages_to_relayout": [{"target_slide": 2, "message": "展示项目实践成果的证据",
                               "role": "evidence_results"}],
    })
    assert plan["primary_slide_count"] == 2
    assert plan["primary_page_count_fit"] == "within_range"
    assert plan["planned_page_roles"] == ["cover", "evidence_results"]
    assert plan["pages_to_relayout"][0]["target_slide"] == 2
    request = {"primary_source_path": primary_rel, "primary_source_sha256": _sha(primary),
               "reviewed_primary_pages": [_receipt(tmp_path, primary, 1, "cover"),
                                          _receipt(tmp_path, primary, 2, "body")],
               "selection_note": "保留整套竞赛模板，只调整第二页的项目成果证据排版。",
               "source_content_replacement_plan": "替换供体页的原项目文案，并核对主模板未改页的旧事实。",
               "rights_handling_note": "本地测试不交付，供体和主模板素材仅用于结构验证。",
               "page_overrides": [{"target_slide": 2, "donor_source_path": donor_rel,
                                   "donor_source_sha256": _sha(donor), "donor_slide": 1,
                                   "target_page": {"message": "项目实践成果的证据与方法",
                                                   "role": "evidence_results"},
                                   "selection_note": "供体页以双栏方法和成果组织同类内容，适合第二页局部换版。",
                                   "reviewed_donor_page": _receipt(tmp_path, donor, 1, "donor")}],
    }
    output = tmp_path / "hybrid.pptx"
    result = build_hybrid_candidate(library, request, output)
    assert result["slide_count"] == 2
    assert result["page_origin_map"][1]["source_path"] == donor_rel
    assert "实践成果替代版式" in Presentation(output).slides[1].shapes[0].text
    assert "项目封面" in Presentation(output).slides[0].shapes[0].text
    assert result["visual_review_pages"] == [1, 2]


def test_portrait_role_rejects_a_full_page_background_picture(tmp_path: Path) -> None:
    image = tmp_path / "background.png"
    Image.new("RGB", (40, 30), "blue").save(image)
    deck = tmp_path / "background.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture = slide.shapes.add_picture(str(image), 0, 0,
                                       width=presentation.slide_width,
                                       height=presentation.slide_height)
    presentation.save(deck)
    with pytest.raises(StyleHybridWorkflowError, match="portrait target covers too much"):
        _validate_image_intents(deck, [{"slide": 1, "shape_id": str(picture.shape_id),
                                        "original_media_sha256": _sha(image),
                                        "purpose": "替换人物头像，保持背景不动。", "visual_role": "portrait"}])


def test_case_page_layout_requires_every_old_text_disposition(tmp_path: Path) -> None:
    library, primary_rel, _, primary, _ = _library(tmp_path)
    donor_rel = "创赛/案例/已填写项目案例.pptx"
    donor = library / donor_rel
    _deck(donor, ["旧项目成果"])
    digest = _sha(donor)
    detail_ref = f"创赛/分析索引/逐份结构/{digest}.json"
    detail = library / detail_ref
    detail.parent.mkdir(parents=True, exist_ok=True)
    detail.write_text(json.dumps(inspect_style_deck(donor, donor_rel), ensure_ascii=False), encoding="utf-8")
    analysis = library / "创赛/分析索引"
    with (analysis / "STYLE-CATALOG.jsonl").open("a", encoding="utf-8") as stream:
        stream.write(json.dumps({"path": donor_rel, "source_sha256": digest,
                                 "source_format": "pptx", "status": "parsed", "slide_count": 1,
                                 "detail_ref": detail_ref, "detail_sha256": _sha(detail),
                                 "style_profile": {}}, ensure_ascii=False) + "\n")
    with (analysis / "STYLE-PAGES.jsonl").open("a", encoding="utf-8") as stream:
        stream.write(json.dumps({"source_sha256": digest, "slide": 1, "title": "旧项目成果",
                                 "text_excerpt": "旧项目成果 证据与方法占位", "role_evidence": [
                                     {"role": "evidence_results"}],
                                 "composition": {"text_object_count": 2,
                                                 "object_counts": {"text_shape": 2}}}, ensure_ascii=False) + "\n")
    with (analysis / "creativecup-curation.jsonl").open("a", encoding="utf-8") as stream:
        stream.write(json.dumps({"path": "模板库/" + donor_rel, "sha256": digest,
                                 "material_type": "已填写项目稿/参赛案例候选"}, ensure_ascii=False) + "\n")
    layer_reviews = [{"part": part,
                      "finding": "已检查来源母版与版式的文字及图像层，当前只借版式。",
                      "rights_handling_note": "此测试源的继承层仅用于本地结构验证，不作客户内容或授权结论。"}
                     for part, layer in json.loads(detail.read_text())["shared_layers"].items()
                     if any(str(obj.get("text") or "").strip() or obj.get("image_fills")
                            for obj in layer.get("objects", []))]
    common = {"primary_source_path": primary_rel, "primary_source_sha256": _sha(primary),
              "reviewed_primary_pages": [_receipt(tmp_path, primary, 1, "base-cover"),
                                         _receipt(tmp_path, primary, 2, "base-evidence")],
              "selection_note": "完整保留原竞赛模板，仅从完成案例借其成果页版式。",
              "source_content_replacement_plan": "案例页每条旧项目事实须改写为当前项目证据。",
              "rights_handling_note": "案例文案只作版式参考，不带入客户演示内容。",
              "page_overrides": [{"target_slide": 2, "donor_source_path": donor_rel,
                                  "donor_source_sha256": digest, "donor_slide": 1,
                                  "target_page": {"message": "展示项目成果证据与实施方法", "role": "evidence_results",
                                                  "required_evidence": "项目成果与方法证据"},
                                  "selection_note": "案例页有两条可编辑文字，结构适合承载当前项目成果。",
                                  "reviewed_donor_page": _receipt(tmp_path, donor, 1, "case-page"),
                                  "shared_layer_reviews": layer_reviews,
                                  "layout_only_reference": True}]}
    text_edits = [{"slide": 2, "shape_id": str(shape.shape_id),
                   "original_text_sha256": hashlib.sha256(shape.text.encode()).hexdigest(),
                   "new_text": "当前项目成果" if index == 0 else "当前证据摘要"}
                  for index, shape in enumerate(Presentation(donor).slides[0].shapes)]
    with pytest.raises(StyleHybridWorkflowError, match="disposition for every source text"):
        build_hybrid_candidate(library, {**common, "text_replacements": text_edits[:1]},
                               tmp_path / "incomplete.pptx")
    assert not (tmp_path / "incomplete.pptx").exists()
    bad_role = json.loads(json.dumps(common))
    bad_role["page_overrides"][0]["target_page"]["role"] = "self_assessment"
    with pytest.raises(StyleHybridWorkflowError, match="same page role"):
        build_hybrid_candidate(library, {**bad_role, "text_replacements": text_edits},
                               tmp_path / "wrong-role.pptx")
    original_detail = detail.read_text(encoding="utf-8")
    payload = json.loads(original_detail)
    layer = next(iter(payload["shared_layers"].values()))
    layer["objects"].append({"id": "999", "name": "Inherited old logo", "image_fills": [{"media_part": "ppt/media/old-logo.png"}]})
    detail.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    catalog_rows = [json.loads(line) for line in (analysis / "STYLE-CATALOG.jsonl").read_text().splitlines()]
    catalog_rows[-1]["detail_sha256"] = _sha(detail)
    _jsonl(analysis / "STYLE-CATALOG.jsonl", catalog_rows)
    with pytest.raises(StyleHybridWorkflowError, match="inherited media"):
        build_hybrid_candidate(library, {**common, "text_replacements": text_edits},
                               tmp_path / "old-logo.pptx")
    detail.write_text(original_detail, encoding="utf-8")
    catalog_rows[-1]["detail_sha256"] = _sha(detail)
    _jsonl(analysis / "STYLE-CATALOG.jsonl", catalog_rows)
    incomplete = json.loads(original_detail)
    incomplete["slides"][0]["objects"] = [obj for obj in incomplete["slides"][0]["objects"]
                                          if obj.get("id") != "3"]
    detail.write_text(json.dumps(incomplete, ensure_ascii=False), encoding="utf-8")
    catalog_rows[-1]["detail_sha256"] = _sha(detail)
    _jsonl(analysis / "STYLE-CATALOG.jsonl", catalog_rows)
    with pytest.raises(StyleHybridWorkflowError, match="detail omits"):
        build_hybrid_candidate(library, {**common, "text_replacements": text_edits},
                               tmp_path / "incomplete-detail.pptx")
    detail.write_text(original_detail, encoding="utf-8")
    catalog_rows[-1]["detail_sha256"] = _sha(detail)
    _jsonl(analysis / "STYLE-CATALOG.jsonl", catalog_rows)
    date_payload = json.loads(original_detail)
    date_item = next(obj for layer in date_payload["shared_layers"].values()
                     for obj in layer["objects"] if obj.get("placeholder_type") == "dt")
    date_item["text"] = "8/15/2025"
    detail.write_text(json.dumps(date_payload, ensure_ascii=False), encoding="utf-8")
    catalog_rows[-1]["detail_sha256"] = _sha(detail)
    _jsonl(analysis / "STYLE-CATALOG.jsonl", catalog_rows)
    with pytest.raises(StyleHybridWorkflowError, match="inherited case text"):
        build_hybrid_candidate(library, {**common, "text_replacements": text_edits},
                               tmp_path / "case-date.pptx")
    detail.write_text(original_detail, encoding="utf-8")
    catalog_rows[-1]["detail_sha256"] = _sha(detail)
    _jsonl(analysis / "STYLE-CATALOG.jsonl", catalog_rows)
    result = build_hybrid_candidate(library, {**common, "text_replacements": text_edits},
                                    tmp_path / "case-layout.pptx")
    assert result["page_overrides"][0]["reuse_mode"] == "layout_only_reference"
    assert result["page_overrides"][0]["reference_content_coverage"]["source_text_objects"] == 2
    assert result["reference_output_readback"][0]["verified_text_shape_ids"] == ["2", "3"]
    assert "当前项目成果" in Presentation(tmp_path / "case-layout.pptx").slides[1].shapes[0].text
