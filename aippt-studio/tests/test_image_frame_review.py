from aippt_studio.presentation_quality import image_frame_geometry_signals


def objects(image_geometry="rect", frame_geometry="roundRect", **picture_overrides):
    return [{"object_id": "2", "kind": "shape", "text": "", "geometry": frame_geometry,
             "bounds_pt": [30, 50, 250, 180], "z_order": 0},
            {"object_id": "3", "kind": "picture", "opacity": "opaque", "geometry": image_geometry,
             "bounds_pt": [30, 50, 250, 180], "z_order": 1, **picture_overrides}]


def test_unbound_rectangular_image_over_curved_frame_needs_specific_context():
    for geometry in ("roundRect", "ellipse"):
        signals = image_frame_geometry_signals(objects(frame_geometry=geometry))
        assert len(signals) == 1
        assert signals[0]["object_id"] == "3" and signals[0]["other_object_id"] == "2"
        assert signals[0]["kind"] == "context_required"


def test_real_clipping_and_independent_backdrops_are_not_assumed_broken():
    assert image_frame_geometry_signals(objects("roundRect")) == []
    assert image_frame_geometry_signals(objects(opacity="transparent_or_partial")) == []
    assert image_frame_geometry_signals(objects(bounds_pt=[50, 70, 180, 160])) == []
    assert image_frame_geometry_signals(objects(z_order=-1)) == []


def test_rotation_keeps_uncertainty_instead_of_claiming_a_visible_defect():
    result = image_frame_geometry_signals(objects(rotation_or_shear=True))
    assert result[0]["rotation_or_shear"] is True
    assert "actual image edges" in result[0]["review_question"]


def test_curved_backing_with_opaque_rectangular_photo_keeps_final_qa_open(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from aippt_studio.presentation_quality import audit_presentation

    photo = tmp_path / "photo.png"
    Image.new("RGB", (320, 180), "#405060").save(photo)
    for curve in (MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.OVAL):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        frame = slide.shapes.add_shape(curve, Inches(1), Inches(1), Inches(4), Inches(2))
        picture = slide.shapes.add_picture(str(photo), Inches(1), Inches(1), Inches(4), Inches(2))
        source = tmp_path / f"{curve}.pptx"
        deck.save(source)

        result = audit_presentation(source)

        signal = next(item for item in result["issues"] if item["code"] == "IMAGE_FRAME_CONTOUR_MISMATCH")
        assert result["status"] == "needs_review"
        assert signal["severity"] == "review"
        assert signal["object_id"] == str(picture.shape_id)
        assert signal["details"]["other_object_id"] == str(frame.shape_id)
        assert signal["details"]["rotation_or_shear"] is False


def test_separate_curved_decoration_does_not_hold_final_qa(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from aippt_studio.presentation_quality import audit_presentation

    photo = tmp_path / "photo.png"
    Image.new("RGB", (320, 180), "#405060").save(photo)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    slide.shapes.add_picture(str(photo), Inches(1.5), Inches(1.2), Inches(3), Inches(1.6))
    source = tmp_path / "separate.pptx"
    deck.save(source)

    result = audit_presentation(source)

    assert result["status"] == "passed"
    assert not any(item["code"] == "IMAGE_FRAME_CONTOUR_MISMATCH" for item in result["issues"])


def test_clipped_picture_does_not_claim_to_occlude_text_in_its_empty_corners():
    from io import BytesIO
    from types import SimpleNamespace
    from lxml import etree
    from PIL import Image
    from aippt_studio.presentation_quality import _opaque
    data = BytesIO()
    Image.new("RGB", (2, 2), "white").save(data, format="PNG")
    for preset, expected in (("rect", "opaque"), ("ellipse", "transparent_or_partial"), ("roundRect", "transparent_or_partial")):
        element = etree.fromstring(f'<p:pic xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:spPr><a:prstGeom prst="{preset}"/></p:spPr></p:pic>')
        shape = SimpleNamespace(shape_type=13, _element=element, image=SimpleNamespace(blob=data.getvalue()))
        assert _opaque(shape) == expected


def test_contain_fill_source_outset_and_native_alpha_are_not_full_surface_occlusion():
    from io import BytesIO
    from types import SimpleNamespace
    from lxml import etree
    from PIL import Image
    from aippt_studio.presentation_quality import _opaque
    data = BytesIO()
    Image.new("RGB", (2, 2), "white").save(data, format="PNG")
    cases = [
        ('<a:blip/><a:stretch><a:fillRect t="25000" b="25000"/></a:stretch>', "transparent_or_partial"),
        ('<a:blip/><a:srcRect l="-25000" r="-25000"/><a:stretch><a:fillRect/></a:stretch>', "transparent_or_partial"),
        ('<a:blip><a:alphaModFix amt="25000"/></a:blip><a:stretch><a:fillRect/></a:stretch>', "transparent_or_partial"),
        ('<a:blip><a:alphaModFix amt="0"/></a:blip>', "transparent_or_partial"),
        ('<a:blip><a:alphaModFix amt="100000"/></a:blip><a:stretch><a:fillRect l="0" t="0" r="0" b="0"/></a:stretch>', "opaque"),
        ('<a:blip/><a:srcRect l="25000" r="25000"/><a:stretch><a:fillRect/></a:stretch>', "opaque"),
        ('<a:blip/><a:stretch><a:fillRect t="unsupported"/></a:stretch>', "unknown"),
        ('<a:blip><a:alphaModFix amt="unsupported"/></a:blip>', "unknown"),
    ]
    for fill, expected in cases:
        xml = f'<p:pic xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:blipFill>{fill}</p:blipFill><p:spPr><a:prstGeom prst="rect"/></p:spPr></p:pic>'
        shape = SimpleNamespace(shape_type=13, _element=etree.fromstring(xml), image=SimpleNamespace(blob=data.getvalue()))
        assert _opaque(shape) == expected, fill
