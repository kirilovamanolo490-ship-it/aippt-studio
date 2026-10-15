from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import aippt_studio.local_preview_officecli as preview
from aippt_studio.local_preview_officecli import OfficeCliPreviewError, render_officecli_slide_preview


def _deck(path: Path) -> None:
    presentation = Presentation()
    for label in ("Target page", "Should not be rendered"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        shape.text = label
    presentation.save(path)


def test_uses_managed_officecli_and_writes_source_bound_receipt(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source, output, receipt_path = tmp_path / "deck.pptx", tmp_path / "one.png", tmp_path / "one.json"
    _deck(source)
    original = source.read_bytes()
    monkeypatch.setattr(preview, "PROJECT_ROOT", tmp_path)
    tool = tmp_path / ".tools/officecli"
    tool.parent.mkdir()
    tool.write_text("managed stub", encoding="utf-8")
    tool.chmod(0o755)
    commands: list[list[str]] = []

    def run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        commands.append(command)
        args = dict(zip(command, command[1:]))
        Image.new("RGB", (800, 450), "white").save(args["--out"])
        return subprocess.CompletedProcess(command, 0, json.dumps({"success": True, "data": args["--out"]}).encode(), b"")

    monkeypatch.setattr(preview.subprocess, "run", run)
    result = render_officecli_slide_preview(source, 1, output, receipt_path=receipt_path, width=800, height=450)

    assert commands[0][0] == str(tool)
    assert commands[0][1:] == [
        "view", str(source.resolve()), "screenshot", "--page", "1", "--render", "html",
        "--screenshot-width", "800", "--screenshot-height", "450", "--out", commands[0][commands[0].index("--out") + 1], "--json",
    ]
    assert result["status"] == "rendered_advisory"
    assert result["renderer"] == "officecli-html"
    assert result["slide_count"] == 2
    assert result["approximate_renderer"] is True
    assert result["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert result["output_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert json.loads(receipt_path.read_text(encoding="utf-8")) == result
    assert source.read_bytes() == original
    assert not list(tmp_path.glob("*.pdf"))


def test_infers_preview_height_from_actual_slide_ratio(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source, output = tmp_path / "deck.pptx", tmp_path / "one.png"
    _deck(source)
    monkeypatch.setattr(preview, "PROJECT_ROOT", tmp_path)
    tool = tmp_path / ".tools/officecli"
    tool.parent.mkdir()
    tool.write_text("managed stub", encoding="utf-8")
    tool.chmod(0o755)

    def run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        options = dict(zip(command, command[1:]))
        height = int(options["--screenshot-height"])
        Image.new("RGB", (800, height), "white").save(options["--out"])
        return subprocess.CompletedProcess(command, 0, b'{"success": true}', b"")

    monkeypatch.setattr(preview.subprocess, "run", run)
    result = render_officecli_slide_preview(source, 1, output, width=800)
    assert result["height"] == round(800 * Presentation(source).slide_height / Presentation(source).slide_width)


def test_rejects_out_of_range_slide_before_running_renderer(tmp_path: Path) -> None:
    source = tmp_path / "deck.pptx"
    _deck(source)
    with pytest.raises(OfficeCliPreviewError, match="exceeds the 2-slide"):
        render_officecli_slide_preview(source, 3, tmp_path / "missing.png")


@pytest.mark.parametrize("existing", ["png", "receipt"])
def test_refuses_to_overwrite_any_existing_output(tmp_path: Path, existing: str) -> None:
    source, output, receipt_path = tmp_path / "deck.pptx", tmp_path / "one.png", tmp_path / "one.json"
    _deck(source)
    occupied = output if existing == "png" else receipt_path
    occupied.write_bytes(b"keep")
    with pytest.raises(OfficeCliPreviewError, match="must not already exist"):
        render_officecli_slide_preview(source, 1, output, receipt_path=receipt_path)
    assert occupied.read_bytes() == b"keep"


@pytest.mark.parametrize("slide_number", [0, -1, True, 1.5])
def test_rejects_invalid_slide_number(tmp_path: Path, slide_number: object) -> None:
    with pytest.raises(OfficeCliPreviewError, match="positive 1-based integer"):
        render_officecli_slide_preview(tmp_path / "missing.pptx", slide_number, tmp_path / "one.png")  # type: ignore[arg-type]


def test_bad_officecli_json_does_not_leave_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source, output = tmp_path / "deck.pptx", tmp_path / "one.png"
    _deck(source)
    monkeypatch.setattr(preview, "PROJECT_ROOT", tmp_path)
    tool = tmp_path / ".tools/officecli"
    tool.parent.mkdir()
    tool.write_text("managed stub", encoding="utf-8")
    tool.chmod(0o755)
    monkeypatch.setattr(preview.subprocess, "run", lambda *args, **kwargs: subprocess.CompletedProcess([], 0, b"not json", b""))
    with pytest.raises(OfficeCliPreviewError, match="invalid JSON"):
        render_officecli_slide_preview(source, 1, output)
    assert not output.exists()


def test_officecli_failure_does_not_remove_competing_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source, output = tmp_path / "deck.pptx", tmp_path / "one.png"
    _deck(source)
    monkeypatch.setattr(preview, "PROJECT_ROOT", tmp_path)
    tool = tmp_path / ".tools/officecli"
    tool.parent.mkdir()
    tool.write_text("managed stub", encoding="utf-8")
    tool.chmod(0o755)

    def fail_after_competitor(*args: object, **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        output.write_bytes(b"someone else's output")
        return subprocess.CompletedProcess([], 1, b'{"success": false, "error": {"error": "render failed"}}', b"")

    monkeypatch.setattr(preview.subprocess, "run", fail_after_competitor)
    with pytest.raises(OfficeCliPreviewError, match="could not render"):
        render_officecli_slide_preview(source, 1, output)
    assert output.read_bytes() == b"someone else's output"


def test_officecli_refuses_dangling_output_symlink(tmp_path: Path) -> None:
    source, output = tmp_path / "deck.pptx", tmp_path / "one.png"
    _deck(source)
    output.symlink_to(tmp_path / "absent.png")
    with pytest.raises(OfficeCliPreviewError, match="must not already exist"):
        render_officecli_slide_preview(source, 1, output)
    assert output.is_symlink()


@pytest.mark.skipif(not (Path(__file__).parents[1] / ".tools/officecli").is_file(), reason="managed OfficeCLI unavailable")
def test_real_officecli_renders_only_one_slide(tmp_path: Path) -> None:
    source, output = tmp_path / "deck.pptx", tmp_path / "one.png"
    _deck(source)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    result = render_officecli_slide_preview(source, 1, output, width=800, height=450)
    assert result["source_sha256"] == original_hash
    assert result["slide_number"] == 1
    assert result["rendered_slide_count"] == 1
    assert Image.open(output).size == (800, 450)
    assert not list(tmp_path.glob("*.pdf"))
    assert not (tmp_path / "slide-2.png").exists()
