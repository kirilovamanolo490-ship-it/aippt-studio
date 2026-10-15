import hashlib
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET
from PIL import Image, ImageCms
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Mm, Pt
import pytest

from aippt_studio import pptx_edit as pe


def test_picture_reencoding_keeps_jpeg_color_profile(tmp_path):
    image_path = tmp_path / "replacement.jpg"
    profile = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    Image.new("RGB", (80, 60), (120, 45, 90)).save(image_path, "JPEG", icc_profile=profile)
    encoded = pe._replacement_raster(str(image_path), "ppt/media/image1.jpeg", image_path.read_bytes())
    with Image.open(BytesIO(encoded)) as image:
        assert image.format == "JPEG"
        assert image.info["icc_profile"] == profile


def make_deck(folder: Path) -> Path:
    """Real, editable two-slide deck with mixed Chinese/Latin runs and a photo."""
    folder.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Mm(254)
    presentation.slide_height = Mm(142.875)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Mm(10), Mm(15), Mm(68), Mm(30))
    shape.name = "明确选择的中文正文"
    shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    paragraph = shape.text_frame.paragraphs[0]
    first = paragraph.add_run()
    first.text = "中文正文："
    first.font.size = Pt(12)
    first.font.name = "Microsoft YaHei"
    first.font.bold = True
    first.font.color.rgb = RGBColor(32, 72, 128)
    ea = OxmlElement("a:ea")
    ea.set("typeface", "Microsoft YaHei")
    first._r.get_or_add_rPr().append(ea)
    second = paragraph.add_run()
    second.text = "保留多段格式 Native text"
    second.font.size = Pt(14)
    second.font.italic = True
    second.font.color.rgb = RGBColor(180, 72, 35)
    shape2 = slide.shapes.add_textbox(Mm(95), Mm(25), Mm(35), Mm(20))
    shape2.text = "第二对象\n保持内容"
    shape2.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    shape2.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    shape3 = slide.shapes.add_textbox(Mm(175), Mm(75), Mm(45), Mm(25))
    shape3.text = "第三对象"
    shape3.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    image = folder / "image.png"
    Image.new("RGB", (160, 100), (224, 236, 245)).save(image)
    picture = slide.shapes.add_picture(str(image), Mm(12), Mm(92), width=Mm(42))
    picture.name = "保留原图"
    untouched = presentation.slides.add_slide(presentation.slide_layouts[6])
    untouched.shapes.add_textbox(Mm(12), Mm(15), Mm(100), Mm(30)).text = "这一页及母版、图片与所有关系均保持字节不变"
    path = folder / "source.pptx"
    presentation.save(path)
    return path


@pytest.fixture
def source(tmp_path):
    return make_deck(tmp_path / "输入 中文")


def req(source, *edits, **options):
    return {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "edits": list(edits), **options}


def parts(path):
    with ZipFile(path) as archive:
        return {i.filename: archive.read(i) for i in archive.infolist()}


def object_at(path, shape_id="2", slide=1):
    report = pe.inspect_pptx(path, slides=[slide])
    return next(o for o in report["slides"][0]["objects"] if o["shape_id"] == shape_id)


def set_edit(**fields):
    return {"action": "set", "slide": 1, "shape_id": "2", **fields}


def test_inspect_lists_exact_ids_geometry_runs_and_slide_subset(source):
    report = pe.inspect_pptx(source, slides=[1])
    assert len(report["slides"]) == 1
    obj = report["slides"][0]["objects"][0]
    assert obj["shape_id"] == "2" and obj["type"] == "sp" and obj["editable"]
    assert obj["geometry"]["mm"]["x"] == 10
    assert obj["font"]["explicit_sizes_pt"] == [12, 14]
    assert obj["font"]["runs"][0]["fonts"]["ea"] == "Microsoft YaHei"
    assert obj["font"]["runs"][1]["properties"]["i"] == "1"
    assert report["source_sha256"] == req(source)["source_sha256"]
    assert report["slide_size_emu"]["width"] == 254 * 36000
    inherited = pe.inspect_pptx(source, slides=[2])["slides"][0]["objects"][0]
    assert inherited["font"]["runs"][0]["size_origin"] == "inherited_unresolved"
    with pytest.raises(ValueError, match="slide numbers"):
        pe.inspect_pptx(source, slides=[True])


def test_replace_exact_text_run_keeps_style_and_non_target_parts(source, tmp_path):
    original_parts = parts(source)
    original = object_at(source)
    edit = {"action": "replace_text", "slide": 1, "shape_id": "2",
            "paragraph": 1, "run": 2,
            "expected_text": "保留多段格式 Native text", "replacement_text": "改成正式的汇报文字"}
    output = tmp_path / "text-patch"
    result = pe.patch_pptx(req(source, edit), output)
    candidate = output / "candidate.pptx"
    changed = parts(candidate)
    after = object_at(candidate)
    assert after["text"] == "中文正文：改成正式的汇报文字"
    assert after["font"]["runs"][0] == original["font"]["runs"][0]
    assert after["font"]["runs"][1]["properties"] == original["font"]["runs"][1]["properties"]
    assert {name for name in original_parts if original_parts[name] != changed[name]} == {"ppt/slides/slide1.xml"}
    assert result["native_visual_verified"] is False


def test_replace_text_requires_exact_unambiguous_run_and_single_line(source, tmp_path):
    duplicate = Presentation(source)
    paragraph = duplicate.slides[0].shapes[0].text_frame.paragraphs[0]
    paragraph.add_run().text = "中文正文："
    duplicate.save(source)
    edit = {"action": "replace_text", "slide": 1, "shape_id": "2",
            "expected_text": "中文正文：", "replacement_text": "核心结论："}
    with pytest.raises(ValueError, match="exactly one run"):
        pe.patch_pptx(req(source, edit), tmp_path / "ambiguous")
    assert not (tmp_path / "ambiguous").exists()
    selected = pe.patch_pptx(req(source, edit | {"paragraph": 1, "run": 1}), tmp_path / "selected")
    assert selected["status"] == "candidate"
    assert object_at(tmp_path / "selected/candidate.pptx")["text"].count("中文正文：") == 1
    with pytest.raises(ValueError, match="single XML-safe"):
        pe.patch_pptx(req(source, edit | {"paragraph": 1, "run": 1, "replacement_text": "两行\n内容"}),
                      tmp_path / "multiline")


def test_add_textbox_preserves_every_existing_xml_object_and_package_member(source, tmp_path):
    before_parts = parts(source)
    before_objects = pe.inspect_pptx(source, slides=[1])["slides"][0]["objects"]
    edit = {"action": "add_textbox", "slide": 1, "text": "新增的可编辑文本",
            "geometry": {"unit": "mm", "x": 30, "y": 105, "width": 80, "height": 14},
            "font_size_pt": 18, "font_name": "Aptos"}
    output = tmp_path / "add-text"
    receipt = pe.patch_pptx(req(source, edit), output)
    candidate = output / "candidate.pptx"
    after_parts = parts(candidate)
    after_objects = pe.inspect_pptx(candidate, slides=[1])["slides"][0]["objects"]
    added = after_objects[-1]
    assert added["type"] == "sp" and added["editable"]
    assert added["text"] == "新增的可编辑文本"
    assert added["geometry"]["mm"] == {"x": 30, "y": 105, "width": 80, "height": 14}
    assert added["font"]["explicit_sizes_pt"] == [18]
    assert added["font"]["runs"][0]["fonts"]["latin"] == "Aptos"
    assert added["shape_id"] not in {obj["shape_id"] for obj in before_objects}
    assert after_objects[:-1] == before_objects
    assert {name for name in before_parts if before_parts[name] != after_parts[name]} == {"ppt/slides/slide1.xml"}
    original_xml = before_parts["ppt/slides/slide1.xml"]
    for start, end in pe._child_spans(original_xml):
        assert original_xml[start:end] in after_parts["ppt/slides/slide1.xml"]
    assert receipt["operations"][0]["before"] == []
    assert receipt["operations"][0]["after"][0]["text"] == "新增的可编辑文本"
    assert receipt["native_visual_verified"] is False
    deck = Presentation(candidate)
    assert deck.slides[0].shapes[-1].text == "新增的可编辑文本"


def test_add_textbox_preserves_intentional_edge_spaces(source, tmp_path):
    edit = {"action": "add_textbox", "slide": 1, "text": " 重点结论 ",
            "geometry": {"unit": "mm", "x": 30, "y": 105, "width": 80, "height": 14}}
    candidate = tmp_path / "spaced" / "candidate.pptx"
    pe.patch_pptx(req(source, edit), candidate.parent)
    with ZipFile(candidate) as package:
        root = pe._xml(package.read("ppt/slides/slide1.xml"))
    text_nodes = root.xpath(".//a:t[text()=' 重点结论 ']", namespaces=pe.NS)
    assert len(text_nodes) == 1
    assert text_nodes[0].get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
    assert Presentation(candidate).slides[0].shapes[-1].text == " 重点结论 "


@pytest.mark.parametrize("edit,match", [
    ({"action": "add_textbox", "slide": 1, "text": "x", "geometry": {"unit": "mm", "x": 1, "y": 1, "width": 10}}, "explicit x, y, width, height"),
    ({"action": "add_textbox", "slide": 1, "text": "two\nlines", "geometry": {"unit": "mm", "x": 1, "y": 1, "width": 10, "height": 10}}, "single-line"),
    ({"action": "add_textbox", "slide": 1, "text": "x", "geometry": {"unit": "mm", "x": 250, "y": 1, "width": 10, "height": 10}}, "slide bounds"),
    ({"action": "add_textbox", "slide": 1, "text": "x", "geometry": {"unit": "inch", "x": 1, "y": 1, "width": 100, "height": 100}}, "slide bounds"),
])
def test_add_textbox_rejects_incomplete_unsafe_or_out_of_bounds_requests(source, tmp_path, edit, match):
    original = source.read_bytes()
    output = tmp_path / "bad-add"
    with pytest.raises(ValueError, match=match):
        pe.patch_pptx(req(source, edit), output)
    assert not output.exists()
    assert source.read_bytes() == original


def test_move_resize_font_preserves_runs_unselected_objects_and_every_other_zip_member(source, tmp_path):
    original = source.read_bytes()
    original_parts = parts(source)
    before_root = pe._xml(original_parts["ppt/slides/slide1.xml"])
    output = tmp_path / "精控 候选"
    result = pe.patch_pptx(req(source, set_edit(geometry={"unit": "mm", "x": 18, "y": 22, "width": 108, "height": 42},
                                                  font_size_pt=24)), output)
    candidate = output / "candidate.pptx"
    candidate_parts = parts(candidate)
    assert source.read_bytes() == original
    assert {k for k in original_parts if original_parts[k] != candidate_parts[k]} == {"ppt/slides/slide1.xml"}
    assert result["changed_parts"] == [{"part": "ppt/slides/slide1.xml",
                                        "before_sha256": pe._sha(original_parts["ppt/slides/slide1.xml"]),
                                        "after_sha256": pe._sha(candidate_parts["ppt/slides/slide1.xml"])}]
    for node in before_root.findall("p:cSld/p:spTree/p:sp", pe.NS)[1:]:
        # Check lexical object bytes, not merely its XML semantics.
        original_raw = original_parts["ppt/slides/slide1.xml"]
        tree = before_root.find("p:cSld/p:spTree", pe.NS)
        index = list(tree).index(node)
        start, end = pe._child_spans(original_raw)[index]
        assert original_raw[start:end] in candidate_parts["ppt/slides/slide1.xml"]
    after_root = pe._xml(candidate_parts["ppt/slides/slide1.xml"])
    for before_run, after_run in zip(before_root.xpath(".//p:sp[1]/p:txBody/a:p/a:r", namespaces=pe.NS),
                                     after_root.xpath(".//p:sp[1]/p:txBody/a:p/a:r", namespaces=pe.NS)):
        assert before_run.find("a:t", pe.NS).text == after_run.find("a:t", pe.NS).text
        before_prop, after_prop = before_run.find("a:rPr", pe.NS), after_run.find("a:rPr", pe.NS)
        assert {k: v for k, v in before_prop.attrib.items() if k != "sz"} == {k: v for k, v in after_prop.attrib.items() if k != "sz"}
        assert [ET.tostring(c) for c in before_prop] == [ET.tostring(c) for c in after_prop]
        assert after_prop.get("sz") == "2400"
    obj = object_at(candidate)
    assert obj["geometry"]["mm"] == {"x": 18, "y": 22, "width": 108, "height": 42}
    assert obj["font"]["explicit_sizes_pt"] == [24]
    assert [r["run"] for r in obj["font"]["runs"]] == [1, 2]
    assert obj["font"]["runs"][0]["fonts"]["ea"] == "Microsoft YaHei"
    assert len(Presentation(candidate).slides) == 2
    assert Presentation(candidate).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt == 24
    assert result["non_target_parts_byte_preserved"] and result["unselected_object_xml_byte_preserved"]
    assert not result["native_visual_verified"] and result["native_visual_review_required"]
    assert json.loads((output / "receipt.json").read_text())["candidate_sha256"] == pe._sha(candidate.read_bytes())


def test_inch_geometry_and_picture_edit_preserve_image_payload(source, tmp_path):
    output = tmp_path / "candidate"
    pe.patch_pptx(req(source, set_edit(shape_id=5, geometry={"unit": "inch", "x": 1.25, "width": 2.5})), output)
    obj = object_at(output / "candidate.pptx", "5")
    assert obj["geometry"]["emu"]["x"] == 1143000
    assert obj["geometry"]["inch"]["width"] == 2.5
    before, after = parts(source), parts(output / "candidate.pptx")
    assert all(before[k] == after[k] for k in before if k.startswith("ppt/media/"))


def test_replace_picture_changes_only_unique_media_and_preserves_picture_geometry(source, tmp_path):
    original_parts = parts(source)
    original_picture = object_at(source, "5")
    replacement = tmp_path / "replacement.jpg"
    Image.new("RGB", (240, 90), (176, 42, 36)).save(replacement, format="JPEG")
    output = tmp_path / "picture-replacement"
    edit = {"action": "replace_picture", "slide": 1, "shape_id": "5",
            "replacement_image": str(replacement)}
    result = pe.patch_pptx(req(source, edit), output)
    candidate = output / "candidate.pptx"
    candidate_parts = parts(candidate)
    changed = {name for name in original_parts if original_parts[name] != candidate_parts[name]}
    media_part = next(name for name in original_parts if name.startswith("ppt/media/"))
    assert changed == {media_part}
    assert object_at(candidate, "5") == original_picture
    assert candidate_parts["ppt/slides/slide2.xml"] == original_parts["ppt/slides/slide2.xml"]
    assert source.read_bytes() == Path(result["source"]).read_bytes()
    assert result["changed_parts"][0]["part"] == media_part
    with Image.open(__import__("io").BytesIO(candidate_parts[media_part])) as image:
        assert image.format == "PNG"
        assert image.size == (240, 90)
        assert all(abs(actual - expected) <= 2 for actual, expected in zip(image.getpixel((0, 0)), (176, 42, 36)))
    assert result["native_visual_verified"] is False


def test_replace_picture_rejects_shared_media_without_candidate(source, tmp_path):
    presentation = Presentation(source)
    image_path = tmp_path / "shared.png"
    media_part = next(name for name in parts(source) if name.startswith("ppt/media/"))
    image_path.write_bytes(parts(source)[media_part])
    presentation.slides[1].shapes.add_picture(str(image_path), Mm(80), Mm(90), width=Mm(20))
    presentation.save(source)
    picture_ids = [str(shape.shape_id) for shape in Presentation(source).slides[0].shapes
                   if shape.shape_type == 13]
    assert len(picture_ids) == 1
    replacement = tmp_path / "replacement.png"
    Image.new("RGB", (12, 12), (200, 100, 50)).save(replacement)
    output = tmp_path / "shared-replacement"
    edit = {"action": "replace_picture", "slide": 1, "shape_id": picture_ids[0],
            "replacement_image": str(replacement)}
    original = source.read_bytes()
    with pytest.raises(ValueError, match="shared or ambiguously referenced"):
        pe.patch_pptx(req(source, edit), output)
    assert not output.exists()
    assert source.read_bytes() == original


def test_dry_run_and_source_bound_plan_produce_same_applied_plan(source, tmp_path):
    args = req(source, set_edit(font_size_pt=28))
    planned = pe.patch_pptx(args | {"dry_run": True}, tmp_path / "plan")
    assert planned["status"] == "planned" and planned["candidate"] is None
    assert not (tmp_path / "plan" / "candidate.pptx").exists()
    assert json.loads((tmp_path / "plan" / "after.json").read_text())["slides"][0]["objects"][0]["font"]["explicit_sizes_pt"] == [28]
    applied = pe.patch_pptx(args | {"expected_plan_sha256": planned["plan_sha256"]}, tmp_path / "applied")
    assert applied["plan_sha256"] == planned["plan_sha256"]
    with pytest.raises(ValueError, match="expected_plan_sha256"):
        pe.patch_pptx(args | {"expected_plan_sha256": "0" * 64}, tmp_path / "bad-plan")
    assert not (tmp_path / "bad-plan").exists()


@pytest.mark.parametrize("position,expected", [("front", ["3", "4", "5", "2"]),
                                              ("back", ["2", "3", "4", "5"]),
                                              ("forward", ["3", "2", "4", "5"]),
                                              ("backward", ["2", "3", "4", "5"])])
def test_z_order_reorders_whole_objects_preserving_their_exact_bytes(source, tmp_path, position, expected):
    if position in {"back", "backward"}:
        selected, expected = "3", ["3", "2", "4", "5"]
    else:
        selected = "2"
    before = parts(source)["ppt/slides/slide1.xml"]
    output = tmp_path / "stack"
    result = pe.patch_pptx(req(source, {"action": "z_order", "slide": 1, "shape_id": selected, "position": position}), output)
    after = parts(output / "candidate.pptx")["ppt/slides/slide1.xml"]
    assert [o["shape_id"] for o in pe.inspect_pptx(output / "candidate.pptx", slides=[1])["slides"][0]["objects"]] == expected
    assert all(before[start:end] in after for start, end in pe._child_spans(before))
    assert result["operations"][0]["before"][0]["z_order"] != result["operations"][0]["after"][0]["z_order"]


@pytest.mark.parametrize("axis,coordinate,dimension,factor", [("left", "x", "width", 0), ("center", "x", "width", .5),
                                                             ("right", "x", "width", 1), ("top", "y", "height", 0),
                                                             ("middle", "y", "height", .5), ("bottom", "y", "height", 1)])
def test_align_uses_selected_anchor_without_changing_other_axis_or_size(source, tmp_path, axis, coordinate, dimension, factor):
    output = tmp_path / "aligned"
    pe.patch_pptx(req(source, {"action": "align", "slide": 1, "shape_ids": ["2", "3", "4"],
                              "anchor_shape_id": "3", "axis": axis}), output)
    before = {i: object_at(source, i)["geometry"]["emu"] for i in ["2", "3", "4"]}
    after = {i: object_at(output / "candidate.pptx", i)["geometry"]["emu"] for i in before}
    assert after["3"] == before["3"]
    target = before["3"][coordinate] + before["3"][dimension] * factor
    for identifier, geometry in after.items():
        assert geometry[coordinate] + geometry[dimension] * factor == target
        assert {k: v for k, v in geometry.items() if k != coordinate} == {k: v for k, v in before[identifier].items() if k != coordinate}


def test_distribution_equal_gaps_preserves_outer_edges_and_selection_order(source, tmp_path):
    output = tmp_path / "distributed"
    pe.patch_pptx(req(source, {"action": "distribute", "slide": 1, "shape_ids": ["4", "2", "3"], "axis": "horizontal"}), output)
    after = [object_at(output / "candidate.pptx", i)["geometry"]["emu"] for i in ["2", "3", "4"]]
    assert after[0]["x"] == 10 * 36000
    assert after[2]["x"] + after[2]["width"] == 220 * 36000
    assert after[1]["x"] - after[0]["x"] - after[0]["width"] == after[2]["x"] - after[1]["x"] - after[1]["width"]
    assert object_at(output / "candidate.pptx", "5") == object_at(source, "5")


def test_ordered_batch_set_then_align_uses_updated_geometry(source, tmp_path):
    output = tmp_path / "batch"
    result = pe.patch_pptx(req(source, set_edit(geometry={"unit": "mm", "x": 32}),
                              {"action": "align", "slide": 1, "shape_ids": ["2", "3"], "anchor_shape_id": "2", "axis": "left"}), output)
    assert object_at(output / "candidate.pptx", "3")["geometry"]["mm"]["x"] == 32
    assert result["operations"][1]["before"][0]["geometry"]["mm"]["x"] == 32


@pytest.mark.parametrize("edit,match", [
    (set_edit(shape_id="999", font_size_pt=24), "missing or ambiguous"),
    (set_edit(shape_id=True, font_size_pt=24), "exact positive"),
    (set_edit(slide=0, font_size_pt=24), "1-based"),
    (set_edit(slide=True, font_size_pt=24), "1-based"),
    (set_edit(geometry={"x": 5}), "unit"),
    (set_edit(geometry={"unit": "px", "x": 5}), "unit"),
    (set_edit(geometry={"unit": "mm", "height": 0}), "non-positive"),
    (set_edit(geometry={"unit": "mm", "x": float("nan")}), "finite"),
    (set_edit(font_size_pt=.5), "1–4000"),
    (set_edit(font_size_pt=4001), "1–4000"),
    (set_edit(font_size_pt=24.005), "increments"),
    (set_edit(font_size_pt=True), "finite"),
    (set_edit(shape_id="5", font_size_pt=24), "native text body"),
    (set_edit(font_size_pt=24, color="red"), "only these fields"),
    (set_edit(), "requires geometry"),
    ({"action": "align", "slide": 1, "shape_ids": ["2", "3"], "anchor_shape_id": "4", "axis": "left"}, "anchor_shape_id"),
    ({"action": "align", "slide": 1, "shape_ids": ["2", "2"], "anchor_shape_id": "2", "axis": "left"}, "distinct"),
    ({"action": "distribute", "slide": 1, "shape_ids": ["2", "3"], "axis": "horizontal"}, "3–100"),
    (set_edit(autofit="none"), "with font_size_pt"),
    (set_edit(autofit="shape", font_size_pt=24), "with font_size_pt"),
])
def test_invalid_request_never_creates_candidate_or_changes_source(source, tmp_path, edit, match):
    original = source.read_bytes()
    with pytest.raises(ValueError, match=match):
        pe.patch_pptx(req(source, edit), tmp_path / "bad")
    assert not (tmp_path / "bad").exists()
    assert source.read_bytes() == original


def test_stale_source_hash_missing_hash_and_existing_output_rejected(source, tmp_path):
    args = req(source, set_edit(font_size_pt=24))
    for hash_value in [None, "stale", "0" * 64]:
        with pytest.raises(ValueError, match="source changed"):
            pe.patch_pptx(args | {"source_sha256": hash_value}, tmp_path / "bad")
    with pytest.raises(ValueError, match="never overwritten"):
        pe.patch_pptx(args, source)
    existing = tmp_path / "exists"
    existing.mkdir()
    (existing / "candidate.pptx").write_bytes(b"preserved")
    with pytest.raises(ValueError, match="never overwritten"):
        pe.patch_pptx(args, existing)
    assert (existing / "candidate.pptx").read_bytes() == b"preserved"


def test_grouped_table_rotated_and_autofit_objects_fail_explicitly(source, tmp_path):
    presentation = Presentation(source)
    slide = presentation.slides[0]
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_textbox(Mm(2), Mm(2), Mm(20), Mm(10))
    nested.text = "组内对象"
    nested_id = nested.shape_id
    table = slide.shapes.add_table(2, 2, Mm(20), Mm(60), Mm(40), Mm(20))
    table_id = table.shape_id
    slide.shapes[0].rotation = 10
    slide.shapes[1].text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    presentation.save(source)
    for edit, match in [
        (set_edit(shape_id=nested_id, geometry={"unit": "mm", "x": 5}), "grouped or nested"),
        (set_edit(shape_id=table_id, geometry={"unit": "mm", "x": 5}), "only ordinary"),
        (set_edit(shape_id="3", font_size_pt=24), "autofit"),
        ({"action": "align", "slide": 1, "shape_ids": ["2", "3"], "anchor_shape_id": "3", "axis": "left"}, "unrotated"),
    ]:
        with pytest.raises(ValueError, match=match):
            pe.patch_pptx(req(source, edit), tmp_path / "unsupported")
        assert not (tmp_path / "unsupported").exists()
    inspection = pe.inspect_pptx(source, slides=[1])["slides"][0]["objects"]
    assert not next(o for o in inspection if o["shape_id"] == str(nested_id))["editable"]


@pytest.mark.parametrize("mode", [MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE])
def test_explicit_disable_autofit_allows_font_edit_preserving_body_attributes_and_box(source, tmp_path, mode):
    presentation = Presentation(source)
    shape = presentation.slides[0].shapes[0]
    shape.text_frame.auto_size = mode
    shape.text_frame.margin_top = Mm(4)
    shape.text_frame.word_wrap = True
    presentation.save(source)
    before = object_at(source)
    output = tmp_path / "font"
    result = pe.patch_pptx(req(source, set_edit(font_size_pt=24, autofit="none")), output)
    after = object_at(output / "candidate.pptx")
    assert after["geometry"] == before["geometry"]
    assert after["text_autofit"]["mode"] == "noAutofit"
    assert after["font"]["explicit_sizes_pt"] == [24]
    assert result["operations"][0]["before"][0]["text_autofit"]["mode"] != "noAutofit"
    assert result["operations"][0]["after"][0]["text_autofit"]["mode"] == "noAutofit"
    native = Presentation(output / "candidate.pptx").slides[0].shapes[0]
    assert native.text_frame.margin_top == Mm(4) and native.text_frame.word_wrap
    assert native.text_frame.auto_size == MSO_AUTO_SIZE.NONE


def test_distribution_rejects_insufficient_space_and_equal_positions(source, tmp_path):
    presentation = Presentation(source)
    presentation.slides[0].shapes[2].top = Mm(40)
    presentation.save(source)
    with pytest.raises(ValueError, match="insufficient space"):
        pe.patch_pptx(req(source, {"action": "distribute", "slide": 1, "shape_ids": ["2", "3", "4"], "axis": "vertical"}), tmp_path / "bad")
    presentation.slides[0].shapes[2].top = Mm(25)
    presentation.save(source)
    with pytest.raises(ValueError, match="equal positions"):
        pe.patch_pptx(req(source, {"action": "distribute", "slide": 1, "shape_ids": ["2", "3", "4"], "axis": "vertical"}), tmp_path / "bad")
    assert not (tmp_path / "bad").exists()


def test_atomic_failure_during_package_verification_removes_output(source, tmp_path, monkeypatch):
    original = source.read_bytes()
    output = tmp_path / "interrupted"

    def fail(candidate, expected):
        assert candidate.exists()
        raise ValueError("injected package verification failure")

    monkeypatch.setattr(pe, "_verify_candidate", fail)
    with pytest.raises(ValueError, match="injected"):
        pe.patch_pptx(req(source, set_edit(font_size_pt=24)), output)
    assert source.read_bytes() == original
    assert not output.exists()


def test_late_invalid_operation_rolls_back_whole_plan(source, tmp_path):
    original = source.read_bytes()
    with pytest.raises(ValueError, match="missing or ambiguous"):
        pe.patch_pptx(req(source, set_edit(font_size_pt=24), set_edit(shape_id="999", font_size_pt=28)), tmp_path / "bad")
    assert source.read_bytes() == original and not (tmp_path / "bad").exists()


def test_source_changed_during_write_never_publishes_candidate(source, tmp_path, monkeypatch):
    original_write = pe._write_candidate

    def change_source(data, candidate, expected):
        original_write(data, candidate, expected)
        source.write_bytes(source.read_bytes() + b"externally changed")

    monkeypatch.setattr(pe, "_write_candidate", change_source)
    with pytest.raises(ValueError, match="while preparing"):
        pe.patch_pptx(req(source, set_edit(font_size_pt=24)), tmp_path / "stale")
    assert source.read_bytes().endswith(b"externally changed")
    assert not (tmp_path / "stale").exists()


def test_slide_number_follows_presentation_order_not_part_filename(source, tmp_path):
    presentation = Presentation(source)
    ids = presentation.slides._sldIdLst
    ids.insert(0, ids[-1])
    presentation.save(source)
    output = tmp_path / "reordered"
    result = pe.patch_pptx(req(source, set_edit(geometry={"unit": "mm", "x": 32})), output)
    assert result["changed_parts"][0]["part"] == "ppt/slides/slide2.xml"
    assert object_at(output / "candidate.pptx")["geometry"]["mm"]["x"] == 32
    assert parts(source)["ppt/slides/slide1.xml"] == parts(output / "candidate.pptx")["ppt/slides/slide1.xml"]


def rewrite_parts(path, replacements):
    original = parts(path)
    with ZipFile(path, "w") as archive:
        for name, data in (original | replacements).items():
            archive.writestr(name, data)


def test_xml_comments_and_sibling_whitespace_survive_exact_patch(source, tmp_path):
    raw = parts(source)["ppt/slides/slide1.xml"]
    raw = raw.replace(b"</p:sp>", b"</p:sp>\n  <!-- literal > separator -->\n", 1)
    rewrite_parts(source, {"ppt/slides/slide1.xml": raw})
    output = tmp_path / "comments"
    pe.patch_pptx(req(source, set_edit(font_size_pt=24)), output)
    after = parts(output / "candidate.pptx")["ppt/slides/slide1.xml"]
    start, end = pe._child_spans(raw)[2]
    new_start, new_end = pe._child_spans(after)[2]
    assert raw[:start] == after[:new_start] and raw[end:] == after[new_end:]


def test_duplicate_object_ids_fail_instead_of_selecting_first(source, tmp_path):
    raw = parts(source)["ppt/slides/slide1.xml"].replace(b'id="3"', b'id="2"', 1)
    rewrite_parts(source, {"ppt/slides/slide1.xml": raw})
    with pytest.raises(ValueError, match="missing or ambiguous"):
        pe.patch_pptx(req(source, set_edit(font_size_pt=24)), tmp_path / "bad")
    assert not (tmp_path / "bad").exists()


def test_nonidentity_slide_group_transform_is_not_assumed_to_be_slide_coordinates(source, tmp_path):
    raw = parts(source)["ppt/slides/slide1.xml"]
    root = pe._xml(raw)
    transform = ET.SubElement(root.find("p:cSld/p:spTree/p:grpSpPr", pe.NS), f"{{{pe.NS['a']}}}xfrm")
    for tag, attributes in [("off", {"x": "0", "y": "0"}), ("ext", {"cx": "200", "cy": "200"}),
                            ("chOff", {"x": "0", "y": "0"}), ("chExt", {"cx": "100", "cy": "100"})]:
        ET.SubElement(transform, f"{{{pe.NS['a']}}}{tag}", attributes)
    rewrite_parts(source, {"ppt/slides/slide1.xml": ET.tostring(root)})
    with pytest.raises(ValueError, match="non-identity group transform"):
        pe.patch_pptx(req(source, set_edit(geometry={"unit": "mm", "x": 25})), tmp_path / "bad")


def test_invalid_pptx_and_nonboolean_dry_run_have_actionable_errors(source, tmp_path):
    with pytest.raises(ValueError, match="boolean"):
        pe.patch_pptx(req(source, set_edit(font_size_pt=24), dry_run="yes"), tmp_path / "bad")
    source.write_bytes(b"not a zip")
    with pytest.raises(ValueError, match="invalid or encrypted PPTX"):
        pe.patch_pptx(req(source, set_edit(font_size_pt=24)), tmp_path / "bad")
    assert not (tmp_path / "bad").exists()
