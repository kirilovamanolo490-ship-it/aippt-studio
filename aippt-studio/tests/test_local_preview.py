from __future__ import annotations

import hashlib
import json
from pathlib import Path
import subprocess

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from aippt_studio.local_preview import LocalPreviewError, render_slide_preview


def _deck(path: Path) -> None:
    presentation = Presentation()
    for label in ("First slide", "Second slide should remain unrendered"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        shape.text = label
    presentation.save(path)


def test_renders_one_slide_of_python_pptx_and_writes_bound_receipt(tmp_path: Path) -> None:
    source = tmp_path / "two-slides.pptx"
    output = tmp_path / "preview.png"
    receipt_path = tmp_path / "preview.json"
    _deck(source)

    receipt = render_slide_preview(source, 1, output, receipt_path=receipt_path, width=640)

    assert receipt["status"] == "rendered_advisory"
    assert receipt["source_sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert receipt["slide_number"] == 1
    assert receipt["slide_count"] == 2
    assert receipt["rendered_slide_count"] == 1
    assert receipt["approximate_renderer"] is True
    assert receipt["renderer_version"] == "0.11.0"
    assert "not Office-native" in " ".join(receipt["limitations"])
    assert receipt["output_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert json.loads(receipt_path.read_text(encoding="utf-8")) == receipt
    assert Image.open(output).size[0] == 640
    assert not list(tmp_path.glob("*.pdf"))
    assert not list(tmp_path.glob("slide-2*.png"))


def test_rejects_slide_outside_deck_without_claiming_preview(tmp_path: Path) -> None:
    source = tmp_path / "two-slides.pptx"
    _deck(source)
    with pytest.raises(LocalPreviewError, match="could not render"):
        render_slide_preview(source, 3, tmp_path / "invalid.png")
    assert not (tmp_path / "invalid.png").exists()


@pytest.mark.parametrize("slide_number", [0, -1, True, 1.5])
def test_rejects_invalid_slide_number(tmp_path: Path, slide_number: object) -> None:
    with pytest.raises(LocalPreviewError, match="positive 1-based integer"):
        render_slide_preview(tmp_path / "missing.pptx", slide_number, tmp_path / "invalid.png")  # type: ignore[arg-type]


@pytest.mark.parametrize("existing", ["png", "receipt"])
def test_refuses_to_overwrite_outputs(tmp_path: Path, existing: str) -> None:
    source = tmp_path / "two-slides.pptx"
    output = tmp_path / "preview.png"
    receipt = tmp_path / "preview.json"
    _deck(source)
    occupied = output if existing == "png" else receipt
    occupied.write_text("keep existing bytes", encoding="utf-8")
    with pytest.raises(LocalPreviewError, match="must not already exist"):
        render_slide_preview(source, 1, output, receipt_path=receipt)
    assert occupied.read_text(encoding="utf-8") == "keep existing bytes"
    assert not output.exists() if existing == "receipt" else True


def test_missing_dependency_has_actionable_optional_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import aippt_studio.local_preview as preview

    source = tmp_path / "two-slides.pptx"
    _deck(source)
    monkeypatch.setattr(preview, "PROJECT_ROOT", tmp_path)
    with pytest.raises(LocalPreviewError, match="npm ci"):
        render_slide_preview(source, 1, tmp_path / "preview.png")


def test_renderer_failure_does_not_delete_competing_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import aippt_studio.local_preview as preview

    source, output = tmp_path / "two-slides.pptx", tmp_path / "preview.png"
    _deck(source)
    monkeypatch.setattr(preview, "probe_node_runtime", lambda: {"status": "passed", "path": "fake-node"})

    def fail_after_competitor(*args: object, **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        output.write_bytes(b"someone else's output")
        return subprocess.CompletedProcess([], 1, b"", b"simulated renderer error")

    monkeypatch.setattr(preview.subprocess, "run", fail_after_competitor)
    with pytest.raises(LocalPreviewError, match="could not render"):
        render_slide_preview(source, 1, output)
    assert output.read_bytes() == b"someone else's output"


def test_refuses_dangling_output_symlink(tmp_path: Path) -> None:
    source, output = tmp_path / "two-slides.pptx", tmp_path / "preview.png"
    _deck(source)
    output.symlink_to(tmp_path / "absent.png")
    with pytest.raises(LocalPreviewError, match="must not already exist"):
        render_slide_preview(source, 1, output)
    assert output.is_symlink()
