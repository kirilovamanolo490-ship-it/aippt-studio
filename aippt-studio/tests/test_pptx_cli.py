"""The quick local edit path should avoid scripts and large inspection output."""

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
import pytest
from typer.testing import CliRunner

from aippt_studio import cli


def _deck(path: Path) -> Path:
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    first.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "不要改动的第一页"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "原文"
    deck.save(path)
    return path


def test_scoped_brief_inspect_and_direct_text_edit(tmp_path: Path):
    source = _deck(tmp_path / "source.pptx")
    original = source.read_bytes()
    report = tmp_path / "inspection.json"
    inspection = CliRunner().invoke(cli.app, ["pptx-inspect", str(source), "--slide", "2", "--brief",
                                               "--output", str(report), "--json"])
    assert inspection.exit_code == 0, inspection.output
    summary = json.loads(inspection.output)
    assert [slide["slide"] for slide in summary["slides"]] == [2]
    assert summary["slides"][0]["objects"][0]["text_preview"] == "原文"
    assert "font" not in summary["slides"][0]["objects"][0]
    assert json.loads(report.read_text())["slides"][0]["objects"][0]["font"]["runs"]

    output = tmp_path / "local-edit"
    result = CliRunner().invoke(cli.app, ["pptx-text", str(source), "--source-sha256", summary["source_sha256"],
                                           "--slide", "2", "--shape-id", "2", "--expected-text", "原文",
                                           "--replacement-text", "面向听众的新文案", "--output", str(output), "--json"])
    assert result.exit_code == 0, result.output
    edited = json.loads(result.output)
    assert edited["status"] == "candidate" and edited["native_visual_verified"] is False
    assert Presentation(output / "candidate.pptx").slides[1].shapes[0].text == "面向听众的新文案"
    assert source.read_bytes() == original
    with ZipFile(source) as before, ZipFile(output / "candidate.pptx") as after:
        assert before.read("ppt/slides/slide1.xml") == after.read("ppt/slides/slide1.xml")

    qa_path = tmp_path / "revision-qa.json"
    qa = CliRunner().invoke(cli.app, ["revision-check", "--receipt", str(output / "receipt.json"),
                                       "--output", str(qa_path), "--json"])
    assert qa.exit_code == 0, qa.output
    qa_summary = json.loads(qa.output)
    assert qa_summary["status"] == "passed_scoped"
    assert qa_summary["scope"]["selected_slides"] == [2]
    assert qa_summary["visual_review"]["scope"] == "target_slides_only"
    assert qa_path.is_file()
    assert not list(output.glob("*.pdf"))


def test_add_text_cli_creates_one_editable_shape_with_scoped_qa(tmp_path: Path):
    source = _deck(tmp_path / "source.pptx")
    original = source.read_bytes()
    inspected = CliRunner().invoke(cli.app, ["pptx-inspect", str(source), "--slide", "2", "--brief", "--json"])
    assert inspected.exit_code == 0, inspected.output
    source_hash = json.loads(inspected.output)["source_sha256"]
    output = tmp_path / "add-text"
    added = CliRunner().invoke(cli.app, ["pptx-add-text", str(source), "--source-sha256", source_hash,
                                          "--slide", "2", "--text", "新增的听众文案",
                                          "--unit", "inch", "--x", "1", "--y", "2", "--width", "4",
                                          "--height", "1", "--font-size-pt", "18", "--output", str(output), "--json"])
    assert added.exit_code == 0, added.output
    edit_receipt = json.loads(added.output)
    assert edit_receipt["status"] == "candidate"
    assert edit_receipt["new_shape_id"] == "3"
    candidate = Presentation(output / "candidate.pptx")
    assert len(candidate.slides[1].shapes) == 2
    assert candidate.slides[1].shapes[1].text == "新增的听众文案"
    assert source.read_bytes() == original
    with ZipFile(source) as before, ZipFile(output / "candidate.pptx") as after:
        assert before.read("ppt/slides/slide1.xml") == after.read("ppt/slides/slide1.xml")
    report_path = tmp_path / "add-text-qa.json"
    checked = CliRunner().invoke(cli.app, ["revision-check", "--receipt", str(output / "receipt.json"),
                                           "--output", str(report_path), "--json"])
    assert checked.exit_code == 0, checked.output
    summary = json.loads(checked.output)
    assert summary["status"] == "passed_scoped"
    assert summary["visual_review"]["required"] is True
    assert summary["visual_review"]["assistant_review"] == "required_for_new_page_copy_and_hierarchy"


def test_direct_text_edit_rejects_stale_hash_without_output(tmp_path: Path):
    source = _deck(tmp_path / "source.pptx")
    output = tmp_path / "stale"
    result = CliRunner().invoke(cli.app, ["pptx-text", str(source), "--source-sha256", "0" * 64,
                                           "--slide", "2", "--shape-id", "2", "--expected-text", "原文",
                                           "--replacement-text", "新文案", "--output", str(output), "--json"])
    assert result.exit_code == 2
    assert not output.exists()
    assert hashlib.sha256(source.read_bytes()).hexdigest() != "0" * 64


def test_nudge_moves_one_object_without_rebuilding_other_slides(tmp_path: Path):
    source = _deck(tmp_path / "source.pptx")
    original = source.read_bytes()
    inspection = CliRunner().invoke(cli.app, ["pptx-inspect", str(source), "--slide", "2", "--json"])
    assert inspection.exit_code == 0, inspection.output
    report = json.loads(inspection.output)
    shape = report["slides"][0]["objects"][0]
    output = tmp_path / "nudge"
    edited = CliRunner().invoke(cli.app, ["pptx-nudge", str(source), "--source-sha256", report["source_sha256"],
                                           "--slide", "2", "--shape-id", shape["shape_id"],
                                           "--dx-mm", "2.5", "--dy-mm", "-1", "--dwidth-mm", "3",
                                           "--output", str(output), "--json"])
    assert edited.exit_code == 0, edited.output
    result = json.loads(edited.output)
    assert result["status"] == "candidate"
    assert result["native_visual_verified"] is False
    assert source.read_bytes() == original
    changed = Presentation(output / "candidate.pptx").slides[1].shapes[0]
    assert changed.left - shape["geometry"]["emu"]["x"] == 90000
    assert changed.top - shape["geometry"]["emu"]["y"] == -36000
    assert changed.width - shape["geometry"]["emu"]["width"] == 108000
    with ZipFile(source) as before, ZipFile(output / "candidate.pptx") as after:
        assert before.read("ppt/slides/slide1.xml") == after.read("ppt/slides/slide1.xml")
    qa = CliRunner().invoke(cli.app, ["revision-check", "--receipt", str(output / "receipt.json"),
                                       "--output", str(tmp_path / "nudge-qa.json"), "--json"])
    assert qa.exit_code == 0, qa.output
    assert json.loads(qa.output)["status"] == "passed_scoped"


def test_nudge_rejects_stale_hash_or_invalid_size_without_output(tmp_path: Path):
    source = _deck(tmp_path / "source.pptx")
    stale = tmp_path / "stale-nudge"
    result = CliRunner().invoke(cli.app, ["pptx-nudge", str(source), "--source-sha256", "0" * 64,
                                           "--slide", "2", "--shape-id", "2", "--dx-mm", "1",
                                           "--output", str(stale), "--json"])
    assert result.exit_code == 2 and not stale.exists()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    invalid = tmp_path / "invalid-nudge"
    result = CliRunner().invoke(cli.app, ["pptx-nudge", str(source), "--source-sha256", source_hash,
                                           "--slide", "2", "--shape-id", "2", "--dwidth-mm", "-200",
                                           "--output", str(invalid), "--json"])
    assert result.exit_code == 2 and not invalid.exists()
    offslide = tmp_path / "offslide-nudge"
    result = CliRunner().invoke(cli.app, ["pptx-nudge", str(source), "--source-sha256", source_hash,
                                           "--slide", "2", "--shape-id", "2", "--dx-mm", "-30",
                                           "--output", str(offslide), "--json"])
    assert result.exit_code == 2 and not offslide.exists()


def test_host_nudge_uses_staged_copy_and_verifies_readback(tmp_path: Path, monkeypatch):
    from aippt_studio import host_nudge
    import subprocess

    source = _deck(tmp_path / "source.pptx")
    original = source.read_bytes()
    inspection = json.loads(CliRunner().invoke(
        cli.app, ["pptx-inspect", str(source), "--slide", "2", "--json"]
    ).output)

    def fake_host(command):
        candidate = Path(command[3])
        slide_index = int(command[4])
        shape_name = command[5]
        target = [float(value) for value in command[-4:]]
        deck = Presentation(candidate)
        shapes = [shape for shape in deck.slides[slide_index - 1].shapes if shape.name == shape_name]
        assert len(shapes) == 1
        shape = shapes[0]
        shape.left, shape.top, shape.width, shape.height = [int(round(value * 12700)) for value in target]
        deck.save(candidate)
        return subprocess.CompletedProcess(command, 0, stdout="saved", stderr="")

    monkeypatch.setattr(host_nudge.platform, "system", lambda: "Darwin")
    out = tmp_path / "host-edit"
    result = host_nudge.nudge_with_powerpoint(
        source, source_sha256=inspection["source_sha256"], slide=2, shape_id="2", output=out,
        dx_pt=2.5, dy_pt=-1.25, executor=fake_host, osascript=Path("/mock/osascript"),
    )
    assert result["status"] == "candidate"
    assert result["native_save_readback"] == "passed"
    assert result["native_visual_verified"] is False
    assert source.read_bytes() == original
    assert (out / "host-nudge-receipt.json").is_file()
    candidate = Presentation(result["candidate"])
    original_shape = Presentation(source).slides[1].shapes[0]
    changed_shape = candidate.slides[1].shapes[0]
    assert changed_shape.left - original_shape.left == Pt(2.5)
    assert changed_shape.top - original_shape.top == Pt(-1.25)


def test_host_nudge_rejects_non_mac_and_stale_source_before_staging(tmp_path: Path):
    from aippt_studio.host_nudge import nudge_with_powerpoint

    source = _deck(tmp_path / "source.pptx")
    report = json.loads(CliRunner().invoke(cli.app, ["pptx-inspect", str(source), "--json"]).output)
    output = tmp_path / "not-created"
    with pytest.raises(ValueError, match="macOS"):
        nudge_with_powerpoint(source, source_sha256=report["source_sha256"], slide=1,
                              shape_id="2", output=output, dx_pt=1, host_platform="Windows")
    with pytest.raises(ValueError, match="source changed"):
        nudge_with_powerpoint(source, source_sha256="0" * 64, slide=1,
                              shape_id="2", output=output, dx_pt=1, host_platform="Darwin")
    assert not output.exists()


def test_host_nudge_permission_failure_keeps_attempt_out_of_candidate_state(tmp_path: Path):
    from aippt_studio.host_nudge import nudge_with_powerpoint
    import subprocess

    source = _deck(tmp_path / "source.pptx")
    report = json.loads(CliRunner().invoke(cli.app, ["pptx-inspect", str(source), "--json"]).output)
    output = tmp_path / "host-permission-failure"

    def denied(command):
        return subprocess.CompletedProcess(command, 1, stdout="", stderr="PowerPoint file access denied")

    with pytest.raises(RuntimeError, match="host edit failed"):
        nudge_with_powerpoint(source, source_sha256=report["source_sha256"], slide=1,
                              shape_id="2", output=output, dx_pt=1,
                              executor=denied, osascript=Path("/mock/osascript"), host_platform="Darwin")
    failure = json.loads((output / "host-nudge-failure.json").read_text())
    assert failure["status"] == "host_failed" and failure["candidate_accepted"] is False
    assert not list(output.glob("host-candidate-*.pptx"))


def test_image_check_reports_local_advisories_without_editing_asset(tmp_path: Path):
    source = tmp_path / "asset.png"
    Image.new("RGB", (400, 200), (50, 90, 160)).save(source)
    original = source.read_bytes()
    output = tmp_path / "image-qa.json"
    result = CliRunner().invoke(cli.app, ["image-check", str(source), "--width", "8", "--height", "4",
                                           "--unit", "in", "--output", str(output), "--json"])
    assert result.exit_code == 0, result.output
    report = json.loads(result.output)
    assert report["assessment"] == "advisory_only"
    assert report["decode_status"] == "decoded"
    assert report["effective_ppi"] == {"width": 50, "height": 50}
    assert json.loads(output.read_text()) == report
    assert source.read_bytes() == original


def test_image_prepare_cli_writes_local_copy_without_modifying_source(tmp_path: Path):
    source = tmp_path / "source.png"
    Image.new("RGB", (900, 450), (50, 90, 160)).save(source)
    original = source.read_bytes()
    output = tmp_path / "prepared.png"
    result = CliRunner().invoke(cli.app, ["image-prepare", str(source), "--max-edge", "450",
                                           "--output", str(output), "--json"])
    assert result.exit_code == 0, result.output
    receipt = json.loads(result.output)
    assert receipt["status"] == "prepared"
    assert receipt["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert receipt["output_size_px"] == {"width": 450, "height": 225}
    assert receipt["source_modified"] is False
    assert source.read_bytes() == original
    with Image.open(output) as image:
        assert image.size == (450, 225)


def test_picture_cli_replaces_only_unshared_media_and_scoped_qa(tmp_path: Path):
    original_image = tmp_path / "original.png"
    replacement_image = tmp_path / "replacement.png"
    Image.new("RGB", (600, 400), (50, 90, 160)).save(original_image)
    Image.new("RGB", (600, 400), (160, 90, 50)).save(replacement_image)
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "不改的页"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_picture(str(original_image), Inches(1), Inches(1), width=Inches(4))
    source = tmp_path / "source.pptx"
    deck.save(source)
    before = source.read_bytes()
    inspect = CliRunner().invoke(cli.app, ["pptx-inspect", str(source), "--slide", "2", "--brief", "--json"])
    assert inspect.exit_code == 0, inspect.output
    summary = json.loads(inspect.output)
    picture_id = summary["slides"][0]["objects"][0]["shape_id"]
    output = tmp_path / "picture-edit"
    edit = CliRunner().invoke(cli.app, ["pptx-picture", str(source), "--source-sha256", summary["source_sha256"],
                                        "--slide", "2", "--shape-id", str(picture_id),
                                        "--replacement-image", str(replacement_image),
                                        "--output", str(output), "--json"])
    assert edit.exit_code == 0, edit.output
    receipt = json.loads(edit.output)
    assert receipt["status"] == "candidate"
    assert any(part["part"].startswith("ppt/media/") for part in receipt["changed_parts"])
    assert source.read_bytes() == before
    with ZipFile(source) as old, ZipFile(output / "candidate.pptx") as new:
        assert old.read("ppt/slides/slide1.xml") == new.read("ppt/slides/slide1.xml")
        assert old.read("ppt/slides/slide2.xml") == new.read("ppt/slides/slide2.xml")
    report_path = tmp_path / "picture-qa.json"
    checked = CliRunner().invoke(cli.app, ["revision-check", "--receipt", str(output / "receipt.json"),
                                           "--output", str(report_path), "--json"])
    assert checked.exit_code == 0, checked.output
    report = json.loads(checked.output)
    assert report["status"] == "passed_scoped"
    assert report["visual_review"]["required"] is True
    assert report["visual_review"]["assistant_review"] == "required_for_image_to_claim_and_crop"


def test_preview_slide_cli_renders_one_advisory_page_without_pdf(tmp_path: Path):
    source = _deck(tmp_path / "source.pptx")
    output = tmp_path / "page-2.png"
    result = CliRunner().invoke(cli.app, ["preview-slide", str(source), "--slide", "2",
                                           "--width", "640", "--output", str(output),
                                           "--no-system-fonts", "--json"])
    assert result.exit_code == 0, result.output
    summary = json.loads(result.output)
    assert summary["status"] == "rendered_advisory"
    assert summary["slide_number"] == 2
    assert summary["slide_count"] == 2
    assert summary["approximate_renderer"] is True
    assert output.is_file()
    assert (tmp_path / "page-2.preview.json").is_file()
    assert not list(tmp_path.glob("*.pdf"))
    assert len(list(tmp_path.glob("*.png"))) == 1


def test_preview_slide_cli_supports_managed_officecli_backend(tmp_path: Path):
    source = _deck(tmp_path / "source.pptx")
    output = tmp_path / "officecli-page-2.png"
    result = CliRunner().invoke(cli.app, ["preview-slide", str(source), "--slide", "2",
                                           "--backend", "officecli", "--width", "640",
                                           "--output", str(output), "--json"])
    assert result.exit_code == 0, result.output
    summary = json.loads(result.output)
    assert summary["renderer"] == "officecli-html"
    assert summary["slide_count"] == 2
    assert summary["approximate_renderer"] is True
    assert output.is_file()
    assert not list(tmp_path.glob("*.pdf"))
