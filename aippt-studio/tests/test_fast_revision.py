from __future__ import annotations

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.util import Inches
from typer.testing import CliRunner

from aippt_studio import cli
from aippt_studio.fast_revision import revise_pptx
from aippt_studio.pptx_edit import inspect_pptx


def _deck(path: Path) -> Path:
    deck = Presentation()
    for copy in ("keep", "change"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = copy
    deck.save(path)
    return path


def _request(source: Path) -> dict:
    inspected = inspect_pptx(source, slides=[2])
    shape_id = inspected["slides"][0]["objects"][0]["shape_id"]
    return {"source": str(source), "source_sha256": inspected["source_sha256"],
            "edits": [{"action": "replace_text", "slide": 2, "shape_id": shape_id,
                       "expected_text": "change", "replacement_text": "revised"}]}


def test_fast_revision_reuses_deck_and_checks_only_changed_page(tmp_path: Path):
    source = _deck(tmp_path / "original.pptx")
    original = source.read_bytes()
    result = revise_pptx(_request(source), tmp_path / "revision", preview_backend="none")

    assert result["status"] == "needs_visual_review"
    assert result["changed_slides"] == [2]
    assert result["revision_qa_status"] == "passed_scoped"
    assert result["native_visual_verified"] is False
    assert result["final_deck_verified"] is False
    assert source.read_bytes() == original
    candidate = tmp_path / "revision/candidate.pptx"
    assert hashlib.sha256(candidate.read_bytes()).hexdigest() == result["candidate_sha256"]
    assert Presentation(candidate).slides[1].shapes[0].text == "revised"
    assert json.loads((tmp_path / "revision/revision-qa.json").read_text())["scope"]["selected_slides"] == [2]
    assert not list((tmp_path / "revision").glob("*.pdf"))
    with ZipFile(source) as before, ZipFile(candidate) as after:
        assert before.read("ppt/slides/slide1.xml") == after.read("ppt/slides/slide1.xml")


def test_fast_revision_cli_produces_one_scoped_candidate(tmp_path: Path):
    source = _deck(tmp_path / "original.pptx")
    request = tmp_path / "request.json"
    request.write_text(json.dumps(_request(source)), encoding="utf-8")
    result = CliRunner().invoke(cli.app, ["fast-revise", "--request", str(request),
                                           "--output", str(tmp_path / "revision"),
                                           "--preview", "none", "--json"])
    assert result.exit_code == 0, result.output
    payload = json.loads(result.output)
    assert payload["changed_slides"] == [2]
    assert payload["revision_qa_status"] == "passed_scoped"
    assert (tmp_path / "revision/fast-revision.json").is_file()


def test_fast_revision_previews_only_changed_page_and_records_unavailable_renderer(tmp_path: Path, monkeypatch):
    from aippt_studio import local_preview

    source = _deck(tmp_path / "original.pptx")
    called: list[int] = []

    def preview(candidate, slide, output, *, width, load_system_fonts):
        called.append(slide)
        raise local_preview.LocalPreviewError("renderer unavailable")

    monkeypatch.setattr(local_preview, "render_slide_preview", preview)
    output = tmp_path / "revision"
    result = revise_pptx(_request(source), output)

    assert called == [2]
    assert result["status"] == "needs_visual_review"
    assert result["revision_qa_status"] == "passed_scoped"
    assert result["previews"] == [{"slide": 2, "status": "unavailable", "reason": "renderer unavailable"}]
    assert (output / "candidate.pptx").is_file()
    assert json.loads((output / "fast-revision.json").read_text())["previews"] == result["previews"]


def test_fast_revision_rejects_full_deck_change_before_writing(tmp_path: Path):
    source = _deck(tmp_path / "original.pptx")
    request = _request(source)
    request["edits"] = [dict(request["edits"][0]) for _ in range(5)]
    for index, edit in enumerate(request["edits"], 1):
        edit["slide"] = index
    output = tmp_path / "revision"
    try:
        revise_pptx(request, output, preview_backend="none")
    except ValueError as error:
        assert "four changed pages" in str(error)
    else:
        raise AssertionError("full-deck edit should not use fast revision")
    assert not output.exists()


def test_fast_revision_stays_in_active_work_workspace(tmp_path: Path, monkeypatch):
    work = tmp_path / "work" / "37"
    workspace = work / "workspace"
    workspace.mkdir(parents=True)
    source = _deck(workspace / "original.pptx")
    monkeypatch.setenv("AIPPT_WORK_DIR", str(work))

    with pytest.raises(ValueError, match="inside the active Work workspace"):
        revise_pptx(_request(source), tmp_path / "outside" / "revision", preview_backend="none")
    assert not (tmp_path / "outside" / "revision").exists()

    outside = tmp_path / "outside"
    outside.mkdir()
    (workspace / "escape").symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match="inside the active Work workspace"):
        revise_pptx(_request(source), workspace / "escape" / "revision", preview_backend="none")
    assert not (outside / "revision").exists()


def test_fast_revision_rejects_preview_of_a_changed_candidate(tmp_path: Path, monkeypatch):
    from aippt_studio import local_preview

    source = _deck(tmp_path / "original.pptx")

    def changed_preview(candidate, slide, output, *, width, load_system_fonts):
        candidate.write_bytes(b"replaced while rendering")
        return {"source_sha256": "other", "output_png": str(output),
                "output_sha256": "other", "renderer": "mock"}

    monkeypatch.setattr(local_preview, "render_slide_preview", changed_preview)
    output = tmp_path / "revision"
    with pytest.raises(ValueError, match="preview source differs"):
        revise_pptx(_request(source), output)
    assert not (output / "fast-revision.json").exists()
