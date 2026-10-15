from __future__ import annotations

import subprocess
from pathlib import Path

import pytest
from pptx import Presentation

from aippt_studio import ingest
from aippt_studio.ingest import extract_text, ingest_sources
from aippt_studio.models import EvidenceLedger, JobState
from aippt_studio.storage import job_dir, read_model, save_job
from aippt_studio.workflow import create_job


def _powerpoint_source(tmp_path: Path) -> Path:
    path = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "客户年度经营复盘"
    slide.placeholders[1].text = "收入增长 28%，续约率达到 93%"
    presentation.save(path)
    return path


def test_ingest_hashes_normalizes_and_extracts_locked_facts(tmp_path: Path) -> None:
    source = tmp_path / "facts.txt"
    source.write_text("2026 年收入增长 28%。\n交付周期从 21 天缩短到 14 天。", encoding="utf-8")
    job = create_job("事实追踪", job_id="facts")
    records = ingest_sources(job.id, [str(source)])
    assert len(records) == 1
    assert len(records[0].sha256) == 64
    assert records[0].normalized_path
    ledger = read_model(job_dir(job.id) / "evidence-ledger.json", EvidenceLedger)
    assert len(ledger.locked_facts) == 2
    assert all(fact.source_id == records[0].id for fact in ledger.locked_facts)


def test_inline_chat_text_uses_normal_ingest_provenance() -> None:
    job = create_job("聊天文字", job_id="inline-text")

    records = ingest_sources(job.id, [], inline_texts=[("customer-chat.md", "客户要求保留 2026 年收入增长 28% 的事实。")])

    assert records[0].original == "chat:inline-text"
    assert (job_dir(job.id) / records[0].stored_path).read_text(encoding="utf-8").startswith("客户要求")
    ledger = read_model(job_dir(job.id) / "evidence-ledger.json", EvidenceLedger)
    assert ledger.sources[0].sha256 == records[0].sha256


def test_multi_source_ingest_rolls_back_when_a_later_source_fails(tmp_path: Path) -> None:
    valid = tmp_path / "valid.txt"
    valid.write_text("2026 年收入增长 28%。", encoding="utf-8")
    job = create_job("事务导入", job_id="transactional-ingest")

    with pytest.raises(FileNotFoundError):
        ingest_sources(job.id, [str(valid), str(tmp_path / "missing.txt")])

    assert not any((job_dir(job.id) / "sources").iterdir())
    ledger = read_model(job_dir(job.id) / "evidence-ledger.json", EvidenceLedger)
    assert ledger.sources == []


def test_gb18030_text_is_decoded_without_replacement_characters(tmp_path: Path) -> None:
    source = tmp_path / "gbk-notes.txt"
    source.write_bytes("客户经营复盘：收入增长 28%。".encode("gb18030"))
    job = create_job("中文编码", job_id="gb18030")

    record = ingest_sources(job.id, [str(source)])[0]

    normalized = (job_dir(job.id) / (record.normalized_path or "")).read_text(encoding="utf-8")
    assert normalized == "客户经营复盘：收入增长 28%。"
    assert "�" not in normalized
    assert record.metadata["encoding"] == "gb18030"


@pytest.mark.parametrize("state", [JobState.FINAL_REVIEW, JobState.DELIVERED])
def test_ingest_rejects_jobs_after_build_is_complete(tmp_path: Path, state: JobState) -> None:
    source = tmp_path / "late-material.txt"
    source.write_text("这份材料不应在终稿阶段被追加。", encoding="utf-8")
    job = create_job("禁止追加材料", job_id=f"late-{state.value}")
    job.state = state
    save_job(job)

    with pytest.raises(ValueError):
        ingest_sources(job.id, [str(source)])

    assert not any((job_dir(job.id) / "sources").iterdir())


def test_markitdown_runs_in_isolated_subprocess_when_cli_is_not_on_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _powerpoint_source(tmp_path)
    real_run = subprocess.run
    commands: list[list[str]] = []
    monkeypatch.setenv("PATH", "")
    monkeypatch.setattr(ingest.shutil, "which", lambda name: None)

    def capture_subprocess(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
        commands.append(command)
        return real_run(command, **kwargs)

    monkeypatch.setattr(ingest.subprocess, "run", capture_subprocess)

    text = extract_text(source, "powerpoint")

    expected = [ingest.sys.executable, "-m", "markitdown", str(source)]
    assert commands == [expected]
    assert "客户年度经营复盘" in text
    assert "收入增长 28%" in text


def test_markitdown_timeout_is_explicit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _powerpoint_source(tmp_path)

    def time_out(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
        raise subprocess.TimeoutExpired(command, kwargs["timeout"])

    monkeypatch.setattr(ingest.subprocess, "run", time_out)

    with pytest.raises(ValueError, match=r"MarkItDown timed out after 180s for source\.pptx"):
        extract_text(source, "powerpoint")


def test_markitdown_nonzero_exit_is_explicit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = _powerpoint_source(tmp_path)

    def fail_conversion(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
        return subprocess.CompletedProcess(command, 2, stdout="", stderr="converter exploded")

    monkeypatch.setattr(ingest.subprocess, "run", fail_conversion)

    with pytest.raises(ValueError, match=r"MarkItDown failed for source\.pptx: converter exploded"):
        extract_text(source, "powerpoint")
