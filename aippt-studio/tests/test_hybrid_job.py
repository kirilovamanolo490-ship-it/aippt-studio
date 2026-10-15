from __future__ import annotations

import copy
import hashlib
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
import pytest

from adapter_test_helpers import write_contract_job
from test_style_hybrid_workflow import _library, _receipt
from aippt_studio import hybrid_job
from aippt_studio.adapter_contracts import AdapterArtifact, AdapterBuildRequest, AdapterStatus
from aippt_studio.adapters.template import TemplateAdapter
from aippt_studio.models import DeckPlan, JobRecord, JobState, ProductionRoute
from aippt_studio.renderer import sha256
from aippt_studio.storage import load_job, read_json, read_model, save_job, write_json
from aippt_studio.workflow import _ensure_adapter_previews, _promote_adapter_result, _template_application_artifact_reason, _validate_external_output


@pytest.fixture
def bound_job(tmp_path, monkeypatch):
    _library_root, rel, _, primary, _ = _library(tmp_path)
    monkeypatch.setattr(hybrid_job, "WORKSPACE_ROOT", tmp_path)
    monkeypatch.setenv("AIPPT_JOBS_DIR", str(tmp_path / "jobs"))
    root = tmp_path / "jobs/fixture"
    request_path, _ = write_contract_job(root, route="template", role="template", source=primary)
    write_json(root / "evidence-ledger.json", {"job_id": "fixture", "sources": []})
    job = JobRecord(id="fixture", name="hybrid integration fixture", state=JobState.FULL_BUILD,
                    route=ProductionRoute.TEMPLATE, source_ids=["source-1"], input_roles={"source-1": "template"},
                    approvals={"brief": True, "prototype": True, "final": True})
    save_job(job)
    packet = {"primary_source_path": rel, "primary_source_sha256": sha256(primary),
              "reviewed_primary_pages": [_receipt(tmp_path, primary, 1, "cover"),
                                         _receipt(tmp_path, primary, 2, "body")],
              "selection_note": "保留真实测试模板的页面对象、主题和相邻页结构。",
              "source_content_replacement_plan": "所有来源文字都替换为当前计划中的项目标题和内容。",
              "rights_handling_note": "仅用于内部结构回归的自建素材，不作为客户成品。",
              "text_replacements": []}
    for number, slide in enumerate(Presentation(primary).slides, 1):
        for index, shape in enumerate(slide.shapes):
            packet["text_replacements"].append({"slide": number, "shape_id": str(shape.shape_id),
                "original_text_sha256": hashlib.sha256(shape.text.encode()).hexdigest(),
                "new_text": f"Fixture title {number}" if index == 0 else f"Point {number}"})
    return root, job, AdapterBuildRequest.model_validate(read_json(request_path)), packet


def build(root, request):
    # The registered adapter must choose the internal binding without invoking
    # an unconfigured subprocess/Automizer runtime.
    adapter = TemplateAdapter(project_root=root, upstream_root=root / "missing", reviewed_commit=None)
    return adapter.build(request, root)


def test_registered_adapter_build_promotes_and_rechecks_actual_package(bound_job):
    root, job, request, packet = bound_job
    previous = root / "build/deck.pptx"
    previous.write_bytes(b"prior editable deck stays until promotion")
    hybrid_job.bind_hybrid_template(job.id, packet)
    assert previous.read_bytes() == b"prior editable deck stays until promotion"
    assert load_job(job.id).approvals == {"brief": True, "prototype": False, "final": False}
    result = build(root, request)
    assert result.status == AdapterStatus.SUCCEEDED, result.failure
    assert result.receipt.engine == hybrid_job.ENGINE
    candidate = root / next(a.path for a in result.artifacts if a.kind == "pptx")
    plan = read_model(root / "planning/deck-plan.json", DeckPlan)
    # Fixtures prove artifact promotion, not visual quality. Production renders
    # actual candidate pages through the standard _ensure_adapter_previews.
    for slide in plan.slides:
        preview = candidate.parent / f"{slide.id}.png"
        Image.new("RGB", (20, 12), "blue").save(preview)
        result.artifacts.append(AdapterArtifact(kind="preview", path=preview.relative_to(root).as_posix(), sha256=sha256(preview)))
    _promote_adapter_result(job, root, result)
    assert not candidate.exists(), "successful promotion must not retain a duplicate staged PPTX"
    manifest, previews = _validate_external_output(job, root, plan)
    assert manifest.engine == hybrid_job.ENGINE and len(previews) == 2
    assert not read_json(root / hybrid_job.RECEIPT)["visual_approved"]
    # Repacking same package bytes (e.g. ZIP timestamps) does not invalidate
    # source identity, but changing an untouched theme does.
    deck = root / "build/deck.pptx"
    with ZipFile(deck) as z:
        members = {name: z.read(name) for name in z.namelist()}
    members["ppt/theme/theme1.xml"] = members["ppt/theme/theme1.xml"].replace(b"Office Theme", b"Changed Theme")
    with ZipFile(deck, "w") as z:
        for name, data in members.items():
            z.writestr(name, data)
    assert "differs" in _template_application_artifact_reason(job, root, plan, manifest.template_application,
                                                               manifest.build_source_sha256, deck)


def test_missing_current_copy_blocks_without_replacing_current_deck(bound_job):
    root, job, request, packet = bound_job
    packet["text_replacements"] = packet["text_replacements"][:1]
    hybrid_job.bind_hybrid_template(job.id, packet)
    current = root / "build/deck.pptx"
    current.write_bytes(b"current")
    result = build(root, request)
    assert result.status == AdapterStatus.BLOCKED
    assert "mapping is incomplete" in result.failure.message
    record = read_json(root / hybrid_job.RECEIPT)
    assert record["content_check"]["blockers"] and not record["visual_approved"]
    assert current.read_bytes() == b"current"


@pytest.mark.parametrize("target", ["plan", "source", "preview"])
def test_stale_binding_rejects_real_input_drift(bound_job, target):
    root, job, request, packet = bound_job
    hybrid_job.bind_hybrid_template(job.id, packet)
    binding = read_json(root / hybrid_job.BINDING)
    if target == "plan":
        path = root / "planning/deck-plan.json"
        value = read_json(path)
        value["slides"][0]["title"] = "Changed request"
        write_json(path, value)
    elif target == "source":
        path = next((root / "sources").iterdir())
        path.write_bytes(b"changed source")
    else:
        path = root / binding["request"]["bound_dependencies"][0]["path"]
        path.write_bytes(b"changed preview")
    result = build(root, request)
    assert result.status == AdapterStatus.BLOCKED
    assert not (root / hybrid_job.RECEIPT).exists()


def test_bad_primary_or_route_never_mutates_job(bound_job):
    root, job, _, packet = bound_job
    before = (root / "job.json").read_bytes()
    wrong = {**packet, "primary_source_sha256": "0" * 64}
    with pytest.raises(ValueError, match="selected template"):
        hybrid_job.bind_hybrid_template(job.id, wrong)
    assert (root / "job.json").read_bytes() == before
    job.route = ProductionRoute.NATIVE
    save_job(job)
    with pytest.raises(ValueError, match="selected template job"):
        hybrid_job.bind_hybrid_template(job.id, packet)
    assert not (root / hybrid_job.BINDING).exists()


def test_page_proof_cannot_be_changed_independently(bound_job):
    root, job, request, packet = bound_job
    hybrid_job.bind_hybrid_template(job.id, packet)
    result = build(root, request)
    assert result.status == AdapterStatus.SUCCEEDED
    proof = copy.deepcopy(result.metadata["template_application"])
    proof["pages"][0]["deck_slide_id"] = "other-slide"
    deck = root / next(a.path for a in result.artifacts if a.kind == "pptx")
    assert "page proof" in hybrid_job.verify_hybrid_output(job, root,
        read_model(root / "planning/deck-plan.json", DeckPlan), proof, sha256(root / hybrid_job.BINDING), deck)


def test_replacement_image_is_stored_and_hash_bound(bound_job, tmp_path):
    root, _, _, packet = bound_job
    pic = tmp_path / "replacement.png"
    Image.new("RGB", (32, 20), "red").save(pic)
    packet["image_replacements"] = [{"slide": 1, "shape_id": "3", "path": str(pic)}]
    stored = hybrid_job._materialize_inputs(root, packet)
    image_path = Path(stored["image_replacements"][0]["image_path"])
    assert image_path.is_relative_to(root) and sha256(image_path) == sha256(pic)
    assert "path" not in stored["image_replacements"][0]
    assert str(pic) == packet["image_replacements"][0]["path"]


def test_planned_copy_cannot_be_satisfied_by_substrings_or_cross_shape_text(bound_job):
    root, job, request, packet = bound_job
    packet["text_replacements"][0]["new_text"] = "Prefix Fixture title 1 suffix"
    hybrid_job.bind_hybrid_template(job.id, packet)
    result = build(root, request)
    assert result.status == AdapterStatus.BLOCKED
    blockers = read_json(root / hybrid_job.RECEIPT)["content_check"]["blockers"]
    assert any(row.get("text") == "Fixture title 1" for row in blockers)


def test_generated_previews_share_engine_stage_and_cleanup_errors_do_not_undo_promotion(bound_job, monkeypatch):
    root, job, request, packet = bound_job
    hybrid_job.bind_hybrid_template(job.id, packet)
    result = build(root, request)
    monkeypatch.setenv("AIPPT_FAKE_RENDER", "1")
    result = _ensure_adapter_previews(job, root, result)
    deck = next(a for a in result.artifacts if a.kind == "pptx")
    assert all(Path(a.path).parent.parent == Path(deck.path).parent for a in result.artifacts if a.kind == "preview")
    original_unlink = Path.unlink
    def denied_staging(path, *args, **kwargs):
        if path.name == "deck.pptx" and path.parent.name.startswith("hybrid-"):
            raise OSError("simulated file in use")
        return original_unlink(path, *args, **kwargs)
    monkeypatch.setattr(Path, "unlink", denied_staging)
    _promote_adapter_result(job, root, result)
    _validate_external_output(job, root, read_model(root / "planning/deck-plan.json", DeckPlan))
    assert (root / "build/deck.pptx").is_file()
