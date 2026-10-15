from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Mm, Pt
import pytest

from aippt_studio.pptx_edit import inspect_pptx, patch_pptx
from aippt_studio.revision_qa import audit_revision


def _source(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    first = deck.slides.add_slide(deck.slide_layouts[6])
    for x, text in ((1, "左侧说明"), (4, "重叠说明")):
        shape = first.shapes.add_textbox(Inches(x), Inches(2), Inches(5), Inches(1))
        shape.text = text
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    second = deck.slides.add_slide(deck.slide_layouts[6])
    untouched = second.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    untouched.text = "没有修改的第二页"
    deck.save(path)
    return path


def _candidate(source: Path, output: Path, *, x_mm: int):
    obj = inspect_pptx(source, slides=[1])["slides"][0]["objects"][1]
    request = {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
               "edits": [{"action": "set", "slide": 1, "shape_id": obj["shape_id"],
                          "geometry": {"unit": "mm", "x": x_mm}}]}
    receipt = patch_pptx(request, output)
    return output / "candidate.pptx", receipt


def test_local_revision_qa_distinguishes_cleared_baseline_issue_and_scopes_page(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    candidate, receipt = _candidate(source, tmp_path / "patch", x_mm=200)

    result = audit_revision(source, candidate, receipt)

    assert result["status"] == "passed_scoped"
    assert result["scope"] == {"kind": "selected_slides", "selected_slides": [1],
                               "full_deck_verified": False, "whole_deck_pass_claimed": False}
    assert result["quality"]["source_open_issue_count"] == 1
    assert result["quality"]["candidate_open_issue_count"] == 0
    assert [issue["code"] for issue in result["quality"]["cleared_issues"]] == ["TEXT_OVERLAP"]
    assert not result["quality"]["new_issues"]
    assert result["visual_review"]["required"] is True
    assert result["visual_review"]["scope"] == "target_slides_only"


def test_local_revision_qa_flags_new_out_of_bounds_issue_and_recommends_context_review(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    candidate, receipt = _candidate(source, tmp_path / "patch", x_mm=350)

    result = audit_revision(source, candidate, receipt)

    assert result["status"] == "needs_review"
    assert [issue["code"] for issue in result["quality"]["new_issues"]] == ["OUT_OF_BOUNDS"]
    assert "assistant_review_changed_quality_findings" in result["visual_review"]["recommended"]


def test_local_revision_qa_retains_unresolved_baseline_issue_as_persistent(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    candidate, receipt = _candidate(source, tmp_path / "patch", x_mm=120)

    result = audit_revision(source, candidate, receipt)

    assert result["status"] == "passed_scoped"
    assert result["quality"]["new_issues"] == []
    assert [issue["code"] for issue in result["quality"]["persistent_issues"]] == ["TEXT_OVERLAP"]
    assert result["quality"]["candidate_open_issue_count"] == 1


def test_local_revision_qa_marks_increased_overlap_as_worsened_not_persistent(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    candidate, receipt = _candidate(source, tmp_path / "patch", x_mm=50)

    result = audit_revision(source, candidate, receipt)

    assert result["status"] == "needs_review"
    assert result["quality"]["persistent_issues"] == []
    worsened, = result["quality"]["worsened_issues"]
    assert worsened["code"] == "TEXT_OVERLAP"
    assert worsened["worsening"]["metric"] == "overlap_ratio"
    assert worsened["worsening"]["candidate_value"] > worsened["worsening"]["baseline_value"]


def test_local_revision_qa_marks_decreased_font_size_as_worsened(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    obj = inspect_pptx(source, slides=[1])["slides"][0]["objects"][0]
    request = {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
               "edits": [{"action": "set", "slide": 1, "shape_id": obj["shape_id"],
                          "font_size_pt": 18, "autofit": "none"}]}
    receipt = patch_pptx(request, tmp_path / "patch")
    candidate = tmp_path / "patch" / "candidate.pptx"
    policy = {"typography_mode": "explicit", "minimum_font_sizes": {"body": 30}}

    result = audit_revision(source, candidate, receipt, policy=policy)

    assert result["status"] == "needs_review"
    worsened, = result["quality"]["worsened_issues"]
    assert worsened["code"] == "SMALL_FONT"
    assert worsened["worsening"]["baseline_value"] == 24
    assert worsened["worsening"]["candidate_value"] == 18


def test_local_revision_qa_marks_increased_overflow_as_worsened(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    deck = Presentation(source)
    deck.slides[0].shapes[0].left = Mm(300)
    deck.save(source)
    obj = inspect_pptx(source, slides=[1])["slides"][0]["objects"][0]
    request = {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
               "edits": [{"action": "set", "slide": 1, "shape_id": obj["shape_id"],
                          "geometry": {"unit": "mm", "x": 320}}]}
    receipt = patch_pptx(request, tmp_path / "patch")

    result = audit_revision(source, tmp_path / "patch" / "candidate.pptx", receipt)

    worsened, = result["quality"]["worsened_issues"]
    assert result["status"] == "needs_review"
    assert worsened["code"] == "OUT_OF_BOUNDS"
    assert worsened["worsening"]["metric"] == "overflow_pt"
    assert worsened["worsening"]["candidate_value"] > worsened["worsening"]["baseline_value"]


def test_replace_text_requires_target_preview_and_limits_assistant_review_to_meaning_or_hierarchy(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    obj = inspect_pptx(source, slides=[1])["slides"][0]["objects"][1]
    request = {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
               "edits": [{"action": "replace_text", "slide": 1, "shape_id": obj["shape_id"],
                          "expected_text": "重叠说明", "replacement_text": "更新后的说明文字"}]}
    receipt = patch_pptx(request, tmp_path / "patch")

    result = audit_revision(source, tmp_path / "patch" / "candidate.pptx", receipt)

    assert result["actions"] == ["replace_text"]
    assert result["visual_review"]["required"] is True
    assert result["visual_review"]["recommended"] == ["render_target_page_preview", "inspect_actual_page_image"]
    assert result["visual_review"]["assistant_review"] == "only_if_copy_meaning_or_hierarchy_changed"
    assert "PDF" not in json.dumps(result["visual_review"])


def test_local_revision_qa_rejects_stale_source_candidate_and_mismatched_package_receipt(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    candidate, receipt = _candidate(source, tmp_path / "patch", x_mm=200)

    altered_source = tmp_path / "changed-source.pptx"
    shutil.copy2(source, altered_source)
    altered_source.write_bytes(altered_source.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="source SHA-256"):
        audit_revision(altered_source, candidate, receipt)
    altered_candidate = tmp_path / "other-candidate.pptx"
    shutil.copy2(candidate, altered_candidate)
    altered_candidate.write_bytes(altered_candidate.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="candidate SHA-256"):
        audit_revision(source, altered_candidate, {**receipt, "candidate": str(altered_candidate)})
    altered_receipt = {**receipt, "changed_parts": []}
    with pytest.raises(ValueError, match="changed_parts"):
        audit_revision(source, candidate, altered_receipt)


def test_local_revision_qa_accepts_receipt_file_but_rejects_a_different_candidate_path(tmp_path: Path):
    source = _source(tmp_path / "source.pptx")
    candidate, receipt = _candidate(source, tmp_path / "patch", x_mm=200)
    receipt_file = tmp_path / "patch" / "receipt.json"
    receipt_file.write_text(json.dumps(receipt), encoding="utf-8")

    result = audit_revision(source, candidate, receipt_file)
    assert result["candidate_sha256"] == receipt["candidate_sha256"]
    with pytest.raises(ValueError, match="candidate path"):
        audit_revision(source, source, receipt_file)
