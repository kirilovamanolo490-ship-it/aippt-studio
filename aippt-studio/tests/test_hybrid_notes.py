from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from lxml import etree as LET
from pptx import Presentation

from aippt_studio.hybrid_notes import update_hybrid_notes
from aippt_studio.models import DeckPlan, SlideSpec


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"


def _plan(*notes: str) -> DeckPlan:
    return DeckPlan(
        title="Notes test",
        slides=[
            SlideSpec(
                id=f"page-{index}",
                title=f"Page {index}",
                purpose="test notes",
                layout="content",
                speaker_notes=note,
            )
            for index, note in enumerate(notes, 1)
        ],
    )


def _make_deck(path: Path, *, existing_notes: list[str] | None) -> None:
    presentation = Presentation()
    for index in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(100000, 100000, 3000000, 400000)
        shape.text = f"untouched slide body {index + 1}"
        if existing_notes is not None:
            slide.notes_slide.notes_text_frame.text = existing_notes[index]
    presentation.save(path)


def _notes_text_by_slide(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
        rels = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
        rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        rel_by_id = {rel.get("Id"): rel for rel in rels.findall(f"{{{rels_ns}}}Relationship")}
        slides = presentation.find(f"{{{P_NS}}}sldIdLst")
        slide_parts = []
        for slide_id in slides.findall(f"{{{P_NS}}}sldId"):
            rel = rel_by_id[slide_id.get(f"{{{R_NS}}}id")]
            slide_parts.append("ppt/" + rel.get("Target").removeprefix("ppt/"))
        result = []
        for slide_part in slide_parts:
            rels_part = slide_part.rsplit("/", 1)[0] + "/_rels/" + slide_part.rsplit("/", 1)[1] + ".rels"
            slide_rels = ET.fromstring(archive.read(rels_part))
            note_rel = next(
                rel for rel in slide_rels.findall(f"{{{rels_ns}}}Relationship")
                if rel.get("Type", "").endswith("/notesSlide")
            )
            notes_part = "ppt/" + re.sub(r"^\.\./", "", note_rel.get("Target"))
            notes_root = ET.fromstring(archive.read(notes_part))
            body_shapes = []
            for shape in notes_root.findall(f".//{{{P_NS}}}sp"):
                placeholder = shape.find(f"./{{{P_NS}}}nvSpPr/{{{P_NS}}}nvPr/{{{P_NS}}}ph")
                if placeholder is not None and placeholder.get("type") == "body":
                    body = shape.find(f"{{{P_NS}}}txBody")
                    body_shapes.append("\n".join(
                        "".join(node.text or "" for node in paragraph.iter(f"{{{A_NS}}}t"))
                        for paragraph in body.findall(f"{{{A_NS}}}p")
                    ).strip())
            result.append(next((text for text in body_shapes if text), ""))
        return result


R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _allowed_changed_part(name: str) -> bool:
    return (
        name == "[Content_Types].xml"
        or name == "ppt/presentation.xml"
        or name == "ppt/_rels/presentation.xml.rels"
        or bool(re.fullmatch(r"ppt/slides/_rels/slide\d+\.xml\.rels", name))
        or name.startswith("ppt/notesSlides/")
        or name.startswith("ppt/notesMasters/")
    )


def _assert_non_notes_parts_identical(before: dict[str, bytes], after: dict[str, bytes]) -> None:
    names = (set(before) | set(after)) - {name for name in set(before) | set(after) if _allowed_changed_part(name)}
    assert {name: before.get(name) for name in names} == {name: after.get(name) for name in names}


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {item.filename: archive.read(item.filename) for item in archive.infolist()}


def _rewrite_package(path: Path, updates: dict[str, bytes], remove: set[str] | None = None) -> None:
    remove = remove or set()
    temp = path.with_suffix(".rewritten.pptx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w") as target:
        for info in source.infolist():
            if info.filename not in remove:
                target.writestr(info, updates.get(info.filename, source.read(info.filename)))
    temp.replace(path)


def _ignorable_xml(data: bytes) -> bytes:
    root = LET.fromstring(data)
    nsmap = {**root.nsmap, "mc": MC_NS, "p14": P14_NS}
    updated = LET.Element(root.tag, nsmap=nsmap)
    for name, value in root.attrib.items():
        updated.set(name, value)
    updated.set(f"{{{MC_NS}}}Ignorable", "p14")
    updated.text = root.text
    updated.tail = root.tail
    for child in root:
        updated.append(child)
    return LET.tostring(updated, encoding="utf-8", xml_declaration=True)


def test_update_replaces_old_notes_and_preserves_existing_notes_master_and_non_notes_parts(tmp_path: Path) -> None:
    deck = tmp_path / "existing-notes.pptx"
    _make_deck(
        deck,
        existing_notes=[
            "旧项目名和事实。\n[Sources]\n/Users/old/private/source.pptx",
            "旧文件位置 C:\\secret\\draft.pptx",
        ],
    )
    before = _parts(deck)
    masters_before = {name: data for name, data in before.items() if name.startswith("ppt/notesMasters/")}

    result = update_hybrid_notes(deck, _plan("新版讲稿：说明当前结论。", "解释新版数据来源。"))

    after = _parts(deck)
    assert result["status"] == "updated"
    assert _notes_text_by_slide(deck) == ["新版讲稿：说明当前结论。", "解释新版数据来源。"]
    assert not any("旧项目名" in data.decode("utf-8", errors="ignore") for name, data in after.items() if name.startswith("ppt/notesSlides/"))
    assert {name: data for name, data in after.items() if name.startswith("ppt/notesMasters/")} == masters_before
    _assert_non_notes_parts_identical(before, after)


def test_update_adds_notes_to_deck_without_notes_and_sanitizes_plan_notes(tmp_path: Path) -> None:
    deck = tmp_path / "no-notes.pptx"
    _make_deck(deck, existing_notes=None)
    original_parts = _parts(deck)
    _rewrite_package(
        deck,
        {
            "ppt/presentation.xml": _ignorable_xml(original_parts["ppt/presentation.xml"]),
            "ppt/slides/_rels/slide1.xml.rels": _ignorable_xml(
                original_parts["ppt/slides/_rels/slide1.xml.rels"]
            ),
        },
    )
    before = _parts(deck)
    assert not any(name.startswith("ppt/notesSlides/") for name in before)

    result = update_hybrid_notes(
        deck,
        _plan(
            "面向听众解释结论。\n[Sources]\n内部来源",
            "保留本页说明。\n路径：/Users/person/private/file.pptx",
        ),
    )

    after = _parts(deck)
    assert result["slide_count"] == 2
    assert _notes_text_by_slide(deck) == ["面向听众解释结论。", "保留本页说明。"]
    assert any(name.startswith("ppt/notesMasters/") for name in after)
    for part in ("ppt/presentation.xml", "ppt/slides/_rels/slide1.xml.rels"):
        root = LET.fromstring(after[part])
        assert root.get(f"{{{MC_NS}}}Ignorable") == "p14"
        assert root.nsmap.get("p14") == P14_NS
    _assert_non_notes_parts_identical(before, after)


def test_slide_count_mismatch_does_not_modify_candidate(tmp_path: Path) -> None:
    deck = tmp_path / "mismatch.pptx"
    _make_deck(deck, existing_notes=None)
    before = deck.read_bytes()

    with pytest.raises(ValueError, match="Plan has 1 slides but candidate has 2 slides"):
        update_hybrid_notes(deck, _plan("Only one note."))

    assert deck.read_bytes() == before


def test_rejects_multiple_slides_sharing_one_notes_part(tmp_path: Path) -> None:
    deck = tmp_path / "shared-notes.pptx"
    _make_deck(deck, existing_notes=["First notes", "Second notes"])
    with zipfile.ZipFile(deck) as archive:
        rels = LET.fromstring(archive.read("ppt/slides/_rels/slide2.xml.rels"))
    notes_rel = next(rel for rel in rels if rel.get("Type", "").endswith("/notesSlide"))
    notes_rel.set("Target", "../notesSlides/notesSlide1.xml")
    _rewrite_package(
        deck,
        {"ppt/slides/_rels/slide2.xml.rels": LET.tostring(rels, encoding="utf-8", xml_declaration=True)},
    )
    before = deck.read_bytes()

    with pytest.raises(ValueError, match="Multiple slides share one notes part"):
        update_hybrid_notes(deck, _plan("New first", "New second"))

    assert deck.read_bytes() == before


def test_rejects_missing_notes_master_part(tmp_path: Path) -> None:
    deck = tmp_path / "missing-notes-master.pptx"
    _make_deck(deck, existing_notes=["Old one", "Old two"])
    before_parts = _parts(deck)
    master = next(name for name in before_parts if re.fullmatch(r"ppt/notesMasters/notesMaster\d+\.xml", name))
    _rewrite_package(deck, {}, remove={master})
    before = deck.read_bytes()

    with pytest.raises(ValueError, match="notes master relationship points to missing part"):
        update_hybrid_notes(deck, _plan("New one", "New two"))

    assert deck.read_bytes() == before
