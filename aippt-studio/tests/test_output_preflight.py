from __future__ import annotations

import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from typer.testing import CliRunner

from aippt_studio import cli
from aippt_studio.output_preflight import output_preflight


def _deck(path: Path, count: int = 1) -> Path:
    deck = Presentation()
    for number in range(count):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        box.text = f"结论页 {number + 1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)
    return path


def _write(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def test_current_output_identity_and_local_audit_are_bound_to_exact_bytes(tmp_path: Path):
    work = tmp_path / "work" / "1"
    output = _deck(work / "output" / "final.pptx", 2)
    before = output.read_bytes()

    report = output_preflight(work)

    assert report["status"] == "passed"
    assert report["identity"] == {
        "sha256": hashlib.sha256(before).hexdigest(),
        "bytes": len(before),
        "slide_count": 2,
    }
    assert report["quality"]["source_sha256"] == report["identity"]["sha256"]
    assert report["quality"]["summary"]["total_deck_slides"] == 2
    assert output.read_bytes() == before


def test_old_manifest_and_quality_report_are_marked_stale_after_output_changes(tmp_path: Path):
    work = tmp_path / "work" / "2"
    old = _deck(work / "output" / "old.pptx", 1)
    old_sha = hashlib.sha256(old.read_bytes()).hexdigest()
    old.rename(work / "output" / "final.pptx")
    qa = work / "workspace" / "qa"
    _write(qa / "final-output-manifest.json", {
        "outputs": [{"path": "output/final.pptx", "sha256": old_sha}],
        "qa": {"rendered_pages": 1},
    })
    _write(qa / "presentation-quality.json", {
        "source_sha256": old_sha,
        "status": "passed",
        "scope": {"full_deck_verified": True},
        "summary": {"total_deck_slides": 1},
    })
    _deck(work / "output" / "final.pptx", 2)

    report = output_preflight(work)

    assert report["status"] == "passed"
    assert report["evidence_summary"]["stale"] == 2
    assert {row["status"] for row in report["evidence"]} == {"stale"}
    assert all("differs" in row["reason"] for row in report["evidence"])


def test_matching_hash_with_wrong_page_count_is_stale_and_unhashed_claim_is_unbound(tmp_path: Path):
    work = tmp_path / "work" / "3"
    output = _deck(work / "output" / "final.pptx", 2)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    qa = work / "workspace" / "qa"
    _write(qa / "presentation-quality.json", {
        "source_sha256": digest,
        "status": "passed",
        "scope": {"full_deck_verified": True},
        "summary": {"total_deck_slides": 1},
    })
    _write(qa / "final-output-manifest.json", {
        "outputs": [{"path": "output/final.pptx", "bytes": output.stat().st_size}],
    })

    report = output_preflight(work)

    by_name = {Path(row["path"]).name: row["status"] for row in report["evidence"]}
    assert by_name == {"presentation-quality.json": "stale", "final-output-manifest.json": "unbound"}


def test_stale_terminal_powerpoint_proofs_block_preflight(tmp_path: Path):
    work = tmp_path / "work" / "proof-stale"
    output = _deck(work / "output" / "final.pptx", 2)
    old_sha = hashlib.sha256(_deck(tmp_path / "old.pptx", 1).read_bytes()).hexdigest()
    qa = work / "workspace" / "qa"
    _write(qa / "native-roundtrip-v8.json", {
        "status": "passed", "sha256": old_sha, "slides": 1,
        "roundtrip": "open_save_close_reopen",
    })
    _write(qa / "package-audit-v9.json", {
        "status": "passed", "sha256": old_sha, "slides": 1,
    })
    _write(qa / "final-delivery-v9.json", {
        "status": "ready_for_user_review", "slides": 1,
        "deliverables": [{"kind": "editable_powerpoint", "path": "output/final.pptx",
                          "sha256": old_sha}],
    })

    report = output_preflight(work)

    assert report["status"] == "needs_review"
    assert {issue["family"] for issue in report["delivery_proof_issues"]} == {
        "native-roundtrip", "package-audit", "final-delivery",
    }
    assert report["identity"]["sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert report["identity"]["slide_count"] == 2


def test_current_terminal_proof_supersedes_historical_receipts(tmp_path: Path):
    work = tmp_path / "work" / "proof-current"
    output = _deck(work / "output" / "final.pptx", 2)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    qa = work / "workspace" / "qa"
    _write(qa / "native-roundtrip-v8.json", {
        "status": "passed", "sha256": "0" * 64, "slides": 1,
    })
    _write(qa / "native-roundtrip-current.json", {
        "status": "passed", "sha256": digest, "slides": 2,
    })

    report = output_preflight(work)

    assert report["status"] == "passed"
    assert report["delivery_proof_issues"] == []
    assert report["evidence_summary"]["stale"] == 1


def test_fresh_terminal_receipts_allow_rebuilding_a_stale_closure_manifest(tmp_path: Path):
    work = tmp_path / "work" / "reclose"
    output = _deck(work / "output" / "final.pptx", 2)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    old_sha = "0" * 64
    qa = work / "workspace" / "qa"
    _write(qa / "final-output-manifest.json", {
        "work_id": "reclose", "status": "review",
        "outputs": [{"path": "output/final.pptx", "sha256": old_sha}],
    })
    _write(qa / "native-roundtrip-current.json", {
        "status": "passed", "sha256": digest, "slides": 2,
    })
    _write(qa / "package-audit-current.json", {
        "status": "passed", "sha256": digest, "slides": 2,
    })
    _write(qa / "final-delivery-current.json", {
        "status": "ready_for_user_review", "slides": 2,
        "deliverables": [{"kind": "editable_powerpoint", "path": "output/final.pptx",
                          "sha256": digest}],
    })

    report = output_preflight(work)

    assert report["status"] == "passed"
    assert report["delivery_proof_issues"] == []
    assert {row["status"] for row in report["evidence"]} == {"stale", "current"}


def test_single_delivery_row_for_another_output_cannot_borrow_top_level_identity(tmp_path: Path):
    work = tmp_path / "work" / "wrong-single-row"
    output = _deck(work / "output" / "final.pptx", 2)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    qa = work / "workspace" / "qa"
    _write(qa / "final-delivery-current.json", {
        "status": "ready_for_user_review", "source_sha256": digest, "slides": 2,
        "deliverables": [{"path": "output/another.pptx", "sha256": digest, "slides": 2}],
    })

    blocked = output_preflight(work)

    assert blocked["status"] == "needs_review"
    assert blocked["delivery_proof_issues"][0]["family"] == "final-delivery"
    assert blocked["evidence"][0]["status"] == "unbound"

    _write(qa / "final-delivery-current.json", {
        "status": "ready_for_user_review", "source_sha256": digest, "slides": 2,
        "deliverables": [{"kind": "editable_powerpoint"}],
    })
    legacy = output_preflight(work)
    assert legacy["status"] == "passed"
    assert legacy["delivery_proof_issues"] == []


def test_matching_hash_is_insufficient_when_page_count_or_proof_status_fails(tmp_path: Path):
    work = tmp_path / "work" / "proof-incomplete"
    output = _deck(work / "output" / "final.pptx", 2)
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    qa = work / "workspace" / "qa"
    _write(qa / "native-roundtrip-current.json", {
        "status": "passed", "sha256": digest, "slides": 1,
    })
    _write(qa / "package-audit-current.json", {
        "status": "failed", "sha256": digest, "slides": 2,
    })

    report = output_preflight(work)

    assert report["status"] == "needs_review"
    assert {issue["family"] for issue in report["delivery_proof_issues"]} == {
        "native-roundtrip", "package-audit",
    }
    assert {row["status"] for row in report["evidence"]} == {"stale", "failed"}


def test_legacy_work_without_terminal_receipts_is_not_blocked_by_new_proof_gate(tmp_path: Path):
    work = tmp_path / "work" / "legacy-no-proof"
    _deck(work / "output" / "final.pptx", 1)

    report = output_preflight(work)

    assert report["status"] == "passed"
    assert report["delivery_proof_issues"] == []


def test_two_pptx_outputs_use_only_the_matching_manifest_and_roundtrip_rows(tmp_path: Path):
    work = tmp_path / "work" / "two-decks"
    first = _deck(work / "output" / "first.pptx", 1)
    second = _deck(work / "output" / "second.pptx", 2)
    first_sha = hashlib.sha256(first.read_bytes()).hexdigest()
    second_sha = hashlib.sha256(second.read_bytes()).hexdigest()
    qa = work / "workspace" / "qa"
    _write(qa / "final-delivery-current.json", {
        "status": "ready_for_user_review",
        "deliverables": [
            {"kind": "editable_powerpoint", "path": "output/first.pptx",
             "sha256": first_sha, "slides": 1},
            {"kind": "editable_powerpoint", "path": "output/second.pptx",
             "sha256": second_sha, "slides": 2},
        ],
    })
    _write(qa / "native-roundtrip-first.json", {
        "status": "passed", "sha256": first_sha, "slides": 1,
    })
    _write(qa / "native-roundtrip-second.json", {
        "status": "passed", "sha256": second_sha, "slides": 2,
    })

    first_report = output_preflight(work, "output/first.pptx")
    second_report = output_preflight(work, "output/second.pptx")

    assert first_report["status"] == second_report["status"] == "passed"
    for report, digest, pages in ((first_report, first_sha, 1), (second_report, second_sha, 2)):
        current_delivery = [row for row in report["evidence"]
                            if row["path"].endswith("final-delivery-current.json")]
        current_roundtrips = [row for row in report["evidence"]
                              if row["path"].endswith(f"native-roundtrip-{'first' if pages == 1 else 'second'}.json")]
        assert len(current_delivery) == len(current_roundtrips) == 1
        assert current_delivery[0]["status"] == current_roundtrips[0]["status"] == "current"
        assert report["identity"]["sha256"] == digest
        assert report["identity"]["slide_count"] == pages

    # If the second deck's roundtrip is replaced with the first deck's identity,
    # the first proof cannot unlock the second output.
    _write(qa / "native-roundtrip-second.json", {
        "status": "passed", "sha256": first_sha, "slides": 1,
    })
    blocked_second = output_preflight(work, "output/second.pptx")
    assert blocked_second["status"] == "needs_review"
    assert blocked_second["delivery_proof_issues"] == [{
        "code": "STALE_DELIVERY_PROOF",
        "family": "native-roundtrip",
        "message": "native-roundtrip receipts exist, but none is bound to the current PPTX SHA-256 and slide count",
    }]


def test_multi_deck_final_delivery_never_falls_back_to_top_level_identity_or_pages(tmp_path: Path):
    for case in ("missing-second-sha", "missing-second-pages"):
        work = tmp_path / "work" / case
        first = _deck(work / "output" / "first.pptx", 1)
        second = _deck(work / "output" / "second.pptx", 2)
        first_sha = hashlib.sha256(first.read_bytes()).hexdigest()
        second_sha = hashlib.sha256(second.read_bytes()).hexdigest()
        second_row = {"kind": "editable_powerpoint", "path": "output/second.pptx"}
        if case == "missing-second-sha":
            second_row["slides"] = 2
        else:
            second_row.update(sha256=second_sha)
        qa = work / "workspace" / "qa"
        _write(qa / "final-delivery-current.json", {
            "status": "ready_for_user_review",
            "source_sha256": first_sha,
            "slides": 2,
            "qa": {"rendered_pages": 2},
            "deliverables": [
                {"kind": "editable_powerpoint", "path": "output/first.pptx",
                 "sha256": first_sha, "slides": 1},
                second_row,
            ],
        })
        _write(qa / "native-roundtrip-second.json", {
            "status": "passed", "sha256": second_sha, "slides": 2,
        })

        report = output_preflight(work, "output/second.pptx")

        final_delivery = next(row for row in report["evidence"]
                              if row["path"].endswith("final-delivery-current.json"))
        assert report["status"] == "needs_review"
        assert final_delivery["status"] == "unbound"
        assert "no recognized artifact SHA-256" in final_delivery["reason"] or \
            "no slide count" in final_delivery["reason"]


def test_multiple_or_missing_outputs_fail_closed(tmp_path: Path):
    work = tmp_path / "work" / "4"
    report = output_preflight(work)
    assert report["status"] == "errors"
    _deck(work / "output" / "one.pptx")
    _deck(work / "output" / "two.pptx")
    report = output_preflight(work)
    assert report["status"] == "errors"
    assert report["errors"][0]["code"] == "OUTPUT_SELECTION"


def test_specific_output_path_must_exist_and_be_pptx(tmp_path: Path):
    work = tmp_path / "work" / "5"
    report = output_preflight(work, "workspace/not-there.pptx")
    assert report["status"] == "errors"
    assert report["errors"][0]["code"] == "OUTPUT_MISSING"
    outside = _deck(tmp_path / "outside.pptx")
    scoped = output_preflight(work, outside)
    assert scoped["errors"][0]["code"] == "OUTPUT_SCOPE"


def test_quality_policy_requires_exact_current_source_hash(tmp_path: Path):
    work = tmp_path / "work" / "6"
    output = _deck(work / "output" / "final.pptx")
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    policy = {"exceptions": []}

    missing = output_preflight(work, policy=policy)
    stale = output_preflight(work, policy=policy, policy_source_sha256="0" * 64)
    current = output_preflight(work, policy=policy, policy_source_sha256=digest)

    assert missing["errors"][0]["code"] == "POLICY_SOURCE_SHA_REQUIRED"
    assert stale["errors"][0]["code"] == "POLICY_SOURCE_SHA_MISMATCH"
    assert current["status"] == "passed"
    assert current["quality_policy_source_sha256"] == digest


def test_output_preflight_cli_saves_current_form_review_and_blocks_open_overlap(tmp_path: Path):
    work = tmp_path / "work" / "7"
    _deck(work / "output" / "final.pptx")
    report_path = work / "workspace" / "qa" / "current-output-preflight.json"
    result = CliRunner().invoke(cli.app, ["output-preflight", "--work-dir", str(work),
                                            "--form", "teaching", "--output", str(report_path), "--json"])
    assert result.exit_code == 0, result.output
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["status"] == "passed"
    assert report["form_review"]["presentation_form"] == "teaching"
    assert report["form_review"]["source_sha256"] == report["identity"]["sha256"]

    deck = Presentation(work / "output" / "final.pptx")
    slide = deck.slides[0]
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = "重叠文字"
    deck.save(work / "output" / "final.pptx")
    blocked = CliRunner().invoke(cli.app, ["output-preflight", "--work-dir", str(work),
                                             "--output", str(report_path), "--json"])
    assert blocked.exit_code == 1, blocked.output
    assert json.loads(report_path.read_text(encoding="utf-8"))["status"] == "needs_review"


def test_output_preflight_cli_exit_code_blocks_stale_roundtrip_receipt(tmp_path: Path):
    work = tmp_path / "work" / "stale-roundtrip-cli"
    _deck(work / "output" / "final.pptx", 2)
    qa = work / "workspace" / "qa"
    _write(qa / "native-roundtrip-v8.json", {
        "status": "passed", "sha256": "0" * 64, "slides": 1,
    })
    report_path = qa / "output-preflight.json"

    result = CliRunner().invoke(cli.app, ["output-preflight", "--work-dir", str(work),
                                            "--output", str(report_path), "--json"])

    assert result.exit_code == 1, result.output
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["status"] == "needs_review"
    assert report["delivery_proof_issues"] == [{
        "code": "STALE_DELIVERY_PROOF",
        "family": "native-roundtrip",
        "message": "native-roundtrip receipts exist, but none is bound to the current PPTX SHA-256 and slide count",
    }]


def test_output_preflight_cli_replaces_old_report_with_structured_input_error(tmp_path: Path):
    work = tmp_path / "work" / "8"
    _deck(work / "output" / "final.pptx")
    report_path = work / "workspace" / "qa" / "current.json"
    _write(report_path, {"status": "passed"})
    invalid_policy = tmp_path / "invalid-policy.json"
    invalid_policy.write_text("{broken", encoding="utf-8")

    failed = CliRunner().invoke(cli.app, ["output-preflight", "--work-dir", str(work),
                                            "--policy", str(invalid_policy), "--output", str(report_path), "--json"])

    assert failed.exit_code == 2
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["status"] == "errors"
    assert report["errors"][0]["code"] == "PREFLIGHT_INPUT_ERROR"


def test_output_preflight_cli_reports_invalid_pptx_without_form_traceback(tmp_path: Path):
    work = tmp_path / "work" / "9"
    source = work / "output" / "corrupt.pptx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"not a PowerPoint package")
    report_path = work / "workspace" / "qa" / "current.json"

    failed = CliRunner().invoke(cli.app, ["output-preflight", "--work-dir", str(work),
                                            "--output", str(report_path), "--json"])

    assert failed.exit_code == 2
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["status"] == "errors"
    assert report["quality"]["errors"][0]["code"] == "AUDIT_ERROR"
    assert "Traceback" not in failed.output
