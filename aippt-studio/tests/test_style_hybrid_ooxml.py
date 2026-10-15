from pathlib import Path
import hashlib
import zipfile
import xml.etree.ElementTree as ET
import pytest
from pptx import Presentation
from pptx.util import Inches
from aippt_studio.style_hybrid_ooxml import replace_slides, UnsafeSlideImport, P, REL


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def deck(path, prefix, image=False):
    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = f'{prefix}-{i + 1}'
        if image:
            import base64
            png = path.parent / 'pixel.png'
            png.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII='))
            slide.shapes.add_picture(str(png), Inches(2), Inches(2), Inches(1), Inches(1))
    prs.save(path)
    return path


def entries(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def graft(a, b, out, **kw):
    return replace_slides(a, out, primary_sha256=sha(a), overrides=[{'target_slide': 2, 'donor_pptx': b, 'donor_sha256': sha(b), 'donor_slide': 1}], **kw)


def test_graft_preserves_base_and_imports_master_cycle_and_media(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor', True)
    before = entries(a)
    out = tmp_path / 'out.pptx'
    result = graft(a, b, out)
    after = entries(out)
    assert b'<Types xmlns=' in after['[Content_Types].xml']
    assert b'<Relationships xmlns=' in after['ppt/slides/_rels/slide2.xml.rels']
    changed = {'ppt/slides/slide2.xml', 'ppt/slides/_rels/slide2.xml.rels', 'ppt/presentation.xml', 'ppt/_rels/presentation.xml.rels', '[Content_Types].xml'}
    assert all(after[n] == data for n, data in before.items() if n not in changed)
    prs = Presentation(out)
    assert [s.shapes[0].text for s in prs.slides] == ['base-1', 'donor-1', 'base-3']
    assert prs.slides[1].shapes[1].image.blob == Presentation(b).slides[0].shapes[1].image.blob
    assert len(prs.slide_masters) == 2
    assert len(prs.slides[1].slide_layout.slide_master.slide_layouts) == 11
    assert result['native_review'] == 'not_performed'
    assert entries(a) == before


def test_donor_notes_need_explicit_omission_and_do_not_change_neighbor_notes(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    for path, text in ((a, 'primary note'), (b, 'old donor fact')):
        p = Presentation(path)
        for s in p.slides:
            s.notes_slide.notes_text_frame.text = text
        p.save(path)
    before = entries(a)
    out = tmp_path / 'out.pptx'
    with pytest.raises(UnsafeSlideImport, match='notesSlide'):
        graft(a, b, out)
    assert not out.exists()
    result = replace_slides(a, out, primary_sha256=sha(a), overrides=[{
        'target_slide': 2, 'donor_pptx': b, 'donor_sha256': sha(b), 'donor_slide': 1,
        'omit_donor_notes': True}])
    assert result['overrides'][0]['omitted_donor_notes']
    p = Presentation(out)
    assert not p.slides[1].has_notes_slide
    assert p.slides[0].notes_slide.notes_text_frame.text == 'primary note'
    assert p.slides[2].notes_slide.notes_text_frame.text == 'primary note'
    after = entries(out)
    assert all(after[name] == data for name, data in before.items() if name.startswith('ppt/notes'))
    assert not any(b'old donor fact' in data for data in after.values())


@pytest.mark.parametrize('malformed,tag_rels', [
    (False, None), (True, None), (False, 'empty'), (False, 'wrong_root'), (False, 'linked')])
def test_plain_wps_tags_preserved_but_arbitrary_xml_under_tags_rejected(tmp_path, malformed, tag_rels):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    parts = entries(b)
    tags = ET.Element(f'{{{P}}}tagLst')
    ET.SubElement(tags, f'{{{P}}}' + ('oleObj' if malformed else 'tag'),
                  {'name': 'KSO_WM_BEAUTIFY_FLAG', 'val': ''})
    parts['ppt/tags/tag1.xml'] = ET.tostring(tags)
    if tag_rels:
        metadata_rels = ET.Element(f'{{{REL}}}' + ('WrongRoot' if tag_rels == 'wrong_root' else 'Relationships'))
        if tag_rels == 'linked':
            ET.SubElement(metadata_rels, f'{{{REL}}}Relationship', {'Id': 'rId1',
                'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image',
                'Target': '../media/image1.png'})
        parts['ppt/tags/_rels/tag1.xml.rels'] = ET.tostring(metadata_rels)
    slide = ET.fromstring(parts['ppt/slides/slide1.xml'])
    custom_data = ET.SubElement(slide.find(f'{{{P}}}cSld'), f'{{{P}}}custDataLst')
    ET.SubElement(custom_data, f'{{{P}}}tags', {'{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id': 'rIdTags'})
    parts['ppt/slides/slide1.xml'] = ET.tostring(slide)
    rels = ET.fromstring(parts['ppt/slides/_rels/slide1.xml.rels'])
    ET.SubElement(rels, f'{{{REL}}}Relationship', {'Id': 'rIdTags',
        'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags',
        'Target': '../tags/tag1.xml'})
    parts['ppt/slides/_rels/slide1.xml.rels'] = ET.tostring(rels)
    ct = ET.fromstring(parts['[Content_Types].xml'])
    ET.SubElement(ct, '{http://schemas.openxmlformats.org/package/2006/content-types}Override', {
        'PartName': '/ppt/tags/tag1.xml',
        'ContentType': 'application/vnd.openxmlformats-officedocument.presentationml.tags+xml'})
    parts['[Content_Types].xml'] = ET.tostring(ct)
    with zipfile.ZipFile(b, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)
    out = tmp_path / 'out.pptx'
    if malformed or tag_rels in {'wrong_root', 'linked'}:
        with pytest.raises(UnsafeSlideImport, match='non-plain tags|valid empty relationship'):
            graft(a, b, out)
        assert not out.exists()
    else:
        result = graft(a, b, out)
        copied = result['overrides'][0]['part_map']['ppt/tags/tag1.xml']
        assert entries(out)[copied] == parts['ppt/tags/tag1.xml']
        assert Presentation(out).slides[1].shapes[0].text == 'donor-1'
        output_slide = ET.fromstring(entries(out)['ppt/slides/slide2.xml'])
        assert output_slide.find(f'.//{{{P}}}tags').get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id') == 'rIdTags'
        output_rels = ET.fromstring(entries(out)['ppt/slides/_rels/slide2.xml.rels'])
        tag_rel = next(r for r in output_rels if r.get('Id') == 'rIdTags')
        assert tag_rel.get('Target') == '../' + copied.removeprefix('ppt/')


@pytest.mark.parametrize('mutation', ['timing', 'notes', 'external', 'ole'])
def test_refuses_unsupported_without_output(tmp_path, mutation):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    parts = entries(b)
    if mutation in {'timing', 'transition', 'ole'}:
        tree = ET.fromstring(parts['ppt/slides/slide1.xml'])
        ET.SubElement(tree, f'{{{P}}}' + ('oleObj' if mutation == 'ole' else mutation))
        parts['ppt/slides/slide1.xml'] = ET.tostring(tree)
    else:
        tree = ET.fromstring(parts['ppt/slides/_rels/slide1.xml.rels'])
        ET.SubElement(tree, f'{{{REL}}}Relationship', {'Id': 'bad', 'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + ('notesSlide' if mutation == 'notes' else 'image'), 'Target': 'https://example.com' if mutation == 'external' else '../notesSlides/notesSlide1.xml', **({'TargetMode': 'External'} if mutation == 'external' else {})})
        parts['ppt/slides/_rels/slide1.xml.rels'] = ET.tostring(tree)
    with zipfile.ZipFile(b, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)
    out = tmp_path / 'out.pptx'
    with pytest.raises(UnsafeSlideImport): graft(a, b, out)
    assert not out.exists()


def test_hash_size_existing_and_missing_dependency_guards(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    out = tmp_path / 'out.pptx'
    with pytest.raises(UnsafeSlideImport, match='SHA256'):
        replace_slides(a, out, primary_sha256='0' * 64, overrides=[{}])
    prs = Presentation(b)
    prs.slide_width += 100
    prs.save(b)
    with pytest.raises(UnsafeSlideImport, match='sizes'): graft(a, b, out)
    out.write_bytes(b'existing')
    with pytest.raises(UnsafeSlideImport, match='new isolated'): graft(a, b, out)
    assert out.read_bytes() == b'existing'


def test_two_donors_have_unique_graph_parts(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    out = tmp_path / 'out.pptx'
    replace_slides(a, out, primary_sha256=sha(a), overrides=[{'target_slide': n, 'donor_pptx': b, 'donor_sha256': sha(b), 'donor_slide': n} for n in [1, 3]])
    prs = Presentation(out)
    assert [s.shapes[0].text for s in prs.slides] == ['donor-1', 'base-2', 'donor-3']
    assert len(prs.slide_masters) == 3


def test_neighbor_local_timing_preserved(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    parts = entries(a)
    root = ET.fromstring(parts['ppt/slides/slide3.xml'])
    ET.SubElement(root, f'{{{P}}}timing')
    parts['ppt/slides/slide3.xml'] = ET.tostring(root)
    with zipfile.ZipFile(a, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)
    before = entries(a)
    graft(a, b, tmp_path / 'out.pptx')
    assert entries(tmp_path / 'out.pptx')['ppt/slides/slide3.xml'] == before['ppt/slides/slide3.xml']


def test_missing_dependency_blocks_graft(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    parts = entries(b)
    del parts['ppt/slideLayouts/slideLayout7.xml']
    with zipfile.ZipFile(b, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)
    with pytest.raises(UnsafeSlideImport, match='Missing dependency'): graft(a, b, tmp_path / 'out.pptx')


def add_transition(path, slide, effect='fade', ns=P):
    parts = entries(path)
    root = ET.fromstring(parts[f'ppt/slides/slide{slide}.xml'])
    transition = ET.SubElement(root, f'{{{P}}}transition', {'spd': 'slow', 'advTm': '1200'})
    ET.SubElement(transition, f'{{{ns}}}{effect}')
    parts[f'ppt/slides/slide{slide}.xml'] = ET.tostring(root)
    with zipfile.ZipFile(path, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)


def test_ordinary_neighbor_transition_allowed_target_transition_preserved(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    add_transition(a, 1)
    add_transition(a, 2, 'wipe')
    add_transition(a, 3, 'push')
    add_transition(b, 1, 'zoom')
    before = entries(a)
    result = graft(a, b, tmp_path / 'out.pptx')
    after = entries(tmp_path / 'out.pptx')
    assert after['ppt/slides/slide1.xml'] == before['ppt/slides/slide1.xml']
    assert after['ppt/slides/slide3.xml'] == before['ppt/slides/slide3.xml']
    transition = ET.fromstring(after['ppt/slides/slide2.xml']).find(f'{{{P}}}transition')
    assert transition[0].tag == f'{{{P}}}wipe'
    assert transition.get('advTm') == '1200'
    assert result['overrides'][0]['primary_transition_preserved'] is True
    assert result['overrides'][0]['donor_transition_omitted'] is True


def test_morph_neighbor_rejected(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    add_transition(a, 3, 'morph', 'http://schemas.microsoft.com/office/powerpoint/2015/09/main')
    with pytest.raises(UnsafeSlideImport): graft(a, b, tmp_path / 'out.pptx')


def add_timing(path, slide=1, trigger=None, missing=False):
    parts = entries(path)
    root = ET.fromstring(parts[f'ppt/slides/slide{slide}.xml'])
    shapeid = root.find('.//' + f'{{{P}}}cNvPr').get('id')
    timing = ET.SubElement(root, f'{{{P}}}timing')
    ET.SubElement(timing, f'{{{P}}}cTn', {'id': '1'})
    ET.SubElement(timing, f'{{{P}}}spTgt', {'spid': '99999' if missing else shapeid})
    if trigger: ET.SubElement(timing, f'{{{P}}}cond', {'evt': trigger})
    parts[f'ppt/slides/slide{slide}.xml'] = ET.tostring(root)
    with zipfile.ZipFile(path, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)


def test_explicit_staticize_removes_only_donor_timing(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    add_timing(a, 3)
    add_timing(b)
    before = entries(a)
    with pytest.raises(UnsafeSlideImport): graft(a, b, tmp_path / 'strict.pptx')
    out = tmp_path / 'out.pptx'
    result = replace_slides(a, out, primary_sha256=sha(a), overrides=[dict(target_slide=2, donor_pptx=b, donor_sha256=sha(b), donor_slide=1, staticize_donor_timing=True)])
    assert ET.fromstring(entries(out)['ppt/slides/slide2.xml']).find(f'{{{P}}}timing') is None
    assert entries(out)['ppt/slides/slide3.xml'] == before['ppt/slides/slide3.xml']
    assert result['overrides'][0]['staticized_parts'][0]['removed_node'] == 'p:timing'
    assert result['output_sha256'] == sha(out)


@pytest.mark.parametrize('trigger,missing', [('onClick', False), (None, True)])
def test_staticize_refuses_trigger_or_uncertain_target(tmp_path, trigger, missing):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    add_timing(b, trigger=trigger, missing=missing)
    with pytest.raises(UnsafeSlideImport):
        replace_slides(a, tmp_path / 'out.pptx', primary_sha256=sha(a), overrides=[dict(target_slide=2, donor_pptx=b, donor_sha256=sha(b), donor_slide=1, staticize_donor_timing=True)])


def test_staticize_refuses_layout_timing_and_dynamic_task(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    add_timing(b)
    override = dict(target_slide=2, donor_pptx=b, donor_sha256=sha(b), donor_slide=1, staticize_donor_timing=True, dynamic_required=True)
    with pytest.raises(UnsafeSlideImport, match='requiring dynamic'):
        replace_slides(a, tmp_path / 'dynamic.pptx', primary_sha256=sha(a), overrides=[override])
    parts = entries(b)
    root = ET.fromstring(parts['ppt/slideLayouts/slideLayout7.xml'])
    ET.SubElement(root, f'{{{P}}}timing')
    parts['ppt/slideLayouts/slideLayout7.xml'] = ET.tostring(root)
    with zipfile.ZipFile(b, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)
    override.update(donor_sha256=sha(b), dynamic_required=False)
    with pytest.raises(UnsafeSlideImport, match='timing'):
        replace_slides(a, tmp_path / 'layout.pptx', primary_sha256=sha(a), overrides=[override])


def test_minor_size_delta_is_explicit_bounded_and_unscaled(tmp_path):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    prs = Presentation(b)
    prs.slide_width += 1587
    prs.save(b)
    with pytest.raises(UnsafeSlideImport, match='sizes'): graft(a, b, tmp_path / 'strict.pptx')
    out = tmp_path / 'out.pptx'
    override = dict(target_slide=2, donor_pptx=b, donor_sha256=sha(b), donor_slide=1, allow_minor_size_delta=True)
    result = replace_slides(a, out, primary_sha256=sha(a), overrides=[override])
    assert Presentation(out).slide_width == Presentation(a).slide_width
    assert Presentation(out).slides[1].shapes[0].left == prs.slides[0].shapes[0].left
    assert result['overrides'][0]['size_policy']['absolute_delta_emu'] == [1587, 0]
    assert result['overrides'][0]['size_policy']['potential_edge_clipping'] is True
    prs.slide_width += 20000
    prs.save(b)
    override['donor_sha256'] = sha(b)
    with pytest.raises(UnsafeSlideImport, match='sizes'):
        replace_slides(a, tmp_path / 'too-large.pptx', primary_sha256=sha(a), overrides=[override])


def test_color_overrides_affect_only_imported_xml_and_preserve_media(tmp_path):
    from aippt_studio.style_hybrid_ooxml import A
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor', True)
    parts = entries(b)
    root = ET.fromstring(parts['ppt/slides/slide1.xml'])
    ET.SubElement(root, f'{{{A}}}srgbClr', {'val': 'abcdef'})
    parts['ppt/slides/slide1.xml'] = ET.tostring(root)
    with zipfile.ZipFile(b, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)
    out = tmp_path / 'out.pptx'
    r = replace_slides(a, out, primary_sha256=sha(a), overrides=[dict(target_slide=2, donor_pptx=b, donor_sha256=sha(b), donor_slide=1, theme_accent_overrides={'accent5': '6b8eb6'}, direct_rgb_overrides={'ABCDEF': '4274B0'})])
    adaptation = r['overrides'][0]['color_adaptation']
    output = entries(out)
    assert '4274B0' in output['ppt/slides/slide2.xml'].decode()
    assert output['ppt/slides/slide1.xml'] == entries(a)['ppt/slides/slide1.xml']
    assert output['ppt/theme/theme1.xml'] == entries(a)['ppt/theme/theme1.xml']
    for old, new in r['overrides'][0]['part_map'].items():
        if old.startswith('ppt/media/'): assert output[new] == parts[old]
    for part in adaptation['part_changes']:
        assert sha_bytes(output[part['output_part']]) == part['after_sha256']
    assert sum(p['theme_slot_count'] for p in adaptation['part_changes']) == 1
    assert sum(p['direct_rgb_count'] for p in adaptation['part_changes']) == 1


def sha_bytes(data): return hashlib.sha256(data).hexdigest()


@pytest.mark.parametrize('colors', [{'theme_accent_overrides': {'dk1': '000000'}}, {'theme_accent_overrides': {'accent1': '#123456'}}, {'direct_rgb_overrides': {'12345': '123456'}}, {'direct_rgb_overrides': {'123456': '12345678'}}, {'direct_rgb_overrides': {'ffffff': '000000', 'FFFFFF': '123456'}}])
def test_color_overrides_reject_invalid_options(tmp_path, colors):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    with pytest.raises(UnsafeSlideImport):
        replace_slides(a, tmp_path / 'out.pptx', primary_sha256=sha(a), overrides=[dict(target_slide=2, donor_pptx=b, donor_sha256=sha(b), donor_slide=1, **colors)])


def test_theme_override_resolves_explicit_scheme_fill_preserves_transform_and_style(tmp_path):
    from aippt_studio.style_hybrid_ooxml import A
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor')
    parts = entries(b)
    root = ET.fromstring(parts['ppt/slides/slide1.xml'])
    shape = root.find(f'.//{{{P}}}sp')
    props = shape.find(f'{{{P}}}spPr')
    fill = ET.SubElement(props, f'{{{A}}}solidFill')
    scheme = ET.SubElement(fill, f'{{{A}}}schemeClr', {'val': 'accent3'})
    ET.SubElement(scheme, f'{{{A}}}lumMod', {'val': '80000'})
    style = ET.SubElement(shape, f'{{{P}}}style')
    fillref = ET.SubElement(style, f'{{{A}}}fillRef', {'idx': '1'})
    ET.SubElement(fillref, f'{{{A}}}schemeClr', {'val': 'accent3'})
    parts['ppt/slides/slide1.xml'] = ET.tostring(root)
    with zipfile.ZipFile(b, 'w') as z:
        for name, data in parts.items(): z.writestr(name, data)
    out = tmp_path / 'out.pptx'
    replace_slides(a, out, primary_sha256=sha(a), overrides=[dict(target_slide=2, donor_pptx=b, donor_sha256=sha(b), donor_slide=1, theme_accent_overrides={'accent3': '6B8EB6'})])
    result = ET.fromstring(entries(out)['ppt/slides/slide2.xml'])
    color = result.find(f'.//{{{A}}}solidFill/{{{A}}}srgbClr')
    assert color.get('val') == '6B8EB6'
    assert color[0].get('val') == '80000'
    assert result.find(f'.//{{{A}}}fillRef/{{{A}}}schemeClr').get('val') == 'accent3'


def test_repeated_build_hash_is_stable_across_wall_clock_dates(tmp_path, monkeypatch):
    a, b = deck(tmp_path / 'a.pptx', 'base'), deck(tmp_path / 'b.pptx', 'donor', True)
    before = entries(a)
    first, second = tmp_path / 'first.pptx', tmp_path / 'second.pptx'
    monkeypatch.setattr(zipfile.time, 'localtime', lambda *_: (2026, 10, 8, 10, 20, 30, 3, 281, -1))
    one = graft(a, b, first)
    monkeypatch.setattr(zipfile.time, 'localtime', lambda *_: (2027, 11, 9, 11, 21, 32, 1, 313, -1))
    two = graft(a, b, second)
    assert one['output_sha256'] == two['output_sha256'] == sha(first) == sha(second)
    assert first.read_bytes() == second.read_bytes()
    assert entries(a) == before
    assert entries(first) == entries(second)
    with zipfile.ZipFile(first) as z:
        assert z.namelist() == sorted(z.namelist())
        assert all(info.date_time == (1980, 1, 1, 0, 0, 0) for info in z.infolist())
        assert all(info.create_system == 3 and info.external_attr == 0o600 << 16 for info in z.infolist())
    with pytest.raises(UnsafeSlideImport, match='new isolated'):
        graft(a, b, first)
