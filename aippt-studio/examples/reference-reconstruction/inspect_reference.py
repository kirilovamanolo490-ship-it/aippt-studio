#!/usr/bin/env python3
"""Inventory source slide geometry/styles and optionally extract media by hash."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE

EMU = 914400
SUPPORTED_NAMES = {
    MSO_SHAPE_TYPE.TEXT_BOX: "text_box",
    MSO_SHAPE_TYPE.PICTURE: "picture",
    MSO_SHAPE_TYPE.PLACEHOLDER: "placeholder",
    MSO_SHAPE_TYPE.GROUP: "group",
    MSO_SHAPE_TYPE.TABLE: "table",
    MSO_SHAPE_TYPE.CHART: "chart",
    MSO_SHAPE_TYPE.LINE: "line",
}
SIMPLE_RECTS = {MSO_AUTO_SHAPE_TYPE.RECTANGLE, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE}


def color_rgb(color) -> str | None:
    try:
        return str(color.rgb) if color and color.rgb else None
    except (AttributeError, TypeError, ValueError):
        return None


def shape_color(get_color) -> str | None:
    try:
        return color_rgb(get_color())
    except (AttributeError, TypeError, ValueError):
        return None


def text_metrics(shape) -> dict[str, Any]:
    sizes = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text and run.font.size is not None:
                sizes.append(round(run.font.size.pt, 1))
    return {"text_chars": len(shape.text), "paragraph_count": len(shape.text_frame.paragraphs),
            "explicit_font_sizes_pt": sorted(set(sizes))}


def inventory(source: Path, media_dir: Path | None, include_text: bool) -> dict:
    presentation = Presentation(source)
    pages = []
    unsupported = []
    for page_no, slide in enumerate(presentation.slides, 1):
        records = []
        textual = []
        for object_no, shape in enumerate(slide.shapes, 1):
            kind = SUPPORTED_NAMES.get(shape.shape_type, f"shape_type_{shape.shape_type}")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                kind = "picture"
            elif shape.has_text_frame:
                kind = "text"
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.auto_shape_type in SIMPLE_RECTS:
                kind = "rect"
            rec: dict[str, Any] = {
                "object": object_no,
                "source_shape_id": shape.shape_id,
                "kind": kind,
                "geometry_in": {"x": round(shape.left / EMU, 4), "y": round(shape.top / EMU, 4),
                                "width": round(shape.width / EMU, 4), "height": round(shape.height / EMU, 4)},
            }
            if shape.has_text_frame:
                metrics = text_metrics(shape)
                if metrics["explicit_font_sizes_pt"]:
                    textual.append(max(metrics["explicit_font_sizes_pt"]))
                role = "body"
                rec.update(metrics)
                rec["role_hint"] = role
                if include_text:
                    rec["text"] = shape.text
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = shape.image.blob
                sha = hashlib.sha256(blob).hexdigest()
                ext = shape.image.ext or "bin"
                filename = f"{sha}.{ext}"
                rec["media"] = {"sha256": sha, "bytes": len(blob), "size_px": list(shape.image.size),
                                 "content_type": shape.image.content_type, "filename_by_hash": filename}
                if media_dir is not None:
                    media_dir.mkdir(parents=True, exist_ok=True)
                    (media_dir / filename).write_bytes(blob)
                try:
                    rec["crop"] = {name: round(float(getattr(shape, f"crop_{name}")), 6)
                                   for name in ("left", "top", "right", "bottom")}
                except (AttributeError, TypeError, ValueError):
                    rec["crop"] = None
            else:
                is_auto = shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                rec["fill_rgb"] = shape_color(lambda: shape.fill.fore_color) if is_auto else None
                rec["line_rgb"] = shape_color(lambda: shape.line.color) if is_auto else None
                if is_auto:
                    rec["auto_shape_type"] = str(shape.auto_shape_type)
            records.append(rec)
            if kind not in {"text", "text_box", "picture", "rect"}:
                unsupported.append({"page": page_no, "object": object_no, "kind": kind})
        if textual:
            median = sorted(textual)[len(textual) // 2]
            for rec in records:
                if rec.get("explicit_font_sizes_pt"):
                    mx = max(rec["explicit_font_sizes_pt"])
                    rec["role_hint"] = "title" if mx >= median * 1.35 else "body"
        pages.append({"page": page_no, "objects": records})
    return {
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "page_count": len(pages),
        "canvas_inches": [round(presentation.slide_width / EMU, 4), round(presentation.slide_height / EMU, 4)],
        "slides": pages,
        "unsupported_objects": unsupported,
        "limitations": [
            "role_hint and geometry are structural heuristics, not visual approval or a finished layout plan",
            "inherited master/layout content, group internals, animation, SmartArt, charts and complex fills need separate handling",
            "explicit font sizes only; inherited fonts/sizes and rendered glyphs may differ",
            "media hashes identify bytes, not meaning, permission, or suitability",
            "source text is omitted unless --include-text is explicitly supplied",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--media-dir", type=Path, help="Optional output directory for independent media")
    parser.add_argument("--include-text", action="store_true", help="Include source copy in the inventory")
    args = parser.parse_args()
    result = inventory(args.source, args.media_dir, args.include_text)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"pages": result["page_count"], "unsupported": len(result["unsupported_objects"]),
                      "output": str(args.output)}))


if __name__ == "__main__":
    main()
