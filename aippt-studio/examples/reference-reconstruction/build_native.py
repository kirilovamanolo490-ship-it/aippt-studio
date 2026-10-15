#!/usr/bin/env python3
"""Create a new PPTX from an authored JSON plan using native basic objects."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def rgb(value: str) -> RGBColor:
    clean = value.strip().lstrip("#")
    if len(clean) != 6 or any(c not in "0123456789abcdefABCDEF" for c in clean):
        raise ValueError(f"expected six-digit RGB color, got {value!r}")
    return RGBColor.from_string(clean.upper())


def set_run(run, spec: dict) -> None:
    run.text = str(spec.get("text", ""))
    font = run.font
    if spec.get("font_family"):
        font.name = spec["font_family"]
        rpr = run._r.get_or_add_rPr()
        ea = OxmlElement("a:ea")
        ea.set("typeface", spec["font_family"])
        rpr.append(ea)
    if spec.get("font_size_pt") is not None:
        font.size = Pt(float(spec["font_size_pt"]))
    if spec.get("bold") is not None:
        font.bold = bool(spec["bold"])
    if spec.get("italic") is not None:
        font.italic = bool(spec["italic"])
    if spec.get("color"):
        font.color.rgb = rgb(spec["color"])


def add_text(slide, spec: dict) -> None:
    box = slide.shapes.add_textbox(Inches(spec["x"]), Inches(spec["y"]),
                                   Inches(spec["width"]), Inches(spec["height"]))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    margin = float(spec.get("margin", 0))
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = VALIGN.get(spec.get("valign", "top"), MSO_ANCHOR.TOP)
    paragraphs = spec.get("paragraphs")
    if not paragraphs:
        paragraphs = [{"runs": [{"text": spec.get("text", ""), **{
            k: spec[k] for k in ("font_family", "font_size_pt", "bold", "italic", "color") if k in spec
        }}]}]
    for index, item in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = ALIGN.get(item.get("align", spec.get("align", "left")), PP_ALIGN.LEFT)
        paragraph.level = int(item.get("level", 0))
        for run_spec in item.get("runs", []):
            run = paragraph.add_run()
            set_run(run, run_spec)


def add_rect(slide, spec: dict) -> None:
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if spec.get("rounded") else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(spec["x"]), Inches(spec["y"]),
                                   Inches(spec["width"]), Inches(spec["height"]))
    if spec.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(spec["fill"])
    else:
        shape.fill.background()
    if spec.get("line"):
        shape.line.color.rgb = rgb(spec["line"])
        shape.line.width = Pt(float(spec.get("line_width_pt", 0.75)))
    else:
        shape.line.fill.background()


def add_picture(slide, spec: dict, media_dir: Path) -> None:
    filename = Path(spec["file"]).name
    path = media_dir / filename
    if not path.is_file():
        raise FileNotFoundError(f"media file not found: {path}")
    expected_hash = Path(filename).stem.lower()
    if len(expected_hash) != 64 or hashlib.sha256(path.read_bytes()).hexdigest() != expected_hash:
        raise ValueError(f"media filename must be its verified SHA-256 plus extension: {filename}")
    # The visible crop and destination must share an aspect ratio. Otherwise
    # PowerPoint stretches the pixels even though width/height look correct.
    from PIL import Image
    with Image.open(path) as image:
        natural_ratio = image.width / image.height
    frame_ratio = float(spec["width"]) / float(spec["height"])
    crop = spec.get("crop")
    if crop is None:
        crop = {"left": 0.0, "top": 0.0, "right": 0.0, "bottom": 0.0}
        if frame_ratio < natural_ratio:
            crop["left"] = crop["right"] = (1 - frame_ratio / natural_ratio) / 2
        else:
            crop["top"] = crop["bottom"] = (1 - natural_ratio / frame_ratio) / 2
    crop_values = {}
    for side in ("left", "top", "right", "bottom"):
        value = float(crop.get(side, 0))
        if not 0 <= value < 1:
            raise ValueError(f"picture crop {side} must be between 0 and 1")
        crop_values[side] = value
    visible_width = 1 - crop_values["left"] - crop_values["right"]
    visible_height = 1 - crop_values["top"] - crop_values["bottom"]
    if visible_width <= 0 or visible_height <= 0:
        raise ValueError("picture crop removes the entire image")
    if abs(natural_ratio * visible_width / visible_height / frame_ratio - 1) > 0.01:
        raise ValueError("authored picture crop would stretch pixels; adjust crop or destination proportions")
    shape = slide.shapes.add_picture(str(path), Inches(spec["x"]), Inches(spec["y"]),
                                     Inches(spec["width"]), Inches(spec["height"]))
    for side, value in crop_values.items():
        setattr(shape, f"crop_{side}", value)


def build(plan_path: Path, output: Path, media_dir: Path) -> dict:
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    width, height = map(float, plan["canvas_inches"])
    if width <= 0 or height <= 0 or not plan.get("slides"):
        raise ValueError("plan needs a positive canvas_inches and at least one slide")
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    blank_layout = presentation.slide_layouts[6]
    for page_no, page in enumerate(plan["slides"], 1):
        slide = presentation.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(page.get("background", "FFFFFF"))
        for object_no, item in enumerate(page.get("elements", []), 1):
            kind = item.get("type")
            if kind == "text":
                add_text(slide, item)
            elif kind == "rect":
                add_rect(slide, item)
            elif kind == "picture":
                add_picture(slide, item, media_dir)
            else:
                raise ValueError(f"page {page_no}, object {object_no}: unsupported type {kind!r}; use text, rect, picture")
            for key in ("x", "y", "width", "height"):
                if float(item[key]) < 0 or (key in ("width", "height") and float(item[key]) == 0):
                    raise ValueError(f"page {page_no}, object {object_no}: invalid geometry {key}")
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    reopened = Presentation(output)
    if len(reopened.slides) != len(plan["slides"]):
        raise RuntimeError("saved output page count does not match the plan")
    return {"output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
            "page_count": len(reopened.slides),
            "object_counts": [len(slide.shapes) for slide in reopened.slides],
            "native_types": [
                {"text": sum(1 for s in slide.shapes if s.has_text_frame),
                 "pictures": sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE),
                 "rectangles": sum(1 for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)}
                for slide in reopened.slides
            ]}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("plan", type=Path, help="Authored JSON plan (not source slide XML)")
    parser.add_argument("output", type=Path)
    parser.add_argument("--media-dir", type=Path, required=True)
    args = parser.parse_args()
    result = build(args.plan, args.output, args.media_dir)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
