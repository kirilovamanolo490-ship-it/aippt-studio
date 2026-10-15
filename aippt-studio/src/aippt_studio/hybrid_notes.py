"""Update presenter notes in a hybrid PPTX without round-tripping slide parts."""

from __future__ import annotations

import posixpath
import re
import tempfile
import zipfile
from pathlib import Path
from lxml import etree as ET

from pptx import Presentation

from .models import DeckPlan
from .notes import sanitize_speaker_notes_text


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
SLIDE_REL = f"{R_NS}/slide"
NOTES_SLIDE_REL = f"{R_NS}/notesSlide"
NOTES_MASTER_REL = f"{R_NS}/notesMaster"
NOTES_SLIDE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml"
)
NOTES_MASTER_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml"
)

def _part_target(source_part: str, target: str) -> str:
    """Resolve an OPC relationship target to a package part name."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))


def _relative_target(source_part: str, target_part: str) -> str:
    return posixpath.relpath(target_part, posixpath.dirname(source_part))


def _rels_part(part: str) -> str:
    directory, name = posixpath.split(part)
    return posixpath.join(directory, "_rels", f"{name}.rels")


def _parse(data: bytes, part: str) -> ET.Element:
    try:
        root = ET.fromstring(data, parser=ET.XMLParser(resolve_entities=False, no_network=True))
        if root.getroottree().docinfo.doctype:
            raise ValueError(f"DTD-bearing XML is unsupported: {part}")
        return root
    except ET.XMLSyntaxError as error:
        raise ValueError(f"Invalid OOXML in {part}: {error}") from error


def _serialize(root: ET.Element) -> bytes:
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _relationship_entries(root: ET.Element) -> list[ET.Element]:
    return root.findall(f"{{{PKG_REL_NS}}}Relationship")


def _next_rid(root: ET.Element) -> str:
    used = {item.get("Id", "") for item in _relationship_entries(root)}
    number = 1
    while f"rId{number}" in used:
        number += 1
    return f"rId{number}"


def _add_relationship(root: ET.Element, rel_id: str, rel_type: str, target: str) -> None:
    ET.SubElement(
        root,
        f"{{{PKG_REL_NS}}}Relationship",
        {"Id": rel_id, "Type": rel_type, "Target": target},
    )


def _slide_parts(archive: zipfile.ZipFile) -> list[str]:
    presentation = _parse(archive.read("ppt/presentation.xml"), "ppt/presentation.xml")
    rels_part = "ppt/_rels/presentation.xml.rels"
    relationships = _parse(archive.read(rels_part), rels_part)
    by_id = {item.get("Id"): item for item in _relationship_entries(relationships)}
    slide_ids = presentation.find(f"{{{P_NS}}}sldIdLst")
    if slide_ids is None:
        return []
    parts: list[str] = []
    for slide_id in slide_ids.findall(f"{{{P_NS}}}sldId"):
        rel_id = slide_id.get(f"{{{R_NS}}}id")
        rel = by_id.get(rel_id)
        if rel is None or rel.get("Type") != SLIDE_REL or not rel.get("Target"):
            raise ValueError(f"Presentation has an unresolved slide relationship: {rel_id}")
        parts.append(_part_target("ppt/presentation.xml", rel.get("Target", "")))
    return parts


def _ensure_override(root: ET.Element, part: str, content_type: str) -> None:
    name = "/" + part
    for item in root.findall(f"{{{CT_NS}}}Override"):
        if item.get("PartName") == name:
            item.set("ContentType", content_type)
            return
    ET.SubElement(root, f"{{{CT_NS}}}Override", {"PartName": name, "ContentType": content_type})


def _notes_master_part(presentation_rels: ET.Element) -> str | None:
    for rel in _relationship_entries(presentation_rels):
        if rel.get("Type") == NOTES_MASTER_REL and rel.get("Target"):
            return _part_target("ppt/presentation.xml", rel.get("Target", ""))
    return None


def _presentation_notes_master_id(presentation: ET.Element, rel_id: str) -> None:
    notes_list = presentation.find(f"{{{P_NS}}}notesMasterIdLst")
    if notes_list is None:
        notes_list = ET.Element(f"{{{P_NS}}}notesMasterIdLst")
        slide_master_list = presentation.find(f"{{{P_NS}}}sldMasterIdLst")
        insert_at = list(presentation).index(slide_master_list) + 1 if slide_master_list is not None else 0
        presentation.insert(insert_at, notes_list)
    notes_list.clear()
    ET.SubElement(notes_list, f"{{{P_NS}}}notesMasterId", {f"{{{R_NS}}}id": rel_id})


def _build_notes_package(plan: DeckPlan) -> tuple[dict[str, bytes], str, bytes, bytes]:
    """Create canonical notes parts in a temporary deck, never from the target deck."""
    with tempfile.TemporaryDirectory(prefix="aippt-hybrid-notes-") as directory:
        temporary = Path(directory) / "notes-template.pptx"
        presentation = Presentation()
        blank_layout = presentation.slide_layouts[6]
        for spec in plan.slides:
            slide = presentation.slides.add_slide(blank_layout)
            slide.notes_slide.notes_text_frame.text = sanitize_speaker_notes_text(spec.speaker_notes)
        presentation.save(temporary)
        with zipfile.ZipFile(temporary, "r") as package:
            slides = {
                f"ppt/notesSlides/notesSlide{index}.xml": package.read(
                    f"ppt/notesSlides/notesSlide{index}.xml"
                )
                for index in range(1, len(plan.slides) + 1)
            }
            rels = {
                f"ppt/notesSlides/_rels/notesSlide{index}.xml.rels": package.read(
                    f"ppt/notesSlides/_rels/notesSlide{index}.xml.rels"
                )
                for index in range(1, len(plan.slides) + 1)
            }
            master = package.read("ppt/notesMasters/notesMaster1.xml")
            master_rels = package.read("ppt/notesMasters/_rels/notesMaster1.xml.rels")
    notes = {}
    for index, part in enumerate(slides, 1):
        notes[part] = slides[part]
        notes[f"ppt/notesSlides/_rels/notesSlide{index}.xml.rels"] = rels[
            f"ppt/notesSlides/_rels/notesSlide{index}.xml.rels"
        ]
    return notes, "ppt/notesMasters/notesMaster1.xml", master, master_rels


def update_hybrid_notes(deck: Path, plan: DeckPlan) -> dict[str, object]:
    """Replace candidate speaker notes in-place while preserving non-notes parts byte-for-byte.

    Existing slide XML and every non-notes part except the explicitly allowed
    presentation/slide relationship and content-type parts remain byte-identical. If the deck
    has no notes master, a minimal notes master generated by python-pptx is
    imported, with its theme relationship redirected to the target deck theme.
    """
    deck = Path(deck)
    if not deck.is_file():
        raise FileNotFoundError(deck)

    generated_notes, generated_master_part, generated_master, generated_master_rels = (
        _build_notes_package(plan)
    )
    with zipfile.ZipFile(deck, "r") as source:
        infos = source.infolist()
        original = {info.filename: source.read(info.filename) for info in infos}

    with zipfile.ZipFile(deck, "r") as source:
        slides = _slide_parts(source)
        if len(slides) != len(plan.slides):
            raise ValueError(
                f"Plan has {len(plan.slides)} slides but candidate has {len(slides)} slides"
            )

    replacements: dict[str, bytes] = {}
    additions: dict[str, bytes] = {}
    removed: set[str] = set()
    presentation_part = "ppt/presentation.xml"
    presentation_rels_part = "ppt/_rels/presentation.xml.rels"
    content_types_part = "[Content_Types].xml"
    presentation = _parse(original[presentation_part], presentation_part)
    presentation_rels = _parse(original[presentation_rels_part], presentation_rels_part)
    content_types = _parse(original[content_types_part], content_types_part)

    notes_master = _notes_master_part(presentation_rels)
    if notes_master is not None and notes_master not in original:
        raise ValueError(f"Presentation notes master relationship points to missing part: {notes_master}")
    if notes_master is None:
        notes_master = generated_master_part
        master_rels_root = _parse(generated_master_rels, "temporary notes master relationships")
        # Use a theme already present in the candidate, preserving its theme bytes.
        for rel in _relationship_entries(master_rels_root):
            if rel.get("Type", "").endswith("/theme"):
                themes = [name for name in original if re.fullmatch(r"ppt/theme/theme\d+\.xml", name)]
                if not themes:
                    raise ValueError("Candidate has no theme part for its notes master")
                rel.set("Target", _relative_target(notes_master, sorted(themes)[0]))
        master_rel_id = _next_rid(presentation_rels)
        _add_relationship(
            presentation_rels,
            master_rel_id,
            NOTES_MASTER_REL,
            _relative_target(presentation_part, notes_master),
        )
        _presentation_notes_master_id(presentation, master_rel_id)
        additions[notes_master] = generated_master
        additions[_rels_part(notes_master)] = _serialize(master_rels_root)
        _ensure_override(content_types, notes_master, NOTES_MASTER_CONTENT_TYPE)
        replacements[presentation_part] = _serialize(presentation)
        replacements[presentation_rels_part] = _serialize(presentation_rels)

    # Locate existing notes for each slide so names can be retained where possible.
    existing_notes_by_slide: dict[str, str] = {}
    slide_relationship_roots: dict[str, ET.Element] = {}
    for slide_part in slides:
        relation_part = _rels_part(slide_part)
        root = _parse(original[relation_part], relation_part) if relation_part in original else ET.Element(
            f"{{{PKG_REL_NS}}}Relationships"
        )
        slide_relationship_roots[slide_part] = root
        note_rels = [
            rel for rel in _relationship_entries(root)
            if rel.get("Type") == NOTES_SLIDE_REL
        ]
        if len(note_rels) > 1:
            raise ValueError(f"Slide has multiple notes relationships: {slide_part}")
        if note_rels:
            rel = note_rels[0]
            if not rel.get("Target"):
                raise ValueError(f"Slide has a notes relationship without a target: {slide_part}")
            notes_part = _part_target(slide_part, rel.get("Target", ""))
            if notes_part not in original:
                raise ValueError(f"Slide notes relationship points to missing part: {notes_part}")
            existing_notes_by_slide[slide_part] = notes_part
    if len(set(existing_notes_by_slide.values())) != len(existing_notes_by_slide):
        raise ValueError("Multiple slides share one notes part; refusing ambiguous note ownership")

    all_notes_parts = {
        name for name in original if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
    }
    used_notes_parts: set[str] = set()
    for index, slide_part in enumerate(slides, 1):
        notes_part = existing_notes_by_slide.get(slide_part)
        if notes_part is None:
            candidate = f"ppt/notesSlides/notesSlide{index}.xml"
            if candidate in original or candidate in additions:
                number = 1
                while f"ppt/notesSlides/notesSlide{number}.xml" in original or f"ppt/notesSlides/notesSlide{number}.xml" in additions:
                    number += 1
                candidate = f"ppt/notesSlides/notesSlide{number}.xml"
            notes_part = candidate
            root = slide_relationship_roots[slide_part]
            _add_relationship(
                root,
                _next_rid(root),
                NOTES_SLIDE_REL,
                _relative_target(slide_part, notes_part),
            )
        used_notes_parts.add(notes_part)

        # Copy the generated notes page structure, retaining the source deck's
        # slide and notes-master relationships only.
        generated_part = f"ppt/notesSlides/notesSlide{index}.xml"
        replacements[notes_part] = generated_notes[generated_part]
        old_generated_rels = generated_notes[f"ppt/notesSlides/_rels/notesSlide{index}.xml.rels"]
        rel_root = _parse(old_generated_rels, "temporary notes slide relationships")
        for rel in _relationship_entries(rel_root):
            if rel.get("Type") == SLIDE_REL:
                rel.set("Target", _relative_target(notes_part, slide_part))
            elif rel.get("Type") == NOTES_MASTER_REL:
                rel.set("Target", _relative_target(notes_part, notes_master))
        replacements[_rels_part(notes_part)] = _serialize(rel_root)
        if notes_part not in original:
            additions[notes_part] = replacements.pop(notes_part)
            additions[_rels_part(notes_part)] = replacements.pop(_rels_part(notes_part))
        _ensure_override(content_types, notes_part, NOTES_SLIDE_CONTENT_TYPE)

    # Retire orphaned notes so old provenance cannot remain attached to the
    # candidate as an unreferenced note page.
    for stale in all_notes_parts - used_notes_parts:
        removed.add(stale)
        removed.add(_rels_part(stale))
    for item in list(content_types.findall(f"{{{CT_NS}}}Override")):
        part_name = item.get("PartName", "").lstrip("/")
        if part_name.startswith("ppt/notesSlides/") and part_name not in used_notes_parts:
            content_types.remove(item)

    # Save each slide relationship part after notes relationships are updated.
    for slide_part, root in slide_relationship_roots.items():
        replacements[_rels_part(slide_part)] = _serialize(root)
    replacements[content_types_part] = _serialize(content_types)

    # Preserve every original ZipInfo and the exact payload for untouched parts.
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{deck.stem}-notes-", suffix=".pptx", dir=deck.parent, delete=False
        ) as stream:
            temporary_path = Path(stream.name)
        with zipfile.ZipFile(temporary_path, "w") as target:
            for info in infos:
                if info.filename in removed:
                    continue
                target.writestr(info, replacements.get(info.filename, original[info.filename]))
            for part, data in additions.items():
                target.writestr(part, data, compress_type=zipfile.ZIP_DEFLATED)
        temporary_path.replace(deck)
    except Exception:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise

    return {
        "status": "updated",
        "slide_count": len(slides),
        "notes_parts": len(used_notes_parts),
        "added_parts": len(additions),
        "removed_orphan_notes": len(all_notes_parts - used_notes_parts),
    }
