"""Contextual, local signals for presentation form and page expression.

This helper locates concrete review candidates in the current PPTX. It never
scores visual quality, deletes copy, imposes font minima, or decides whether a
number or phrase is legitimate. The form selects questions for a human review.
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation

from .expression_review import _objects, copy_signals, decorative_ordinal_signals, density_signals
from .models import PresentationForm
from .renderer import sha256


_FORM_REVIEW_LENSES: dict[PresentationForm, tuple[str, ...]] = {
    PresentationForm.GENERAL: (
        "Read the page as the intended audience: does its claim follow from its evidence?",
        "Check hierarchy, proportions, image meaning, and neighboring-page continuity on the whole page.",
    ),
    PresentationForm.GOVERNMENT_BRIEFING: (
        "Check that claims, responsibility, place, and timing are supported by the supplied evidence.",
        "Check that the wording is formal and source-faithful for this briefing occasion.",
    ),
    PresentationForm.BUSINESS_REPORT: (
        "Check that the conclusion, evidence, action, and owner form a clear decision path.",
        "Check that repeated page structures still fit each page's actual content.",
    ),
    PresentationForm.COMPETITION_DEFENSE: (
        "Check that the page answers a judging question with comparison, evidence, and a feasible claim.",
        "Check that claimed innovation is supported by the supplied material.",
    ),
    PresentationForm.ACADEMIC_DEFENSE: (
        "Check that the research question, method, evidence, contribution, and limitations agree.",
        "Check terminology, data basis, and conclusions against the supplied sources.",
    ),
    PresentationForm.TEACHING: (
        "Check that each page advances one understandable learning step and gives a concrete example where needed.",
        "Check that the sequence supports comprehension without hiding essential content.",
    ),
    PresentationForm.TECHNICAL_EXPLANATION: (
        "Check that the concept, mechanism, evidence, and operating boundary are explicit and accurate.",
        "Check that diagrams show real relationships rather than decorative interface elements.",
    ),
    PresentationForm.CEREMONIAL_STORY: (
        "Check that the page contributes to the occasion and the narrative's emotional progression.",
        "Check continuity of identity and visual rhythm across neighboring pages.",
    ),
    PresentationForm.PREMIUM_BRAND_PORTFOLIO: (
        "Check that authentic work and imagery carry the page, with their important details intact.",
        "Check editorial hierarchy and restraint against the actual portfolio material.",
    ),
}


def select_work_form(work_dir: Path, requested: str | None = None) -> dict:
    """Use an explicit form or a unanimous local brief, never a style guess.

    A Work may contain several historical Studio jobs. Conflicting briefs do
    not identify the current output, so they fall back to general review and
    are surfaced for the assistant to resolve.
    """
    if requested is not None:
        form = PresentationForm(requested)
        return {"form": form.value, "source": "explicit_cli", "briefs": []}
    rows = []
    root = Path(work_dir) / "workspace" / "studio-jobs"
    for path in sorted(root.glob("*/planning/brief.json")):
        try:
            document = json.loads(path.read_text(encoding="utf-8-sig"))
            raw = document["presentation_intent"]["presentation_form"]
            form = PresentationForm(raw)
        except (OSError, UnicodeError, json.JSONDecodeError, KeyError, TypeError, ValueError):
            continue
        rows.append({"path": path.relative_to(work_dir).as_posix(), "form": form.value})
    choices = {row["form"] for row in rows}
    if len(choices) == 1:
        return {"form": next(iter(choices)), "source": "unanimous_work_briefs_unbound_to_output", "briefs": rows}
    return {"form": PresentationForm.GENERAL.value,
            "source": "conflicting_work_briefs" if choices else "no_work_brief", "briefs": rows}


def audit_form_quality(pptx: Path, presentation_form: PresentationForm | str = PresentationForm.GENERAL) -> dict:
    """Return local review signals and a form-specific review lens for a PPTX.

    All returned detections have ``severity='advisory'`` and
    ``disposition='context_required'``. They are not part of an automatic pass
    or fail decision. Raster/master-only content is outside the detectors.
    """
    try:
        form = presentation_form if isinstance(presentation_form, PresentationForm) else PresentationForm(presentation_form)
    except (TypeError, ValueError) as error:
        raise ValueError(f"unsupported presentation form: {presentation_form!r}") from error

    path = Path(pptx).expanduser().resolve()
    source_hash = sha256(path)
    presentation = Presentation(path)
    signals: list[dict] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        objects = list(_objects(slide.shapes, presentation=presentation, slide=slide))
        for obj in objects:
            for candidate in copy_signals(obj["text"]):
                signals.append({**candidate, "severity": "advisory", "disposition": "context_required",
                                "slide": slide_number, "object_id": obj["object_id"],
                                "excerpt": candidate["excerpt"][:240],
                                "guidance": "Check whether this is audience-facing copy or a production instruction in this context. Preserve genuine examples and source wording when appropriate."})
        for candidate in decorative_ordinal_signals(objects, presentation.slide_width / 12700,
                                                     presentation.slide_height / 12700,
                                                     page_number=slide_number):
            signals.append({**candidate, "severity": "advisory", "disposition": "context_required",
                            "slide": slide_number,
                            "guidance": "Determine its audience-facing function on the rendered page and neighboring pages. Preserve useful page navigation, chapters, steps, and real quantities; remove unexplained decoration only after review."})
        for candidate in density_signals(objects, presentation.slide_width / 12700,
                                         presentation.slide_height / 12700):
            signals.append({**candidate, "severity": "advisory", "disposition": "context_required",
                            "slide": slide_number,
                            "guidance": "Inspect this explanatory copy at presentation scale alongside the full page. The measured geometry and type contrast locate a review candidate; they do not decide readability or meaning."})

    # Detect source mutation while reading; every finding is tied to current bytes.
    if sha256(path) != source_hash:
        raise ValueError("PPTX changed during form quality audit; rerun against an immutable candidate")
    return {
        "schema_version": 1,
        "status": "advisory_review",
        "source": str(path),
        "source_sha256": source_hash,
        "presentation_form": form.value,
        "slide_count": len(presentation.slides),
        "review_lens": list(_FORM_REVIEW_LENSES[form]),
        "signals": signals,
        "summary": {"signals": len(signals), "decorative_ordinal_candidates": sum(s["code"] == "DECORATIVE_ORDINAL" for s in signals),
                    "production_copy_candidates": sum(s["code"] in {"PRODUCTION_INSTRUCTION", "ASSISTANT_REPLY", "REVIEW_REMINDER", "INTERNAL_LABEL"} for s in signals),
                    "density_context_candidates": sum(s["code"] == "TEXT_DENSITY_CONTEXT" for s in signals),
                    "automatic_removals": 0, "automatic_visual_score": None},
        "coverage": "Slide-level editable text only. Rasterized text, master-only ornaments, glyph collision, page balance, image meaning, and overall visual quality require inspection of current rendered pages.",
    }
