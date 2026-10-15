"""Small, file-oriented PowerPoint operations that do not create Studio Works."""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


def _pptx_path(value: Path, role: str) -> Path:
    path = Path(value).expanduser().absolute()
    if path.suffix.lower() != ".pptx":
        raise ValueError(f"{role} must use the .pptx extension: {path}")
    return path


def _publish_new_file(temporary: Path, output: Path) -> None:
    """Publish without replacing a file created concurrently by another process."""
    try:
        os.link(temporary, output)
    except FileExistsError as error:
        raise FileExistsError(f"output already exists: {output}") from error
    temporary.unlink()


def create_blank_pptx(output: Path) -> dict[str, object]:
    """Create a one-slide, blank 16:9 editable deck without launching Office."""
    output = _pptx_path(output, "output")
    if output.exists() or output.is_symlink():
        raise FileExistsError(f"output already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(40 / 3)
    presentation.slide_height = Inches(7.5)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output.stem}.", suffix=".pptx", dir=output.parent, delete=False
        ) as handle:
            temporary = Path(handle.name)
        presentation.save(temporary)
        _publish_new_file(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return {
        "status": "created",
        "path": str(output),
        "slides": 1,
        "width_inches": 40 / 3,
        "height_inches": 7.5,
    }


def save_as_pptx(source: Path, output: Path, *, replace: bool = False) -> dict[str, object]:
    """Copy saved PPTX bytes to a new path, optionally updating that path atomically."""
    source = _pptx_path(source, "source").resolve()
    output = _pptx_path(output, "output")
    if not source.is_file():
        raise FileNotFoundError(f"source PPTX does not exist: {source}")
    if source == output.resolve(strict=False):
        raise ValueError("output must be a different path from source")
    if output.is_symlink():
        raise ValueError(f"output must not be a symlink: {output}")
    if output.exists() and not replace:
        raise FileExistsError(f"output already exists: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)

    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output.stem}.", suffix=".pptx", dir=output.parent, delete=False
        ) as handle:
            temporary = Path(handle.name)
        shutil.copyfile(source, temporary)
        if temporary.stat().st_size == 0:
            raise ValueError("source PPTX is empty")
        if replace:
            temporary.replace(output)
        else:
            _publish_new_file(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return {"status": "saved_as", "source": str(source), "path": str(output),
            "bytes": output.stat().st_size, "replaced": replace}


def request_powerpoint_open(source: Path) -> dict[str, object]:
    """Ask macOS Launch Services to open a PPTX; this does not verify native review."""
    if sys.platform != "darwin":
        raise RuntimeError("pptx-open is available only on macOS")
    source = _pptx_path(source, "source").resolve()
    if not source.is_file():
        raise FileNotFoundError(f"source PPTX does not exist: {source}")
    subprocess.run(
        ["/usr/bin/open", "-a", "Microsoft PowerPoint", str(source)],
        check=True,
        capture_output=True,
        text=True,
        timeout=30,
    )
    return {"status": "open_requested", "path": str(source), "native_review": "not_verified"}
