"""Bind the current assistant's page-expression judgment to actual candidates.

This module collects evidence and validates review coverage. It does not pretend
that regex or a geometry score can decide whether a presentation communicates.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import DeckPlan, JobState, select_prototype_slides
from .renderer import sha256, powerpoint_roundtrip, natural_sort_paths
from .storage import job_dir, job_lock, load_job, read_json, read_model, save_job, write_json


REVIEW_LENS = [
    "面向本次听众，这页究竟说明了什么，标题与证据是否对应？",
    "用演讲者或汇报人的口吻读一遍：是否具体、准确、符合场合？有无助手对话、制作提醒、空泛套话？",
    "看完整页面：大小、位置、宽度和留白是否体现信息的重要性与关系？是否有巨大孤立文字旁挤窄密说明？",
    "图、文、数据是否共同说明一件事？叠放是有意构图还是遮挡？",
    "每张图是否对应本页论点，是否重复、变形或裁掉关键信息？证照、专利、人物按其证据/身份作用是否看得清？截图是否真的支持标题？",
    "图片实际可见边缘与边框、底形是否协调？圆角或椭圆是否裁在图片本身，还是外面一个曲框、里面仍露直角？检查不同长宽比、留边、白边黑边和人物/证据被裁掉的部分。",
    "精简后是否仍保留必要事实、名称、完整指标、单位和比较基线？百分比与百分点有没有混淆？强调是否落在真正重要的关系或异常上？",
    "参考中的元素关系、间距、人物位置和已认可视觉体系是否得到体现？不能只套颜色、加框或重复卡片。局部改稿是否保护最新版已调好的对象，修改在实际PPTX里是否真的可见？",
    "页角的01、02等独立数字传达了什么？保留有作用的页码、章节导航、步骤或数据；没有信息作用的装饰编号应先从作品中移除，而不是补一段解释文字。",
    "结合前后页：这页补充了什么，是否缺前提、重复结论或突然换话题？",
    "检查后先修作品；审阅意见、制作过程和QA提醒不得作为修复文字放到PPT里。",
]

# Only precise production-language signals. Quoted examples in a training deck
# may be legitimate; the assistant must inspect the context and record why.
_COPY_SIGNALS = [
    ("PRODUCTION_INSTRUCTION", re.compile(r"(?:这[一张]?页|本页|该页|此页).{0,24}(?:需要|应当|应该|建议|注意|记得|待).{0,28}(?:排版|字号|留白|对齐|遮挡|配图|前[一]?页|后[一]?页|其他页面|衔接|关联|调整|补充)", re.S)),
    ("ASSISTANT_REPLY", re.compile(r"(?:我会|我将|我可以|我帮你|为您|为你).{0,24}(?:制作|生成|调整|修改|优化|检查|排版).{0,16}(?:PPT|幻灯片|页面|这一页|这页|演示文稿)", re.I | re.S)),
    ("ASSISTANT_REPLY", re.compile(r"(?:以下是|这是)(?:我)?(?:为您|为你|给您|给你).{0,16}(?:准备|制作|整理|生成).{0,12}(?:的)?[^。！？；\n]{0,32}(?:PPT|幻灯片|演示文稿)", re.I)),
    ("PRODUCTION_INSTRUCTION", re.compile(r"(?:这一页|这页|本页|该页|此页).{0,12}(?:可以|可|建议|应该|需要).{0,10}(?:放|加|配|写|排|补).{0,18}(?:标题|图片|配图|文字|内容|项目背景|流程|图表|照片|说明)", re.S)),
    ("REVIEW_REMINDER", re.compile(r"(?:请|需要|记得|注意).{0,12}(?:检查|确认|调整).{0,24}(?:页面关联|与前页|与后页|图文遮挡|字号大小|排版效果|占位文字)", re.S)),
    ("INTERNAL_LABEL", re.compile(r"(?:^|\n)\s*(?:制作说明|排版建议|设计提示|待办事项|QA检查|修改备注|助手回复)\s*[:：]", re.I)),
]


def copy_signals(text: str) -> list[dict]:
    # Quoted training examples and reported speech remain reviewable content.
    # Suppress only matches fully enclosed by explicit quotation marks.
    quoted = []
    for quote_open, quote_close in (("“", "”"), ("「", "」"), ("\"", "\""), ("'", "'")):
        start = 0
        while (left := text.find(quote_open, start)) >= 0:
            right = text.find(quote_close, left + len(quote_open))
            if right < 0:
                break
            quoted.append((left, right + len(quote_close)))
            start = right + len(quote_close)
    result = []
    for code, pattern in _COPY_SIGNALS:
        for match in pattern.finditer(text):
            if any(left <= match.start() and match.end() <= right for left, right in quoted):
                continue
            result.append({"code": code, "excerpt": match.group(0), "kind": "context_required"})
    return result


def density_signals(objects: list[dict], width_pt: float, height_pt: float) -> list[dict]:
    """Locate long explanatory copy whose measured page geometry may impede reading.

    This is a contextual prompt only: it does not judge meaning, actual wrapping,
    contrast, or audience fit, and never applies a universal font-size cutoff.
    """
    if not isinstance(width_pt, (int, float)) or not isinstance(height_pt, (int, float)) or width_pt <= 0 or height_pt <= 0:
        return []
    sizes = [size for obj in objects for size in obj.get("font_sizes_pt", [])
             if isinstance(size, (int, float)) and size > 0]
    largest = max(sizes, default=None)
    signals = []
    for obj in objects:
        text = obj.get("text", "")
        if not isinstance(text, str) or len(text.strip()) < 70:
            continue
        bounds = obj.get("bounds_pt", {})
        width = bounds.get("width") if isinstance(bounds, dict) else None
        height = bounds.get("height") if isinstance(bounds, dict) else None
        if not isinstance(width, (int, float)) or not isinstance(height, (int, float)) or width <= 0:
            continue
        sizes_here = [v for v in obj.get("font_sizes_pt", []) if isinstance(v, (int, float)) and v > 0]
        local = max(sizes_here, default=None)
        contrast = largest / local if largest and local else None
        lines = max(text.count("\n") + 1, int(obj.get("paragraph_count", 1) or 1))
        width_ratio = width / width_pt
        # A narrow measure is a candidate only when copy is substantial. A
        # large type contrast is supporting evidence, not a sufficient trigger.
        narrow = width_ratio <= .28 and (len(text) >= 90 or lines >= 4)
        contrast_case = width_ratio <= .62 and len(text) >= 100 and contrast is not None and contrast >= 2.2
        if not (narrow or contrast_case):
            continue
        signals.append({"code": "TEXT_DENSITY_CONTEXT", "kind": "context_required",
                        "object_id": obj.get("object_id"), "excerpt": text[:180],
                        "components": {"characters": len(text), "line_count_or_paragraphs": lines,
                                       "width_pt": round(width, 2), "width_ratio": round(width_ratio, 4),
                                       "height_ratio": round(height / height_pt, 4),
                                       "max_font_pt_in_object": local,
                                       "largest_font_pt_on_page": largest,
                                       "font_contrast_ratio": round(contrast, 3) if contrast else None,
                                       "narrow_region_candidate": narrow,
                                       "large_type_contrast_candidate": contrast_case},
                        "limitations": "This signal does not judge meaning, actual rendered line wrapping, contrast, or audience needs; font sizes are evidence for review, not a pass/fail threshold."})
    return signals


def _digest(value: dict) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _objects(shapes, parent=None, group_ids=(), presentation=None, slide=None):
    from .presentation_quality import IDENTITY, _group_transform, _transform
    parent = IDENTITY if parent is None else parent
    for shape in shapes:
        if shape.element.xpath('./*/p:cNvPr[@hidden="1" or @hidden="true"]'):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _objects(shape.shapes, _group_transform(shape, parent), (*group_ids, str(shape.shape_id)), presentation, slide)
            continue
        paragraphs = shape.element.xpath('.//a:p')
        text = '\n'.join(''.join(p.xpath('.//a:t/text()')) for p in paragraphs).strip()
        if text:
            _, rect, _, _ = _transform(shape, parent)
            hints = []
            if (shape.element.xpath('./p:nvSpPr/p:nvPr/p:ph[@type="sldNum"]')
                    or shape.element.xpath('.//a:fld[@type="slidenum"]')):
                hints.append('ooxml_page_number')
            paragraphs_count = len([p for p in paragraphs if ''.join(p.xpath('.//a:t/text()')).strip()])
            font_sizes = {int(v) / 100 for v in shape.element.xpath('.//a:rPr/@sz')}
            if presentation is not None and slide is not None and getattr(shape, "has_text_frame", False):
                from .presentation_quality import _text_samples
                samples = _text_samples(presentation, slide, shape, shape.text_frame._txBody, "text")
                font_sizes.update(sample["font_pt"] for sample in samples if sample.get("font_pt") is not None)
            yield {"object_id": str(shape.shape_id), "text": text,
                   "bounds_pt": dict(zip(("left", "top", "width", "height"),
                                         (round(v, 2) for v in (rect[0], rect[1], rect[2]-rect[0], rect[3]-rect[1])))),
                   "group_ids": list(group_ids), "semantic_hints": hints,
                   "paragraph_count": paragraphs_count,
                   "font_sizes_pt": sorted(font_sizes)}


def decorative_ordinal_signals(objects: list[dict], width_pt: float, height_pt: float, *, page_number: int) -> list[dict]:
    """Locate isolated corner numerals for contextual review, never ban numbers.

    The outer quarters only locate candidate corners; they do not score design.
    A real page number, chapter, step or quantity can be legitimate. Text in an
    image or a master is outside this detector and remains part of image review.
    """
    signals = []
    for obj in objects:
        text = obj['text'].strip()
        if not re.fullmatch(r'[0-9]{1,3}', text):
            continue
        bounds = obj['bounds_pt']
        x, y = bounds['left'] + bounds['width']/2, bounds['top'] + bounds['height']/2
        horizontal = 'left' if x <= width_pt/4 else 'right' if x >= width_pt*3/4 else None
        vertical = 'top' if y <= height_pt/4 else 'bottom' if y >= height_pt*3/4 else None
        if not horizontal or not vertical:
            continue
        hints = list(obj.get('semantic_hints', []))
        if int(text) == page_number:
            hints.append('matches_deck_page_number')
        nearby = []
        for other in objects:
            if other['object_id'] == obj['object_id']:
                continue
            box = other['bounds_pt']
            distance = ((box['left'] + box['width']/2-x)**2 + (box['top'] + box['height']/2-y)**2)**.5
            nearby.append((distance, other))
        signals.append({'code': 'DECORATIVE_ORDINAL', 'kind': 'context_required',
                        'object_id': obj['object_id'], 'excerpt': text, 'bounds_pt': dict(bounds),
                        'location': f'{vertical}_{horizontal}', 'semantic_hints': hints,
                        'nearby_text': [{'object_id': other['object_id'], 'text': other['text'][:160],
                                         'center_distance_pt': round(distance, 2)}
                                        for distance, other in sorted(nearby, key=lambda item: item[0])[:3]],
                        'guidance': 'Identify the number\'s audience-facing function in the actual page and adjacent pages. Page numbers, chapter navigation, real steps and quantities can be legitimate; remove unexplained decoration from the candidate, not genuine source numbering.'})
    return signals


def review_required(root: Path) -> bool:
    policy = root / 'planning/authoring-policy.json'
    return policy.is_file() and read_json(policy).get('expression_review_required') is True


def _current_previews(root: Path, gate: str, deck: Path) -> dict:
    path = root / 'reviews' / f'{gate}-expression-render.json'
    if not path.is_file():
        raise ValueError('Expression page images are missing or stale; prepare expression-review to render the current PPTX')
    binding = read_json(path)
    if binding.get('source_sha256') != sha256(deck):
        raise ValueError('Expression page images are stale; prepare expression-review for the current PPTX')
    for relative, digest in binding.get('files', {}).items():
        p = (root / relative).resolve()
        if not p.is_relative_to(root.resolve()) or not p.is_file() or sha256(p) != digest:
            raise ValueError('Expression page images are stale or changed; prepare expression-review again')
    if not binding.get('previews'):
        raise ValueError('Expression render has no page images')
    return binding


def _refresh_previews(root: Path, gate: str) -> None:
    """Render a copy of the actual PPTX; never pair it with old HTML previews."""
    deck = root / 'build' / ('prototype.pptx' if gate == 'prototype' else 'deck.pptx')
    try:
        _current_previews(root, gate, deck)
        return
    except ValueError:
        pass
    source_hash = sha256(deck)
    # Keep Office staging paths short; the manifest holds source identity.
    directory = root / 'reviews' / f'view-{uuid.uuid4().hex[:8]}'
    directory.mkdir(parents=True)
    candidate = directory / 'candidate.pptx'
    shutil.copy2(deck, candidate)
    report = powerpoint_roundtrip(candidate, directory)
    previews = natural_sort_paths(list((directory / 'slides').glob('*.png')))
    count = len(Presentation(deck).slides)
    if len(previews) != count or sha256(deck) != source_hash:
        raise ValueError('Expression render is incomplete or candidate changed during rendering')
    write_json(directory / 'roundtrip.json', report)
    files = [candidate, directory / 'roundtrip.json', *previews]
    binding = {'source_sha256': source_hash,
               'previews': [p.relative_to(root).as_posix() for p in previews],
               'files': {p.relative_to(root).as_posix(): sha256(p) for p in files},
               'method': 'Native Office render of a copied current PPTX; original candidate preserved'}
    write_json(root / 'reviews' / f'{gate}-expression-render.json', binding)


def _packet(root: Path, gate: str) -> dict:
    if gate not in {'prototype', 'final'}:
        raise ValueError('expression review gate must be prototype or final')
    deck = root / 'build' / ('prototype.pptx' if gate == 'prototype' else 'deck.pptx')
    binding = _current_previews(root, gate, deck)
    plan = read_model(root / 'planning/deck-plan.json', DeckPlan)
    selected = select_prototype_slides(plan) if gate == 'prototype' else plan.slides
    presentation = Presentation(deck)
    if len(presentation.slides) != len(selected):
        raise ValueError('Actual deck and selected plan pages differ; rebuild before expression review')
    files = [deck, root / 'planning/deck-plan.json', root / 'reviews' / f'{gate}-expression-render.json']
    files.extend(root / relative for relative in binding['files'])
    context = {}
    for relative in ('planning/brief.json', 'planning/authoring-policy.json', 'evidence-ledger.json', 'planning/editorial-inputs.json', 'planning/design-context.json'):
        path = root / relative
        if path.is_file():
            files.append(path)
            context[relative] = read_json(path)
    design_context = context.get('planning/design-context.json', {})
    design_binding_current = bool(design_context) and (
        design_context.get('plan_binding', {}).get('plan_sha256') == sha256(root / 'planning/deck-plan.json')
        and design_context.get('scene_sha256') == (
            sha256(root / 'planning/native-scene.json') if (root / 'planning/native-scene.json').is_file() else None)
        and design_context.get('plan_binding', {}).get('brief_sha256') == (
            sha256(root / 'planning/brief.json') if (root / 'planning/brief.json').is_file() else None))
    pages = []
    for index, (slide, spec) in enumerate(zip(presentation.slides, selected)):
        preview = root / binding['previews'][index]
        if not preview.is_file():
            raise ValueError(f'Actual page preview missing: {spec.id}; render before reviewing')
        files.append(preview)
        objects = list(_objects(slide.shapes, presentation=presentation, slide=slide))
        plan_index = next(i for i, row in enumerate(plan.slides) if row.id == spec.id)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ''
        signals = []
        for obj in [*objects, {"object_id": "speaker_notes", "text": notes}]:
            for signal in copy_signals(obj['text']):
                item = {**signal, 'object_id': obj['object_id'], 'slide_id': spec.id}
                item['signal_id'] = _digest(item)[:20]
                signals.append(item)
        for signal in decorative_ordinal_signals(objects, presentation.slide_width/12700,
                                                 presentation.slide_height/12700, page_number=plan_index + 1):
            item = {**signal, 'slide_id': spec.id}
            item['signal_id'] = _digest(item)[:20]
            signals.append(item)
        for signal in density_signals(objects, presentation.slide_width/12700, presentation.slide_height/12700):
            item = {**signal, 'slide_id': spec.id}
            item['signal_id'] = _digest(item)[:20]
            signals.append(item)
        from .presentation_quality import slide_image_geometry_signals
        for signal in slide_image_geometry_signals(presentation, slide, plan_index + 1):
            item = {**signal, 'slide_id': spec.id}
            item['signal_id'] = _digest(item)[:20]
            signals.append(item)
        pages.append({'slide_id': spec.id, 'page_number': plan_index + 1,
                      'preview': preview.relative_to(root).as_posix(), 'preview_sha256': sha256(preview),
                      'objects': objects, 'speaker_notes': notes, 'signals': signals,
                      'internal_plan': spec.model_dump(mode='json'),
                      'previous_page': plan.slides[plan_index - 1].model_dump(mode='json') if plan_index else None,
                      'next_page': plan.slides[plan_index + 1].model_dump(mode='json') if plan_index + 1 < len(plan.slides) else None})
    inputs = {p.relative_to(root).as_posix(): sha256(p) for p in files}
    from .lesson_application import project_lesson_context, plan_query
    project_lessons = project_lesson_context(root, plan_query(plan, context))
    # Unrelated projects entering/leaving an error state do not invalidate a
    # reviewed deck. Selected lesson content/receipts still bind the judgment.
    project_lessons = {"lessons": project_lessons["lessons"],
                       "instruction": project_lessons.get("instruction", "")}
    packet = {'schema_version': 1, 'gate': gate, 'status': 'awaiting_assistant_review',
              'inputs': inputs, 'context': context, 'slides': pages,
              'project_lessons': project_lessons,
              'design_context_status': 'current_advisory' if design_binding_current else 'stale_or_not_recorded',
              'review_lens': REVIEW_LENS,
              'ordinal_signal_coverage': 'Isolated numeric text in slide-level corner objects, including group transforms; master-only objects and raster text require the actual page image. Signals are context prompts, not automatic removal or aesthetic failure.',
              'image_geometry_signal_coverage': 'Opaque pictures over curved shapes with coincident transformed bounds are candidates for edge review. Objects are not automatically paired or cropped; other sizes, rasterized frames, alpha edges and master content still require the actual page image.',
              'instruction': 'Open every referenced page image at presentation scale and compare adjacent pages. Read the actual copy and source context. Record a concrete audience takeaway and whole-page assessment. Automatic signals are prompts, not judgments. Do not fill ready before looking, and do not ask the user to perform this review for the assistant.'}
    packet['packet_sha256'] = _digest(packet)
    return packet


def prepare_expression_review(job_id: str, gate: str) -> dict:
    with job_lock(job_id):
        if gate not in {'prototype', 'final'}:
            raise ValueError('expression review gate must be prototype or final')
        root = job_dir(job_id)
        _refresh_previews(root, gate)
        packet = _packet(root, gate)
        path = root / 'reviews' / f'{gate}-expression-packet.json'
        write_json(path, packet)
        return {'status': packet['status'], 'packet': str(path), 'packet_sha256': packet['packet_sha256'],
                'slides': len(packet['slides']), 'review_lens': REVIEW_LENS,
                'design_context_status': packet['design_context_status'],
                'design_review_focus': 'Inspect narrative intent, page relationships and actual reference/output pairs in packet.context; discovery or object links do not establish design quality.'}


def _nonempty(value, field: str) -> None:
    if not isinstance(value, str) or not value.strip() or value.strip().lower() in {'passed', 'ok', 'ready', '通过', '无', '正常'}:
        raise ValueError(f'{field} requires a concrete observation, not a pass label')


def _validate_review(packet: dict, request: dict) -> dict:
    allowed = {'packet_sha256', 'reviewer', 'narrative_assessment', 'slides', 'lesson_dispositions'}
    if not isinstance(request, dict) or set(request) - allowed or request.get('reviewer') != 'current_assistant':
        raise ValueError('Review must contain current_assistant judgment and supported fields only')
    if request.get('packet_sha256') != packet['packet_sha256']:
        raise ValueError('Expression review is stale; prepare and inspect the current candidate again')
    _nonempty(request.get('narrative_assessment'), 'narrative_assessment')
    rows = request.get('slides')
    if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
        raise ValueError('slides must contain one assessment per actual page')
    expected = {page['slide_id']: page for page in packet['slides']}
    if len(rows) != len(expected) or {row.get('slide_id') for row in rows} != set(expected):
        raise ValueError('Every actual page must be reviewed exactly once')
    issues = []
    for row in rows:
        if set(row) - {'slide_id', 'preview_sha256', 'audience_takeaway', 'page_assessment', 'decision', 'signal_dispositions'}:
            raise ValueError('Unsupported page review fields')
        page = expected[row['slide_id']]
        if row.get('preview_sha256') != page['preview_sha256']:
            raise ValueError('Page review must identify the preview that was inspected')
        for field in ('audience_takeaway', 'page_assessment'):
            _nonempty(row.get(field), field)
        if row.get('decision') not in {'ready', 'revise'}:
            raise ValueError('Page decision must be ready or revise')
        dispositions = row.get('signal_dispositions', [])
        signals = {item['signal_id'] for item in page['signals']}
        if (not isinstance(dispositions, list) or not all(isinstance(d, dict) for d in dispositions)
                or len(dispositions) != len(signals) or {d.get('signal_id') for d in dispositions} != signals):
            raise ValueError('Address each expression signal in its actual context')
        for disposition in dispositions:
            if set(disposition) != {'signal_id', 'decision', 'reason'} or disposition.get('decision') not in {'legitimate', 'revise'}:
                raise ValueError('Signal disposition must be legitimate or revise with a reason')
            _nonempty(disposition.get('reason'), 'signal reason')
            if disposition['decision'] == 'revise' and row['decision'] == 'ready':
                raise ValueError('A page with unresolved copy cannot be ready')
        if row['decision'] == 'revise':
            issues.append({'code': 'PAGE_EXPRESSION_REQUIRES_REVISION', 'slide_id': row['slide_id'],
                           'assessment': row['page_assessment']})
    from .lesson_application import validate_lesson_dispositions
    lesson_rows, lesson_issues = validate_lesson_dispositions(packet, request)
    issues.extend(lesson_issues)
    result = {'schema_version': 1, 'gate': packet['gate'], 'packet_sha256': packet['packet_sha256'],
            'inputs': packet['inputs'], 'status': 'ready' if not issues else 'needs_revision',
            'reviewer': 'current_assistant', 'basis': 'Assistant judgment of the bound actual pages; not an automated aesthetic score or user acceptance',
            'narrative_assessment': request['narrative_assessment'], 'slides': rows, 'issues': issues}
    if lesson_rows or 'lesson_dispositions' in request:
        result['lesson_dispositions'] = lesson_rows
    return result


def _record_project_learning_expression_qa(root: Path, gate: str, review_path: Path, report: dict) -> None:
    """Write a compact final-expression QA receipt when its Work proof is live.

    The expression-review document intentionally keeps its domain status
    (``ready``/``needs_revision``).  This sibling QA receipt expresses the
    narrower archival fact that a ready review covered the current canonical
    deck, while preserving the original review format and its validation path.
    """

    if gate != "final":
        return
    receipt_path = root / "qa" / "expression-review-final.json"
    receipt = {
        "schema_version": 1,
        "qa_kind": "current_assistant_expression_review",
        "status": "needs_revision",
        "review": {
            "path": review_path.relative_to(root).as_posix(),
            "sha256": sha256(review_path),
            "packet_sha256": report.get("packet_sha256"),
            "status": report.get("status"),
        },
    }
    if report.get("status") == "ready":
        try:
            from .project_learning import current_qa_provenance_for_output

            proof = current_qa_provenance_for_output(root, root / "build" / "deck.pptx")
            inputs = report.get("inputs")
            if proof is not None and isinstance(inputs, dict) and inputs.get("build/deck.pptx") == proof["output_sha256"]:
                receipt.update(status="passed", project_learning_provenance=proof)
            else:
                receipt["reason"] = "current_work_bound_output_unavailable"
        except (OSError, ValueError, TypeError, AttributeError):
            receipt["reason"] = "current_work_bound_output_unavailable"
    else:
        receipt["reason"] = "expression_review_not_ready"
    write_json(receipt_path, receipt)


def record_expression_review(job_id: str, gate: str, request: dict) -> dict:
    with job_lock(job_id):
        root = job_dir(job_id)
        job = load_job(job_id)
        expected = JobState.PROTOTYPE_REVIEW if gate == 'prototype' else JobState.FINAL_REVIEW
        recover = False
        if job.state == JobState.BLOCKED and job.resume_state == expected:
            from .reconcile import audit_job
            issues = audit_job(job_id).issues
            candidate_artifacts = {'expression_review', 'native_scene', 'deck_pptx', 'deck_html', 'previews', 'montage',
                                   'prototype_pptx', 'prototype_html', 'prototype_previews', 'prototype_montage'}
            recover = bool(issues) and all(
                issue.code in {'approval_hash_mismatch', 'approval_artifact_missing'}
                and issue.message.startswith(f'{gate} approval artifact ')
                and issue.message.rsplit(': ', 1)[-1] in candidate_artifacts for issue in issues)
        if gate not in {'prototype', 'final'} or (job.state != expected and not recover):
            raise ValueError('Record expression review only in the corresponding candidate review stage')
        packet = _packet(root, gate)
        report = _validate_review(packet, request)
        path = root / 'reviews' / f'{gate}-expression-review.json'
        if path.is_file() and read_json(path) == report and not recover:
            _record_project_learning_expression_qa(root, gate, path, report)
            return report
        if path.is_file():
            write_json(root / 'reviews/expression-history' / f'{sha256(path)}.json', read_json(path))
        write_json(path, report)
        _record_project_learning_expression_qa(root, gate, path, report)
        job.approvals[gate] = False
        job.approval_hashes.pop(gate, None)
        if recover:
            job.state = expected
            job.resume_state = None
            job.blocked_reason = None
            job.failure = None
        if gate == 'prototype':
            job.approvals['final'] = False
            job.approval_hashes.pop('final', None)
        save_job(job)
        if report['issues']:
            from .authoring_policy import record_quality_findings
            record_quality_findings(root, {**report, 'status': 'needs_review'})
        return report


def require_expression_review(root: Path, gate: str) -> dict:
    if not review_required(root):
        return {'status': 'legacy_not_required', 'visual_reviewed': False}
    if os.environ.get('AIPPT_FAKE_RENDER') == '1':
        return {'status': 'skipped-test-render', 'visual_reviewed': False}
    path = root / 'reviews' / f'{gate}-expression-review.json'
    if not path.is_file():
        raise ValueError(f'Assistant whole-page expression review is missing; run expression-review for {gate}, inspect the pages and record the assessment')
    packet = _packet(root, gate)
    report = read_json(path)
    request = {key: report.get(key) for key in ('packet_sha256', 'reviewer', 'narrative_assessment', 'slides')}
    if 'lesson_dispositions' in report:
        request['lesson_dispositions'] = report['lesson_dispositions']
    validated = _validate_review(packet, request)
    if validated != report:
        raise ValueError('Expression review record differs from its validated content')
    if report['status'] != 'ready':
        raise ValueError('Page expression still needs revision; fix and inspect the current pages before approval')
    return report
