"""Live platform-aware dependency selection and bounded end-to-end runtime QA.

Directory presence is never a capability result. Historical Work caches are
diagnostic evidence, not a source of native dependencies for a new task.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys

from .config import PROJECT_ROOT
from .runtime import binary_identity, compatible_binary, historical_runtime_path, host_identity, resolve_executable
from .python_runtime import current_process_identity, validate_identity


def _write(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _artifact_fingerprint(package: Path) -> dict:
    return {name: _file_hash(package / name) for name in ("package.json", "dist/artifact_tool.mjs", "node_modules/skia-canvas/lib/skia.node")}


def historical_cache(path: Path) -> bool:
    return historical_runtime_path(path)


def _run(command: list[str], *, cwd: Path | None = None, timeout: int = 30) -> dict:
    env = dict(os.environ)
    # An inherited --require/--import or NODE_PATH can select old native caches
    # before our explicit resolution. These probes use only their chosen roots.
    env.pop("NODE_OPTIONS", None)
    env.pop("NODE_PATH", None)
    try:
        result = subprocess.run(command, cwd=cwd, env=env, capture_output=True, text=True,
                                encoding="utf-8", errors="replace", timeout=timeout, check=False)
        return {"status": "passed" if result.returncode == 0 else "failed", "exit_code": result.returncode,
                "stdout": result.stdout[-5000:], "stderr": result.stderr[-3000:]}
    except (OSError, subprocess.TimeoutExpired) as error:
        return {"status": "failed", "error": str(error)}


def probe_node_runtime(node: Path | None = None) -> dict:
    node = Path(node) if node is not None else resolve_executable("node")
    if node is None or not node.is_file():
        return {"status": "failed", "reason": "A compatible Node.js executable was not found"}
    if historical_cache(node) or not compatible_binary(node):
        return {"status": "failed", "path": str(node), "reason": "Historical Work cache or foreign native executable is not a runtime source"}
    result = _run([str(node), "-p", "JSON.stringify({platform:process.platform,architecture:process.arch,version:process.version,napi:process.versions.napi,executable:process.execPath})"])
    if result["status"] == "passed":
        try:
            identity = json.loads(result["stdout"].strip())
            native_arch = {"x86_64": "x64", "amd64": "x64", "aarch64": "arm64"}.get(host_identity()["architecture"], host_identity()["architecture"])
            # Puppeteer 25 in the locked core requires Node >=22.12, not merely
            # the old package's >=20 declaration. Probe the actual chosen host.
            version = tuple(int(v) for v in identity["version"].lstrip("v").split(".")[:2])
            if identity["platform"] != sys.platform or identity["architecture"] != native_arch or version < (22, 12):
                raise ValueError("Node must match the current platform/architecture and be >=22.12 for the locked core")
            return {"status": "passed", "path": str(node.resolve()), **identity}
        except (ValueError, KeyError, TypeError) as error:
            return {"status": "failed", "path": str(node), "reason": str(error), "probe": result}
    return {"status": "failed", "path": str(node), "probe": result}


def artifact_candidates() -> list[Path]:
    values = [os.environ.get("AIPPT_ARTIFACT_NODE_MODULES"), os.environ.get("RUNTIME_NODE_MODULES"),
              str(Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules"),
              str(PROJECT_ROOT / "node_modules")]
    return list(dict.fromkeys(Path(v).expanduser() for v in values if v))


_ARTIFACT_PROBE = r"""
const {createRequire}=require('node:module');
const {pathToFileURL}=require('node:url');
const fs=require('node:fs'); const path=require('node:path');
(async()=>{
  const root=process.argv[1], output=process.argv[2];
  const packageRoot=path.join(root,'@oai/artifact-tool');
  const req=createRequire(path.join(packageRoot,'package.json'));
  const {Canvas}=req('skia-canvas');
  const canvas=new Canvas(64,48); const ctx=canvas.getContext('2d');
  ctx.fillStyle='#ffffff';ctx.fillRect(0,0,64,48);
  ctx.fillStyle='#1759a7';ctx.fillRect(4,4,28,24);
  ctx.fillStyle='#b03030';ctx.font='14px sans-serif';ctx.fillText('图 A',4,43);
  const png=await canvas.toBuffer('png');fs.writeFileSync(output,png);
  const meta=JSON.parse(fs.readFileSync(path.join(packageRoot,'package.json'),'utf8'));
  const entry=typeof meta.exports?.['.']==='string'?meta.exports['.']:meta.main;
  const artifact=await import(pathToFileURL(path.join(packageRoot,entry)).href);
  if(typeof artifact.Presentation?.create!=='function'||typeof artifact.PresentationFile?.exportPptx!=='function') throw Error('Presentation API missing');
  const deck=artifact.Presentation.create({slideSize:{width:640,height:360}});deck.slides.add();
  console.log(JSON.stringify({version:meta.version,png_bytes:png.length,presentation_api:true,platform:process.platform,architecture:process.arch}));
})().catch(e=>{console.error(e.stack||e);process.exitCode=1});
"""


def select_artifact_runtime(output_dir: Path, candidates: list[Path] | None = None) -> dict:
    """Select only a live, host-native runtime; optional failure never blocks core."""
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    node = probe_node_runtime()
    result = {"schema_version": 1, "status": "unavailable", "host": host_identity(), "node": node,
              "environment": {}, "selected_node_modules": None, "capabilities": [],
              "rejected_candidates": [], "optional": True, "core_blocking": False}
    if node["status"] == "passed":
        result["environment"]["AIPPT_NODE"] = node["path"]
        for index, candidate in enumerate(candidates if candidates is not None else artifact_candidates()):
            candidate = Path(candidate).expanduser()
            item = {"path": str(candidate), "resolved": str(candidate.resolve())}
            package = candidate / "@oai/artifact-tool"
            native = package / "node_modules/skia-canvas/lib/skia.node"
            if historical_cache(candidate):
                item["reason"] = "historical_work_cache_rejected"
                if native.is_file():
                    item["native_binary"] = binary_identity(native)
                result["rejected_candidates"].append(item)
                continue
            if not (package / "package.json").is_file() or not native.is_file():
                item["reason"] = "artifact_package_or_native_addon_missing"
                result["rejected_candidates"].append(item)
                continue
            item["native_binary"] = binary_identity(native)
            if not compatible_binary(native, architecture=node["architecture"]):
                item["reason"] = "native_addon_platform_architecture_mismatch"
                result["rejected_candidates"].append(item)
                continue
            png = output / f"artifact-probe-{index}.png"
            probe = _run([node["path"], "-e", _ARTIFACT_PROBE, str(candidate.resolve()), str(png.resolve())], timeout=30)
            try:
                if probe["status"] != "passed":
                    raise ValueError("native drawing/import probe failed")
                from PIL import Image
                with Image.open(png) as image:
                    image.load()
                    if image.size != (64, 48) or len(image.convert("RGB").getcolors(4096) or []) < 2:
                        raise ValueError("drawing probe did not produce the expected nonblank raster")
                identity = json.loads(probe["stdout"].splitlines()[-1])
                fingerprint = _artifact_fingerprint(package)
            except (OSError, ValueError, IndexError) as error:
                item.update({"reason": str(error), "probe": probe})
                result["rejected_candidates"].append(item)
                continue
            result.update({"status": "ready", "selected_node_modules": str(candidate.resolve())})
            result["environment"]["RUNTIME_NODE_MODULES"] = str(candidate.resolve())
            result["capabilities"].append({"name": "artifact-tool", "available": True, "required_for": "optional",
                "verification": "native raster drawing and Presentation API construction", "path": str(package.resolve()),
                "native_binary": item["native_binary"], "probe": identity,
                "runtime_fingerprint": fingerprint,
                "evidence": str(png.resolve()), "sha256": hashlib.sha256(png.read_bytes()).hexdigest()})
            break
    if result["status"] != "ready":
        result["capabilities"].append({"name": "artifact-tool", "available": False, "required_for": "optional",
                                       "remedy": "Use the verified Studio PptxGenJS/HTML route or repair the managed runtime; never borrow an old Work node_modules cache."})
    _write(output / "runtime-selection.json", result)
    return result


def validated_environment(report: dict) -> dict[str, str]:
    """Recheck saved Work environment before a wrapper makes it authoritative."""
    if not isinstance(report, dict) or report.get("schema_version") != 1:
        raise ValueError("runtime environment has an unsupported schema")
    host = report.get("host", {})
    current = host_identity()
    if any(host.get(key) != current[key] for key in ("platform", "architecture")):
        raise ValueError("runtime environment belongs to another platform or architecture")
    environment = report.get("environment", {})
    if not isinstance(environment, dict) or set(environment) - {"AIPPT_NODE", "RUNTIME_NODE_MODULES"}:
        raise ValueError("runtime environment contains unsupported keys")
    node = probe_node_runtime(Path(environment["AIPPT_NODE"])) if environment.get("AIPPT_NODE") else None
    if node is None or node["status"] != "passed":
        raise ValueError("saved Node runtime no longer matches the current host")
    approved_node = report.get("node", {})
    if any(approved_node.get(key) != node.get(key) for key in ("path", "platform", "architecture", "version", "napi")):
        raise ValueError("Node runtime changed since the dependency probe; rerun runtime selection")
    if environment.get("RUNTIME_NODE_MODULES"):
        root = Path(environment["RUNTIME_NODE_MODULES"])
        native = root / "@oai/artifact-tool/node_modules/skia-canvas/lib/skia.node"
        if report.get("status") != "ready" or historical_cache(root) or not native.is_file() or not compatible_binary(native, architecture=node["architecture"]):
            raise ValueError("saved Artifact runtime is stale, historical or incompatible")
        proof = next((c for c in report.get("capabilities", []) if c.get("name") == "artifact-tool"), {})
        if proof.get("runtime_fingerprint") != _artifact_fingerprint(root / "@oai/artifact-tool"):
            raise ValueError("Artifact runtime changed since its drawing probe; rerun runtime selection")
    return dict(environment)


def runtime_doctor(output_dir: Path, *, include_powerpoint: bool = True) -> dict:
    """Build a tiny real deck through both core routes and optionally native Office."""
    output = Path(output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    capabilities = []
    result = {"schema_version": 1, "status": "unavailable", "host": host_identity(),
              "core_ready": False, "export_ready": False, "capabilities": capabilities,
              "output_dir": str(output), "visual_quality_approved": False}
    if os.environ.get("AIPPT_FAKE_RENDER") == "1":
        result["reason"] = "AIPPT_FAKE_RENDER cannot establish live runtime readiness"
        _write(output / "runtime-doctor.json", result)
        return result

    def check(name, required_for, operation):
        try:
            details = operation()
            capabilities.append({"name": name, "available": True, "status": "passed", "required_for": required_for, "details": details})
            return True
        except Exception as error:
            capabilities.append({"name": name, "available": False, "status": "failed", "required_for": required_for,
                                 "error": f"{type(error).__name__}: {str(error)[-2500:]}"})
            return False

    def python_drawing():
        from PIL import Image, ImageDraw
        from pptx import Presentation
        from pptx.util import Inches
        import pymupdf
        import lxml.etree
        image = Image.new("RGB", (80, 50), "white")
        ImageDraw.Draw(image).rectangle((2, 2, 40, 30), fill="#1759a7")
        image.save(output / "python-drawing.png")
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Runtime probe 原生文字"
        deck.save(output / "python-probe.pptx")
        reopened = Presentation(output / "python-probe.pptx")
        if len(reopened.slides) != 1: raise ValueError("python-pptx reopen failed")
        return {"pillow": Image.__version__, "pymupdf": pymupdf.VersionBind, "lxml": lxml.etree.LXML_VERSION, "native_text": True}

    check("python-native-packages", "all", python_drawing)
    node = probe_node_runtime()
    capabilities.append({"name": "node-runtime", "available": node["status"] == "passed", "status": node["status"], "required_for": "native", "details": node})
    if node["status"] == "passed":
        def pptxgenjs():
            modules = PROJECT_ROOT / "node_modules"
            if historical_cache(modules): raise ValueError("core modules point to an old Work cache")
            script = r"const{createRequire}=require('node:module');const req=createRequire(process.argv[1]+'/package.json');const PptxGenJS=req('pptxgenjs');const ppt=new PptxGenJS();ppt.layout='LAYOUT_WIDE';const s=ppt.addSlide();s.addText('Mac runtime 中文绘图',{x:.5,y:.5,w:10,h:.8,fontSize:28});s.addShape(ppt.ShapeType.rect,{x:1,y:2,w:3,h:1,fill:{color:'1759A7'}});ppt.writeFile({fileName:process.argv[2]}).then(()=>console.log('written')).catch(e=>{console.error(e);process.exitCode=1});"
            probe = _run([node["path"], "-e", script, str(PROJECT_ROOT), str(output / "pptxgenjs-probe.pptx")], timeout=45)
            if probe["status"] != "passed": raise ValueError(json.dumps(probe))
            # Exercise the same bounded empty-notes normalization as production.
            from .renderer import normalize_ooxml
            normalize_ooxml(output / "pptxgenjs-probe.pptx")
            from pptx import Presentation
            deck = Presentation(output / "pptxgenjs-probe.pptx")
            if len(deck.slides) != 1 or len(deck.slides[0].shapes) != 2: raise ValueError("PptxGenJS native objects missing")
            return {"native_shapes": 2, "normalized_empty_notes": True, "pptx": str(output / "pptxgenjs-probe.pptx")}
        built = check("pptxgenjs-native-build", "native", pptxgenjs)
        if built:
            from .renderer import officecli_validate
            check("officecli-actual-pptx", "all", lambda: officecli_validate(output / "pptxgenjs-probe.pptx"))

        def html_render():
            from .renderer import normalize_ooxml
            from pptx import Presentation
            html = output / "html-probe.html"
            html.write_text('<!doctype html><meta charset="utf-8"><style>body{margin:0}.slide{width:1280px;height:720px;background:white;position:relative}.title{position:absolute;left:80px;top:80px;font:40px Arial;color:#1759a7}.body{position:absolute;left:80px;top:200px;font:32px Arial}</style><section class="slide" data-slide-id="runtime-01"><div class="title">Runtime probe</div><div class="body">中文正文 Native editable text</div></section>', encoding="utf-8")
            pptx = output / "html-probe.pptx"
            previews = output / "html-previews"
            probe = _run([node["path"], str(PROJECT_ROOT / "node/render.mjs"), str(html), str(pptx), str(previews)], cwd=PROJECT_ROOT, timeout=90)
            if probe["status"] != "passed": raise ValueError(json.dumps(probe))
            normalize_ooxml(pptx)
            if len(Presentation(pptx).slides) != 1 or not (previews / "runtime-01.png").is_file():
                raise ValueError("HTML route did not produce both a one-page deck and its actual PNG")
            return {"pptx": str(pptx), "preview": str(previews / "runtime-01.png"), "browser_layout_and_export": True}
        check("html-chromium-actual-render", "native", html_render)
        if include_powerpoint and built:
            def native_office():
                from .renderer import powerpoint_roundtrip
                report = powerpoint_roundtrip(output / "pptxgenjs-probe.pptx", output / "powerpoint")
                if report.get("status") != "passed": raise ValueError(json.dumps(report))
                return report
            check("powerpoint-save-reopen-slide-preview", "export", native_office)
    result["artifact_runtime"] = select_artifact_runtime(output / "artifact")
    capabilities.extend(result["artifact_runtime"]["capabilities"])
    required = [item for item in capabilities if item.get("required_for") in {"all", "native"}]
    result["core_ready"] = bool(required) and all(item["available"] for item in required)
    result["export_ready"] = any(item["name"] == "powerpoint-save-reopen-slide-preview" and item["available"] for item in capabilities)
    result["status"] = "ready" if result["core_ready"] and (result["export_ready"] or not include_powerpoint) else "unavailable"
    _write(output / "runtime-doctor.json", result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("operation", choices=["select", "doctor", "validate-environment"])
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--skip-powerpoint", action="store_true")
    args = parser.parse_args()
    # This module is also a documented direct entrypoint.  Do not allow a
    # Rosetta Python to create a Work runtime receipt before its x64 process
    # architecture is corrected by a supported launcher.  Parse first so
    # ``--help`` remains available for recovery instructions.
    runtime_reason = validate_identity(current_process_identity())
    if runtime_reason is not None:
        parser.error("Studio runtime selection requires a host-native Python: " + runtime_reason)
    if args.operation == "validate-environment":
        result = validated_environment(json.loads(args.report.read_text(encoding="utf-8-sig")))
    elif args.operation == "select":
        result = select_artifact_runtime(args.output_dir)
    else:
        result = runtime_doctor(args.output_dir, include_powerpoint=not args.skip_powerpoint)
    print(json.dumps(result, ensure_ascii=False))
    if args.operation == "doctor" and result.get("status") != "ready":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
