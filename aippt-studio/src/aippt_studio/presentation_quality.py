"""Read-only, object-level PPTX readability and geometry checks.

This is not a renderer. Explicit/inherited font sizes and transformed rectangles
are inspectable; real line wrapping, glyph ink, contrast and design quality are
not proved by a structural pass. A policy exception names one slide and object,
has a bounded allowance and retains its rationale and evidence in the report.
Font sizes are diagnostic references by default, not a verdict on a page's
expression. Only explicitly supplied project typography constraints can block.
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import math
from pathlib import Path
import re

from PIL import Image
from pptx import Presentation

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
EMU = 12700.0  # points
FONT_SIZE_REFERENCES = {"title": 30.0, "body": 24.0, "label": 18.0, "caption": 18.0,
                        "source": 16.0, "page_number": 16.0}
EXCEPTION_CODES = {"SMALL_FONT", "OUT_OF_BOUNDS", "TEXT_OVERLAP", "IMAGE_TEXT_OVERLAP", "TEXT_OCCLUDED"}
IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


def _find(node, path):
    return node.find(path, NS) if node is not None else None


def _all(node, path):
    return node.findall(path, NS) if node is not None else []


def _finite(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _policy(value: dict | None) -> dict:
    value = {} if value is None else value
    if not isinstance(value, dict):
        raise ValueError("policy must be an object")
    unknown = set(value) - {"mode", "typography_mode", "minimum_font_sizes", "object_roles", "exceptions"}
    if unknown:
        raise ValueError(f"unsupported policy keys: {sorted(unknown)}; global skip flags are not supported")
    mode = value.get("mode", "new")
    if mode not in {"new", "legacy_audit"}:
        raise ValueError("mode must be new or legacy_audit")
    typography_mode = value.get("typography_mode", "advisory")
    if typography_mode not in {"advisory", "explicit"}:
        raise ValueError("typography_mode must be advisory or explicit")
    # Never silently turn reference values into a user's project constraints.
    minimums = {}
    overrides = value.get("minimum_font_sizes", {})
    if not isinstance(overrides, dict):
        raise ValueError("minimum_font_sizes must be an object")
    for role, size in overrides.items():
        if role not in FONT_SIZE_REFERENCES or not _finite(size) or not 1 <= size <= 400:
            raise ValueError("font minima need a known role and a finite size between 1 and 400pt")
        minimums[role] = float(size)
    roles, exceptions = [], []
    for key, target in (("object_roles", roles), ("exceptions", exceptions)):
        rows = value.get(key, [])
        if not isinstance(rows, list) or len(rows) > 500:
            raise ValueError(f"{key} must be a list of at most 500 explicit object records")
        seen = set()
        for index, row in enumerate(rows):
            if not isinstance(row, dict):
                raise ValueError(f"{key}[{index}] must be an object")
            slide, object_id = row.get("slide"), row.get("object_id")
            if not isinstance(slide, int) or isinstance(slide, bool) or slide < 1 or not str(object_id).isdigit():
                raise ValueError(f"{key}[{index}] needs a positive slide and numeric OOXML object_id; wildcards are forbidden")
            if not isinstance(row.get("reason"), str) or not row["reason"].strip():
                raise ValueError(f"{key}[{index}] needs a reason")
            item = {**row, "object_id": str(object_id)}
            if key == "object_roles":
                if row.get("role") not in FONT_SIZE_REFERENCES or set(row) - {"slide", "object_id", "role", "reason"}:
                    raise ValueError(f"invalid object role record at {index}")
                identity = (slide, str(object_id))
            else:
                code = row.get("code")
                if code not in EXCEPTION_CODES or not isinstance(row.get("evidence"), str) or not row["evidence"].strip():
                    raise ValueError(f"exceptions[{index}] needs a supported code and evidence locator or current approval quote")
                allowed = {"slide", "object_id", "code", "reason", "evidence"}
                if code == "SMALL_FONT":
                    allowed.add("min_font_pt")
                    if not _finite(row.get("min_font_pt")) or not 1 <= row["min_font_pt"] <= 400:
                        raise ValueError("font exceptions need a finite min_font_pt between 1 and 400; an unlimited waiver is forbidden")
                elif code == "OUT_OF_BOUNDS":
                    allowed.add("max_overflow_pt")
                    if not _finite(row.get("max_overflow_pt")) or not 0 < row["max_overflow_pt"] <= 36:
                        raise ValueError("boundary exceptions need 0 < max_overflow_pt <= 36")
                else:
                    allowed.update({"other_object_id", "max_overlap_ratio"})
                    if not str(row.get("other_object_id")).isdigit() or not _finite(row.get("max_overlap_ratio")) or not 0 < row["max_overlap_ratio"] <= 1:
                        raise ValueError("overlap exceptions need one other_object_id and 0 < max_overlap_ratio <= 1")
                    item["other_object_id"] = str(row["other_object_id"])
                if set(row) - allowed:
                    raise ValueError(f"unsupported exception fields: {sorted(set(row) - allowed)}")
                identity = (slide, str(object_id), code, item.get("other_object_id"))
            if identity in seen:
                raise ValueError(f"duplicate {key} entry: {identity}")
            seen.add(identity)
            target.append(item)
    return {"mode": mode, "typography_mode": typography_mode, "minimum_font_sizes": minimums,
            "object_roles": roles, "exceptions": exceptions}


def _mul(a, b):
    return (a[0]*b[0]+a[2]*b[1], a[1]*b[0]+a[3]*b[1],
            a[0]*b[2]+a[2]*b[3], a[1]*b[2]+a[3]*b[3],
            a[0]*b[4]+a[2]*b[5]+a[4], a[1]*b[4]+a[3]*b[5]+a[5])


def _translate(x, y):
    return (1, 0, 0, 1, x, y)


def _rotate(angle):
    angle = math.radians(angle)
    return (math.cos(angle), math.sin(angle), -math.sin(angle), math.cos(angle), 0, 0)


def _point(matrix, x, y):
    return (matrix[0]*x + matrix[2]*y + matrix[4], matrix[1]*x + matrix[3]*y + matrix[5])


def _transform(shape, parent):
    x, y, w, h = [float(value or 0) / EMU for value in (shape.left, shape.top, shape.width, shape.height)]
    xfrm = _find(shape._element, "p:spPr/a:xfrm")
    if xfrm is None:
        xfrm = _find(shape._element, "p:xfrm")
    rotation = float(xfrm.get("rot", "0"))/60000 if xfrm is not None else float(shape.rotation or 0)
    flip_h = xfrm is not None and xfrm.get("flipH") in {"1", "true"}
    flip_v = xfrm is not None and xfrm.get("flipV") in {"1", "true"}
    center = _mul(_translate(x+w/2, y+h/2), _mul(_rotate(rotation), (-1 if flip_h else 1, 0, 0, -1 if flip_v else 1, 0, 0)))
    matrix = _mul(parent, _mul(center, _translate(-w/2, -h/2)))
    points = [_point(matrix, xx, yy) for xx, yy in ((0, 0), (w, 0), (w, h), (0, h))]
    rect = [min(p[0] for p in points), min(p[1] for p in points), max(p[0] for p in points), max(p[1] for p in points)]
    return matrix, rect, w, h


def _group_transform(shape, parent):
    xfrm = _find(shape._element, "p:grpSpPr/a:xfrm")
    if xfrm is None:
        return parent
    off, ext, child_off, child_ext = [_find(xfrm, "a:" + name) for name in ("off", "ext", "chOff", "chExt")]
    def pair(node, keys, fallback):
        return tuple(float(node.get(key, default)) / EMU if node is not None else default/EMU for key, default in zip(keys, fallback))
    x, y = pair(off, ("x", "y"), (0, 0))
    w, h = pair(ext, ("cx", "cy"), (float(shape.width or 0), float(shape.height or 0)))
    cx, cy = pair(child_off, ("x", "y"), (0, 0))
    cw, ch = pair(child_ext, ("cx", "cy"), (w*EMU, h*EMU))
    if cw <= 0 or ch <= 0:
        raise ValueError(f"group {shape.shape_id} has zero/invalid child coordinate extent")
    local = _mul(_translate(x+w/2, y+h/2), _rotate(float(xfrm.get("rot", "0"))/60000))
    local = _mul(local, (-1 if xfrm.get("flipH") in {"1", "true"} else 1, 0, 0, -1 if xfrm.get("flipV") in {"1", "true"} else 1, 0, 0))
    local = _mul(local, _mul(_translate(-w/2, -h/2), _mul((w/cw, 0, 0, h/ch, 0, 0), _translate(-cx, -cy))))
    return _mul(parent, local)


def _placeholder(element):
    return _find(element, "p:nvSpPr/p:nvPr/p:ph")


def _placeholder_chain(shape, slide):
    ph = _placeholder(shape._element)
    if ph is None:
        return []
    idx, kind = ph.get("idx", "0"), ph.get("type", "obj")
    chain = []
    layout_shapes = list(slide.slide_layout.shapes)
    layout = next((candidate for candidate in layout_shapes if _placeholder(candidate._element) is not None
                   and _placeholder(candidate._element).get("idx", "0") == idx), None)
    if layout is not None:
        chain.append(("layout_placeholder", layout._element))
        kind = _placeholder(layout._element).get("type", kind)
    master_kind = "title" if kind in {"title", "ctrTitle"} else "body" if kind in {"body", "obj", "subTitle"} else kind
    for candidate in slide.slide_layout.slide_master.shapes:
        inherited_ph = _placeholder(candidate._element)
        if inherited_ph is not None and inherited_ph.get("type", "obj") == master_kind:
            chain.append(("master_placeholder", candidate._element))
            break
    return chain


def _font_size(presentation, slide, shape, txbody, paragraph, run=None, *, chart=False):
    level = int((_find(paragraph, "a:pPr").get("lvl", "0")) if _find(paragraph, "a:pPr") is not None else 0)
    level = min(8, max(0, level))
    candidates = []
    if run is not None:
        candidates.append(("run.rPr", _find(run, "a:rPr")))
    candidates += [("paragraph.defRPr", _find(paragraph, "a:pPr/a:defRPr")),
                   ("shape.lstStyle", _find(txbody, f"a:lstStyle/a:lvl{level+1}pPr/a:defRPr")),
                   ("shape.defaultStyle", _find(txbody, "a:lstStyle/a:defPPr/a:defRPr"))]
    if chart:
        # A chart's global txPr is inherited by labels that have no local size.
        chart_space = shape.chart._chartSpace
        candidates += [("chart.txPr", _find(chart_space, "c:txPr/a:p/a:pPr/a:defRPr"))]
    else:
        for name, element in _placeholder_chain(shape, slide):
            body = _find(element, "p:txBody")
            candidates.extend([(name + ".paragraph", _find(body, "a:p/a:pPr/a:defRPr")),
                               (name + ".lstStyle", _find(body, f"a:lstStyle/a:lvl{level+1}pPr/a:defRPr"))])
        ph = _placeholder(shape._element)
        kind = ph.get("type", "obj") if ph is not None else None
        style = "titleStyle" if kind in {"title", "ctrTitle"} else "bodyStyle" if ph is not None and kind in {"body", "obj", "subTitle"} else "otherStyle"
        candidates.append((f"master.{style}.level{level+1}", _find(slide.slide_layout.slide_master._element,
                           f"p:txStyles/p:{style}/a:lvl{level+1}pPr/a:defRPr")))
        candidates.append((f"presentation.defaultTextStyle.level{level+1}", _find(presentation._element,
                           f"p:defaultTextStyle/a:lvl{level+1}pPr/a:defRPr")))
    for source, node in candidates:
        if node is not None and node.get("sz") is not None:
            size = float(node.get("sz")) / 100.0
            if not math.isfinite(size) or size <= 0:
                raise ValueError(f"invalid font size on object {shape.shape_id}")
            norm = _find(txbody, "a:bodyPr/a:normAutofit")
            scale = float(norm.get("fontScale", "100000"))/100000.0 if norm is not None else 1.0
            return {"declared_font_pt": size, "font_pt": size*scale, "autofit_scale": scale,
                    "font_source": source, "paragraph_level": level}
    return {"font_pt": None, "font_source": "unresolved_chart_default" if chart else "unresolved_inheritance", "paragraph_level": level}


def _text_samples(presentation, slide, shape, txbody, target, *, chart=False):
    samples = []
    for index, paragraph in enumerate(_all(txbody, "a:p"), 1):
        runs = [child for child in paragraph if child.tag.rsplit("}", 1)[-1] in {"r", "fld"}]
        for run in runs:
            text = "".join(node.text or "" for node in _all(run, "a:t"))
            if text.strip():
                samples.append({"text": text, "text_target": target, "paragraph": index,
                                **_font_size(presentation, slide, shape, txbody, paragraph, run, chart=chart)})
    return samples


def _chart_samples(presentation, slide, shape):
    space = shape.chart._chartSpace
    samples = []
    for label in ("title", "catAx", "valAx", "legend", "dLbls"):
        for index, container in enumerate(_all(space, f".//c:{label}"), 1):
            deleted = _find(container, "c:delete")
            if deleted is not None and deleted.get("val") in {"1", "true"}:
                continue
            tick_position = _find(container, "c:tickLblPos")
            if label in {"catAx", "valAx"} and tick_position is not None and tick_position.get("val") == "none":
                continue
            if label == "dLbls":
                visibility = [node for tag in ("showVal", "showCatName", "showSerName", "showPercent", "showBubbleSize")
                              for node in _all(container, f".//c:{tag}")]
                # Dormant formats can retain 12pt while label contents are off.
                custom_text = _all(container, ".//a:t")
                if visibility and not custom_text and all(node.get("val", "1") in {"0", "false"} for node in visibility):
                    continue
            body = _find(container, "c:tx/c:rich") if label == "title" else None
            if body is None:
                body = _find(container, "c:txPr")
            target = f"chart:{label}:{index}"
            written = _text_samples(presentation, slide, shape, body, target, chart=True) if body is not None else []
            if written:
                samples.extend(written)
            else:
                paragraph = _find(body, "a:p")
                sample = _font_size(presentation, slide, shape, body, paragraph, chart=True)
                samples.append({"text": f"[{label} labels]", "text_target": target, **sample})
    # Charts without explicit axes/legend/title may still show data labels.
    return samples


def _area(rect):
    return max(0, rect[2]-rect[0]) * max(0, rect[3]-rect[1])


def _intersection(a, b):
    return [max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])]


def _role(shape, text, rect, slide_height, override):
    ph = _placeholder(shape._element)
    kind = ph.get("type", "obj") if ph is not None else ""
    page_semantic = (kind == "sldNum" or any(field.get("type") == "slidenum" for field in _all(shape._element, ".//a:fld"))
                     or (bool(re.fullmatch(r"\s*(?:第\s*)?\d{1,3}(?:\s*/\s*\d{1,3}|\s*页)?\s*", text))
                         and rect[1] >= slide_height*.85 and rect[2]-rect[0] <= 72))
    source_semantic = bool(re.match(r"\s*(?:数据来源|图片来源|资料来源|来源|参考文献|参考资料|sources?|references?)\s*[:：]\s*\S", text, re.I)) and len(text) <= 350
    if override:
        role = override["role"]
        if role == "page_number" and not page_semantic:
            raise ValueError(f"object {shape.shape_id} cannot be labelled page_number without page-number semantics")
        if role == "source" and not source_semantic:
            raise ValueError(f"object {shape.shape_id} cannot be labelled source without source/reference semantics")
        return role, "explicit_object_role"
    if page_semantic:
        return "page_number", "page_number_semantics"
    if source_semantic:
        return "source", "source_prefix_semantics"
    if kind in {"title", "ctrTitle"}:
        return "title", "title_placeholder"
    if getattr(shape, "has_chart", False):
        return "label", "chart_labels"
    return "body", "default_body"


def _opaque(shape):
    if shape.shape_type == 13:  # picture
        geometry = _find(shape._element, "p:spPr/a:prstGeom")
        if _find(shape._element, "p:spPr/a:custGeom") is not None:
            return "unknown"
        if geometry is not None and geometry.get("prst") != "rect":
            # A clipped picture does not cover its whole bounding rectangle,
            # even when every pixel in the original JPEG is opaque.
            return "transparent_or_partial"
        # The source raster can be fully opaque while its picture surface is
        # not: contain uses an inset fill rectangle, and native transparency
        # applies to the blip independently of PNG/JPEG alpha.
        try:
            fill_rect = _find(shape._element, "p:blipFill/a:stretch/a:fillRect")
            if fill_rect is not None and any(int(fill_rect.get(side, "0")) > 0 for side in ("l", "t", "r", "b")):
                return "transparent_or_partial"
            source_rect = _find(shape._element, "p:blipFill/a:srcRect")
            if source_rect is not None and any(int(source_rect.get(side, "0")) < 0 for side in ("l", "t", "r", "b")):
                # Some existing generators express contain as source outset.
                return "transparent_or_partial"
            for alpha in _all(shape._element, "p:blipFill/a:blip/a:alphaModFix"):
                if int(alpha.get("amt", "100000")) < 100000:
                    return "transparent_or_partial"
        except (ValueError, TypeError):
            return "unknown"
        try:
            with Image.open(BytesIO(shape.image.blob)) as image:
                if "A" in image.getbands():
                    return "opaque" if image.getchannel("A").getextrema() == (255, 255) else "transparent_or_partial"
                return "unknown" if "transparency" in image.info else "opaque"
        except Exception:
            return "unknown"
    fill = _find(shape._element, "p:spPr/a:solidFill")
    geometry = _find(shape._element, "p:spPr/a:prstGeom")
    if fill is not None and geometry is not None and geometry.get("prst") == "rect":
        alphas = _all(fill, ".//a:alpha")
        return "opaque" if all(int(item.get("val", "100000")) == 100000 for item in alphas) else "transparent_or_partial"
    return "unknown"


def _objects(presentation, slide, slide_number, policy):
    rows, hidden = [], []
    role_map = {(row["slide"], row["object_id"]): row for row in policy["object_roles"]}
    def visit(shapes, parent=IDENTITY, groups=()):
        for shape in shapes:
            own_properties = [node for tag in ("nvSpPr", "nvPicPr", "nvGraphicFramePr", "nvCxnSpPr", "nvGrpSpPr")
                              for node in _all(shape._element, f"p:{tag}/p:cNvPr")]
            if any(node.get("hidden") in {"1", "true"} for node in own_properties):
                hidden.append(str(shape.shape_id))
                continue
            if shape.shape_type == 6:  # group
                visit(shape.shapes, _group_transform(shape, parent), groups+(str(shape.shape_id),))
                continue
            matrix, rect, width, height = _transform(shape, parent)
            samples = []
            kind = "picture" if shape.shape_type == 13 else "shape"
            if getattr(shape, "has_table", False):
                kind = "table"
                for r, row in enumerate(shape.table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        if not cell.is_spanned:
                            samples.extend(_text_samples(presentation, slide, shape, cell._tc.txBody, f"cell:{r}:{c}"))
            elif getattr(shape, "has_chart", False):
                kind = "chart"
                samples = _chart_samples(presentation, slide, shape)
            elif getattr(shape, "has_text_frame", False):
                samples = _text_samples(presentation, slide, shape, shape.text_frame._txBody, "text")
            text = "\n".join(sample["text"] for sample in samples)
            role, role_source = _role(shape, text, rect, float(presentation.slide_height)/EMU,
                                      role_map.get((slide_number, str(shape.shape_id))))
            scale = min(math.hypot(parent[0], parent[1]), math.hypot(parent[2], parent[3]))
            for sample in samples:
                if sample["font_pt"] is not None:
                    sample["font_pt"] *= scale
                sample["group_scale"] = scale
            body = _find(shape._element, "p:txBody/a:bodyPr")
            left, top, right, bottom = [float(body.get(key, default))/EMU if body is not None else default/EMU
                                        for key, default in (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720))]
            inner = [_point(matrix, x, y) for x, y in ((left, top), (max(left, width-right), top),
                                                       (max(left, width-right), max(top, height-bottom)), (left, max(top, height-bottom)))]
            text_rect = [min(p[0] for p in inner), min(p[1] for p in inner), max(p[0] for p in inner), max(p[1] for p in inner)]
            rotated = abs(matrix[1]) > 1e-8 or abs(matrix[2]) > 1e-8
            geometry = _find(shape._element, "p:spPr/a:prstGeom")
            rows.append({"slide": slide_number, "object_id": str(shape.shape_id), "name": shape.name,
                         "kind": kind, "z_order": len(rows), "group_ids": list(groups), "bounds_pt": rect,
                         "text_bounds_pt": text_rect if kind not in {"table", "chart"} else rect,
                         "rotation_or_shear": rotated, "role": role, "role_source": role_source,
                         "text_preview": text[:160], "text": text, "font_samples": samples,
                         "opacity": _opaque(shape), "background_container": False,
                         "geometry": geometry.get("prst") if geometry is not None else "custom_or_unspecified",
                         "auto_fit": "shape" if _find(body, "a:spAutoFit") is not None else "normal" if _find(body, "a:normAutofit") is not None else "none",
                         "xml_object_type": shape._element.tag.rsplit("}", 1)[-1]})
    visit(slide.shapes)
    return rows, hidden


def image_frame_geometry_signals(rows: list[dict]) -> list[dict]:
    """Locate probable separate frame/photo pairs, without inventing ownership.

    Coincident boxes are geometric evidence only. A backing may be intentional;
    alpha, clipping, rotation and semantic purpose still need actual page review.
    """
    signals = []
    for picture in rows:
        if picture.get("kind") != "picture" or picture.get("opacity") != "opaque":
            continue
        pbox = picture["bounds_pt"]
        for frame in rows:
            if (frame.get("kind") != "shape" or frame.get("text", "").strip()
                    or frame.get("geometry") not in {"roundRect", "ellipse"}
                    or frame["z_order"] >= picture["z_order"]
                    or frame.get("geometry") == picture.get("geometry")):
                continue
            fbox = frame["bounds_pt"]
            # This tolerance identifies a potential shared surface. It is not
            # an aesthetic radius, spacing or image-size rule.
            tolerance = max(.5, min(fbox[2]-fbox[0], fbox[3]-fbox[1]) * .005)
            if max(abs(a-b) for a, b in zip(pbox, fbox)) > tolerance:
                continue
            signals.append({"code": "IMAGE_FRAME_CONTOUR_MISMATCH", "kind": "context_required",
                "object_id": picture["object_id"], "other_object_id": frame["object_id"],
                "image_geometry": picture.get("geometry"), "frame_geometry": frame["geometry"],
                "bounds_pt": pbox, "rotation_or_shear": bool(picture.get("rotation_or_shear") or frame.get("rotation_or_shear")),
                "reason": "An opaque image and an underlying curved shape share their bounds but have different surface geometry.",
                "review_question": "Inspect the actual image edges and decide whether these objects form one frame/photo surface. If so, apply matching geometry to the image itself and preserve required source content. If this is an intentional independent composition, record that contextual judgment; do not auto-pair or crop unrelated objects."})
    return signals


def slide_image_geometry_signals(presentation, slide, slide_number: int) -> list[dict]:
    rows, _ = _objects(presentation, slide, slide_number, _policy(None))
    return image_frame_geometry_signals(rows)


def _inherited_static_text(slide) -> list[dict]:
    """Return non-placeholder text boxes from the active layout and master.

    These objects are visible on the rendered page but are absent from
    ``slide.shapes``. Keep this deliberately narrow: inherited text can be
    compared geometrically, while its rendered ink and contrast still need
    visual review.
    """
    found = []
    sources = (("layout", slide.slide_layout.shapes), ("master", slide.slide_layout.slide_master.shapes))
    for source, shapes in sources:
        for shape in shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            if _placeholder(shape._element) is not None:
                continue
            _, bounds, _, _ = _transform(shape, IDENTITY)
            found.append({"source": source, "object_id": str(shape.shape_id),
                          "name": shape.name, "text": shape.text.strip(), "bounds_pt": bounds})
    return found


def audit_presentation(
    pptx: Path,
    policy: dict | None = None,
    *,
    selected_slides: list[int] | None = None,
) -> dict:
    """Audit structure and explicit constraints without modifying the PPTX.

    ``passed`` means no unresolved measured blockers/review signals. It never
    certifies semantic expression, visual hierarchy, or reading-scale quality.
    Default typography observations have advisory severity and cannot fail it.
    """
    path = Path(pptx).expanduser().resolve()
    if selected_slides is not None and (
        not isinstance(selected_slides, list) or not selected_slides
        or any(type(number) is not int or number < 1 for number in selected_slides)
        or len(set(selected_slides)) != len(selected_slides)
    ):
        raise ValueError("selected_slides must contain distinct positive 1-based slide numbers")
    result = {"schema_version": 1, "generated_at": datetime.now(timezone.utc).isoformat(),
              "source": str(path), "source_sha256": None, "status": "errors", "issues": [], "errors": [], "blocking_issues": [], "advisory_issues": [],
              "slides": [], "exceptions": {"applied": [], "unused": []}, "source_modified": False,
              "scope": {"kind": "full_deck" if selected_slides is None else "selected_slides",
                        "selected_slides": None if selected_slides is None else list(selected_slides),
                        "full_deck_verified": selected_slides is None},
              "visual_quality_verified": False,
              "coverage": {"font_sizes": "explicit and resolved OOXML inheritance; normal autofit and group scale included",
                           "typography_judgment": "font-size references are advisory; page semantics, viewing scale and visual hierarchy require actual slide review",
                           "geometry": "transformed object rectangles, including nested groups; rotated overlaps require rendered review",
                           "text_wrapping": "not rendered; no guarantee of actual line breaks, clipping or text fit",
                           "image_semantics": "background geometry and raster alpha only; no photo meaning or contrast judgment",
                           "scope": "slide objects plus geometric collision prompts for static layout/master text; decorative non-text master objects are not collision-audited"}}
    try:
        resolved_policy = _policy(policy)
        result["policy"] = resolved_policy
        result["typography_mode"] = resolved_policy["typography_mode"]
        result["font_size_references"] = dict(FONT_SIZE_REFERENCES)
        data = path.read_bytes()
        result["source_sha256"] = hashlib.sha256(data).hexdigest()
        presentation = Presentation(BytesIO(data))
        slide_count = len(presentation.slides)
        if selected_slides is not None and any(number > slide_count for number in selected_slides):
            raise ValueError(f"selected_slides must refer to existing slides 1–{slide_count}")
        selected_set = None if selected_slides is None else set(selected_slides)
        if selected_set is not None:
            # A saved whole-deck policy can accompany a one-page revision.
            # Rules naming untouched pages are retained in the project policy,
            # but they are not evaluated or falsely marked unused by this
            # explicitly scoped check.
            excluded_roles = len([row for row in resolved_policy["object_roles"] if row["slide"] not in selected_set])
            excluded_exceptions = len([row for row in resolved_policy["exceptions"] if row["slide"] not in selected_set])
            resolved_policy = {**resolved_policy,
                               "object_roles": [row for row in resolved_policy["object_roles"] if row["slide"] in selected_set],
                               "exceptions": [row for row in resolved_policy["exceptions"] if row["slide"] in selected_set]}
            result["scope"]["policy_records_not_evaluated"] = {
                "object_roles": excluded_roles, "exceptions": excluded_exceptions}
            result["policy"] = resolved_policy
        width, height = float(presentation.slide_width)/EMU, float(presentation.slide_height)/EMU
        result["slide_size_pt"] = {"width": width, "height": height}
        issues = []
        def issue(row, code, severity, **details):
            issues.append({"slide": row["slide"], "object_id": row["object_id"], "object_name": row["name"],
                           "code": code, "severity": severity, "details": {"role": row["role"], "role_source": row["role_source"], **details}})
        all_ids = set()
        for number, slide in enumerate(presentation.slides, 1):
            if selected_set is not None and number not in selected_set:
                continue
            rows, hidden = _objects(presentation, slide, number, resolved_policy)
            if len({row["object_id"] for row in rows}) != len(rows):
                raise ValueError(f"slide {number} contains duplicate OOXML object IDs; scoped exceptions would be ambiguous")
            all_ids.update((number, row["object_id"]) for row in rows)
            texts = [row for row in rows if row["font_samples"]]
            inherited_text = _inherited_static_text(slide)
            row_by_id = {row["object_id"]: row for row in rows}
            for signal in image_frame_geometry_signals(rows):
                details = {key: value for key, value in signal.items() if key not in {"code", "kind", "object_id"}}
                # A matched opaque rectangular photo over a curved backing is
                # a concrete composition question. Keep the final check open
                # until the actual page confirms whether the pair is intended.
                issue(row_by_id[signal["object_id"]], signal["code"], "review", **details)
            for row in rows:
                rect = row["bounds_pt"]
                # A text-bearing shape whose transformed usable inset rectangle
                # has collapsed cannot contain visible text. This is a
                # structural clipping risk; it does not estimate wrapping or
                # judge visual quality.
                if row in texts and row["kind"] == "shape" and _area(row["text_bounds_pt"]) <= .01:
                    issue(row, "TEXT_CONTENT_AREA_COLLAPSED", "review",
                          bounds_pt=rect, text_bounds_pt=row["text_bounds_pt"],
                          certainty="insets_leave_no_text_area",
                          text_preview=row["text_preview"],
                          guidance="The text box margins leave no usable content area. Inspect its inset settings and the rendered page.")
                # A full-slide image below every intersecting text object is a
                # background container, not a foreground text occluder.
                if row["kind"] == "picture" and _area(_intersection(rect, [0, 0, width, height])) >= width*height*.85:
                    overlapping = [text for text in texts if _area(_intersection(rect, text["bounds_pt"])) > 0]
                    row["background_container"] = bool(overlapping) and all(row["z_order"] < text["z_order"] for text in overlapping)
                overflow = max(0, -rect[0], -rect[1], rect[2]-width, rect[3]-height)
                if overflow > .5 and (row["font_samples"] or row["kind"] in {"picture", "table", "chart"}):
                    if not row["background_container"]:
                        issue(row, "OUT_OF_BOUNDS", "blocker", overflow_pt=round(overflow, 3), bounds_pt=rect,
                              certainty="object_geometry", text_preview=row["text_preview"])
                explicit_minimum = (resolved_policy["minimum_font_sizes"].get(row["role"])
                                    if resolved_policy["typography_mode"] == "explicit" else None)
                reference = resolved_policy["minimum_font_sizes"].get(row["role"], FONT_SIZE_REFERENCES[row["role"]])
                threshold = explicit_minimum if explicit_minimum is not None else reference
                small = [sample for sample in row["font_samples"] if sample["font_pt"] is not None and sample["font_pt"] < threshold-.01]
                unknown = [sample for sample in row["font_samples"] if sample["font_pt"] is None]
                if small:
                    typography_details = ({"minimum_font_pt": explicit_minimum, "policy_basis": "explicit_project_constraint"}
                                          if explicit_minimum is not None else
                                          {"reference_font_pt": reference, "policy_basis": "diagnostic_reference"})
                    issue(row, "SMALL_FONT", "blocker" if explicit_minimum is not None else "advisory", **typography_details,
                          actual_min_font_pt=min(sample["font_pt"] for sample in small),
                          text_samples=[{**sample, "text": sample["text"][:100]} for sample in small], certainty="resolved_ooxml_font",
                          guidance=("This object is below a supplied project constraint; check the actual page and the object's role."
                                    if explicit_minimum is not None else
                                    "This is a size reference, not a legibility verdict. Judge the object's meaning, hierarchy and intended viewing scale in the actual page."))
                if unknown:
                    issue(row, "FONT_UNRESOLVED", "review" if explicit_minimum is not None else "advisory",
                          text_targets=sorted({sample["text_target"] for sample in unknown}),
                          certainty="insufficient_font_inheritance", text_preview=row["text_preview"],
                          guidance="The actual font size could not be resolved; inspect this text in the intended viewing context.")
                if row["role_source"] == "default_body" and row["kind"] == "shape" and row["text"].strip() and rect[1] < height*.15 and rect[2]-rect[0] >= width*.35 and len(row["text"]) <= 80:
                    sizes = [sample["font_pt"] for sample in row["font_samples"] if sample["font_pt"] is not None]
                    if sizes and min(sizes) < FONT_SIZE_REFERENCES["title"] and not small:
                        issue(row, "TITLE_ROLE_HINT", "advisory", suggested_role="title", actual_min_font_pt=min(sizes),
                              reference_font_pt=FONT_SIZE_REFERENCES["title"], certainty="position_heuristic",
                              guidance="This may be a title or another semantic role. Position and point size alone do not determine its role or quality.")
            for index, first in enumerate(texts):
                if first["kind"] == "chart":
                    continue
                for second in texts[index+1:]:
                    if second["kind"] == "chart":
                        continue
                    area = _area(_intersection(first["text_bounds_pt"], second["text_bounds_pt"]))
                    denominator = min(_area(first["text_bounds_pt"]), _area(second["text_bounds_pt"]))
                    if area > 4 and denominator and area/denominator > .025:
                        issue(first, "TEXT_OVERLAP", "review", other_object_id=second["object_id"], overlap_ratio=round(area/denominator, 6),
                              certainty="text_container_geometry_only", text_preview=first["text_preview"], other_text_preview=second["text_preview"],
                              guidance="Actual glyph positions and wrapping require rendered review; intersecting boxes alone do not prove colliding ink.")
            for text in texts:
                text_area = _area(text["text_bounds_pt"])
                if not text_area:
                    continue
                for inherited in inherited_text:
                    area = _area(_intersection(text["text_bounds_pt"], inherited["bounds_pt"]))
                    if area > 4 and area / text_area > .025:
                        issue(text, "INHERITED_TEXT_OVERLAP", "review",
                              other_object_id=f"{inherited['source']}:{inherited['object_id']}",
                              other_object_name=inherited["name"], other_text_preview=inherited["text"][:160],
                              overlap_ratio=round(area / text_area, 6), certainty="text_container_geometry_only",
                              guidance="Static layout/master text is rendered with this slide. Check the actual page for colliding ink, intended layering and contrast.")
            for text in texts:
                text_area = _area(text["text_bounds_pt"])
                if not text_area:
                    continue
                for graphic in rows:
                    if graphic is text or graphic["kind"] not in {"picture", "shape"}:
                        continue
                    if graphic["font_samples"] and graphic["opacity"] != "opaque":
                        continue
                    area = _area(_intersection(text["text_bounds_pt"], graphic["bounds_pt"]))
                    ratio = area/text_area
                    if area <= 4 or ratio <= .025 or graphic["background_container"]:
                        continue
                    above = graphic["z_order"] > text["z_order"]
                    if above and graphic["opacity"] == "opaque" and ratio >= .98 and not graphic["rotation_or_shear"] and not text["rotation_or_shear"]:
                        issue(text, "TEXT_OCCLUDED", "blocker", other_object_id=graphic["object_id"], overlap_ratio=round(ratio, 6),
                              certainty="opaque_foreground_covers_text_container", text_preview=text["text_preview"])
                    elif graphic["kind"] == "picture" or above:
                        issue(text, "IMAGE_TEXT_OVERLAP", "review", other_object_id=graphic["object_id"], overlap_ratio=round(ratio, 6),
                              graphic_above_text=above, graphic_opacity=graphic["opacity"], certainty="geometry_only",
                              guidance="Verify intentional text-on-image composition and legibility in the actual slide.")
            result["slides"].append({"slide": number, "hidden_object_ids": hidden,
                                     "objects": [{key: value for key, value in row.items() if key != "text"} for row in rows]})
        for role in resolved_policy["object_roles"]:
            if (role["slide"], role["object_id"]) not in all_ids:
                raise ValueError(f"object role targets a missing/hidden object: slide {role['slide']} object {role['object_id']}")
        applied = set()
        for item in issues:
            for index, exception in enumerate(resolved_policy["exceptions"]):
                if (item["slide"], item["object_id"], item["code"]) != (exception["slide"], exception["object_id"], exception["code"]):
                    continue
                details = item["details"]
                allowed = (item["code"] == "SMALL_FONT" and details["actual_min_font_pt"] >= exception["min_font_pt"]-.01)
                if item["code"] == "OUT_OF_BOUNDS":
                    allowed = details["overflow_pt"] <= exception["max_overflow_pt"]
                if item["code"] in {"TEXT_OVERLAP", "IMAGE_TEXT_OVERLAP", "TEXT_OCCLUDED"}:
                    allowed = details.get("other_object_id") == exception["other_object_id"] and details["overlap_ratio"] <= exception["max_overlap_ratio"]
                if allowed:
                    item["exception"] = dict(exception)
                    item["disposition"] = "accepted_object_exception"
                    applied.add(index)
                    result["exceptions"]["applied"].append({"issue_code": item["code"], **exception})
                    break
            item.setdefault("disposition", "open")
        result["exceptions"]["unused"] = [row for index, row in enumerate(resolved_policy["exceptions"]) if index not in applied]
        result["issues"] = issues
        open_issues = [item for item in issues if item["disposition"] == "open" and item["severity"] in {"blocker", "review"}]
        advisories = [item for item in issues if item["disposition"] == "open" and item["severity"] == "advisory"]
        result["advisory_issues"] = advisories
        result["status"] = "needs_review" if open_issues else "passed" if selected_set is None else "passed_scoped"
        result["blocking_issues"] = [item for item in open_issues if item["severity"] == "blocker"] if resolved_policy["mode"] == "new" else []
        result["summary"] = {"slides": len(result["slides"]), "total_deck_slides": slide_count,
                             "objects": sum(len(slide["objects"]) for slide in result["slides"]),
                             "open_issues": len(open_issues), "blockers": len(result["blocking_issues"]),
                             "review_issues": sum(item["severity"] == "review" for item in open_issues), "advisories": len(advisories),
                             "accepted_exceptions": len(result["exceptions"]["applied"]),
                             "codes": dict(Counter(item["code"] for item in open_issues)),
                             "advisory_codes": dict(Counter(item["code"] for item in advisories)),
                             "typography_enforcement": "supplied_project_constraints" if resolved_policy["typography_mode"] == "explicit" else "diagnostic_references_only",
                             "enforcement": "new_authoring_gate" if resolved_policy["mode"] == "new" else "legacy_read_only_advisory"}
        result["source_modified"] = hashlib.sha256(path.read_bytes()).hexdigest() != result["source_sha256"]
        if result["source_modified"]:
            raise ValueError("source changed during audit; rerun against an immutable candidate")
    except Exception as error:
        result["status"] = "errors"
        result["errors"].append({"code": "AUDIT_ERROR", "type": type(error).__name__, "message": str(error)})
    return result
