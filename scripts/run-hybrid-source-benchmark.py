#!/usr/bin/env python3
"""Run three internal, source-bound template-edit regressions on one real PPTX."""
from __future__ import annotations

import hashlib
import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation

from aippt_studio.local_preview import render_slide_preview
from aippt_studio.motion_clone import (
    _read_xml,
    _safe_entries,
    _slide_objects,
    _slide_order,
    text_sha256,
)
from aippt_studio.pptx_edit import patch_pptx
from aippt_studio.style_hybrid_workflow import build_hybrid_candidate, _sha
from aippt_studio.template_content_replace import (
    _image_refs,
    inspect_page_images,
    replace_content,
)


ROOT = Path(__file__).resolve().parents[1]
LIBRARY = ROOT / "模板库"
OUT = ROOT / "diagnostics/studio-audit-2026-10-08/production-benchmark"
SOURCE_REL = "创赛/04-绿色-乡村振兴、科技/16-柑心科技.pptx"
SOURCE_SHA = "cac2f6cdfa730843453b5bd5931bf6fd3be80844ce7529ba61bf8162f745513c"
DONOR_REL = "创赛/01-主图款/07-保鲜膜.pptx"
DONOR_SHA = "4db535ff097fbaf2549b5d8c415598f6530a99a2cd39501ac3519c22a26786a0"


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def package_parts(path: Path) -> dict[str, bytes]:
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def slides_and_texts(path: Path) -> tuple[list[str], list[list[str]]]:
    texts_by_slide: list[list[str]] = []
    with ZipFile(path) as archive:
        entries = _safe_entries(archive)
        _, slide_parts = _slide_order(archive, entries)
        for part in slide_parts:
            root = _read_xml(archive, entries, part)
            texts = []
            for shape in _slide_objects(root).values():
                if shape.tag.rsplit("}", 1)[-1] != "sp":
                    continue
                text = "".join(shape.xpath(
                    ".//a:t/text()",
                    namespaces={"a": "http://schemas.openxmlformats.org/drawingml/2006/main"},
                )).strip()
                if text:
                    texts.append(text)
            texts_by_slide.append(texts)
    return slide_parts, texts_by_slide


def object_counts(path: Path, slide_number: int) -> dict[str, int]:
    presentation = Presentation(path)
    slide = presentation.slides[slide_number - 1]
    return {
        "top_level_objects": len(slide.shapes),
        "text_objects": sum(1 for item in slide.shapes if item.has_text_frame and item.text.strip()),
        "picture_objects": sum(1 for item in slide.shapes if item.shape_type == 13),
        "editable_text_and_shapes": sum(
            1 for item in slide.shapes if item.has_text_frame or item.shape_type == 1
        ),
    }


def render_target(deck: Path, case_dir: Path, slide_number: int) -> dict:
    image = case_dir / f"slide-{slide_number}.png"
    receipt = render_slide_preview(deck, slide_number, image, width=1280, load_system_fonts=True)
    return {
        "png": str(image.relative_to(ROOT)),
        "png_sha256": receipt["output_sha256"],
        "receipt": str(image.with_suffix(".preview.json").relative_to(ROOT)),
        "renderer": receipt["renderer"],
        "approximate": receipt["approximate_renderer"],
        "system_font_fallback": receipt.get("system_font_fallback"),
    }


def changed_parts(before: dict[str, bytes], after: dict[str, bytes]) -> list[dict[str, str]]:
    names = sorted(set(before) | set(after))
    return [
        {"part": name, "before_sha256": sha(before[name]) if name in before else None,
         "after_sha256": sha(after[name]) if name in after else None}
        for name in names if before.get(name) != after.get(name)
    ]


def preservation(before: dict[str, bytes], after: dict[str, bytes], target_slide: int) -> dict:
    slide_parts = sorted(
        name for name in before if name.startswith("ppt/slides/slide") and name.endswith(".xml")
    )
    untouched = [name for index, name in enumerate(slide_parts, 1) if index != target_slide]
    return {
        "non_target_slide_xml_byte_identical": all(before[name] == after.get(name) for name in untouched),
        "non_target_slide_parts": untouched,
        "non_target_slide_part_sha256": {
            name: sha(before[name]) for name in untouched
        },
    }


def save_receipt(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def fresh_case_dir(path: Path) -> Path:
    resolved = path.resolve()
    if not resolved.is_relative_to(OUT.resolve()) or resolved == OUT.resolve():
        raise RuntimeError(f"Refusing to clear a path outside the benchmark directory: {path}")
    if path.exists() and (not path.is_dir() or any(path.iterdir())):
        raise RuntimeError(f"Existing benchmark evidence must be preserved; choose a fresh case directory: {path}")
    path.mkdir(parents=True, exist_ok=True)
    return path


def extract_image(source: Path, slide_number: int, shape_id: str, output: Path) -> tuple[str, str]:
    with ZipFile(source) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        part = slides[slide_number - 1]
        root = _read_xml(archive, entries, part)
        match = next(
            (item for item in _image_refs(archive, entries, part, root, slide_number)
             if item["shape_id"] == shape_id),
            None,
        )
        if match is None:
            raise RuntimeError(f"No embedded image for slide {slide_number} shape {shape_id}")
        data = archive.read(match["media_part"])
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(data)
        return match["media_part"], sha(data)


def convert_to_png(source: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as image:
        image.convert("RGBA").save(output, format="PNG", optimize=True)


def retry_layout_swap_only() -> dict:
    """Retry only the previously blocked page graft with explicit donor-note handling."""
    source, donor = LIBRARY / SOURCE_REL, LIBRARY / DONOR_REL
    if _sha(source) != SOURCE_SHA or _sha(donor) != DONOR_SHA:
        raise RuntimeError("A benchmark source changed; retry stopped")
    case_dir = OUT / "02-one-page-layout-swap"
    prior = case_dir / "attempt-1-blocked.json"
    if not prior.is_file():
        prior = OUT / "02-attempt-1-blocked.json"
    if not prior.is_file():
        raise RuntimeError("The initial blocked attempt must be preserved before retry")
    candidate = case_dir / "candidate.pptx"
    if candidate.exists():
        raise RuntimeError(f"Refusing to overwrite existing retry candidate: {candidate}")
    inputs = OUT / "assets/hybrid-inputs"
    inputs.mkdir(parents=True, exist_ok=True)
    donor_photo_map = {
        "14": (1, "4", "上游种植：果树挂果实景"),
        "18": (1, "5", "中游采收：果园鲜果直采场景"),
        "20": (4, "383", "下游产品：柑橘果品与包装场景"),
    }
    paths = {}
    for donor_shape_id, (source_slide, source_shape, _) in donor_photo_map.items():
        raw = inputs / f"source-slide-{source_slide}-shape-{source_shape}.jpg"
        extract_image(source, source_slide, source_shape, raw)
        png = inputs / f"donor-card-{donor_shape_id}.png"
        convert_to_png(raw, png)
        paths[donor_shape_id] = png
    text_values = {
        "1514": "产业协同", "1666": "项目概况", "1665": "产业链条",
        "1685": "市场机遇", "1686": "实施路径",
        "22": "内部示例：以标准化果品和产地管理稳定原料品质。",
        "38": "上游：优质果园", "43": "内部示例：鲜果直供周边渠道。",
        "55": "中游：产地直采", "1800": "内部示例：按场景组织产品并提升渠道触达。",
        "1803": "下游：品牌与渠道", "1818": "协同增值",
    }
    donor_images = {item["shape_id"]: item for item in inspect_page_images(donor, 4)["objects"]
                    if item["layer"] == "slide"}
    image_edits = [{
        "slide": 2, "shape_id": shape_id,
        "original_media_sha256": donor_images[shape_id]["media"][0]["media_sha256"],
        "image_path": str(paths[shape_id]), "purpose": purpose, "visual_role": "illustration",
    } for shape_id, (_, _, purpose) in donor_photo_map.items()]
    request = {
        "primary_source_path": SOURCE_REL, "primary_source_sha256": SOURCE_SHA,
        "project": {"occasion": "创新创业竞赛答辩", "track": "农业科技与乡村振兴",
                    "topic": "柑橘产业链结构示例", "style": "绿色",
                    "collection": "creative", "desired_slide_count": 5},
        "reviewed_primary_pages": [
            {"slide": 1, "finding": "绿色大字封面以项目名为唯一视觉焦点。",
             "preview_receipt": str((OUT / "source-inspection/source-slide-1.preview.json").resolve())},
            {"slide": 2, "finding": "市场页将增长曲线与上中下游关系并置。",
             "preview_receipt": str((OUT / "source-inspection/source-slide-2.preview.json").resolve())},
        ],
        "selection_note": "用三栏驱动页替换拥挤的市场图表页，验证一页供体及关系闭包导入。",
        "source_content_replacement_plan": "仅填入显式内部示例标题、分段说明和柑橘场景图；不是完整客户事实改写。",
        "rights_handling_note": "只在内部候选中复用两份模板包内图片；版权授权未核实，不作外部交付。",
        "page_overrides": [{
            "target_slide": 2, "donor_source_path": DONOR_REL,
            "donor_source_sha256": DONOR_SHA, "donor_slide": 4,
            "selection_note": "三列卡片承载市场驱动、核心说明与图片，页面功能可替代原市场分析页。",
            "target_page": {"role": "market-analysis",
                            "message": "以清晰的信息层级呈现市场增长数据、产业链和核心机会"},
            "omit_donor_notes": True,
            "donor_notes_disposition": "仅省略供体旧讲者备注；job-bound后续须按当前plan重建备注。",
            "reviewed_donor_page": {
                "slide": 4,
                "finding": "来源页以三栏呈现市场驱动因素、说明文字和对应场景图片。",
                "preview_receipt": str((OUT / "source-inspection/donor-film-slide-4.preview.json").resolve()),
            },
        }],
        "text_replacements": [
            {**item, "slide": 2} for item in text_map(donor, 4, text_values)
        ],
        "image_replacements": image_edits,
    }
    before = package_parts(source)
    record = {
        "case": "one_page_layout_swap", "status": "blocked",
        "scope": "primary slide 2 replaced by donor slide 4; mapped current page copy and 3 main photo slots; no full-deck factual cleanup",
        "source_path": SOURCE_REL, "source_sha256": SOURCE_SHA,
        "donor_path": DONOR_REL, "donor_sha256": DONOR_SHA, "target_slide": 2,
        "prior_blocked_attempt": str(prior.relative_to(ROOT)),
    }
    try:
        hybrid_receipt = build_hybrid_candidate(LIBRARY, request, candidate)
        after = package_parts(candidate)
        original_donor_texts = slides_and_texts(donor)[1][3]
        output_texts = slides_and_texts(candidate)[1][1]
        record.update({
            "status": "candidate_review_pending", "output": str(candidate.relative_to(ROOT)),
            "output_sha256": _sha(candidate), "editability": object_counts(candidate, 2),
            "changed_parts": changed_parts(before, after),
            "preservation": preservation(before, after, 2),
            "donor_text_remaining_on_target": [text for text in original_donor_texts if text in output_texts],
            "remaining_template_placeholders": any(
                marker in text for text in output_texts for marker in ("请在此处输入", "请用一句话", "关键词", "XXXXX", "标题")
            ),
            "full_page_content_adapted": False,
            "hybrid_build_receipt": hybrid_receipt,
        })
        record["preview"] = render_target(candidate, case_dir, 2)
    except (OSError, ValueError) as error:
        record["blocked_reason"] = str(error)
        record["output"] = None
        record["native_visual_review"] = "not_performed"
    retry_number = 1
    while (case_dir / f"retry-{retry_number:02d}-receipt.json").exists():
        retry_number += 1
    save_receipt(case_dir / f"retry-{retry_number:02d}-receipt.json", record)
    report_path = OUT / "benchmark-report.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    report["status"] = "partial_blocked_review_pending" if record["status"] == "blocked" else "internal_structural_regressions_candidate_review_pending"
    report["cases"] = [record if item.get("case") == "one_page_layout_swap" else item for item in report["cases"]]
    save_receipt(report_path, report)
    return record


def text_map(source: Path, slide_number: int, values: dict[str, str]) -> list[dict]:
    with ZipFile(source) as archive:
        entries = _safe_entries(archive)
        _, slides = _slide_order(archive, entries)
        root = _read_xml(archive, entries, slides[slide_number - 1])
        objects = _slide_objects(root)
        edits = []
        for shape_id, new_text in values.items():
            shape = objects.get(str(shape_id))
            if shape is None or shape.tag.rsplit("}", 1)[-1] != "sp":
                raise RuntimeError(f"Text shape {shape_id} is missing on page {slide_number}")
            original = "".join(shape.xpath(
                ".//a:t/text()",
                namespaces={"a": "http://schemas.openxmlformats.org/drawingml/2006/main"},
            ))
            edits.append({
                "slide": slide_number,
                "shape_id": str(shape_id),
                "original_text_sha256": text_sha256(original),
                "new_text": new_text,
            })
        return edits


def main() -> None:
    source = LIBRARY / SOURCE_REL
    donor = LIBRARY / DONOR_REL
    if _sha(source) != SOURCE_SHA:
        raise RuntimeError("Primary source hash changed; benchmark stopped")
    if _sha(donor) != DONOR_SHA:
        raise RuntimeError("Donor source hash changed; benchmark stopped")
    if not (OUT / "source-inspection/source-inspection.json").is_file():
        raise RuntimeError("Run source inspection and target selection before the benchmark")

    OUT.mkdir(parents=True, exist_ok=True)
    inputs = OUT / "assets/hybrid-inputs"
    inputs.mkdir(parents=True, exist_ok=True)
    source_before = package_parts(source)
    root_before_hash = _sha(source)
    records: list[dict] = []

    # Case 1: copy the actual template and replace only selected market copy/photo.
    case_dir = fresh_case_dir(OUT / "01-template-copy-text-image")
    candidate = case_dir / "candidate.pptx"
    replacement_photo = inputs / "citrus-tree-slide1-shape5.jpg"
    media_part, replacement_sha = extract_image(source, 1, "5", replacement_photo)
    source_images = inspect_page_images(source, 2)["objects"]
    target_image = next(item for item in source_images if item["shape_id"] == "42" and item["layer"] == "slide")
    copy_receipt = replace_content(
        source,
        candidate,
        source_sha256=SOURCE_SHA,
        text_replacements=text_map(source, 2, {
            "67": "市场机会",
            "51": "内部示例：果园标准化管理、采后分级和产地冷藏构成供给改善重点。",
            "59": "柑橘鲜果与加工产品",
        }),
        image_replacements=[{
            "slide": 2,
            "shape_id": "42",
            "original_media_sha256": target_image["media"][0]["media_sha256"],
            "path": str(replacement_photo),
        }],
    )
    case_after = package_parts(candidate)
    case_texts = slides_and_texts(candidate)[1][1]
    original_texts = slides_and_texts(source)[1]
    template_placeholder_markers = ("请在此处输入", "请用一句话", "关键词", "XXXXX", "标题")
    record = {
        "case": "template_copy_text_and_image",
        "scope": "slide 2: 3 text objects and one orchard photo; remaining template example content intentionally retained",
        "source_path": SOURCE_REL,
        "source_sha256": SOURCE_SHA,
        "output": str(candidate.relative_to(ROOT)),
        "output_sha256": _sha(candidate),
        "target_slide": 2,
        "editability": object_counts(candidate, 2),
        "changed_parts": changed_parts(source_before, case_after),
        "preservation": preservation(source_before, case_after, 2),
        "targeted_original_text_still_present": [text for text in original_texts[1] if text in case_texts],
        "remaining_template_placeholders": any(marker in text for text in case_texts for marker in template_placeholder_markers),
        "full_page_content_adapted": False,
        "new_photo": {"source_part": media_part, "sha256": replacement_sha,
                      "path": str(replacement_photo.relative_to(ROOT))},
        "replacement_receipt": copy_receipt,
    }
    record["preview"] = render_target(candidate, case_dir, 2)
    save_receipt(case_dir / "receipt.json", record)
    records.append(record)

    # Case 2: graft one market-driver page from a second indexed creative template.
    case_dir = fresh_case_dir(OUT / "02-one-page-layout-swap")
    candidate = case_dir / "candidate.pptx"
    donor_photo_map = {
        "14": (1, "4", "上游种植：果树挂果实景"),
        "18": (1, "5", "中游采收：果园鲜果直采场景"),
        "20": (4, "383", "下游产品：柑橘果品与包装场景"),
    }
    donor_image_paths: dict[str, Path] = {}
    for donor_shape_id, (source_slide, source_shape, _) in donor_photo_map.items():
        raw = inputs / f"source-slide-{source_slide}-shape-{source_shape}.jpg"
        extract_image(source, source_slide, source_shape, raw)
        png = inputs / f"donor-card-{donor_shape_id}.png"
        convert_to_png(raw, png)
        donor_image_paths[donor_shape_id] = png

    donor_texts = {
        "1514": "产业协同",
        "1666": "项目概况",
        "1665": "产业链条",
        "1685": "市场机遇",
        "1686": "实施路径",
        "22": "内部示例：以标准化果品和产地管理稳定原料品质。",
        "38": "上游：优质果园",
        "43": "内部示例：鲜果直供周边渠道。",
        "55": "中游：产地直采",
        "1800": "内部示例：按场景组织产品并提升渠道触达。",
        "1803": "下游：品牌与渠道",
        "1818": "协同增值",
    }
    text_edits = [{**item, "slide": 2} for item in text_map(donor, 4, donor_texts)]
    image_targets = {item["shape_id"]: item for item in inspect_page_images(donor, 4)["objects"]
                     if item["layer"] == "slide"}
    image_edits = [
        {
            "slide": 2,
            "shape_id": donor_shape_id,
            "original_media_sha256": image_targets[donor_shape_id]["media"][0]["media_sha256"],
            "image_path": str(image_path),
            "purpose": purpose,
            "visual_role": "illustration",
        }
        for donor_shape_id, (_, _, purpose) in donor_photo_map.items()
        for image_path in [donor_image_paths[donor_shape_id]]
    ]
    request = {
        "primary_source_path": SOURCE_REL,
        "primary_source_sha256": SOURCE_SHA,
        "project": {
            "occasion": "创新创业竞赛答辩",
            "track": "农业科技与乡村振兴",
            "topic": "柑橘产业链结构示例",
            "style": "绿色",
            "collection": "creative",
            "desired_slide_count": 5,
        },
        "reviewed_primary_pages": [
            {"slide": 1, "finding": "绿色大字封面以项目名为唯一视觉焦点。",
             "preview_receipt": str((OUT / "source-inspection/source-slide-1.preview.json").resolve())},
            {"slide": 2, "finding": "市场页将增长曲线与上中下游关系并置。",
             "preview_receipt": str((OUT / "source-inspection/source-slide-2.preview.json").resolve())},
        ],
        "selection_note": "用三栏驱动页替换拥挤的市场图表页，验证一页供体及关系闭包导入。",
        "source_content_replacement_plan": "仅为结构回归填入显式内部示例标题、分段说明和柑橘场景图；不是完整客户事实改写。",
        "rights_handling_note": "只在内部候选中复用两份模板包内图片；版权授权未核实，不作外部交付。",
        "page_overrides": [{
            "target_slide": 2,
            "donor_source_path": DONOR_REL,
            "donor_source_sha256": DONOR_SHA,
            "donor_slide": 4,
            "omit_donor_notes": True,
            "donor_notes_disposition": "仅省略供体旧讲者备注；job-bound后续须按当前plan重建备注。",
            "selection_note": "三列卡片承载市场驱动、核心说明与图片，页面功能可替代原市场分析页。",
            "target_page": {"role": "market-analysis",
            "message": "以清晰的信息层级呈现市场增长数据、产业链和核心机会"},
            "reviewed_donor_page": {
                "slide": 4,
                "finding": "来源页以三栏呈现市场驱动因素、说明文字和对应场景图片。",
                "preview_receipt": str((OUT / "source-inspection/donor-film-slide-4.preview.json").resolve()),
            },
        }],
        "text_replacements": text_edits,
        "image_replacements": image_edits,
    }
    record = {
        "case": "one_page_layout_swap",
        "scope": "primary slide 2 replaced by donor slide 4; mapped current page copy and 3 main photo slots; no full-deck factual cleanup",
        "source_path": SOURCE_REL,
        "source_sha256": SOURCE_SHA,
        "donor_path": DONOR_REL,
        "donor_sha256": DONOR_SHA,
        "target_slide": 2,
    }
    try:
        hybrid_receipt = build_hybrid_candidate(LIBRARY, request, candidate)
        case_after = package_parts(candidate)
        donor_texts_original = slides_and_texts(donor)[1][3]
        candidate_texts = slides_and_texts(candidate)[1][1]
        record.update({
            "status": "candidate_review_pending",
            "output": str(candidate.relative_to(ROOT)),
            "output_sha256": _sha(candidate),
            "editability": object_counts(candidate, 2),
            "changed_parts": changed_parts(source_before, case_after),
            "preservation": preservation(source_before, case_after, 2),
            "donor_text_remaining_on_target": [text for text in donor_texts_original if text in candidate_texts],
            "hybrid_build_receipt": hybrid_receipt,
        })
        record["preview"] = render_target(candidate, case_dir, 2)
    except (OSError, ValueError) as error:
        record.update({"status": "blocked", "blocked_reason": str(error),
                       "output": None, "native_visual_review": "not_performed",
                       "targeted_original_text_still_present": None,
                       "remaining_template_placeholders": None,
                       "full_page_content_adapted": False})
    save_receipt(case_dir / "receipt.json", record)
    records.append(record)

    # Case 3: one exact run edit in the final page, preserving all other slide XML.
    case_dir = fresh_case_dir(OUT / "03-one-slide-exact-edit")
    edit_receipt = patch_pptx({
        "source": str(source),
        "source_sha256": SOURCE_SHA,
        "edits": [{
            "action": "replace_text",
            "slide": 5,
            "shape_id": "76",
            "paragraph": 1,
            "run": 1,
            "expected_text": "项目成果",
            "replacement_text": "成果摘要",
        }],
        "dry_run": False,
    }, case_dir / "candidate")
    candidate = Path(edit_receipt["candidate"])
    case_after = package_parts(candidate)
    candidate_texts = slides_and_texts(candidate)[1][4]
    record = {
        "case": "single_slide_exact_edit",
        "scope": "slide 5 shape 76, paragraph 1 run 1 only",
        "source_path": SOURCE_REL,
        "source_sha256": SOURCE_SHA,
        "output": str(candidate.relative_to(ROOT)),
        "output_sha256": _sha(candidate),
        "target_slide": 5,
        "editability": object_counts(candidate, 5),
        "changed_parts": changed_parts(source_before, case_after),
        "preservation": preservation(source_before, case_after, 5),
        "targeted_original_text_still_present": [text for text in original_texts[4] if text in candidate_texts],
        "remaining_template_placeholders": any(marker in text for text in candidate_texts for marker in template_placeholder_markers),
        "full_page_content_adapted": False,
        "patch_receipt": edit_receipt,
    }
    record["preview"] = render_target(candidate, case_dir, 5)
    save_receipt(case_dir / "receipt.json", record)
    records.append(record)

    if _sha(source) != root_before_hash:
        raise RuntimeError("Primary source bytes changed during benchmark")
    report = {
        "schema": "aippt.template-source-benchmark.v1",
        "status": "partial_blocked_review_pending" if any(item.get("status") == "blocked" for item in records)
        else "internal_structural_regressions_candidate_review_pending",
        "source_path": SOURCE_REL,
        "source_sha256": SOURCE_SHA,
        "source_read_only_verified": True,
        "native_powerpoint_review": "not_performed",
        "candidate_acceptance": "not_claimed",
        "cases": records,
    }
    save_receipt(OUT / "benchmark-report.json", report)
    print(json.dumps({
        "status": report["status"],
        "report": str((OUT / "benchmark-report.json").relative_to(ROOT)),
        "cases": [{"case": item["case"], "output": item["output"],
                   "status": item.get("status", "candidate"),
                   "output_sha256": item.get("output_sha256"),
                   "non_target_slide_xml_byte_identical": (item.get("preservation") or {}).get("non_target_slide_xml_byte_identical"),
                   "blocked_reason": item.get("blocked_reason"),
                   "preview": (item.get("preview") or {}).get("png")} for item in records],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--retry-layout-only", action="store_true")
    args = parser.parse_args()
    if args.retry_layout_only:
        print(json.dumps(retry_layout_swap_only(), ensure_ascii=False, indent=2))
    else:
        main()
